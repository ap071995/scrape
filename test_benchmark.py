"""Offline benchmark regression tests; no browsers, network or live user files."""
import csv
import hashlib
import json
from pathlib import Path
import socket
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

import benchmark as b


def lead(number, **values):
    return {b.PROFILE: f"https://www.linkedin.com/in/person-{number}/",
            "First Name": "Person", "Last Name": str(number),
            "Company Name": f"Company {number % 17}",
            "Industry": ("Education", "Technology", "Healthcare")[number % 3],
            "Company Size": ("11-50 employees", "501-1,000 employees", "10,001+ employees")[number % 3],
            "Country": ("India", "United States", "Germany")[number % 3], **values}


class BenchmarkTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="scraper_benchmark_test_")
        self.root = Path(self.temp.name)

    def tearDown(self):
        self.temp.cleanup()

    def write_csv(self, name, rows, headers=None, delimiter=","):
        path = self.root / name
        headers = headers or list(dict.fromkeys(k for r in rows for k in r))
        with path.open("w", encoding="utf-8-sig", newline="") as stream:
            writer = csv.DictWriter(stream, headers, delimiter=delimiter)
            writer.writeheader()
            writer.writerows(rows)
        return path

    def workbook(self, name, rows, metadata=None):
        path = self.root / name
        book = Workbook()
        book.active.title = "Leads"
        for sheet, records in [(book.active, rows)] + [(book.create_sheet(k), v) for k, v in (metadata or {}).items()]:
            headers = list(dict.fromkeys(k for record in records for k in record)) or [b.PROFILE]
            sheet.append(headers)
            for record in records:
                sheet.append([record.get(h) for h in headers])
        book.save(path)
        book.close()
        return path

    def review(self, record, field, state="Value", expected="", **extra):
        return {"Record ID": b.canonical_profile(record[b.PROFILE]), "Field": field,
                "Reviewed": "Yes", "Expected State": state, "Expected Value": expected,
                "False Match": "No", **extra}

    def test_build_never_connects_and_preserves_source(self):
        path = self.workbook("leads.xlsx", [lead(1), lead(2)])
        before = path.read_bytes()
        with patch.object(socket, "socket", side_effect=AssertionError("Network forbidden")):
            result = b.build_benchmark(path, self.root / "out")
        self.assertEqual(path.read_bytes(), before)
        self.assertEqual(result["summary"]["input_sha256"], hashlib.sha256(before).hexdigest())
        self.assertEqual(result["summary"]["sample_size"], 2)
        self.assertEqual(result["summary"]["sample_shortfall"], 148)
        self.assertEqual(result["summary"]["reviewed_known_cells"], 0)
        self.assertIn("Not evaluated", Path(result["report"]).read_text())

    def test_deterministic_diverse_150_from_240_records(self):
        records, _ = b.prepare_records([lead(i) for i in range(240)])
        first = b.select_sample(records)
        second = b.select_sample(list(reversed(records)))
        self.assertEqual([r["_id"] for r in first], [r["_id"] for r in second])
        self.assertEqual(len(first), 150)
        self.assertEqual(len({r["_id"] for r in first}), 150)
        self.assertEqual(len({r["Country"] for r in first}), 3)
        self.assertNotEqual([r["_id"] for r in first], [r["_id"] for r in b.select_sample(records, seed=12)])

    def test_company_cap_prefers_variety_then_fills(self):
        records, _ = b.prepare_records([lead(i, **{"Company Name": "Common" if i < 90 else f"Rare {i}"}) for i in range(100)])
        sample = b.select_sample(records, sample_size=15, company_cap=1)
        self.assertEqual(len(sample), 15)
        self.assertEqual(len({r["Company Name"] for r in sample}), 11)

    def test_invalid_sample_or_cap_rejected(self):
        for size in (0, 201, True, "150", -1):
            with self.subTest(size=size), self.assertRaises(ValueError):
                b.select_sample([], sample_size=size)
        for cap in (0, True, "3"):
            with self.subTest(cap=cap), self.assertRaises(ValueError):
                b.select_sample([], company_cap=cap)

    def test_canonical_profile_dedup_keeps_complete_row_without_merging(self):
        rows = [lead(1, **{"Company Name": "Old", "Revenue": "$1M"}),
                lead(1, **{b.PROFILE: "http://in.linkedin.com/in/PERSON-1/?trk=test", "Company Name": "New", "City": "Pune", "State": "Maharashtra"})]
        records, duplicates = b.prepare_records(rows)
        self.assertEqual(duplicates, 1)
        self.assertEqual(len(records), 1)
        self.assertEqual(records[0]["Company Name"], "New")
        self.assertNotIn("Revenue", records[0])

    def test_invalid_profiles_not_conflated(self):
        records, duplicates = b.prepare_records([lead(1, **{b.PROFILE: "-"}), lead(2, **{b.PROFILE: "https://notlinkedin.com/in/x"})])
        self.assertEqual(duplicates, 0)
        self.assertEqual(len({r["_id"] for r in records}), 2)
        self.assertTrue(all(not r["_profile"] for r in records))

    def test_sales_navigator_context_dedup_and_review_restart(self):
        rows = [lead(1, **{b.PROFILE: "http://uk.linkedin.com/sales/lead/ACwAAbC1,NAME_SEARCH,abc?trk=123"}),
                lead(1, **{b.PROFILE: "https://www.linkedin.com/sales/lead/ACwAAbC1,SEARCH,xyz?test=1", "Revenue": "$1M"})]
        source = self.write_csv("sales.csv", rows)
        first = b.build_benchmark(source, self.root)
        self.assertEqual(first["summary"]["duplicates_removed"], 1)
        key = "https://www.linkedin.com/sales/lead/ACwAAbC1"
        self.assertEqual(first["summary"]["sample_ids"], [key])
        review = self.write_csv("sales-review.csv", [self.review(rows[0], "Revenue", expected="$1M")])
        second = b.build_benchmark(source, self.root, review_path=review)
        self.assertEqual(second["summary"]["reviewed_known_cells"], 1)
        self.assertNotEqual(key, b.canonical_profile("https://www.linkedin.com/sales/lead/acwaabc1"))

    def test_blank_rows_keep_original_source_row(self):
        path = self.write_csv("blank.csv", [lead(1), {}, lead(2)])
        rows, _ = b.read_saved(path)
        records, _ = b.prepare_records(rows)
        self.assertEqual([r["_row"] for r in records], [2, 4])

    def test_header_only_empty_dataset_supported(self):
        path = self.write_csv("empty.csv", [], [b.PROFILE])
        result = b.build_benchmark(path, self.root)
        self.assertEqual(result["summary"]["sample_size"], 0)
        fields = json.loads(Path(result["json"]).read_text())["fields"]
        self.assertTrue(all(f["Coverage"] is None and f["Accuracy"] is None for f in fields))

    def test_wrong_or_duplicate_headers_rejected_even_empty(self):
        for headers in (["Wrong"], [b.PROFILE, b.PROFILE]):
            path = self.write_csv("invalid.csv", [], headers)
            with self.assertRaises(ValueError):
                b.build_benchmark(path, self.root)

    def test_tsv_supported(self):
        path = self.write_csv("leads.tsv", [lead(1)], delimiter="\t")
        rows, _ = b.read_saved(path)
        self.assertEqual(rows[0]["First Name"], "Person")

    def test_unsupported_file_rejected(self):
        path = self.root / "leads.xls"
        path.write_text("not excel")
        with self.assertRaises(ValueError):
            b.read_saved(path)

    def test_coverage_is_not_accuracy(self):
        records, _ = b.prepare_records([lead(1, **{"Revenue": "$1M"}), lead(2)])
        field = next(r for r in b.field_results(records, {}, "sample") if r["Field"] == "Revenue")
        self.assertEqual(field["Coverage"], .5)
        self.assertEqual(field["Reviewed Known"], 0)
        self.assertIsNone(field["Accuracy"])

    def test_explicit_absence_unknown_and_false_match_denominators(self):
        raw = [lead(i) for i in range(5)]
        raw[0]["Revenue"] = " $1M "
        raw[1]["Revenue"] = "$2M"
        records, _ = b.prepare_records(raw)
        reviews = {
            (records[0]["_id"], "Revenue"): self.review(raw[0], "Revenue", expected="$1m"),
            (records[1]["_id"], "Revenue"): self.review(raw[1], "Revenue", expected="$2M", **{"False Match": "Yes"}),
            (records[2]["_id"], "Revenue"): self.review(raw[2], "Revenue", "Confirmed absent"),
            (records[3]["_id"], "Revenue"): self.review(raw[3], "Revenue", "Unknown"),
        }
        field = next(r for r in b.field_results(records, reviews, "sample") if r["Field"] == "Revenue")
        self.assertEqual((field["Reviewed Known"], field["Correct"], field["Incorrect"], field["Unknown"], field["False Matches"]), (3, 2, 1, 2, 1))
        self.assertAlmostEqual(field["Accuracy"], 2 / 3)

    def test_reviews_require_valid_schema_and_explicit_state(self):
        source = lead(1)
        for values in ({"Expected State": "Value", "Expected Value": ""},
                       {"Expected State": "Confirmed absent", "Expected Value": "$1M"},
                       {"Expected State": "Maybe"}, {"Reviewed": "Y"},
                       {"False Match": "perhaps"}, {"Field": "Invented"}):
            with self.subTest(values=values):
                path = self.write_csv("review.csv", [self.review(source, "Revenue", expected="$1M", **values)] if "Expected Value" not in values else [{**self.review(source, "Revenue", expected="$1M"), **values}])
                with self.assertRaises(ValueError):
                    b.read_reviews(path)
        with self.assertRaises(ValueError):
            b.read_reviews(self.workbook("wrong-sheet.xlsx", [source]))

    def test_duplicate_reviews_rejected(self):
        row = self.review(lead(1), "Revenue", expected="$1M")
        with self.assertRaises(ValueError):
            b.read_reviews(self.write_csv("dup.csv", [row, row]))

    def test_review_input_and_earlier_output_preserved(self):
        source = self.write_csv("leads.csv", [lead(1, **{"Revenue": "$1M"})])
        first = b.build_benchmark(source, self.root)
        book = load_workbook(first["workbook"])
        sheet = book["Review"]
        for row in sheet.iter_rows(min_row=2):
            if row[3].value == "Revenue":
                row[5].value, row[6].value, row[7].value = "Yes", "Value", "$1M"
        book.save(first["workbook"])
        book.close()
        before = Path(first["workbook"]).read_bytes()
        second = b.build_benchmark(source, self.root, review_path=first["workbook"])
        self.assertNotEqual(first["output_dir"], second["output_dir"])
        self.assertEqual(Path(first["workbook"]).read_bytes(), before)
        self.assertEqual(second["summary"]["reviewed_known_cells"], 1)

    def test_baseline_changes_and_reviewed_new_values(self):
        old = [lead(1), lead(2, **{"Revenue": "$1M"}), lead(3, **{"Revenue": "$2M"})]
        new = [lead(1, **{"Revenue": "$3M"}), lead(2, **{"Revenue": "$4M"}), lead(3)]
        baseline = self.write_csv("before.csv", old)
        source = self.write_csv("after.csv", new)
        reviews = self.write_csv("reviews.csv", [self.review(new[0], "Revenue", expected="$3M")])
        baseline_bytes = baseline.read_bytes()
        result = b.build_benchmark(source, self.root, baseline_path=baseline, review_path=reviews)
        self.assertEqual(result["summary"]["baseline_change_counts"], {"Newly filled": 1, "Changed value": 1, "Became missing": 1})
        self.assertEqual(result["summary"]["newly_filled_review_results"], {"Correct": 1})
        self.assertEqual(baseline.read_bytes(), baseline_bytes)

    def test_unmatched_reviews_reported(self):
        source = self.write_csv("source.csv", [lead(1)])
        review = self.write_csv("reviews.csv", [self.review(lead(2), "Revenue", expected="$1M")])
        result = b.build_benchmark(source, self.root, review_path=review)
        self.assertEqual(result["summary"]["unmatched_review_cells"], 1)
        self.assertEqual(result["summary"]["reviewed_known_cells"], 0)

    def test_email_provenance_join_requires_profile_and_company(self):
        raw = [lead(1), lead(2)]
        details = [{b.PROFILE: raw[0][b.PROFILE], "Company Name": raw[0]["Company Name"],
                    "Published Email ID": "person@example.com", "Domain/MX Status": "MX found", "Mailbox Verification": "Not verified"},
                   {b.PROFILE: raw[1][b.PROFILE], "Company Name": "Wrong", "Generated Email Candidate": "other@example.com"}]
        records, _ = b.prepare_records(raw, {"Email Details": details})
        self.assertEqual(records[0]["Published Email ID"], "person@example.com")
        self.assertEqual(records[0]["Mailbox Verification"], "Not verified")
        self.assertNotIn("Generated Email Candidate", records[1])
        self.assertNotIn("Generated Email Candidate", records[0])

    def test_legacy_generated_address_stays_separate(self):
        records, _ = b.prepare_records([lead(1, **{"Business Email ID": "p@example.com", "Email Status": "Generated (unverified)"}),
                                        lead(2, **{"Business Email ID": "q@example.com"})])
        self.assertEqual(records[0]["Generated Email Candidate"], "p@example.com")
        self.assertNotIn("Published Email ID", records[0])
        self.assertNotIn("Published Email ID", records[1])
        self.assertNotIn("Generated Email Candidate", records[1])

    def test_legacy_email_status_requires_exact_application_provenance(self):
        from business_email import build_email_details, FOUND, GENERATED
        for status in ("Not published", "Not generated", "Published lookup failed",
                       "Previously generated, now manually provided", "Verified by user",
                       FOUND, GENERATED):
            with self.subTest(status=status):
                raw = lead(1, **{"Business Email ID": "person@example.com", "Email Status": status})
                expected = build_email_details(raw)
                record = b.prepare_records([raw])[0][0]
                for field, detail in (("Published Email ID", "published_email"),
                                      ("Generated Email Candidate", "generated_email")):
                    self.assertEqual(record.get(field, ""), expected[detail])
                    measured = next(row for row in b.field_results([record], {}, "sample")
                                    if row["Field"] == field)
                    self.assertEqual(measured["Filled"], int(bool(expected[detail])))

    def test_changed_lead_email_rejects_stale_companion_address_and_dns(self):
        raw = lead(1, **{"Business Email ID": "new.person@example.com",
                        "Email Status": "Found online (unverified)"})
        for companion in (
                {"Published Email ID": "old.person@old-domain.com", "Email Domain": "old-domain.com"},
                {"Published Email ID": "new.person@example.com", "Email Domain": "old-domain.com"},
                {"Email Domain": "example.com"}):
            with self.subTest(companion=companion):
                detail = {b.PROFILE: raw[b.PROFILE], "Company Name": raw["Company Name"],
                          "Domain/MX Status": "MX records found", **companion}
                record = b.prepare_records([raw], {"Email Details": [detail]})[0][0]
                self.assertEqual(record["Published Email ID"], "new.person@example.com")
                self.assertNotIn("Domain/MX Status", record)

    def test_matching_companion_email_retains_its_dns_context(self):
        raw = lead(1, **{"Business Email ID": "person@example.com",
                        "Email Status": "Found online (unverified)"})
        detail = {b.PROFILE: raw[b.PROFILE], "Company Name": raw["Company Name"],
                  "Published Email ID": "PERSON@EXAMPLE.COM", "Email Domain": "example.com",
                  "Domain/MX Status": "MX records found"}
        record = b.prepare_records([raw], {"Email Details": [detail]})[0][0]
        self.assertEqual(record["Published Email ID"], "PERSON@EXAMPLE.COM")
        self.assertEqual(record["Domain/MX Status"], "MX records found")

    def test_maps_scope_kept_separate_from_person_location(self):
        raw = lead(1, City="Pune", Address="Original")
        entry = {b.PROFILE: raw[b.PROFILE], "Company Name": raw["Company Name"],
                 "Location Type": "Headquarters", "Company Location Address": "123 Main St, Boston",
                 "Company Location Phone": "+16175551234", "Prospect Location": "Pune, India",
                 "Prospect Works Here Verified": False, "Eligible for Merge": "Yes", "Source URL": "https://maps.google.com/?cid=123"}
        records, _ = b.prepare_records([raw], {"Maps Evidence": [entry]})
        self.assertEqual(records[0]["Company Location Type"], "Headquarters")
        self.assertEqual(records[0]["City"], "Pune")
        self.assertEqual(records[0]["Address"], "Original")
        self.assertIs(records[0]["Prospect Works Here Verified"], False)
        self.assertIn("Company Location Address", records[0]["_evidence"])

    def test_held_maps_candidates_do_not_increase_extracted_coverage(self):
        raw = lead(1)
        for status in ("No", "", None):
            with self.subTest(status=status):
                entry = {b.PROFILE: raw[b.PROFILE], "Company Name": raw["Company Name"],
                         "Location Type": "Headquarters", "Company Location Address": "123 Main St, Boston",
                         "Eligible for Merge": status, "Outcome / Reason": "Conflicting prospect locality"}
                records, _ = b.prepare_records([raw], {"Maps Evidence": [entry]})
                self.assertNotIn("Company Location Address", records[0])
                self.assertEqual(b.missing_reason(records[0], "Company Location Address"), "Conflicting prospect locality")
                result = next(r for r in b.field_results(records, {}, "sample") if r["Field"] == "Company Location Address")
                self.assertEqual(result["Filled"], 0)

    def test_company_evidence_cannot_join_same_name_different_website(self):
        raw = lead(1, **{"Company Website": "https://example.com", "Revenue": "$1M"})
        evidence = {"Company Name": raw["Company Name"], "Company Website": "https://other.com", "Field": "Revenue", "Value": "$1M", "Source URL": "https://other.com/report"}
        records, _ = b.prepare_records([raw], {"Company Evidence": [evidence]})
        self.assertNotIn("Revenue", records[0]["_evidence"])

    def test_missing_reason_only_claims_recorded_cause(self):
        raw = lead(1)
        records, _ = b.prepare_records([raw], {"Field Status": [{b.PROFILE: raw[b.PROFILE], "Field": "Revenue", "Reason": "Page blocked"}]})
        self.assertEqual(b.missing_reason(records[0], "Revenue"), "Page blocked")
        self.assertEqual(b.missing_reason(records[0], "NAICS Code"), "No evidence recorded")

    def test_formula_and_illegal_control_values_export_as_safe_literals(self):
        raw = lead(1, **{"Company Name": '=HYPERLINK("https://example.com")', "Phone Number": "+12345678", "Revenue": "a\x01b"})
        result = b.build_benchmark(self.write_csv("source.csv", [raw]), self.root)
        book = load_workbook(result["workbook"], data_only=False)
        try:
            sample = book["Sample"]
            header = {cell.value: cell.column for cell in sample[1]}
            self.assertEqual(sample.cell(2, header["Company Name"]).data_type, "s")
            self.assertEqual(sample.cell(2, header["Phone Number"]).value, "+12345678")
            self.assertEqual(sample.cell(2, header["Revenue"]).value, "ab")
            self.assertTrue(all(c.data_type != "f" for sheet in book for row in sheet for c in row))
        finally:
            book.close()
        with (Path(result["output_dir"]) / "review_template.csv").open(encoding="utf-8-sig", newline="") as stream:
            rows = list(csv.DictReader(stream))
        phone = next(row for row in rows if row["Field"] == "Phone Number")
        self.assertEqual(phone["Actual Value"], "'+12345678")
        phone.update({"Reviewed": "Yes", "Expected State": "Value", "Expected Value": phone["Actual Value"]})
        parsed = b.read_reviews(self.write_csv("review-copy.csv", [phone]))
        self.assertEqual(next(iter(parsed.values()))["Expected Value"], "+12345678")

    def test_all_fields_have_consistent_result_rows(self):
        result = b.build_benchmark(self.write_csv("source.csv", [lead(1)]), self.root)
        data = json.loads(Path(result["json"]).read_text(encoding="utf-8"))
        self.assertEqual(len(data["fields"]), 2 * len(b.FIELDS))
        book = load_workbook(result["workbook"])
        try:
            self.assertEqual(book["Review"].max_row, 1 + len(b.FIELDS))
            self.assertEqual(len(book["Review"].data_validations.dataValidation), 3)
            self.assertEqual(book["Sample"].freeze_panes, "C2")
            self.assertIn("Field Evidence", book.sheetnames)
        finally:
            book.close()


if __name__ == "__main__":
    unittest.main()
