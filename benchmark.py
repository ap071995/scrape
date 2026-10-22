"""Offline sampling and review of saved scraper output. No browser or network calls.

The sample deliberately increases variety; its accuracy is not a population estimate.
Only explicit human reviews enter accuracy denominators. A filled cell is coverage,
not evidence of correctness. The portable application already depends on openpyxl.
"""
from __future__ import annotations

import argparse
from collections import Counter, defaultdict
import csv
from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path
import random
import re
import tempfile
from urllib.parse import unquote, urlsplit

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation


PROFILE = "Prospect LinkedIn Link"
FIELDS = (
    "First Name", "Last Name", "Tenure", "Job Title", "Company Name",
    "Company Website", "Industry", "Company Size", "City", "State", "Country",
    "Address", "Zip Code", "Phone Number", "Revenue", "SIC Code", "NAICS Code",
    "Business Email ID", "Published Email ID", "Generated Email Candidate",
    "Company Location Type", "Company Location Address", "Company Location Phone",
)
CONTEXT_FIELDS = ("Email Status", "Domain/MX Status", "Mailbox Verification",
                  "Existing Email (origin unknown)", "Prospect Location",
                  "Prospect Works Here Verified", "Location Basis", "Classification Evidence",
                  "Maps Outcome / Reason")
MISSING = {"", "-", "none", "null", "nan", "n/a", "na", "unknown", "not available"}
REVIEW_HEADERS = (
    "Record ID", PROFILE, "Company Name", "Field", "Actual Value", "Reviewed",
    "Expected State", "Expected Value", "False Match", "Review Source", "Review Notes",
)
RESULT_HEADERS = (
    "Scope", "Field", "Records", "Filled", "Missing", "Coverage",
    "Reviewed Known", "Correct", "Incorrect", "Unknown", "Accuracy",
    "False Matches", "Missing Reasons",
)
LINK_FIELDS = {
    "Company Website": "Industry Link", "Industry": "Industry Link",
    "Company Size": "Industry Link", "Address": "Address Link",
    "Zip Code": "Address Link", "Phone Number": "Phone Number Link",
    "Revenue": "Revenue Link", "SIC Code": "SIC/NAICS Code Link",
    "NAICS Code": "SIC/NAICS Code Link",
}


def text(value):
    return "" if value is None else str(value).strip()


def missing(value):
    return text(value).casefold() in MISSING


def canonical_profile(value):
    """Country subdomains, http, query parameters and trailing slashes share a key."""
    raw = text(value)
    if missing(raw):
        return ""
    try:
        parsed = urlsplit(raw if "://" in raw else "https://" + raw)
        host = (parsed.hostname or "").casefold().rstrip(".")
        path = unquote(parsed.path).rstrip("/")
        if (parsed.scheme in ("https", "http") and not parsed.username
                and (host == "linkedin.com" or host.endswith(".linkedin.com"))):
            if re.fullmatch(r"/in/[^/\s]+", path, re.I):
                return "https://www.linkedin.com" + path.casefold()
            # Sales Navigator IDs are opaque and case-sensitive. Comma suffixes
            # describe search context, not the prospect's identity.
            if path.startswith("/sales/lead/"):
                path = path.split(",", 1)[0]
                if re.fullmatch(r"/sales/lead/[^/\s,]+", path):
                    return "https://www.linkedin.com" + path
    except ValueError:
        pass
    return ""


def _table(rows, label, required=()):
    rows = iter(rows)
    headers = [text(v) for v in next(rows, ())]
    if not headers or not any(headers):
        raise ValueError(f"{label}: missing header row")
    nonempty = [h for h in headers if h]
    if len(nonempty) != len(set(nonempty)):
        raise ValueError(f"{label}: duplicate column names")
    if any(header not in headers for header in required):
        raise ValueError(f"{label}: missing required columns: {', '.join(required)}")
    return [{**dict(zip(headers, values)), "_source_row": index}
            for index, values in enumerate(rows, 2) if any(text(v) for v in values)]


def read_saved(path, sheet="Leads", required=(), require_sheet=False):
    """Read literals/formula text without evaluating formulas or refreshing links."""
    path = Path(path).expanduser().resolve(strict=True)
    if path.suffix.lower() in (".csv", ".tsv"):
        with path.open("r", encoding="utf-8-sig", newline="") as stream:
            records = _table(csv.reader(stream, delimiter="\t" if path.suffix.lower() == ".tsv" else ","), path.name, required)
        return records, {}
    if path.suffix.lower() != ".xlsx":
        raise ValueError("Input must be a saved .xlsx, .csv or .tsv file")
    book = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    try:
        if require_sheet and sheet not in book.sheetnames:
            raise ValueError(f"{path.name}: missing '{sheet}' worksheet")
        selected = book[sheet] if sheet in book.sheetnames else book.worksheets[0]
        records = _table(selected.values, selected.title, required)
        metadata = {}
        for name in ("Email Details", "Email Evidence", "Company Evidence", "Maps Evidence", "Field Status"):
            if name in book.sheetnames and name != selected.title:
                metadata[name] = _table(book[name].values, name)
        return records, metadata
    finally:
        book.close()


def _normal(value):
    return re.sub(r"\s+", " ", text(value)).casefold()


def _evidence_url(value):
    try:
        parsed = urlsplit(text(value))
        return parsed.scheme in ("https", "http") and bool(parsed.hostname) and not parsed.username
    except ValueError:
        return False


def prepare_records(raw_records, metadata=None):
    """Keep the most complete duplicate; never merge conflicting source rows."""
    metadata = metadata or {}
    details = defaultdict(list)
    for entry in metadata.get("Email Details", ()):
        key = canonical_profile(entry.get(PROFILE))
        if key:
            details[key].append(entry)
    email_evidence = defaultdict(list)
    for entry in metadata.get("Email Evidence", ()):
        key = canonical_profile(entry.get(PROFILE))
        if key:
            email_evidence[key].append(entry)
    company_evidence = defaultdict(list)
    for entry in metadata.get("Company Evidence", ()):
        company = _normal(entry.get("Company Name"))
        if not missing(company):
            company_evidence[company].append(entry)
    maps_evidence = defaultdict(list)
    for entry in metadata.get("Maps Evidence", ()):
        key = canonical_profile(entry.get(PROFILE))
        if key and not missing(entry.get("Company Name")):
            maps_evidence[key].append(entry)
    field_status = defaultdict(dict)
    for entry in metadata.get("Field Status", ()):
        key, field = canonical_profile(entry.get(PROFILE)), text(entry.get("Field"))
        if key and field:
            field_status[key][field] = text(entry.get("Reason") or entry.get("Status"))
    prepared, best = [], {}
    duplicates = 0
    for index, source in enumerate(raw_records, 2):
        record = dict(source)
        key = canonical_profile(record.get(PROFILE))
        index = source.get("_source_row", index)
        record["_row"] = index
        record["_id"] = key or "row:" + str(index) + ":" + hashlib.sha256(
            json.dumps(source, sort_keys=True, default=str).encode()).hexdigest()[:12]
        record["_profile"] = key
        record["_evidence"] = {}
        record["_reasons"] = field_status.get(key, {}).copy()
        for entry in details.get(key, ()):
            if (missing(record.get("Company Name")) or
                    _normal(entry.get("Company Name")) != _normal(record.get("Company Name"))):
                continue
            # Leads may have been corrected in Excel after the companion sheet
            # was written. Do not attach a previous address's origin or DNS.
            current_email = _normal(record.get("Business Email ID"))
            if not missing(current_email):
                recorded_emails = [_normal(entry.get(field)) for field in (
                    "Published Email ID", "Generated Email Candidate", "Existing Email (origin unknown)")
                    if not missing(entry.get(field))]
                if not recorded_emails or any(value != current_email for value in recorded_emails):
                    continue
                recorded_domain = _normal(entry.get("Email Domain"))
                if (not missing(recorded_domain)
                        and recorded_domain != current_email.rsplit("@", 1)[-1]):
                    continue
            for field in ("Published Email ID", "Generated Email Candidate", "Existing Email (origin unknown)",
                          "Domain/MX Status", "Mailbox Verification", "Email Status"):
                if not missing(entry.get(field)):
                    record[field] = entry[field]
        email = text(record.get("Business Email ID"))
        status = _normal(record.get("Email Status"))
        # Legacy files retain their explicit recorded provenance, not a guessed pattern.
        if not missing(email):
            if status == "generated (unverified)" and missing(record.get("Generated Email Candidate")):
                record["Generated Email Candidate"] = email
            elif status == "found online (unverified)" and missing(record.get("Published Email ID")):
                record["Published Email ID"] = email
        for entry in email_evidence.get(key, ()):
            if (_normal(entry.get("Company Name")) == _normal(record.get("Company Name"))
                    and _evidence_url(entry.get("Source URL"))):
                value = entry.get("Business Email ID")
                for field in ("Business Email ID", "Published Email ID", "Generated Email Candidate"):
                    if not missing(value) and _normal(value) == _normal(record.get(field)):
                        record["_evidence"][field] = text(entry["Source URL"])
        for entry in company_evidence.get(_normal(record.get("Company Name")), ()):
            field = text(entry.get("Field"))
            if (not missing(entry.get("Company Website")) and not missing(record.get("Company Website"))
                    and _normal(entry["Company Website"]).rstrip("/") != _normal(record["Company Website"]).rstrip("/")):
                continue
            if (not missing(entry.get("Value"))
                    and _normal(entry.get("Value")) == _normal(record.get(field))
                    and _evidence_url(entry.get("Source URL"))):
                record["_evidence"][field] = text(entry["Source URL"])
        for entry in maps_evidence.get(key, ()):
            if _normal(entry.get("Company Name")) != _normal(record.get("Company Name")):
                continue
            eligible = _normal(entry.get("Eligible for Merge")) == "yes"
            record["Maps Outcome / Reason"] = entry.get("Outcome / Reason", "")
            if not eligible:
                for field in ("Company Location Type", "Company Location Address", "Company Location Phone"):
                    record["_reasons"][field] = text(entry.get("Outcome / Reason")) or "Maps candidate held; not eligible for merge"
            for target, source_field in (("Company Location Type", "Location Type"),
                                         ("Company Location Address", "Company Location Address"),
                                         ("Company Location Phone", "Company Location Phone"),
                                         *[(field, field) for field in CONTEXT_FIELDS if field.startswith("Prospect") or field in ("Location Basis", "Classification Evidence")]):
                if target.startswith("Company Location") and not eligible:
                    continue
                if not missing(entry.get(source_field)):
                    record[target] = entry[source_field]
                    if _evidence_url(entry.get("Source URL")):
                        record["_evidence"][target] = text(entry["Source URL"])
        for field, link in LINK_FIELDS.items():
            if _evidence_url(record.get(link)):
                record["_evidence"].setdefault(field, text(record[link]))
        for field in FIELDS:
            if not missing(record.get(field + " Missing Reason")):
                record["_reasons"][field] = text(record[field + " Missing Reason"])
        if key and key in best:
            duplicates += 1
            old_index = best[key]
            if sum(not missing(record.get(f)) for f in FIELDS) > sum(
                    not missing(prepared[old_index].get(f)) for f in FIELDS):
                prepared[old_index] = record
        else:
            if key:
                best[key] = len(prepared)
            prepared.append(record)
    return prepared, duplicates


def _size_bucket(value):
    if missing(value):
        return "Unknown"
    numbers = [int(v.replace(",", "")) for v in re.findall(r"\d[\d,]*", text(value))]
    if not numbers:
        return "Other"
    size = max(numbers)
    return "1-50" if size <= 50 else "51-200" if size <= 200 else "201-1000" if size <= 1000 else "1001-10000" if size <= 10000 else "10001+"


def strata(record):
    count = sum(not missing(record.get(f)) for f in FIELDS)
    tier = "Low (<50%)" if count < len(FIELDS) / 2 else "Medium (50-79%)" if count < len(FIELDS) * .8 else "High (80%+)"
    roles = text(record.get("Experience Layout") or record.get("Role Layout"))
    # Grouped roles cannot be reliably reconstructed from a flat job-title column.
    return (_normal(record.get("Industry")) if not missing(record.get("Industry")) else "Unknown",
            _size_bucket(record.get("Company Size")),
            _normal(record.get("Country")) if not missing(record.get("Country")) else "Unknown",
            tier, roles if roles else "Not recorded")


def select_sample(records, sample_size=150, seed=42, company_cap=3):
    if isinstance(sample_size, bool) or not isinstance(sample_size, int) or not 1 <= sample_size <= 200:
        raise ValueError("sample_size must be an integer between 1 and 200")
    if isinstance(company_cap, bool) or not isinstance(company_cap, int) or company_cap < 1:
        raise ValueError("company_cap must be a positive integer")
    rng = random.Random(seed)
    buckets = defaultdict(list)
    for record in sorted(records, key=lambda r: r["_id"]):
        buckets[strata(record)].append(record)
    keys = sorted(buckets)
    rng.shuffle(keys)
    for bucket in buckets.values():
        rng.shuffle(bucket)
    selected, deferred, counts = [], [], Counter()
    while keys and len(selected) < sample_size:
        remaining = []
        for key in keys:
            if len(selected) >= sample_size:
                break
            record = buckets[key].pop()
            company = _normal(record.get("Company Name"))
            company = company if not missing(company) else record["_id"]
            if counts[company] < company_cap:
                selected.append(record)
                counts[company] += 1
            else:
                deferred.append(record)
            if buckets[key]:
                remaining.append(key)
        keys = remaining
    # Relax the cap only after every uncapped candidate has been considered.
    selected.extend(deferred[:max(0, sample_size - len(selected))])
    return selected


def read_reviews(path):
    if not path:
        return {}
    raw, _ = read_saved(path, sheet="Review", require_sheet=True,
                        required=("Record ID", "Field", "Reviewed", "Expected State", "Expected Value"))
    reviews = {}
    for row in raw:
        if _normal(row.get("CSV Escaped")) == "yes":
            row = {k: _csv_restore(v) for k, v in row.items()}
        reviewed = _normal(row.get("Reviewed"))
        if reviewed not in ("", "no", "false", "0", "yes", "true", "1"):
            raise ValueError("Reviewed must be Yes or No")
        if reviewed not in ("yes", "true", "1"):
            continue
        key = text(row.get("Record ID")) or canonical_profile(row.get(PROFILE))
        field = text(row.get("Field"))
        state = _normal(row.get("Expected State"))
        if not key or field not in FIELDS or state not in ("value", "confirmed absent", "unknown"):
            raise ValueError("Reviewed rows require a record ID, supported field and Expected State")
        if state == "value" and missing(row.get("Expected Value")):
            raise ValueError("Reviewed Value requires a nonempty Expected Value")
        if state == "confirmed absent" and not missing(row.get("Expected Value")):
            raise ValueError("Confirmed absent must have an empty Expected Value")
        if _normal(row.get("False Match")) not in ("", "no", "false", "0", "yes", "true", "1"):
            raise ValueError("False Match must be Yes or No")
        if (key, field) in reviews:
            raise ValueError(f"Duplicate review for {key}, {field}")
        reviews[key, field] = row
    return reviews


def assessment(record, field, reviews):
    review = reviews.get((record["_id"], field))
    if not review or _normal(review.get("Expected State")) == "unknown":
        return "Unknown", False
    false_match = _normal(review.get("False Match")) in ("yes", "true", "1")
    if _normal(review.get("Expected State")) == "confirmed absent":
        correct = missing(record.get(field))
    else:
        correct = _normal(record.get(field)) == _normal(review.get("Expected Value"))
    return "Correct" if correct and not false_match else "Incorrect", false_match


def missing_reason(record, field):
    if not missing(record.get(field)):
        return ""
    reason = record.get("_reasons", {}).get(field)
    if reason:
        return reason
    if field in record.get("_evidence", {}):
        return "Source recorded; missing reason not recorded"
    return "No evidence recorded"


def field_results(records, reviews, scope):
    results = []
    for field in FIELDS:
        outcomes = Counter()
        reasons = Counter()
        filled = false_matches = 0
        for record in records:
            filled += not missing(record.get(field))
            status, false_match = assessment(record, field, reviews)
            outcomes[status] += 1
            false_matches += false_match
            reason = missing_reason(record, field)
            if reason:
                reasons[reason] += 1
        reviewed = outcomes["Correct"] + outcomes["Incorrect"]
        results.append({"Scope": scope, "Field": field, "Records": len(records),
                        "Filled": filled, "Missing": len(records) - filled,
                        "Coverage": filled / len(records) if records else None,
                        "Reviewed Known": reviewed, "Correct": outcomes["Correct"],
                        "Incorrect": outcomes["Incorrect"], "Unknown": outcomes["Unknown"],
                        "Accuracy": outcomes["Correct"] / reviewed if reviewed else None,
                        "False Matches": false_matches, "Missing Reasons": dict(sorted(reasons.items()))})
    return results


def baseline_changes(records, baseline, reviews=None):
    old = {r["_profile"]: r for r in baseline if r["_profile"]}
    changes = []
    for record in records:
        previous = old.get(record["_profile"])
        if previous is None:
            continue
        for field in FIELDS:
            before, after = previous.get(field), record.get(field)
            if (missing(before) and missing(after)) or _normal(before) == _normal(after):
                continue
            kind = "Newly filled" if missing(before) else "Became missing" if missing(after) else "Changed value"
            changes.append({"Record ID": record["_id"], "Field": field,
                            "Before": text(before), "After": text(after), "Change": kind,
                            "Current Assessment": assessment(record, field, reviews or {})[0],
                            "Current False Match": "Yes" if assessment(record, field, reviews or {})[1] else "No"})
    return changes


def _csv_safe(value):
    raw = text(value)
    # CSV does not carry cell types. Neutralize formulas even after whitespace.
    return "'" + raw if raw.lstrip().startswith(("=", "+", "-", "@", "'")) else raw


def _csv_restore(value):
    raw = text(value)
    return raw[1:] if len(raw) > 1 and raw[0] == "'" and raw[1] in "=+-@'" else value


def _csv_write(path, headers, rows):
    with Path(path).open("x", encoding="utf-8-sig", newline="") as stream:
        writer = csv.writer(stream)
        writer.writerow(headers)
        for row in rows:
            writer.writerow([_csv_safe(json.dumps(row.get(h), ensure_ascii=False, sort_keys=True)
                                      if isinstance(row.get(h), dict) else row.get(h)) for h in headers])


def _sheet(book, name, headers, rows, editable=()):
    sheet = book.create_sheet(name)
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "C2"
    for row in [dict(zip(headers, headers))] + list(rows):
        values = []
        for header in headers:
            value = row.get(header)
            if isinstance(value, dict):
                value = json.dumps(value, ensure_ascii=False)
            if isinstance(value, str):
                value = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", value)[:32767]
            values.append(value)
        sheet.append(values)
    for row in sheet:
        for cell in row:
            if isinstance(cell.value, str):
                cell.value = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", cell.value)[:32767]
                cell.data_type = "s"  # Source text never becomes a formula or external link.
            cell.font = Font(name="Arial", size=10)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="243B53")
        cell.font = Font(name="Arial", size=10, color="FFFFFF", bold=True)
    sheet.row_dimensions[1].height = 30
    for index, header in enumerate(headers, 1):
        sheet.column_dimensions[get_column_letter(index)].width = (42 if header in ("Record ID", PROFILE, "Source URL", "Review Source")
                                                                    else 38 if header in ("Actual Value", "Expected Value", "Missing Reasons", "Review Notes", "Missing Reason")
                                                                    else 26 if header in ("Field", "Company Name") else 18)
        if header in ("Coverage", "Accuracy"):
            for row in sheet.iter_rows(min_row=2, min_col=index, max_col=index):
                row[0].number_format = "0.0%"
        if header in editable:
            for row in sheet.iter_rows(min_row=2, min_col=index, max_col=index):
                row[0].fill = PatternFill("solid", fgColor="FFF2CC")
    if sheet.max_row > 1:
        sheet.auto_filter.ref = sheet.dimensions
    return sheet


def build_benchmark(input_path, output_dir, sample_size=150, seed=42,
                    baseline_path=None, review_path=None, company_cap=3):
    """Create a unique offline run folder; return paths plus a machine-readable summary.

    Inputs and manual review workbooks are read only. Metadata joins require the same
    profile/company; a populated source URL is evidence availability, not verification.
    """
    source = Path(input_path).expanduser().resolve(strict=True)
    source_hash = hashlib.sha256(source.read_bytes()).hexdigest()
    input_hashes = {source: source_hash}
    for extra in (baseline_path, review_path):
        if extra:
            extra_path = Path(extra).expanduser().resolve(strict=True)
            input_hashes[extra_path] = hashlib.sha256(extra_path.read_bytes()).hexdigest()
    raw, metadata = read_saved(source, required=(PROFILE,))
    records, duplicates = prepare_records(raw, metadata)
    sample = select_sample(records, sample_size, seed, company_cap)
    reviews = read_reviews(review_path)
    baseline, changes = [], []
    if baseline_path:
        baseline_raw, baseline_meta = read_saved(baseline_path, required=(PROFILE,))
        baseline, _ = prepare_records(baseline_raw, baseline_meta)
        changes = baseline_changes(records, baseline, reviews)
    fields = field_results(records, reviews, "Saved records") + field_results(sample, reviews, "Review sample")
    review_rows, evidence_rows = [], []
    for record in sample:
        for field in FIELDS:
            old_review = reviews.get((record["_id"], field), {})
            review_rows.append({"Record ID": record["_id"], PROFILE: record.get(PROFILE),
                                "Company Name": record.get("Company Name"), "Field": field,
                                "Actual Value": record.get(field), "Reviewed": old_review.get("Reviewed", "No"),
                                "Expected State": old_review.get("Expected State", "Unknown"),
                                "Expected Value": old_review.get("Expected Value", ""),
                                "False Match": old_review.get("False Match", "No"),
                                "Review Source": old_review.get("Review Source", ""),
                                "Review Notes": old_review.get("Review Notes", "")})
            outcome, false_match = assessment(record, field, reviews)
            evidence_rows.append({"Record ID": record["_id"], "Field": field, "Actual Value": record.get(field),
                                  "Missing Reason": missing_reason(record, field),
                                  "Source URL": record["_evidence"].get(field, ""),
                                  "Assessment": outcome, "False Match": "Yes" if false_match else "No"})
    generated = datetime.now(timezone.utc).isoformat(timespec="seconds")
    record_keys = {r["_id"] for r in records}
    baseline_keys = {r["_profile"] for r in baseline if r["_profile"]}
    stratum_counts = Counter(strata(r) for r in sample)
    summary = {"generated_at": generated, "input": str(source), "input_sha256": source_hash,
               "raw_records": len(raw), "unique_records": len(records), "duplicates_removed": duplicates,
               "requested_sample_size": sample_size, "sample_size": len(sample), "seed": seed,
               "company_cap": company_cap, "sample_shortfall": max(0, sample_size - len(sample)),
               "sample_ids": [r["_id"] for r in sample],
               "records_without_profile_key": sum(not r["_profile"] for r in records),
               "reviewed_known_cells": sum(r["Reviewed Known"] for r in fields if r["Scope"] == "Review sample"),
               "unmatched_review_cells": sum(k[0] not in record_keys for k in reviews),
               "baseline_matched_records": sum(r["_profile"] in baseline_keys for r in records),
               "baseline_change_counts": dict(Counter(c["Change"] for c in changes)),
               "newly_filled_review_results": dict(Counter(c["Current Assessment"] for c in changes if c["Change"] == "Newly filled")),
               "method": "Deterministic diversity sample across industry, size, country, completeness and recorded role layout; company cap relaxed only to fill remaining places.",
               "limitations": ["Offline only: no new profile visits or increased scraping speed.",
                               "A filled field measures coverage, not verified accuracy.",
                               "Accuracy includes only human-reviewed known expectations. Unreviewed and explicitly unknown cells are excluded.",
                               "This diversity sample is not weighted to estimate population accuracy.",
                               "Grouped roles are unknown unless the source explicitly records the layout.",
                               "Baseline newly filled/changed cells are not automatically correct.",
                               "Missing reasons are reported only when recorded; absence of evidence is not confirmed absence."],
               "strata": [{"industry": k[0], "company_size": k[1], "country": k[2], "completeness": k[3],
                            "role_layout": k[4], "saved_records": n,
                            "sample_records": stratum_counts[k]}
                           for k, n in sorted(Counter(strata(r) for r in records).items())]}
    output_root = Path(output_dir).expanduser().resolve()
    output_root.mkdir(parents=True, exist_ok=True)
    folder = Path(tempfile.mkdtemp(prefix="benchmark_" + datetime.now(timezone.utc).strftime("%Y%m%d_%H%M%S_") , dir=output_root))
    workbook_path = folder / "benchmark_review.xlsx"
    book = Workbook()
    book.remove(book.active)
    _sheet(book, "Field Results", RESULT_HEADERS, fields)
    review_sheet = _sheet(book, "Review", REVIEW_HEADERS, review_rows, editable=REVIEW_HEADERS[5:])
    for column, options in (("F", '"No,Yes"'), ("G", '"Unknown,Value,Confirmed absent"'), ("I", '"No,Yes"')):
        validation = DataValidation(type="list", formula1=options, allow_blank=False)
        validation.errorTitle, validation.error, validation.showErrorMessage = "Choose a listed value", "Use the drop-down list.", True
        review_sheet.add_data_validation(validation)
        if review_sheet.max_row > 1:
            validation.add(f"{column}2:{column}{review_sheet.max_row}")
    sample_headers = ["Record ID", "Source Row", PROFILE] + list(FIELDS) + list(CONTEXT_FIELDS) + ["Industry Stratum", "Size Stratum", "Country Stratum", "Completeness Stratum", "Role Layout Stratum"]
    _sheet(book, "Sample", sample_headers, [{"Record ID": r["_id"], "Source Row": r["_row"],
             **{f: r.get(f) for f in [PROFILE] + list(FIELDS) + list(CONTEXT_FIELDS)},
             **dict(zip(sample_headers[-5:], strata(r)))} for r in sample])
    evidence_headers = ("Record ID", "Field", "Actual Value", "Missing Reason", "Source URL", "Assessment", "False Match")
    _sheet(book, "Field Evidence", evidence_headers, evidence_rows)
    _sheet(book, "Instructions", ("Topic", "Instructions"), [
        {"Topic": "Workflow", "Instructions": "Review only selected fields needed. Edit yellow columns in Review, save this workbook, then rerun benchmark.py with --reviews pointing to this file. Reports are recalculated by the application; editing Excel alone does not refresh them."},
        {"Topic": "Known value", "Instructions": "Set Reviewed=Yes, Expected State=Value, and Expected Value to the independently checked value. Use Review Source/Notes to record supporting evidence."},
        {"Topic": "Confirmed absent", "Instructions": "Use Reviewed=Yes and Expected State=Confirmed absent only when independently confirmed. Leave Expected Value empty. A failed search alone is not confirmed absence."},
        {"Topic": "Unknown", "Instructions": "Leave Reviewed=No or choose Expected State=Unknown when you cannot establish ground truth. It is excluded from accuracy."},
        {"Topic": "False match", "Instructions": "Mark Yes for a wrong person, company or location association. It counts as incorrect on a reviewed known expectation."},
        {"Topic": "Comparison", "Instructions": "Values are compared after trimming, collapsing whitespace and ignoring case. Other spelling/format differences require a reviewer to reconcile; no fuzzy match is used for correctness."},
        *[{"Topic": "Method", "Instructions": s} for s in summary["limitations"]]])
    book["Instructions"].column_dimensions["B"].width = 110
    for index in range(2, book["Instructions"].max_row + 1):
        book["Instructions"].row_dimensions[index].height = 48
    book.save(workbook_path)
    book.close()
    _csv_write(folder / "field_results.csv", RESULT_HEADERS, fields)
    _csv_write(folder / "review_template.csv", REVIEW_HEADERS + ("CSV Escaped",),
               [{**r, "CSV Escaped": "Yes"} for r in review_rows])
    _csv_write(folder / "baseline_changes.csv", ("Record ID", "Field", "Before", "After", "Change", "Current Assessment", "Current False Match"), changes)
    (folder / "field_results.json").write_text(json.dumps({"summary": summary, "fields": fields, "changes": changes}, indent=2, ensure_ascii=False), encoding="utf-8")
    sample_fields = [r for r in fields if r["Scope"] == "Review sample"]
    lines = ["# Saved-results benchmark", "", f"Selected **{len(sample)} of {len(records)} saved unique records** (requested {sample_size}; removed {duplicates} duplicate profile rows).", "",
             "No websites were visited. Source files and earlier review files are unchanged.", "",
             f"Human-reviewed known cells in this sample: **{summary['reviewed_known_cells']}**. Coverage and accuracy are separate.", "",
             "| Field | Filled / sampled | Coverage | Reviewed known | Correct | Incorrect | Accuracy | False matches |",
             "| --- | ---: | ---: | ---: | ---: | ---: | ---: | ---: |"]
    for row in sample_fields:
        percentage = lambda value: "Not evaluated" if value is None else f"{100 * value:.1f}%"
        lines.append(f"| {row['Field']} | {row['Filled']} / {row['Records']} | {percentage(row['Coverage'])} | {row['Reviewed Known']} | {row['Correct']} | {row['Incorrect']} | {percentage(row['Accuracy'])} | {row['False Matches']} |")
    lines += ["", "## Review and rerun", "", "Open benchmark_review.xlsx and edit the yellow Review columns. Save it, then pass it with --reviews on the next run. The workbook stores a report snapshot; the application recalculates the results on rerun.", "",
              "## Baseline comparison", "", f"Matched {summary['baseline_matched_records']} records by normalized LinkedIn profile. Changes: {summary['baseline_change_counts'] or 'none'}. Newly filled cells are not assumed correct. Reviewed outcomes of newly filled cells: {summary['newly_filled_review_results'] or 'none'}.", "", "## Method and limits", "", summary["method"], ""]
    lines += ["- " + value for value in summary["limitations"]]
    if summary["sample_shortfall"]:
        lines += ["", f"Only {len(records)} unique saved records were available. Collect additional records through the existing workflow to reach the target; none were invented or fetched for this benchmark."]
    if summary["unmatched_review_cells"]:
        lines += ["", f"{summary['unmatched_review_cells']} reviewed cells did not match this input and were excluded."]
    (folder / "benchmark_report.md").write_text("\n".join(lines) + "\n", encoding="utf-8")
    for path, expected_hash in input_hashes.items():
        if hashlib.sha256(path.read_bytes()).hexdigest() != expected_hash:
            raise RuntimeError("An input or review changed during benchmark generation; rerun against stable saved files")
    return {"output_dir": str(folder), "workbook": str(workbook_path),
            "report": str(folder / "benchmark_report.md"), "json": str(folder / "field_results.json"),
            "csv": str(folder / "field_results.csv"), "summary": summary}


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", required=True, help="Saved XLSX/CSV/TSV output; read only")
    parser.add_argument("--output-dir", required=True, help="Parent directory for a new benchmark run folder")
    parser.add_argument("--sample-size", type=int, default=150)
    parser.add_argument("--seed", type=int, default=42)
    parser.add_argument("--company-cap", type=int, default=3)
    parser.add_argument("--baseline", help="Optional earlier saved output")
    parser.add_argument("--reviews", help="Earlier reviewed benchmark workbook or review CSV")
    args = parser.parse_args(argv)
    result = build_benchmark(args.input, args.output_dir, args.sample_size, args.seed,
                             args.baseline, args.reviews, args.company_cap)
    print(json.dumps(result, indent=2, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
