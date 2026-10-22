from __future__ import annotations

import csv
import difflib
import json
import os
import queue
import re
import shutil
import threading
import tempfile
import time
import unicodedata
from datetime import datetime, timezone
from pathlib import Path
from typing import Any
from urllib.parse import quote_plus, urljoin, urlsplit, urlunsplit

from google_maps_dom import PLACE_SCRIPT
from company_public_data import public_url as is_public_source_url, host as source_host
from business_email import build_email_details
from run_support import EvidenceCache, EnrichmentSessions, RecoveryJournal, RunMetrics, timed
from linkedin_links import classify_link, extract_link_inputs, LinkResolutionStore, RESOLUTION_DOM

import customtkinter as ctk
import requests
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from playwright.sync_api import sync_playwright, TimeoutError as PlaywrightTimeoutError


# ============================================================
# UNIVERSAL LEAD EXTRACTOR - VERSION 2
# ============================================================
# VERSION 11 CHANGE:
#   STRICT SEARCH-CARD -> PRESENT -> COMPANY MATCH -> COMPANY CLICK FLOW.
#   1. Copy Name, Job Title and Company from the search result card.
#   2. Click the prospect name to open the side panel.
#   3. Search the panel for the literal keyword Present.
#   4. Copy the complete Present line as Tenure.
#   5. No Present => Tenure = Dead.
#   6. Present company mismatch => skip the prospect.
#   7. Present company match => click the company and continue.
#
#   The app no longer discovers all leads first.
#   It starts copying each newly visible result card immediately,
#   writes the row, then continues scrolling for more.
#
#   Speed is improved by:
#     1. Deterministic parsing for obvious card fields.
#     2. Qwen is used only as a fallback for ambiguous cards.
#     3. Excel is checkpointed in batches instead of saving after every row.
#     4. No pre-harvest/discovery phase.
#
# Authentication:
#   The app does NOT ask for or store the LinkedIn password.
#   Sign in manually in the Chrome window.
#
# This program does not bypass MFA, CAPTCHAs, access controls,
# or other website protections.
# ============================================================


APP_NAME = "Universal Lead Extractor"
APP_VERSION = "42.9.0"

# These fields require public company evidence and are never inferred by Qwen.
PUBLIC_COMPANY_FIELDS = (
    "Revenue", "Revenue Link", "SIC Code", "NAICS Code", "SIC/NAICS Code Link",
)
EMAIL_EVIDENCE_HEADERS = (
    'Person', 'Prospect LinkedIn Link', 'Company Name', 'Business Email ID',
    'Method', 'Pattern', 'Source URL', 'Supporting Text', 'Retrieved At',
)
COMPANY_EVIDENCE_HEADERS = (
    "Company Name", "Company Website", "Field", "Value", "Source URL",
    "Period", "Basis", "Excerpt", "Checked At",
)
DETAIL_SHEETS = {
    '_profile_state': ('Profile State', (
        ('profile', 'Prospect LinkedIn Link'), ('company', 'Company Name'),
        ('company_url', 'Company LinkedIn URL'), ('captured_at', 'Profile Captured At'),
    )),
    '_previous_company_location': ('Previous Company Location', (
        ('profile', 'Prospect LinkedIn Link'), ('company', 'Company Name'),
        ('address', 'Previous Company Address'), ('zip', 'Previous Postal Code'),
        ('phone', 'Previous Company Phone'), ('source_url', 'Address Source'),
        ('phone_source', 'Phone Source'), ('basis', 'Reason Preserved'),
    )),
    '_email_details': ('Email Details', (
        ('profile', 'Prospect LinkedIn Link'), ('company', 'Company Name'),
        ('published_email', 'Published Email ID'), ('generated_email', 'Generated Email Candidate'),
        ('existing_email', 'Existing Email (origin unknown)'), ('email_domain', 'Email Domain'),
        ('domain_status', 'Domain/MX Status'), ('checked_at', 'Checked At'),
        ('mailbox_status', 'Mailbox Verification'), ('detail', 'Detail'),
    )),
    '_maps_evidence': ('Maps Evidence', (
        ('profile', 'Prospect LinkedIn Link'), ('company', 'Company Name'),
        ('scope', 'Location Type'), ('address', 'Company Location Address'),
        ('phone', 'Company Location Phone'), ('source_url', 'Source URL'),
        ('prospect_location', 'Prospect Location'), ('location_basis', 'Location Basis'),
        ('basis', 'Classification Evidence'), ('is_mergeable', 'Eligible for Merge'),
        ('reason', 'Outcome / Reason'),
        ('applied_fields', 'Fields Applied'), ('checked_at', 'Checked At'),
        ('prospect_office_verified', 'Prospect Works Here Verified'),
    )),
}

COMPANY_NAME_MATCH_THRESHOLD = 88.0

# LinkedIn displays employee counts either as a single value (``10,001+``)
# or a range (``5,001-10,000``). Keep the complete value so company-size
# buckets are not truncated to the second number in the range.
EMPLOYEE_SIZE_RE = re.compile(
    r"(?i)\b\d[\d,.]*\s*(?:[-–—]\s*\d[\d,.]*)?\s*"
    r"[KMB]?\s*\+?\s*employees\b"
)

# LinkedIn often changes a legal display name after a rebrand. Keep this list
# short and explicit so fuzzy matching does not turn unrelated subsidiaries into
# the same company.
COMPANY_NAME_ALIASES = {
    frozenset({"bhp", "bhp billiton"}),
}

LLM_BASE_URL = "http://127.0.0.1:8080"
LLM_CHAT_URL = f"{LLM_BASE_URL}/v1/chat/completions"
LLM_MODELS_URL = f"{LLM_BASE_URL}/v1/models"
LLM_FALLBACK_MODEL = "qwen3-4b-thinking-2507.Q4_K_M.gguf"

# SMART-mode performance controls.
# Structured validation should not need long free-form generations.
SMART_LLM_READ_TIMEOUT_SECONDS = 60
SMART_LLM_MAX_TOKENS = 700

# Local-LLM resilience. A single failed health check must not disable
# enrichment for an entire scraping run.
LLM_HEALTH_RETRIES = 3
LLM_HEALTH_TIMEOUT_SECONDS = 10
LLM_HEALTH_RETRY_DELAY_SECONDS = 1.25
LLM_RECHECK_INTERVAL_SECONDS = 20

# Ambiguous search-card fallback should be much smaller/faster than the final
# SMART enrichment request.
CARD_LLM_READ_TIMEOUT_SECONDS = 45
CARD_LLM_MAX_TOKENS = 650

# Detailed persistent diagnostics.
ACTIVITY_LOG_RAW_LLM_LIMIT = 12_000
ACTIVITY_LOG_EVIDENCE_LIMIT = 14_000
ACTIVITY_LOG_VALUE_LIMIT = 2_000

# Google Maps enrichment.
# Used only when company Address / Zip / Phone are still missing after
# LinkedIn/company-page extraction.
GOOGLE_MAPS_PAGE_TIMEOUT_MS = 25_000
GOOGLE_MAPS_RESULT_WAIT_MS = 15_000
GOOGLE_MAPS_SETTLE_MS = 2_500
GOOGLE_MAPS_MAX_QUERIES_PER_COMPANY = 2

# Place-details panel is lazy rendered by Google Maps. Give the page time to
# settle, then scroll the left details pane in controlled steps so Address /
# Phone / Website elements have a chance to render.
GOOGLE_MAPS_DETAIL_INITIAL_WAIT_MS = 2_200
GOOGLE_MAPS_DETAIL_SCROLL_WAIT_MS = 650
GOOGLE_MAPS_DETAIL_SCROLL_STEP = 720
GOOGLE_MAPS_DETAIL_MAX_SCROLLS = 10
GOOGLE_MAPS_DETAIL_STAGNANT_LIMIT = 2

# Immediately after clicking the first Maps result, move the place-details
# panel slightly so lazy Address / Phone cards begin rendering before the
# normal deeper scrolling loop starts.
GOOGLE_MAPS_POST_CLICK_SCROLL_PX = 360
GOOGLE_MAPS_POST_CLICK_WAIT_MS = 1_200

# Search-result matching. Google Maps may display several similarly named
# businesses; inspect multiple visible cards before deciding what to click.
GOOGLE_MAPS_RESULT_SCAN_ROUNDS = 5
GOOGLE_MAPS_RESULT_SCAN_LIMIT = 30
GOOGLE_MAPS_RESULT_SCROLL_PX = 900

# Fields whose final enrichment state is audited at the end of every run.
ENRICHMENT_AUDIT_FIELDS = [
    "Job Level",
    "Job Function",
    "Company Website",
    "Industry",
    "Address",
    "City",
    "State",
    "Zip Code",
    "Country",
    "Revenue",
    "SIC Code",
    "NAICS Code",
    "SIC/NAICS Code Link",
]

PAGE_TIMEOUT_MS = 60_000
PAGE_SETTLE_MS = 2_500

# Keep each result card compact. This is enough for the visible result data.
MAX_CARD_TEXT_CHARS = 4_500

# Search-result harvesting settings.
SCROLL_PIXELS = 1_250
SCROLL_WAIT_MS = 650
MAX_STAGNANT_SCROLLS = 5
MAX_SCROLL_ROUNDS = 250

# Dynamic load detection after scrolling.
NEW_RESULT_POLL_MS = 120
NEW_RESULT_TIMEOUT_MS = 2400

# Current-role side-panel extraction.
PROFILE_PANEL_OPEN_TIMEOUT_MS = 8000
PROFILE_PANEL_SCROLL_STEP = 850
PROFILE_PANEL_SCROLL_WAIT_MS = 220
PROFILE_PANEL_MAX_SCROLLS = 20

# Company page/panel extraction.
COMPANY_PAGE_TIMEOUT_MS = 15000
COMPANY_PAGE_ELEMENT_WAIT_MS = 12_000
COMPANY_PAGE_SETTLE_MS = 3_000

# Standard LinkedIn profile/company flow used by Link Scraper.
LINK_PROFILE_SETTLE_MS = 2_500
LINK_PROFILE_MAX_SCROLLS = 24
LINK_COMPANY_SETTLE_MS = 3_000
LINK_COMPANY_MAX_SCROLLS = 36

# Fast-mode controls.
# These improve throughput by reducing unnecessary local work.
# They do not use stealth, CAPTCHA bypass, fingerprint spoofing,
# or other mechanisms intended to defeat site controls.
USE_LLM_FALLBACK = True

# Saving XLSX is relatively expensive because the workbook is rewritten.
# UI runs journal every row, then save XLSX in batches.
EXCEL_CHECKPOINT_EVERY = 10
EXCEL_CHECKPOINT_SECONDS = 45

HEADERS = [
    "Date",
    "Campaign Name",
    "Agent Name",
    "Calling Agent Name",
    "Comment",
    "Dispo",
    "Notes",
    "Email Status",
    "Tenure",
    "First Name",
    "Last Name",
    "Business Email ID",
    "Prospect LinkedIn Link",
    "Job Title",
    "Job Level",
    "Job Function",
    "Company Name",
    "Company Website",
    "Industry",
    "Company Size",
    "See All Company Size",
    "Industry Link",
    "Address",
    "City",
    "State",
    "Zip Code",
    "Address Link",
    "Country",
    "Phone Number",
    "Phone Number Link",
    "Revenue",
    "Revenue Link",
    "SIC Code",
    "NAICS Code",
    "SIC/NAICS Code Link",
    "Audited Date",
    "QA Name",
    "QA Status",
    "Primary Reason",
    "Description",
    "Bounce Email Status",
    "Source File Name",
]

# Fields that can reasonably be extracted/classified from the visible result card.
CARD_LLM_FIELDS = [
    "Tenure",
    "First Name",
    "Last Name",
    "Job Title",
    "Job Level",
    "Job Function",
    "Company Name",
    "Industry",
    "City",
    "State",
    "Country",
    "Description",
]


def app_data_dir() -> Path:
    root = Path(os.getenv("LOCALAPPDATA", str(Path.home())))
    path = root / "UniversalLeadExtractor"
    path.mkdir(parents=True, exist_ok=True)
    return path


CHROME_PROFILE_DIR = app_data_dir() / "chrome_profile"


def normalize_url(url: str) -> str:
    """
    Canonical URL used for duplicate detection.
    Query and fragment are intentionally removed.
    """
    url = (url or "").strip().rstrip(".,;")
    if url.casefold() in {'', '-', 'none', 'null', 'nan', 'n/a', 'unknown'}:
        return ""

    try:
        parts = urlsplit(url)
        scheme = (parts.scheme or "https").lower()
        netloc = parts.netloc.lower()
        host = (parts.hostname or '').lower()
        path = parts.path.rstrip("/")
        if host == "linkedin.com" or host.endswith(".linkedin.com"):
            netloc = "www.linkedin.com"
            scheme = "https"
            if path.startswith("/sales/lead/"):
                path = path.split(",", 1)[0]
        return urlunsplit((scheme, netloc, path, "", ""))
    except Exception:
        return url


def clean_cell_value(value: Any) -> str:
    if value is None:
        return "-"
    if isinstance(value, (int, float)):
        return str(value)
    value = str(value).strip()
    return value if value else "-"


def excel_safe(value: Any) -> Any:
    """
    Prevent scraped text from being interpreted as an Excel formula.
    """
    if value is None:
        return "-"
    text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", str(value))[:32767]
    if text == "-":
        return text
    if text.startswith(("=", "+", "-", "@")):
        return "'" + text
    return text



class CompanyPageMismatchError(RuntimeError):
    """Raised when the opened company page does not match the search-card company."""


class ProfileAccessError(RuntimeError):
    """Authentication or verification requires user action before continuing."""


def canonical_profile_url(value: str) -> str:
    parts = urlsplit(value.strip())
    host = (parts.hostname or '').lower()
    if not (host == 'linkedin.com' or host.endswith('.linkedin.com')):
        raise ValueError('Profile URL must belong to linkedin.com.')
    if not parts.path.startswith('/in/'):
        raise ValueError('Expected a LinkedIn /in/ profile URL.')
    return urlunsplit(('https', 'www.linkedin.com', parts.path.rstrip('/'), '', ''))


class LocalLLMClient:
    def __init__(
        self,
        base_url: str = LLM_BASE_URL,
        trace_callback=None,
    ):
        self.base_url = base_url.rstrip("/")
        self.chat_url = f"{self.base_url}/v1/chat/completions"
        self.models_url = f"{self.base_url}/v1/models"
        self.session = requests.Session()
        self.model_id = LLM_FALLBACK_MODEL
        self.trace_callback = trace_callback

    def _trace(self, event: str, **payload) -> None:
        """
        Send detailed Local-LLM diagnostics to the application's persistent
        activity log without flooding the on-screen Activity Log.
        """
        callback = self.trace_callback

        if not callback:
            return

        safe_payload: dict[str, Any] = {}

        for key, value in payload.items():
            if isinstance(value, str):
                limit = (
                    ACTIVITY_LOG_RAW_LLM_LIMIT
                    if key in {
                        "raw_response",
                        "response_text",
                        "prompt",
                    }
                    else ACTIVITY_LOG_EVIDENCE_LIMIT
                    if key in {
                        "evidence",
                        "website_evidence",
                    }
                    else ACTIVITY_LOG_VALUE_LIMIT
                )

                safe_payload[key] = (
                    value
                    if len(value) <= limit
                    else value[:limit]
                    + f"... <truncated {len(value) - limit} chars>"
                )
            else:
                safe_payload[key] = value

        try:
            callback(event, safe_payload)
        except Exception:
            # Diagnostics must never interrupt scraping/enrichment.
            pass

    def health_check(self) -> tuple[bool, str]:
        """
        Verify that the endpoint actually used for enrichment works.

        Some llamafile/OpenAI-compatible servers can answer
        /v1/chat/completions even when /v1/models is unavailable or behaves
        differently. Therefore /v1/models is only used to discover the model
        id; the lightweight chat request is the real health check.
        """
        discovered_model = None

        self._trace(
            "health_check_start",
            base_url=self.base_url,
            models_url=self.models_url,
            chat_url=self.chat_url,
            configured_model=self.model_id,
        )

        # Optional model discovery. Failure here does NOT mean the LLM is down.
        try:
            model_started = time.perf_counter()

            model_response = self.session.get(
                self.models_url,
                timeout=LLM_HEALTH_TIMEOUT_SECONDS,
            )
            model_response.raise_for_status()

            model_data = model_response.json()
            models = model_data.get("data") or []

            if (
                models
                and isinstance(models[0], dict)
                and models[0].get("id")
            ):
                discovered_model = str(models[0]["id"])
                self.model_id = discovered_model

            self._trace(
                "models_endpoint_ok",
                status_code=model_response.status_code,
                elapsed_seconds=round(
                    time.perf_counter() - model_started,
                    3,
                ),
                discovered_model=discovered_model or "",
                model_count=len(models),
            )

        except Exception as exc:
            self._trace(
                "models_endpoint_error",
                error_type=type(exc).__name__,
                error=str(exc),
            )

        last_error = "Unknown local LLM error"

        for attempt in range(1, LLM_HEALTH_RETRIES + 1):
            try:
                payload = {
                    "model": self.model_id,
                    "messages": [
                        {
                            "role": "user",
                            "content": (
                                "/no_think\n"
                                "Reply with exactly: OK"
                            ),
                        }
                    ],
                    "temperature": 0.0,
                    "max_tokens": 6,
                }

                self._trace(
                    "health_chat_attempt",
                    attempt=attempt,
                    model=self.model_id,
                    max_tokens=payload.get("max_tokens"),
                )

                chat_started = time.perf_counter()

                response = self.session.post(
                    self.chat_url,
                    json=payload,
                    timeout=(
                        5,
                        LLM_HEALTH_TIMEOUT_SECONDS,
                    ),
                )
                response.raise_for_status()

                data = response.json()

                self._trace(
                    "health_chat_http_ok",
                    attempt=attempt,
                    status_code=response.status_code,
                    elapsed_seconds=round(
                        time.perf_counter() - chat_started,
                        3,
                    ),
                )

                content = self._assistant_text(data)

                if not content:
                    raise ValueError(
                        "Chat endpoint returned no assistant text."
                    )

                detail = (
                    discovered_model
                    or self.model_id
                )

                self._trace(
                    "health_check_success",
                    model=detail,
                    assistant_text=content,
                    attempt=attempt,
                )

                return True, detail

            except Exception as exc:
                self._trace(
                    "health_chat_error",
                    attempt=attempt,
                    error_type=type(exc).__name__,
                    error=str(exc),
                )

                last_error = (
                    f"chat attempt {attempt}/{LLM_HEALTH_RETRIES}: "
                    f"{type(exc).__name__}: {exc}"
                )

                if attempt < LLM_HEALTH_RETRIES:
                    time.sleep(
                        LLM_HEALTH_RETRY_DELAY_SECONDS
                    )

        self._trace(
            "health_check_failed",
            error=last_error,
        )

        return False, last_error

    @staticmethod
    def _assistant_text(data: dict[str, Any]) -> str:
        """
        Extract assistant output from common OpenAI-compatible response shapes.

        Supports:
        - choices[0].message.content
        - choices[0].message.reasoning_content
        - choices[0].text
        - list-style message content items
        """
        try:
            choice = (data.get("choices") or [{}])[0] or {}
        except Exception:
            choice = {}

        message = choice.get("message") or {}

        candidates = [
            message.get("content"),
            message.get("reasoning_content"),
            choice.get("text"),
        ]

        for candidate in candidates:
            if isinstance(candidate, str) and candidate.strip():
                return candidate.strip()

            if isinstance(candidate, list):
                parts: list[str] = []

                for item in candidate:
                    if isinstance(item, str):
                        if item.strip():
                            parts.append(item.strip())
                        continue

                    if isinstance(item, dict):
                        value = (
                            item.get("text")
                            or item.get("content")
                            or item.get("value")
                        )

                        if isinstance(value, str) and value.strip():
                            parts.append(value.strip())

                if parts:
                    return "\n".join(parts)

        return ""

    @staticmethod
    def _strip_thinking_and_fences(text: str) -> str:
        text = re.sub(
            r"<think>.*?</think>",
            "",
            text,
            flags=re.DOTALL | re.IGNORECASE,
        ).strip()

        if text.startswith("```"):
            text = re.sub(r"^```(?:json)?\s*", "", text, flags=re.IGNORECASE)
            text = re.sub(r"\s*```$", "", text)

        return text.strip()

    @staticmethod
    def _extract_json_object(text: str) -> dict[str, Any]:
        cleaned = LocalLLMClient._strip_thinking_and_fences(text)

        try:
            obj = json.loads(cleaned)
            if isinstance(obj, dict):
                return obj
        except Exception:
            pass

        start = cleaned.find("{")
        end = cleaned.rfind("}")

        if start >= 0 and end > start:
            obj = json.loads(cleaned[start : end + 1])
            if isinstance(obj, dict):
                return obj

        raise ValueError("The local model did not return a valid JSON object.")

    def extract_search_card(
        self,
        card_text: str,
        lead_url: str,
    ) -> dict[str, str]:
        schema = {field: "-" for field in CARD_LLM_FIELDS}

        prompt = f"""
/no_think
Return JSON immediately. Do not provide reasoning.\n\nYou are a structured lead-data extraction engine.

The text below comes from ONE visible LinkedIn Sales Navigator search-result card.

Extract only values supported by the supplied card text.
Do not use outside knowledge to invent factual information.

You MAY classify:
- Job Level from the visible Job Title.
- Job Function from the visible Job Title.
- Industry only if the card text directly supports it. Otherwise return "-".

Important interpretation rules:
- The person's display name maps to First Name and Last Name.
- If a display name is abbreviated (example: "Yoni K."), preserve what is visible:
  First Name = "Yoni"
  Last Name = "K."
- A line such as:
    "Senior Tax Manager · EY"
  generally means:
    Job Title = "Senior Tax Manager"
    Company Name = "EY"
- A line such as:
    "5 years 3 months in role | 5 years 3 months in company"
  should be captured in Tenure exactly or nearly exactly as shown.
- Location text may contain City, State/Region and Country.
- An "About:" snippet may be stored in Description.
- If a field is not available, return "-".

Return exactly one JSON object.
Return no markdown and no explanation.
Use exactly these keys:
{json.dumps(list(schema.keys()), ensure_ascii=False)}

PROSPECT LINK:
{lead_url}

SEARCH RESULT CARD TEXT:
{card_text[:MAX_CARD_TEXT_CHARS]}
""".strip()

        payload = {
            "model": self.model_id,
            "messages": [
                {
                    "role": "system",
                    "content": (
                        "You map one Sales Navigator result card to structured lead fields. "
                        "Return JSON only and never fabricate unsupported factual values."
                    ),
                },
                {"role": "user", "content": prompt},
            ],
            "temperature": 0.1,
            "max_tokens": CARD_LLM_MAX_TOKENS,
        }

        self._trace(
            "search_card_request",
            model=self.model_id,
            prompt_chars=len(prompt),
            card_text=card_text,
            requested_fields=list(CARD_LLM_FIELDS),
        )

        request_started = time.perf_counter()

        r = self.session.post(
            self.chat_url,
            json=payload,
            timeout=(5, CARD_LLM_READ_TIMEOUT_SECONDS),
        )
        r.raise_for_status()

        data = r.json()
        content = self._assistant_text(data)

        self._trace(
            "search_card_response",
            status_code=r.status_code,
            elapsed_seconds=round(
                time.perf_counter() - request_started,
                3,
            ),
            raw_response=content,
        )

        if not content:
            raise ValueError(
                "The local model returned an empty search-card response."
            )

        try:
            parsed = self._extract_json_object(content)
        except Exception as exc:
            self._trace(
                "search_card_parse_error",
                error_type=type(exc).__name__,
                error=str(exc),
                raw_response=content,
            )
            raise

        self._trace(
            "search_card_parsed",
            parsed=parsed,
        )

        result: dict[str, str] = {}
        for field in CARD_LLM_FIELDS:
            value = parsed.get(field, "-")

            if isinstance(value, (list, dict)):
                value = json.dumps(value, ensure_ascii=False)

            result[field] = clean_cell_value(value)

        return result


    @staticmethod
    def _evidence_normalized(value: str) -> str:
        return re.sub(
            r"[^a-z0-9]+",
            " ",
            (value or "").lower(),
        ).strip()

    @staticmethod
    def _evidence_grounded_confidence(
        field: str,
        value: str,
        evidence: dict[str, str],
    ) -> int:
        """
        Give Address/Zip fallback confidence only when the proposed value is
        actually present in direct official-site evidence.

        This lets local models omit the confidence object without forcing a
        factual street address/postal code back to Unknown, while still
        preventing unsupported address fabrication.
        """
        if not value or value.lower() == "unknown":
            return 0

        website_evidence = str(
            evidence.get(
                "Official Website Evidence",
                "",
            )
            or ""
        )

        if not website_evidence:
            return 0

        if field == "Zip Code":
            token = value.strip()

            if not token:
                return 0

            if re.search(
                r"(?<![A-Za-z0-9])"
                + re.escape(token)
                + r"(?![A-Za-z0-9])",
                website_evidence,
                flags=re.IGNORECASE,
            ):
                return 95

            return 0

        if field == "Address":
            proposed = (
                LocalLLMClient
                ._evidence_normalized(value)
            )
            source = (
                LocalLLMClient
                ._evidence_normalized(
                    website_evidence
                )
            )

            if (
                proposed
                and len(proposed) >= 12
                and proposed in source
            ):
                return 95

        return 0

    def enrich_missing_fields(
        self,
        record: dict[str, Any],
        requested_fields: list[str],
    ) -> dict[str, Any]:
        """
        Ask the local LLM to enrich or validate only the requested fields.

        The caller decides which fields are missing, invalid, or explicitly
        selected for validation. Existing valid scraped fields are supplied as
        evidence but are not requested for replacement.

        Returns:
            {
                "values": {<Excel field>: <value>, ...},
                "confidence": {<Excel field>: <0-100>, ...}
            }
        """
        field_to_json = {
            "Job Level": "job_level",
            "Job Function": "job_function",
            "Company Website": "company_website",
            "Industry": "industry",
            "Address": "address",
            "City": "city",
            "State": "state",
            "Zip Code": "zip_code",
            "Country": "country",
            "Revenue": "revenue",
            "SIC Code": "sic_code",
            "NAICS Code": "naics_code",
            "SIC/NAICS Code Link": "sic_naics_code_link",
        }

        requested_fields = [
            field for field in requested_fields
            if field in field_to_json and field not in PUBLIC_COMPANY_FIELDS
        ]

        if not requested_fields:
            return {
                "values": {},
                "confidence": {},
            }

        requested_json_keys = [
            field_to_json[field]
            for field in requested_fields
        ]

        # Include every useful scraped value we currently have. Extra temporary
        # keys such as Company LinkedIn URL are intentionally accepted even
        # though they are not workbook columns.
        evidence_fields = [
            "First Name",
            "Last Name",
            "Job Title",
            "Job Level",
            "Job Function",
            "Company Name",
            "Company LinkedIn URL",
            "Company Website",
            "Industry",
            "Company Size",
            "Description",
            "Address",
            "City",
            "State",
            "Zip Code",
            "Country",
            "Phone Number",
            "Revenue",
            "SIC Code",
            "NAICS Code",
            "SIC/NAICS Code Link",
            "Prospect LinkedIn Link",
            "Tenure",
            "Industry Link",
            "Address Link",
            "Revenue Link",
            "Official Website Evidence",
        ]

        # Send only useful evidence. Omitting "-" / Unknown fields reduces
        # prompt processing time on the local model.
        evidence = {}

        for field in evidence_fields:
            value = clean_cell_value(record.get(field, "-"))

            if value.lower() in {
                "",
                "-",
                "unknown",
                "n/a",
                "na",
                "none",
                "null",
                "unavailable",
                "not available",
                "not found",
            }:
                continue

            evidence[field] = value

        requested_schema = {
            key: "Unknown"
            for key in requested_json_keys
        }

        confidence_schema = {
            key: 0
            for key in requested_json_keys
        }

        prompt = f"""
/no_think
Return JSON immediately. Do not provide reasoning.

Validate or enrich ONLY REQUESTED FIELDS using SCRAPED EVIDENCE.
Direct scraped evidence is preferred. Never fabricate a value.
Use "Unknown" when evidence is insufficient.

REQUESTED FIELDS:
{json.dumps(requested_json_keys, ensure_ascii=False)}

SCRAPED EVIDENCE:
{json.dumps(evidence, ensure_ascii=False)}

STANDARDIZATION RULES:
- job_level: C-Level, VP, Director, Manager, Senior,
  Individual Contributor, Owner/Founder, Partner, Unknown.
- job_function: Finance, Sales, Marketing, Operations,
  Information Technology, Human Resources, Procurement, Supply Chain,
  Engineering, Legal, Customer Service, Administration, Product,
  Security, Healthcare, Other, Unknown.
- company_website: official company site only; never LinkedIn/social/
  directory/Crunchbase. Use Unknown if not confident.
- industry: one primary standardized industry.
- address/city/state/zip_code/country: must be geographically consistent.
  "Official Website Evidence" is direct scraped website text and may be used
  to extract an exact street address/postal code when it clearly appears there.
  Never invent a street address or postal code when it is not in the evidence.
  Country should be a full commonly accepted country name.
- revenue: exact value or defensible range only; otherwise Unknown.
- sic_code: 4 digits when Company Industry/business activity gives enough
  evidence. SIC/NAICS classification is independent of street-address data.
- naics_code: 6 digits when Company Industry/business activity gives enough
  evidence. Do not return Unknown merely because Address/Zip are unavailable.
- sic_naics_code_link: a relevant http/https reference matching the returned
  SIC or NAICS code; otherwise Unknown.

CONFIDENCE:
Return integer confidence 0-100 for every requested field.

Return exactly one flat JSON object:
{{
  {json.dumps(requested_schema, ensure_ascii=False)[1:-1]},
  "confidence": {json.dumps(confidence_schema, ensure_ascii=False)}
}}
Only requested fields plus "confidence". No markdown. No explanation.
""".strip()

        self._trace(
            "enrichment_prepare",
            model=self.model_id,
            requested_fields=requested_fields,
            requested_json_keys=requested_json_keys,
            evidence=json.dumps(
                evidence,
                ensure_ascii=False,
            ),
            evidence_keys=list(evidence.keys()),
            current_requested_values={
                field: clean_cell_value(
                    record.get(field, "-")
                )
                for field in requested_fields
            },
            prompt_chars=len(prompt),
        )

        payload = {
            "model": self.model_id,
            "messages": [
                {
                    "role": "system",
                    "content": (
                        "You enrich scraped lead records conservatively. "
                        "Use Unknown when evidence is insufficient. "
                        "Return strict JSON only."
                    ),
                },
                {
                    "role": "user",
                    "content": prompt,
                },
            ],
            "temperature": 0.0,
            "max_tokens": SMART_LLM_MAX_TOKENS,
        }

        parsed = None
        last_parse_error = None

        for response_attempt in range(2):
            request_payload = dict(payload)

            if response_attempt == 1:
                request_payload["messages"] = [
                    {
                        "role": "system",
                        "content": (
                            "Return one valid JSON object only. "
                            "No reasoning, no markdown, no commentary."
                        ),
                    },
                    {
                        "role": "user",
                        "content": (
                            prompt
                            + "\n\nIMPORTANT: Your previous response could "
                              "not be parsed. Return ONLY valid JSON now."
                        ),
                    },
                ]
                request_payload["max_tokens"] = min(
                    SMART_LLM_MAX_TOKENS,
                    550,
                )

            self._trace(
                "enrichment_request",
                attempt=response_attempt + 1,
                model=self.model_id,
                max_tokens=request_payload.get("max_tokens"),
                requested_fields=requested_fields,
            )

            request_started = time.perf_counter()

            response = self.session.post(
                self.chat_url,
                json=request_payload,
                timeout=(
                    5,
                    SMART_LLM_READ_TIMEOUT_SECONDS,
                ),
            )
            response.raise_for_status()

            response_data = response.json()
            content = self._assistant_text(
                response_data
            )

            self._trace(
                "enrichment_response",
                attempt=response_attempt + 1,
                status_code=response.status_code,
                elapsed_seconds=round(
                    time.perf_counter() - request_started,
                    3,
                ),
                raw_response=content,
            )

            try:
                if not content:
                    raise ValueError(
                        "The local model returned an empty enrichment response."
                    )

                parsed = self._extract_json_object(
                    content
                )
                break
            except Exception as exc:
                last_parse_error = exc

                self._trace(
                    "enrichment_parse_error",
                    attempt=response_attempt + 1,
                    error_type=type(exc).__name__,
                    error=str(exc),
                    raw_response=content,
                )

        if parsed is None:
            raise ValueError(
                "The local model failed to return valid JSON "
                f"after retry: {last_parse_error}"
            )

        # V22 expects the requested fields at the top level, matching the
        # application's documented JSON contract. For compatibility with older
        # prompts/models, also accept V21's nested {"values": {...}} shape.
        nested_values = parsed.get("values")

        if isinstance(nested_values, dict):
            raw_values = nested_values
        else:
            raw_values = {
                key: parsed.get(key, "Unknown")
                for key in requested_json_keys
            }

        raw_confidence = parsed.get("confidence", {})

        if not isinstance(raw_values, dict):
            raise ValueError("LLM response fields are not a JSON object.")

        if not isinstance(raw_confidence, dict):
            raw_confidence = {}

        # The user's response example groups geography confidence under
        # "location". Accept that as a fallback for the individual geographic
        # fields when field-level confidence is not present.
        location_confidence = raw_confidence.get("location", 0)

        json_to_field = {
            json_key: excel_field
            for excel_field, json_key in field_to_json.items()
        }

        values: dict[str, str] = {}
        confidence: dict[str, int] = {}

        for json_key in requested_json_keys:
            excel_field = json_to_field[json_key]

            value = raw_values.get(json_key, "Unknown")

            if isinstance(value, (dict, list)):
                value = "Unknown"

            value = str(value).strip() if value is not None else "Unknown"
            if not value:
                value = "Unknown"

            # Prevent verbose prose from entering Excel.
            if len(value) > 500:
                value = "Unknown"

            raw_score_found = json_key in raw_confidence
            raw_score = raw_confidence.get(json_key)

            if (
                not raw_score_found
                and json_key in {
                    "address",
                    "city",
                    "state",
                    "zip_code",
                    "country",
                }
                and "location" in raw_confidence
            ):
                raw_score = location_confidence
                raw_score_found = True

            score_parsed = False

            if raw_score_found:
                try:
                    if isinstance(raw_score, str):
                        raw_score = (
                            raw_score
                            .strip()
                            .rstrip("%")
                            .strip()
                        )

                    score = int(float(raw_score))
                    score_parsed = True
                except Exception:
                    score = 0
            else:
                score = 0

            score = max(0, min(100, score))

            if not score_parsed:
                score = (
                    LocalLLMClient
                    ._evidence_grounded_confidence(
                        excel_field,
                        value,
                        evidence,
                    )
                )

                if score <= 0:
                    score = (
                        SearchWorker
                        ._fallback_llm_confidence(
                            excel_field,
                            value,
                        )
                    )

            values[excel_field] = value
            confidence[excel_field] = score

        self._trace(
            "enrichment_parsed",
            requested_fields=requested_fields,
            parsed=parsed,
            normalized_values=values,
            normalized_confidence=confidence,
        )

        return {
            "values": values,
            "confidence": confidence,
        }


class WorkbookManager:
    def __init__(self, output_path: str, checkpoint_every: int = 1):
        self.output_path = Path(output_path)
        self.headers = list(HEADERS)
        self.checkpoint_every = max(1, int(checkpoint_every))
        self.metrics = RunMetrics()
        self.journal = RecoveryJournal(self.output_path)
        self._replaying = False
        self._link_rows = {}
        self.seen_links: set[str] = set()
        self._dirty_rows = 0
        self._last_save_ts = time.time()
        self._company_evidence_rows: dict[tuple[str, str, str, str], int] | None = None
        self._email_evidence_rows = None
        self._detail_row_indexes: dict[str, dict[tuple[str, str], int]] = {}
        if self.output_path.exists():
            self.wb = load_workbook(self.output_path)
            if "Leads" not in self.wb.sheetnames:
                raise ValueError("Existing workbook has no Leads sheet; choose another output file.")
            self.ws = self.wb["Leads"]
            existing_headers = [c.value for c in self.ws[1]]
            if existing_headers[:len(HEADERS)] != HEADERS:
                raise ValueError("Existing Leads headers do not match; choose another output file.")
            self._validate_columns(existing_headers[len(HEADERS):])
            self.headers = existing_headers
            column = HEADERS.index("Prospect LinkedIn Link") + 1
            for row in self.ws.iter_rows(min_row=2, min_col=column, max_col=column):
                value = row[0].value
                key = normalize_url(str(value or ''))
                if key:
                    self.seen_links.add(key)
                    self._link_rows.setdefault(key, row[0].row)
        else:
            self.wb = Workbook()
            self.ws = self.wb.active
            self.ws.title = "Leads"
            self._create_header()
            # Do not truncate a recovery journal before it has replayed.
            self._replaying = True
            self.save(force=True)
            self._replaying = False
        pending = self.journal.read()
        if pending:
            self._replaying = True
            try:
                for entry in pending:
                    record = entry['record']
                    if '_custom_columns' in record:
                        self.add_columns(record['_custom_columns'])
                        continue
                    if '_link_resolution' in record:
                        self.record_link_resolution(**record['_link_resolution'])
                        continue
                    row = entry.get('row')
                    if row is None or row > self.ws.max_row:
                        row = self.find_link_row(record.get('Prospect LinkedIn Link', ''))
                    if row is not None:
                        self.update_record(row, record)
                    else:
                        self.add_record(record)
            finally:
                self._replaying = False
            self.save(force=True)

    def _validate_columns(self, names):
        seen = {name.casefold() for name in self.headers}
        for name in names:
            if (not isinstance(name, str) or not name.strip() or name != name.strip()
                    or len(name) > 120 or name.startswith(('=', '+', '-', '@', '_'))
                    or any(ord(char) < 32 for char in name) or name.casefold() in seen):
                raise ValueError('Column names must be unique, plain text (1–120 characters).')
            seen.add(name.casefold())
        if len(seen) > 16384:
            raise ValueError('Excel supports at most 16,384 columns.')

    def add_columns(self, names):
        names = list(names)
        # Schema events are replayable after a checkpoint or interrupted save.
        names = [name for name in names if name not in self.headers] if self._replaying else names
        self._validate_columns(names)
        if not names:
            return
        if not self._replaying:
            self.journal.append(None, {'_custom_columns': names})
        for name in names:
            self.headers.append(name)
            col = len(self.headers)
            cell = self.ws.cell(1, col, name)
            cell.fill = PatternFill('solid', fgColor='1F4E78')
            cell.font = Font(color='FFFFFF', bold=True)
            cell.alignment = Alignment(wrap_text=True)
            self.ws.column_dimensions[get_column_letter(col)].width = 25
        self._dirty_rows += 1

    def add_manual_rows(self, count):
        if type(count) is not int or not 1 <= count <= 1000:
            raise ValueError('Add between 1 and 1,000 rows at a time.')
        if self.ws.max_row + count > 1048576:
            raise ValueError('Excel worksheet row limit reached.')
        previous_batch = getattr(self, '_batching', False)
        self._batching = True
        try:
            rows = [self.add_record({}) for _ in range(count)]
        finally:
            self._batching = previous_batch
        self.save(force=True)
        return rows

    def record_link_resolution(self, original, resolved='', kind='Invalid',
                               resolution='Unresolved', processing='Not processed', detail='', checked_at=None):
        """Separate audit sheet keeps the 42-column Leads schema compatible."""
        title = 'Link Resolution'
        headers = ['Original Input URL', 'Resolved LinkedIn URL', 'Input Type',
                   'Resolution Status', 'Processing Status', 'Detail', 'Checked At']
        if title in self.wb.sheetnames and [cell.value for cell in self.wb[title][1]] != headers:
            raise ValueError('Existing Link Resolution sheet has different headers; choose another output workbook.')
        entry = dict(original=original, resolved=resolved, kind=kind, resolution=resolution,
                     processing=processing, detail=detail,
                     checked_at=checked_at or datetime.now(timezone.utc).isoformat())
        if not self._replaying:
            self.journal.append(None, {'_link_resolution': entry})
        if title not in self.wb.sheetnames:
            sheet = self.wb.create_sheet(title)
            sheet.append(headers)
            sheet.freeze_panes = 'A2'
            for cell in sheet[1]:
                cell.font = Font(bold=True, color='FFFFFF')
                cell.fill = PatternFill('solid', fgColor='1F4E78')
            for col, width in zip('ABCDEFG', (65, 60, 24, 22, 22, 65, 30)):
                sheet.column_dimensions[col].width = width
        sheet = self.wb[title]
        if not hasattr(self, '_resolution_rows'):
            self._resolution_rows = {row[0].value: row[0].row for row in sheet.iter_rows(min_row=2)}
        row = self._resolution_rows.setdefault(excel_safe(original), sheet.max_row + 1)
        for col, value in enumerate(entry.values(), 1):
            cell = sheet.cell(row, col, excel_safe(value))
            cell.alignment = Alignment(vertical='top', wrap_text=True)
            cell.hyperlink = None
            if col in (1, 2) and str(value).startswith(('https://', 'http://')):
                cell.hyperlink = value
                cell.style = 'Hyperlink'
        sheet.auto_filter.ref = f'A1:G{sheet.max_row}'
        self._dirty_rows += 1
        self.checkpoint()

    def _create_header(self):
        fill = PatternFill("solid", fgColor="1F4E78")
        font = Font(color="FFFFFF", bold=True)

        for col_idx, header in enumerate(self.headers, 1):
            cell = self.ws.cell(row=1, column=col_idx, value=header)
            cell.fill = fill
            cell.font = font
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )

        self.ws.freeze_panes = "A2"
        self.ws.auto_filter.ref = f"A1:{get_column_letter(len(self.headers))}1"
        self.ws.row_dimensions[1].height = 35

        widths = {
            "A": 14, "B": 24, "C": 18, "D": 22, "E": 22, "F": 16,
            "G": 24, "H": 18, "I": 28, "J": 18, "K": 18, "L": 30,
            "M": 44, "N": 34, "O": 20, "P": 24, "Q": 30, "R": 32,
            "S": 26, "T": 20, "U": 24, "V": 34, "W": 38, "X": 20,
            "Y": 20, "Z": 16, "AA": 34, "AB": 20, "AC": 22, "AD": 34,
            "AE": 20, "AF": 34, "AG": 16, "AH": 18, "AI": 34, "AJ": 18,
            "AK": 18, "AL": 18, "AM": 26, "AN": 45, "AO": 22, "AP": 28,
        }

        for letter, width in widths.items():
            self.ws.column_dimensions[letter].width = width

    def is_duplicate(self, linkedin_url: str) -> bool:
        key = normalize_url(linkedin_url)
        return bool(key) and key in self.seen_links

    def find_link_row(self, url: str) -> int | None:
        """Find a saved profile for an enrichment-only update."""
        key = normalize_url(url)
        if not key:
            return None
        return self._link_rows.get(key)

    def read_record(self, row_idx):
        record = dict(zip(self.headers, [clean_cell_value(c.value) for c in self.ws[row_idx]]))
        profile = normalize_url(record['Prospect LinkedIn Link'])
        company = record.get('Company Name', '-').casefold()
        for key, (name, columns) in DETAIL_SHEETS.items():
            if name not in self.wb.sheetnames:
                continue
            sheet = self.wb[name]
            if name not in self._detail_row_indexes:
                self._detail_row_indexes[name] = {
                    (normalize_url(clean_cell_value(row[0].value)), clean_cell_value(row[1].value).casefold()): row[0].row
                    for row in sheet.iter_rows(min_row=2)}
            detail_row = self._detail_row_indexes[name].get((profile, company))
            if detail_row:
                record[key] = {field: sheet.cell(detail_row, col).value
                               for col, (field, _) in enumerate(columns, 1)}
        if record.get('_profile_state', {}).get('company_url'):
            record['Company LinkedIn URL'] = record['_profile_state']['company_url']
        return record

    def profile_needs_visit(self, url, refresh_days=30):
        row = self.find_link_row(url)
        if row is None:
            return True
        record = self.read_record(row)
        if any(SearchWorker._is_unknown_value(record.get(field)) or record.get(field) == 'Dead'
               for field in ('First Name', 'Job Title', 'Company Name', 'Tenure')):
            return True
        if record.get('Company Name', '').casefold() in {
                'full-time', 'full time', 'part-time', 'part time', 'contract',
                'temporary', 'self-employed', 'freelance', 'internship', 'volunteer'}:
            return True
        captured = record.get('_profile_state', {}).get('captured_at')
        if captured:
            try:
                return (datetime.now(timezone.utc) - datetime.fromisoformat(captured)).total_seconds() > refresh_days * 86400
            except (ValueError, TypeError):
                return True
        return False  # Legacy saved rows have unknown age; refresh remains an explicit option.

    def checkpoint(self):
        if (self._dirty_rows >= self.checkpoint_every or
                self._dirty_rows and time.time() - self._last_save_ts >= EXCEL_CHECKPOINT_SECONDS):
            self.save(force=True)

    def _prepare_record_evidence(self, record: dict[str, Any]) -> dict[str, Any]:
        """Validate optional evidence before changing any lead cells or indexes."""
        prepared = dict(record)
        from business_email import build_email_details
        if not isinstance(prepared.get('_email_details'), dict):
            prepared['_email_details'] = build_email_details(record)
        else:
            expected = build_email_details(record)
            provided = prepared['_email_details']
            identity_fields = ('published_email', 'generated_email', 'existing_email', 'email_domain')
            if any(clean_cell_value(provided.get(field)).casefold()
                   != clean_cell_value(expected.get(field)).casefold() for field in identity_fields):
                prepared['_email_details'] = expected
            else:
                prepared['_email_details'] = {**provided, 'mailbox_status': 'Not verified'}
        for key, (sheet_name, columns) in DETAIL_SHEETS.items():
            item = prepared.get(key)
            if not isinstance(item, dict):
                prepared.pop(key, None)
                continue
            profile = normalize_url(clean_cell_value(record.get('Prospect LinkedIn Link')))
            if (not profile or normalize_url(clean_cell_value(item.get('profile'))) != profile
                    or clean_cell_value(item.get('company')).casefold()
                    != clean_cell_value(record.get('Company Name')).casefold()):
                prepared.pop(key, None)
                continue
            if sheet_name in self.wb.sheetnames:
                if [cell.value for cell in self.wb[sheet_name][1]] != [label for _, label in columns]:
                    raise ValueError(f'{sheet_name} sheet headers do not match the expected schema.')
        for key, sheet_name, headers in (
            ('_company_evidence', 'Company Evidence', COMPANY_EVIDENCE_HEADERS),
            ('_email_evidence', 'Email Evidence', EMAIL_EVIDENCE_HEADERS),
        ):
            entries = record.get(key, [])
            accepted = []
            for entry in entries if isinstance(entries, list) else []:
                if not isinstance(entry, dict):
                    continue
                source = clean_cell_value(entry.get('url'))
                if not is_public_source_url(source):
                    continue
                if key == '_company_evidence':
                    if entry.get('field') not in PUBLIC_COMPANY_FIELDS:
                        continue
                else:
                    from business_email import valid_email
                    if not valid_email(entry.get('email')):
                        continue
                    expected_person = ' '.join(clean_cell_value(record.get(field))
                                               for field in ('First Name', 'Last Name'))
                    if any(clean_cell_value(entry.get(field)).casefold() != clean_cell_value(value).casefold()
                           for field, value in (('person', expected_person),
                                                ('company', record.get('Company Name')),
                                                ('email', record.get('Business Email ID')))):
                        continue
                    profile = normalize_url(clean_cell_value(entry.get('profile')))
                    if not profile or profile != normalize_url(clean_cell_value(record.get('Prospect LinkedIn Link'))):
                        continue
                accepted.append(entry)
            prepared[key] = accepted
            if accepted and sheet_name in self.wb.sheetnames:
                if [cell.value for cell in self.wb[sheet_name][1]] != list(headers):
                    raise ValueError(f'{sheet_name} sheet headers do not match the expected schema.')
        return prepared

    def _store_record_details(self, record: dict[str, Any]) -> None:
        """Store current email/office metadata without changing the Leads schema."""
        for key, (sheet_name, columns) in DETAIL_SHEETS.items():
            item = record.get(key)
            if not isinstance(item, dict):
                continue
            if sheet_name not in self.wb.sheetnames:
                sheet = self.wb.create_sheet(sheet_name)
                sheet.append([label for _, label in columns])
                sheet.freeze_panes = 'C2'
                for cell in sheet[1]:
                    cell.fill = PatternFill('solid', fgColor='1F4E78')
                    cell.font = Font(color='FFFFFF', bold=True)
                for index in range(1, len(columns) + 1):
                    sheet.column_dimensions[get_column_letter(index)].width = 32 if index > 1 else 48
            sheet = self.wb[sheet_name]
            profile = normalize_url(clean_cell_value(item.get('profile')))
            company = clean_cell_value(item.get('company')).casefold()
            if sheet_name not in self._detail_row_indexes:
                self._detail_row_indexes[sheet_name] = {
                    (normalize_url(clean_cell_value(row[0].value)), clean_cell_value(row[1].value).casefold()): row[0].row
                    for row in sheet.iter_rows(min_row=2)}
            index = self._detail_row_indexes[sheet_name]
            row_idx = index.get((profile, company), sheet.max_row + 1)
            if key == '_email_details' and (profile, company) in index:
                old = {field: sheet.cell(row_idx, column).value
                       for column, (field, _) in enumerate(columns, 1)}
                if (item.get('domain_status') == 'Not checked (disabled)'
                        and clean_cell_value(old.get('checked_at')) != '-'
                        and all(clean_cell_value(old.get(field)).casefold()
                                == clean_cell_value(item.get(field)).casefold() for field in (
                                    'published_email', 'generated_email', 'existing_email', 'email_domain'))):
                    item = {**item, **{field: old[field] for field in ('domain_status', 'checked_at', 'detail')}}
            index[(profile, company)] = row_idx
            for column, (field, _) in enumerate(columns, 1):
                value = item.get(field, '-')
                if isinstance(value, bool):
                    value = 'Yes' if value else 'No'
                safe_value = re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', '', clean_cell_value(value))[:32767]
                cell = sheet.cell(row_idx, column, safe_value)
                cell.data_type = 's'
                cell.number_format = '@'
                cell.alignment = Alignment(vertical='top', wrap_text=True)
                cell.hyperlink = None
                if field in ('profile', 'source_url') and is_public_source_url(str(value)):
                    cell.hyperlink = str(value)
                    cell.style = 'Hyperlink'
            sheet.auto_filter.ref = f'A1:{get_column_letter(len(columns))}{sheet.max_row}'

    def _store_company_evidence(self, record: dict[str, Any]) -> None:
        """Keep per-field public sources without changing the Leads schema."""
        entries = record.get("_company_evidence", [])
        if not isinstance(entries, list) or not entries:
            return
        name = "Company Evidence"
        if name not in self.wb.sheetnames:
            sheet = self.wb.create_sheet(name)
            sheet.append(list(COMPANY_EVIDENCE_HEADERS))
            sheet.freeze_panes = "A2"
            for cell in sheet[1]:
                cell.fill = PatternFill("solid", fgColor="1F4E78")
                cell.font = Font(color="FFFFFF", bold=True)
            for column, width in zip("ABCDEFGHI", (30, 34, 20, 24, 60, 20, 24, 95, 27)):
                sheet.column_dimensions[column].width = width
        sheet = self.wb[name]
        if [cell.value for cell in sheet[1]] != list(COMPANY_EVIDENCE_HEADERS):
            raise ValueError("Company Evidence sheet headers do not match the expected schema.")
        if self._company_evidence_rows is None:
            self._company_evidence_rows = {}
            for row in sheet.iter_rows(min_row=2):
                values = [clean_cell_value(cell.value) for cell in row]
                key = (values[0].casefold(), values[2], values[3], values[4])
                self._company_evidence_rows[key] = row[0].row
        for entry in entries:
            if not isinstance(entry, dict):
                continue
            values = [clean_cell_value(entry.get(field, "-")) for field in (
                "company", "website", "field", "value", "url", "period",
                "basis", "excerpt", "checked_at",
            )]
            if values[2] not in PUBLIC_COMPANY_FIELDS or not values[4].startswith(("https://", "http://")):
                continue
            key = (values[0].casefold(), values[2], values[3], values[4])
            row_idx = self._company_evidence_rows.get(key, sheet.max_row + 1)
            self._company_evidence_rows[key] = row_idx
            for col_idx, value in enumerate(values, 1):
                cell = sheet.cell(row=row_idx, column=col_idx, value=excel_safe(value))
                cell.number_format = "@"
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if col_idx in (2, 5) and value.startswith(("https://", "http://")):
                    cell.hyperlink = value
                    cell.style = "Hyperlink"
        sheet.auto_filter.ref = f"A1:I{sheet.max_row}"

    def _store_email_evidence(self, record: dict[str, Any]) -> None:
        entries = record.get('_email_evidence', [])
        if not isinstance(entries, list) or not entries:
            return
        name = 'Email Evidence'
        if name not in self.wb.sheetnames:
            sheet = self.wb.create_sheet(name)
            sheet.append(list(EMAIL_EVIDENCE_HEADERS))
            sheet.freeze_panes = 'A2'
            for cell in sheet[1]:
                cell.fill = PatternFill('solid', fgColor='1F4E78')
                cell.font = Font(color='FFFFFF', bold=True)
            for column, width in zip('ABCDEFGHI', (26, 45, 30, 36, 45, 24, 60, 95, 28)):
                sheet.column_dimensions[column].width = width
        sheet = self.wb[name]
        if [cell.value for cell in sheet[1]] != list(EMAIL_EVIDENCE_HEADERS):
            raise ValueError('Email Evidence sheet headers do not match the expected schema.')
        if self._email_evidence_rows is None:
            self._email_evidence_rows = {
                (str(row[1].value), str(row[2].value), str(row[3].value), str(row[6].value)): row[0].row
                for row in sheet.iter_rows(min_row=2)}
        existing = self._email_evidence_rows
        for entry in entries:
            if not isinstance(entry, dict):
                continue
            values = [clean_cell_value(entry.get(key, '-')) for key in (
                'person', 'profile', 'company', 'email', 'method', 'pattern', 'url', 'excerpt', 'checked_at')]
            if '@' not in values[3] or not values[6].startswith(('http://', 'https://')):
                continue
            key = (values[1], values[2], values[3], values[6])
            row_number = existing.get(key, sheet.max_row + 1)
            existing[key] = row_number
            for column, value in enumerate(values, 1):
                cell = sheet.cell(row_number, column, excel_safe(value))
                cell.number_format = '@'
                cell.alignment = Alignment(vertical='top', wrap_text=True)
                if column in (2, 7) and value.startswith(('http://','https://')):
                    cell.hyperlink = value
                    cell.style = 'Hyperlink'
        sheet.auto_filter.ref = f'A1:I{sheet.max_row}'

    def incomplete_link_row(self, url: str) -> int | None:
        """Allow a later run to repair a previously saved incomplete profile."""
        key = normalize_url(url)
        required = (
            'Tenure', 'Job Title', 'Company Name', 'Company Website',
            'Industry', 'Company Size', 'City', 'State', 'Country',
        )
        invalid_company_values = {
            'full-time', 'full time', 'part-time', 'part time',
            'contract', 'temporary', 'self-employed', 'freelance',
            'internship', 'volunteer', 'apprenticeship', 'seasonal',
        }
        for row in self.ws.iter_rows(min_row=2):
            record = dict(zip(HEADERS, [cell.value for cell in row]))
            if normalize_url(str(record.get('Prospect LinkedIn Link') or '')) != key:
                continue
            company_value = str(record.get('Company Name') or '').strip().casefold()
            if (
                any(record.get(field) in (None, '', '-', 'Dead') for field in required)
                or company_value in invalid_company_values
            ):
                return row[0].row
        return None

    def add_record(self, record: dict[str, Any]):
        """
        Fast row append.

        Instead of constructing/styling 42 cells one at a time, append the row
        in one operation and touch only URL cells afterwards.
        """
        link = normalize_url(str(record.get("Prospect LinkedIn Link", "-")))

        if link and link in self.seen_links:
            raise ValueError("Duplicate prospect URL: " + link)

        record = self._prepare_record_evidence(record)

        if not self._replaying:
            self.journal.append(self.ws.max_row + 1, record)

        raw_values = [
            re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', '', clean_cell_value(record.get(header, '-')))[:32767]
            for header in self.headers
        ]

        values = [
            (
                actual
                if header == "Phone Number"
                else excel_safe(actual)
            )
            for header, actual in zip(
                self.headers,
                raw_values,
            )
        ]

        self.ws.append(values)
        row_idx = self.ws.max_row

        # Only URL fields need additional cell work. Phone Number is explicitly
        # stored as text so a leading "+" is preserved exactly.
        for col_idx, (header, actual) in enumerate(
            zip(self.headers, raw_values),
            1,
        ):
            cell = self.ws.cell(
                row=row_idx,
                column=col_idx,
            )

            if header == "Phone Number":
                cell.number_format = "@"
                cell.data_type = "s"

            if actual.startswith(("http://", "https://", "tel:", "mailto:")):
                cell.hyperlink = actual
                cell.style = "Hyperlink"

        if link and link != "-":
            self.seen_links.add(link)
            self._link_rows[link] = row_idx
        self._store_company_evidence(record)
        self._store_email_evidence(record)
        self._store_record_details(record)
        self._dirty_rows += 1
        now = time.time()

        if not getattr(self, '_batching', False) and (
            self._dirty_rows >= self.checkpoint_every
            or (now - self._last_save_ts) >= EXCEL_CHECKPOINT_SECONDS
        ):
            self.save(force=True)

        return row_idx

    def update_record(
        self,
        row_idx: int,
        record: dict[str, Any],
    ) -> None:
        """
        Update an existing data row, preserving fields omitted by the caller.
        """
        if not isinstance(row_idx, int) or row_idx < 2 or row_idx > self.ws.max_row:
            raise ValueError('Only an existing Leads data row can be updated.')
        previous = dict(zip(self.headers, [cell.value for cell in self.ws[row_idx]]))
        record = {**previous, **record}
        old_link = normalize_url(clean_cell_value(previous.get('Prospect LinkedIn Link')))
        new_link = normalize_url(clean_cell_value(record.get('Prospect LinkedIn Link')))
        if new_link and new_link != old_link and new_link in self.seen_links:
            raise ValueError('Duplicate prospect URL: ' + new_link)
        record = self._prepare_record_evidence(record)
        if not self._replaying:
            self.journal.append(row_idx, record)
        raw_values = [
            re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', '', clean_cell_value(record.get(header, '-')))[:32767]
            for header in self.headers
        ]

        for col_idx, (header, actual) in enumerate(
            zip(self.headers, raw_values),
            1,
        ):
            value = (
                actual
                if header == "Phone Number"
                else excel_safe(actual)
            )

            cell = self.ws.cell(
                row=row_idx,
                column=col_idx,
                value=value,
            )

            # Remove stale hyperlink first, then re-add when appropriate.
            cell.hyperlink = None

            if header == "Phone Number":
                cell.number_format = "@"
                cell.data_type = "s"

            if actual.startswith(
                ("http://", "https://", "tel:", "mailto:")
            ):
                cell.hyperlink = actual
                cell.style = "Hyperlink"

        self._store_company_evidence(record)
        self._store_email_evidence(record)
        self._store_record_details(record)
        if old_link != new_link:
            self._link_rows.pop(old_link, None)
            if old_link:
                self.seen_links.discard(old_link)
            if new_link:
                self.seen_links.add(new_link)
                self._link_rows[new_link] = row_idx
        self._dirty_rows += 1

    def enrichment_audit(self) -> dict[str, Any]:
        """
        Inspect every saved row and report unresolved enrichment fields.
        """
        unresolved_rows: list[dict[str, Any]] = []
        field_counts = {
            field: 0
            for field in ENRICHMENT_AUDIT_FIELDS
        }

        header_map = {
            self.ws.cell(row=1, column=col).value: col
            for col in range(1, self.ws.max_column + 1)
        }

        for row_idx in range(2, self.ws.max_row + 1):
            missing: list[str] = []

            for field in ENRICHMENT_AUDIT_FIELDS:
                col = header_map.get(field)

                if not col:
                    missing.append(field)
                    field_counts[field] += 1
                    continue

                value = clean_cell_value(
                    self.ws.cell(
                        row=row_idx,
                        column=col,
                    ).value
                )

                if SearchWorker._is_unknown_value(value):
                    missing.append(field)
                    field_counts[field] += 1

            if missing:
                unresolved_rows.append(
                    {
                        "excel_row": row_idx,
                        "first_name": clean_cell_value(
                            self.ws.cell(
                                row=row_idx,
                                column=header_map.get("First Name", 1),
                            ).value
                        ),
                        "last_name": clean_cell_value(
                            self.ws.cell(
                                row=row_idx,
                                column=header_map.get("Last Name", 1),
                            ).value
                        ),
                        "company": clean_cell_value(
                            self.ws.cell(
                                row=row_idx,
                                column=header_map.get("Company Name", 1),
                            ).value
                        ),
                        "missing_fields": missing,
                    }
                )

        return {
            "total_rows": max(self.ws.max_row - 1, 0),
            "rows_with_unresolved_fields": len(unresolved_rows),
            "field_missing_counts": field_counts,
            "unresolved_rows": unresolved_rows,
        }

    def save(self, force: bool = False):
        if not force and self._dirty_rows <= 0:
            return
        self.output_path.parent.mkdir(parents=True, exist_ok=True)
        self.ws.auto_filter.ref = f"A1:{get_column_letter(len(self.headers))}{self.ws.max_row}"
        # Replace only after the complete ZIP has been written successfully.
        fd, temporary = tempfile.mkstemp(suffix=".xlsx", dir=self.output_path.parent)
        os.close(fd)
        try:
            with self.metrics.measure('excel_save'):
                self.wb.save(temporary)
                os.replace(temporary, self.output_path)
            if not self._replaying:
                self.journal.clear()
        except PermissionError as exc:
            raise PermissionError(
                f"Cannot save {self.output_path}. Close it in Excel and retry."
            ) from exc
        finally:
            Path(temporary).unlink(missing_ok=True)
        self._dirty_rows = 0
        self._last_save_ts = time.time()


class SearchWorker(threading.Thread):
    """
    Owns Playwright and Chrome in one thread.
    UI commands are passed through a queue.
    """

    def __init__(self, callbacks: dict[str, Any]):
        super().__init__(daemon=True)

        self.callbacks = callbacks
        self.commands: queue.Queue = queue.Queue()

        self.pause_event = threading.Event()
        self.stop_event = threading.Event()

        self.pw = None
        self.context = None
        self.page = None
        self.google_maps_page = None

        self.llm = LocalLLMClient(
            trace_callback=self._local_llm_trace,
        )
        self.llm_available: bool | None = None
        self.llm_last_check_ts = 0.0

        # Per-run performance counters.
        self.dom_calls = 0
        self.dom_seconds = 0.0
        self.qwen_calls = 0
        self.qwen_seconds = 0.0

        # Reuse company details when multiple prospects work at the same company.
        self.company_cache: dict[str, dict[str, str]] = {}

        # SMART-mode caches prevent repeated local-LLM validation for the same
        # company and the same job title.
        self.smart_company_cache: dict[str, dict[str, str]] = {}
        self.smart_job_cache: dict[str, dict[str, str]] = {}
        self.website_evidence_cache: dict[str, str] = {}
        self.google_maps_cache: dict[str, dict[str, str]] = {}
        self.public_company_enricher = None
        self.business_email_enricher = None
        self.email_domain_checks_enabled = False
        self.email_domain_checker = None
        self._last_profile_locations: list[str] = []
        self.headless = False
        self._browser_headless = None
        self.prefer_local_office = False
        self.metrics = RunMetrics(self.callbacks.get('stage'))
        self.profile_cache = {}
        self.cache_path = None
        self.enrichment_sessions = None
        self.refresh_profiles = False
        self.enrichment_attempts = {}
        self.link_resolutions = None

    def _configure_run(self, payload):
        self._finish_enrichment_sessions()
        self.headless = bool(payload.get('headless', False))
        self.prefer_local_office = bool(payload.get('prefer_local_office', False))
        self.refresh_profiles = bool(payload.get('refresh_profiles', False))
        self.metrics = RunMetrics(self.callbacks.get('stage'))
        payload['workbook'].metrics.on_stage = self.callbacks.get('stage')
        self.cache_path = None
        if payload.get('persistent_cache', False):
            self.cache_path = payload['workbook'].output_path.parent / '.scraper_cache_v42_7.sqlite3'
        if payload.get('enrichment_sessions', 0):
            from enrichment_session import IsolatedEnrichmentSession
            self.enrichment_sessions = EnrichmentSessions(
                lambda stage: IsolatedEnrichmentSession(self, stage),
                concurrent=int(payload['enrichment_sessions']) > 1, on_error=self.log)

    def _install_caches(self):
        if self.cache_path:
            for attr, namespace, days in (
                ('company_cache', 'linkedin_company', 7),
                ('profile_cache', 'profile_capture', 30),
                ('google_maps_cache', 'maps', 7),
                ('enrichment_attempts', 'enrichment_attempts', 15 / 1440),
            ):
                setattr(self, attr, EvidenceCache(self.cache_path, namespace, days * 86400, self.metrics))

    def _attempt_signature(self, record):
        return {'fields': [clean_cell_value(record.get(key)) for key in HEADERS[8:35]],
                'prefer_local_office': self.prefer_local_office,
                'email_domain_checks': self.email_domain_checks_enabled}

    def _remember_enrichment(self, record):
        if str(record.get('Email Status', '')).startswith('Time limit reached'):
            return  # A timed-out lookup must remain eligible for a later retry.
        if self.cache_path and not self.stop_event.is_set():
            key = normalize_url(record.get('Prospect LinkedIn Link', ''))
            if key:
                self.enrichment_attempts[key] = self._attempt_signature(record)

    def _finish_enrichment_sessions(self):
        if self.enrichment_sessions:
            try:
                self.enrichment_sessions.close()
            finally:
                self.enrichment_sessions = None

    def _enrich_offsite(self, record, source_label):
        with self.metrics.measure('offsite_enrichment'):
            if self.enrichment_sessions:
                return self.enrichment_sessions.enrich(record)
            if self._google_maps_needs_enrichment(record) or self._needs_local_office(record):
                record = self._google_maps_enrich_record(record, source_label=source_label)
            if self._public_company_needs_enrichment(record):
                record = self._public_company_enrich_record(record, source_label=source_label)
            if self._business_email_needs_enrichment(record) or self.email_domain_checks_enabled:
                record = self._business_email_enrich_record(record, source_label=source_label)
            return record

    def _needs_local_office(self, record):
        if not self.prefer_local_office or not any(not self._is_unknown_value(record.get(k))
                                                 for k in ('City', 'State', 'Country')):
            return False
        target = {**record, 'Address': '-'}
        address = record.get('Address', '-')
        return (not self._is_unknown_value(address) and
                not self._maps_address_matches_location(target, address))

    def _capture_profile_once(self, profile_url):
        key = normalize_url(profile_url)
        if not self.refresh_profiles and key in self.profile_cache:
            self.metrics.count('profile_visits_avoided')
            return self.profile_cache[key]
        self._launch_browser()
        with self.metrics.measure('linkedin_profile'):
            data = self._capture_standard_profile(profile_url)
        data['_captured_at'] = datetime.now(timezone.utc).isoformat()
        self.profile_cache[key] = data
        return data

    def _report_performance(self, workbook):
        report = {**self.metrics.snapshot(), 'workbook': workbook.metrics.snapshot()}
        path = workbook.output_path.with_suffix('.performance.json')
        path.write_text(json.dumps(report, indent=2), encoding='utf-8')
        self.log('PERFORMANCE | ' + json.dumps(report))

    def _reset_public_company_enricher(self) -> None:
        self.business_email_enricher = None
        if self.public_company_enricher is not None:
            try:
                self.public_company_enricher.close()
            except Exception as exc:
                self.log(f"PUBLIC COMPANY CLOSE WARNING | {type(exc).__name__}: {exc}")
            finally:
                self.public_company_enricher = None

    @staticmethod
    def _public_company_needs_enrichment(record: dict[str, str]) -> bool:
        return any(SearchWorker._field_missing_or_invalid(record, field)
                   for field in PUBLIC_COMPANY_FIELDS)

    @timed('public_company')
    def _public_company_enrich_record(
        self, record: dict[str, str], source_label: str,
    ) -> dict[str, str]:
        """Fetch sourced company facts in both modes; optional failure keeps the row."""
        if not self.context or self.stop_event.is_set():
            return record

        if not self._public_company_needs_enrichment(record):
            return record
        try:
            if self.public_company_enricher is None:
                from company_public_data import PublicCompanyEnricher
                self.public_company_enricher = PublicCompanyEnricher(
                    self.context, company_match=self._link_company_names_match,
                    log=self.log,
                    trace=lambda event, **payload: self.trace("PUBLIC_COMPANY", event, **payload),
                    stop_event=self.stop_event, pause_check=self._wait_if_paused,
                )
                if self.cache_path:
                    self.public_company_enricher.cache = EvidenceCache(
                        self.cache_path, 'public_company', 7 * 86400, self.metrics)
            return self.public_company_enricher.enrich(dict(record))
        except Exception as exc:
            self.log(f"PUBLIC COMPANY ENRICHMENT WARNING | source={source_label!r} | "
                     f"company={record.get('Company Name', '-')!r} | {type(exc).__name__}: {exc}")
            return record

    @staticmethod
    def _business_email_needs_enrichment(record: dict[str, str]) -> bool:
        from business_email import BusinessEmailEnricher
        return BusinessEmailEnricher.needs_enrichment(record)

    def _retain_same_employer_email(self, record: dict, previous: dict) -> dict:
        """Keep known contact data only when the repaired identity still agrees."""
        if self._is_unknown_value(previous.get('Business Email ID')):
            return record
        from business_email import normalized, valid_email
        from company_public_data import host
        fields = ('First Name', 'Last Name', 'Company Name', 'Company Website')
        if any(self._is_unknown_value(row.get(field))
               for row in (record, previous) for field in fields):
            return record
        same_person = all(normalized(clean_cell_value(record[field])) ==
                          normalized(clean_cell_value(previous[field]))
                          for field in ('First Name', 'Last Name'))
        same_company = self._link_company_names_match(
            record['Company Name'], previous['Company Name'])
        try:
            same_domain = host(record['Company Website']) == host(previous['Company Website'])
        except (TypeError, ValueError):
            same_domain = False
        if same_person and same_company and same_domain and valid_email(previous.get('Business Email ID')):
            for field in ('Business Email ID', 'Email Status', 'Bounce Email Status'):
                record[field] = previous.get(field, '-')
        return record

    @timed('email')
    def _business_email_enrich_record(self, record: dict[str, str], source_label: str) -> dict[str, str]:
        if self.stop_event.is_set() or (not self.context and not self.email_domain_checks_enabled):
            return record
        email_started = time.monotonic()
        result = self._lookup_business_email(record, source_label)
        from business_email import build_email_details
        checker = None
        if (self.email_domain_checks_enabled and not self.stop_event.is_set()
                and time.monotonic() - email_started < 14.9):
            if self.email_domain_checker is None:
                from email_domain_check import DomainMXChecker
                self.email_domain_checker = DomainMXChecker()
            checker = self.email_domain_checker
            checker.total_timeout = min(8.0, max(0.1, 15.0 - (time.monotonic() - email_started)))
        return {**result, '_email_details': build_email_details(result, checker)}

    def _lookup_business_email(self, record: dict[str, str], source_label: str) -> dict[str, str]:
        if not self.context or self.stop_event.is_set():
            return record
        try:
            from business_email import BusinessEmailEnricher
            if not BusinessEmailEnricher.needs_enrichment(record):
                return record
            if self.public_company_enricher is None:
                from company_public_data import PublicCompanyEnricher
                self.public_company_enricher = PublicCompanyEnricher(
                    self.context, company_match=self._link_company_names_match,
                    log=self.log, trace=lambda event, **payload: self.trace('COMPANY_WEB', event, **payload),
                    stop_event=self.stop_event, pause_check=self._wait_if_paused)
            if self.business_email_enricher is None:
                self.business_email_enricher = BusinessEmailEnricher(
                    self.public_company_enricher, self._link_company_names_match, log=self.log)
                if self.cache_path:
                    self.business_email_enricher.pattern_cache = EvidenceCache(
                        self.cache_path, 'email_patterns', 7 * 86400, self.metrics)
            return self.business_email_enricher.enrich(record)
        except Exception as exc:
            self.log(f'BUSINESS EMAIL WARNING | source={source_label!r} | {type(exc).__name__}: {exc}')
            return record

    def _enrich_saved_company_row(
        self, workbook: WorkbookManager, profile_url: str, source_label: str,
    ) -> dict[str, str] | None:
        """Retry missing Maps and public company fields on an existing profile."""
        row_idx = workbook.find_link_row(profile_url)
        if row_idx is None:
            return None
        record = workbook.read_record(row_idx)
        original = dict(record)
        if (not self.refresh_profiles and self.cache_path and
                self.enrichment_attempts.get(normalize_url(profile_url)) == self._attempt_signature(record)):
            self.metrics.count('recent_enrichment_attempts_skipped')
            return None
        company_url = record.get('Company LinkedIn URL', '-')
        if company_url == '-':
            cached = self.profile_cache.get(normalize_url(profile_url))
            if cached:
                company_url = self._parse_standard_current_experience(cached).get('Company Profile Link', '-')
        if (company_url != '-' and '/in/' in profile_url and any(
                self._field_missing_or_invalid(record, field) for field in
                ('Company Website', 'Industry', 'Company Size'))):
            details = self._read_standard_company_about(company_url, record.get('Company Name', '-'))
            for key, value in details.items():
                if self._is_unknown_value(record.get(key)) and not self._is_unknown_value(value):
                    record[key] = value
        needs_maps = self._google_maps_needs_enrichment(record) or self._needs_local_office(record)
        needs_public = self._public_company_needs_enrichment(record)
        needs_email = (self._business_email_needs_enrichment(record)
                       or self.email_domain_checks_enabled)
        if not (needs_maps or needs_public or needs_email) and record == original:
            return None
        record = self._enrich_offsite(record, source_label)
        workbook.update_record(row_idx, record)
        workbook.checkpoint()
        self._remember_enrichment(record)
        unresolved = [field for field in PUBLIC_COMPANY_FIELDS
                      if self._field_missing_or_invalid(record, field)]
        self.log(f"COMPANY EXISTING ROW | row={row_idx} | "
                 f"company={record.get('Company Name', '-')!r} | "
                 f"maps_pending={self._google_maps_needs_enrichment(record)} | "
                 f"public_unresolved={', '.join(unresolved) or 'none'}")
        # Metadata can be refreshed without counting an unchanged old row toward
        # the requested new/updated-lead target.
        if self.cache_path and all(clean_cell_value(record.get(key)) == clean_cell_value(original.get(key))
                                   for key in HEADERS[8:35]):
            return None
        return record

    @staticmethod
    def _is_unknown_value(value: Any) -> bool:
        if value is None:
            return True

        s = re.sub(r"\s+", " ", str(value)).strip().lower()

        return s in {
            "",
            "-",
            "unknown",
            "n/a",
            "na",
            "none",
            "null",
            "nan",
            "not available",
            "unavailable",
            "not found",
            "not provided",
        }

    @staticmethod
    def _is_valid_http_url(value: str) -> bool:
        value = (value or "").strip()

        if not value.startswith(("http://", "https://")):
            return False

        try:
            parts = urlsplit(value)
            host = (parts.hostname or "").strip().lower()

            if not host or " " in host:
                return False

            return True
        except Exception:
            return False

    @staticmethod
    def _is_valid_official_website(value: str) -> bool:
        if not SearchWorker._is_valid_http_url(value):
            return False

        host = (urlsplit(value).hostname or "").lower()
        host = host.removeprefix("www.")

        # Reject malformed single-label guesses such as https://acme.
        if "." not in host:
            return False

        blocked_hosts = (
            "linkedin.com",
            "facebook.com",
            "instagram.com",
            "twitter.com",
            "x.com",
            "crunchbase.com",
            "zoominfo.com",
            "apollo.io",
            "bloomberg.com",
            "wikipedia.org",
            "glassdoor.com",
            "indeed.com",
            "dnb.com",
            "pitchbook.com",
        )

        return not any(
            host == blocked
            or host.endswith("." + blocked)
            for blocked in blocked_hosts
        )

    @staticmethod
    def _smart_company_cache_key(record: dict[str, str]) -> str:
        """
        Stable company-level SMART cache key.

        Never include an enriched field such as Company Website in the key,
        otherwise the key changes after enrichment and future leads cannot hit
        the cache.
        """
        company = SearchWorker._normalize_company_for_match(
            clean_cell_value(record.get("Company Name", "-"))
        )

        company_url = normalize_url(
            clean_cell_value(record.get("Company LinkedIn URL", "-"))
        )

        if company_url and company_url != "-":
            return "|".join(
                part
                for part in (company, company_url)
                if part
            )

        website = clean_cell_value(
            record.get("Company Website", "-")
        )

        if SearchWorker._is_valid_official_website(website):
            host = (urlsplit(website).hostname or "").lower()
            host = host.removeprefix("www.")
            return "|".join(
                part
                for part in (company, host)
                if part
            )

        return company


    @staticmethod
    def _smart_job_cache_key(record: dict[str, str]) -> str:
        title = re.sub(
            r"\s+",
            " ",
            clean_cell_value(record.get("Job Title", "-")),
        ).strip().lower()

        if SearchWorker._is_unknown_value(title):
            return ""

        return title

    def _apply_smart_cache(
        self,
        record: dict[str, str],
        requested_fields: list[str],
    ) -> tuple[list[str], list[str]]:
        """
        Apply previously accepted SMART results conservatively.

        Rules:
        - Missing/invalid fields may be filled from cache.
        - A valid scraped company/geography value is never overwritten merely
          because a cache entry exists.
        - Job Level/Function can reuse an accepted result for the exact same
          normalized Job Title.
        - When a valid current value conflicts with a company cache value,
          leave it for normal SMART validation instead of silently replacing it.
        """
        company_fields = {
            "Company Website",
            "Industry",
            "Address",
            "City",
            "State",
            "Zip Code",
            "Country",
            "Revenue",
            "SIC Code",
            "NAICS Code",
            "SIC/NAICS Code Link",
        }

        job_fields = {
            "Job Level",
            "Job Function",
        }

        remaining = list(requested_fields)
        cache_hits: list[str] = []

        company_key = self._smart_company_cache_key(record)
        job_key = self._smart_job_cache_key(record)

        company_cached = (
            self.smart_company_cache.get(company_key, {})
            if company_key
            else {}
        )
        job_cached = (
            self.smart_job_cache.get(job_key, {})
            if job_key
            else {}
        )

        for field in list(remaining):
            cached_value = None
            has_cache = False

            if field in company_fields and field in company_cached:
                cached_value = clean_cell_value(company_cached[field])
                has_cache = True
            elif field in job_fields and field in job_cached:
                cached_value = clean_cell_value(job_cached[field])
                has_cache = True

            if not has_cache:
                continue

            current_value = clean_cell_value(record.get(field, "-"))
            current_bad = self._field_missing_or_invalid(record, field)
            cached_unknown = self._is_unknown_value(cached_value)

            if field in job_fields:
                # Same normalized title -> accepted classification is reusable.
                if not cached_unknown:
                    record[field] = cached_value
                    remaining.remove(field)
                    cache_hits.append(field)
                elif current_bad:
                    record[field] = "Unknown"
                    remaining.remove(field)
                    cache_hits.append(field)
                # If cached Unknown but current value is valid, leave current
                # value for live SMART validation.
                continue

            # Company/geography fields:
            if current_bad:
                record[field] = (
                    "Unknown"
                    if cached_unknown
                    else cached_value
                )
                remaining.remove(field)
                cache_hits.append(field)
                continue

            # Current direct scrape is valid. If cache agrees, no LLM call is
            # needed. If it disagrees, preserve the direct scrape and let SMART
            # validate the discrepancy live.
            if (
                not cached_unknown
                and cached_value == current_value
            ):
                remaining.remove(field)
                cache_hits.append(field)

        return remaining, cache_hits


    def _store_smart_cache(
        self,
        record: dict[str, str],
        fields: list[str],
    ) -> None:
        company_fields = {
            "Company Website",
            "Industry",
            "Address",
            "City",
            "State",
            "Zip Code",
            "Country",
            "Revenue",
            "SIC Code",
            "NAICS Code",
            "SIC/NAICS Code Link",
        }

        job_fields = {
            "Job Level",
            "Job Function",
        }

        company_key = self._smart_company_cache_key(record)
        job_key = self._smart_job_cache_key(record)

        if company_key:
            bucket = self.smart_company_cache.setdefault(company_key, {})

            for field in fields:
                if field in company_fields:
                    bucket[field] = clean_cell_value(
                        record.get(field, "Unknown")
                    )

        if job_key:
            bucket = self.smart_job_cache.setdefault(job_key, {})

            for field in fields:
                if field in job_fields:
                    bucket[field] = clean_cell_value(
                        record.get(field, "Unknown")
                    )

    @staticmethod
    def _field_missing_or_invalid(
        record: dict[str, str],
        field: str,
    ) -> bool:
        value = clean_cell_value(record.get(field, "-"))
        missing = SearchWorker._is_unknown_value(value)

        if missing:
            return True

        if field == "Job Level":
            return value not in {
                "C-Level",
                "VP",
                "Director",
                "Manager",
                "Senior",
                "Individual Contributor",
                "Owner/Founder",
                "Partner",
            }

        if field == "Job Function":
            return value not in {
                "Finance",
                "Sales",
                "Marketing",
                "Operations",
                "Information Technology",
                "Human Resources",
                "Procurement",
                "Supply Chain",
                "Engineering",
                "Legal",
                "Customer Service",
                "Administration",
                "Product",
                "Security",
                "Healthcare",
                "Other",
            }

        if field == "Company Website":
            return not SearchWorker._is_valid_official_website(value)

        if field == "Industry":
            return (
                SearchWorker._looks_like_revenue_line(value)
                or "employees" in value.lower()
                or value.lower().startswith("star the account")
            )

        if field == "Zip Code":
            return not SearchWorker._is_plausible_zip(value)

        if field == "Revenue":
            return not SearchWorker._is_plausible_revenue(value)

        if field == "SIC Code":
            return re.fullmatch(r"\d{4}", value) is None

        if field == "NAICS Code":
            return re.fullmatch(r"\d{6}", value) is None

        if field in {"Revenue Link", "SIC/NAICS Code Link"}:
            # Company source URLs need not contain their reported code digits.
            # The public-data service verifies the value against fetched text.
            return not SearchWorker._is_valid_http_url(value)

        # Industry / Address / City / State / Country are structurally valid
        # when non-empty. SMART mode may still explicitly validate geography.
        return False

    @staticmethod
    def _required_enrichment_fields(
        record: dict[str, str],
        validate_existing: bool = True,
    ) -> list[str]:
        """
        Identify required fields that are missing/invalid.

        SMART mode additionally revalidates Job Level, Job Function and
        geography even when a syntactically valid scraped value exists.
        """
        required = [
            "Job Level",
            "Job Function",
            "Company Website",
            "Industry",
            "Address",
            "City",
            "State",
            "Zip Code",
            "Country",
            "Revenue",
            "SIC Code",
            "NAICS Code",
            "SIC/NAICS Code Link",
        ]

        validation_fields = (
            {
                "Job Level",
                "Job Function",
                "City",
                "State",
                "Zip Code",
                "Country",
            }
            if validate_existing
            else set()
        )

        return [
            field
            for field in required
            if field not in PUBLIC_COMPANY_FIELDS and (
                SearchWorker._field_missing_or_invalid(record, field)
                or field in validation_fields
            )
        ]


    @staticmethod
    def _canonical_job_level(value: str) -> str:
        raw = re.sub(r"\s+", " ", value or "").strip()
        low = raw.lower()

        exact = {
            "c-level": "C-Level",
            "c level": "C-Level",
            "executive": "C-Level",
            "chief": "C-Level",
            "vp": "VP",
            "vice president": "VP",
            "svp": "VP",
            "senior vice president": "VP",
            "evp": "VP",
            "executive vice president": "VP",
            "director": "Director",
            "managing director": "Director",
            "manager": "Manager",
            "senior manager": "Manager",
            "sr manager": "Manager",
            "sr. manager": "Manager",
            "senior": "Senior",
            "individual contributor": "Individual Contributor",
            "individual": "Individual Contributor",
            "owner": "Owner/Founder",
            "founder": "Owner/Founder",
            "co-founder": "Owner/Founder",
            "cofounder": "Owner/Founder",
            "owner/founder": "Owner/Founder",
            "partner": "Partner",
        }

        return exact.get(low, "Unknown")

    @staticmethod
    def _canonical_job_function(value: str) -> str:
        raw = re.sub(r"\s+", " ", value or "").strip()
        low = raw.lower()

        exact = {
            "finance": "Finance",
            "finance / accounting": "Finance",
            "finance/accounting": "Finance",
            "accounting": "Finance",
            "sales": "Sales",
            "sales / business development": "Sales",
            "sales/business development": "Sales",
            "business development": "Sales",
            "marketing": "Marketing",
            "operations": "Operations",
            "information technology": "Information Technology",
            "it": "Information Technology",
            "it / engineering": "Information Technology",
            "it/engineering": "Information Technology",
            "human resources": "Human Resources",
            "hr": "Human Resources",
            "procurement": "Procurement",
            "procurement / supply chain": "Procurement",
            "procurement/supply chain": "Procurement",
            "supply chain": "Supply Chain",
            "engineering": "Engineering",
            "legal": "Legal",
            "customer service": "Customer Service",
            "customer support": "Customer Service",
            "administration": "Administration",
            "administrative": "Administration",
            "product": "Product",
            "security": "Security",
            "healthcare": "Healthcare",
            "other": "Other",
        }

        return exact.get(low, "Unknown")

    @staticmethod
    def _canonical_country(value: str) -> str:
        raw = re.sub(r"\s+", " ", value or "").strip()

        if not raw:
            return "Unknown"

        key = re.sub(r"[.\s]+", " ", raw.lower()).strip()
        key = re.sub(r"\s+", " ", key)

        mapping = {
            "us": "United States",
            "u s": "United States",
            "usa": "United States",
            "u s a": "United States",
            "united states of america": "United States",
            "uk": "United Kingdom",
            "u k": "United Kingdom",
            "gb": "United Kingdom",
            "great britain": "United Kingdom",
            "uae": "United Arab Emirates",
            "u a e": "United Arab Emirates",
            "ar": "Argentina",
            "ca": "Canada",
            "in": "India",
            "sg": "Singapore",
            "au": "Australia",
            "nz": "New Zealand",
            "my": "Malaysia",
            "ph": "Philippines",
            "id": "Indonesia",
            "th": "Thailand",
            "vn": "Vietnam",
            "de": "Germany",
            "fr": "France",
            "it": "Italy",
            "es": "Spain",
            "nl": "Netherlands",
            "ie": "Ireland",
            "no": "Norway",
            "se": "Sweden",
            "dk": "Denmark",
            "fi": "Finland",
            "ch": "Switzerland",
            "at": "Austria",
            "pl": "Poland",
            "cz": "Czech Republic",
            "tr": "Turkey",
            "mx": "Mexico",
            "cn": "China",
            "jp": "Japan",
            "hk": "Hong Kong",
            "sa": "Saudi Arabia",
            "za": "South Africa",
        }

        return mapping.get(key, raw)

    @staticmethod
    def _is_plausible_zip(value: str) -> bool:
        value = re.sub(r"\s+", " ", value or "").strip()

        if SearchWorker._is_unknown_value(value):
            return False

        # Broad international postal-code structure. Exact geographic
        # correctness is delegated to SMART validation.
        return bool(
            3 <= len(value) <= 12
            and re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9 -]*", value)
        )

    @staticmethod
    def _is_plausible_revenue(value: str) -> bool:
        value = re.sub(r"\s+", " ", value or "").strip()

        if SearchWorker._is_unknown_value(value):
            return False

        # Require either an explicit currency or a financial magnitude. This
        # avoids accepting unrelated bare numbers as revenue.
        currency = r"(?:US\$|USD|EUR|GBP|JPY|INR|AUD|CAD|CHF|CNY|[$€£¥₹])"
        amount = r"\d[\d,.]*"
        magnitude = r"(?:K|M|B|T|thousand|million|billion|trillion)"

        exact_or_range = re.compile(
            rf"(?i)^(?:{currency}\s*)?{amount}\s*{magnitude}"
            rf"(?:\s*[-–—]\s*(?:{currency}\s*)?{amount}\s*{magnitude})?$"
        )

        currency_amount = re.compile(
            rf"(?i)^{currency}\s*{amount}"
            rf"(?:\s*[-–—]\s*{currency}?\s*{amount})?$"
        )

        return bool(
            exact_or_range.search(value)
            or currency_amount.search(value)
        )


    @staticmethod
    def _code_link_matches_codes(
        link: str,
        sic_code: str,
        naics_code: str,
    ) -> bool:
        if not SearchWorker._is_valid_http_url(link):
            return False

        link_lower = link.lower()

        codes = [
            code
            for code in (sic_code, naics_code)
            if (
                code
                and not SearchWorker._is_unknown_value(code)
                and re.fullmatch(r"\d{4,6}", code)
            )
        ]

        if not codes:
            return False

        return any(code in link_lower for code in codes)

    @staticmethod
    def _fallback_llm_confidence(
        field: str,
        value: str,
    ) -> int:
        """
        Some local models return the requested value but omit/misformat the
        confidence object. Use conservative defaults only for structurally
        valid non-exact fields.

        Exact street Address and Zip Code never receive fallback confidence;
        they still require explicit model confidence/evidence.
        """
        if SearchWorker._is_unknown_value(value):
            return 0

        defaults = {
            "Job Level": 75,
            "Job Function": 75,
            "Company Website": 85,
            "Industry": 75,
            "City": 75,
            "State": 75,
            "Country": 80,
            "Revenue": 70,
            "SIC Code": 72,
            "NAICS Code": 72,
            "SIC/NAICS Code Link": 75,
        }

        return defaults.get(field, 0)

    @staticmethod
    def _confidence_threshold(field: str) -> int:
        thresholds = {
            "Job Level": 65,
            "Job Function": 65,
            "Company Website": 85,
            "Industry": 70,
            "Address": 85,
            "City": 75,
            "State": 75,
            "Zip Code": 95,
            "Country": 75,
            "Revenue": 70,
            "SIC Code": 70,
            "NAICS Code": 70,
            "SIC/NAICS Code Link": 75,
        }
        return thresholds.get(field, 70)

    @staticmethod
    def _normalize_llm_enrichment_value(
        field: str,
        value: str,
    ) -> str:
        raw = str(value or "").strip()

        if SearchWorker._is_unknown_value(raw):
            return "Unknown"

        if field == "Address":
            # Preserve meaningful multiline addresses while cleaning each line.
            lines = [
                re.sub(r"[ \t]+", " ", line).strip()
                for line in raw.replace("\r\n", "\n").replace("\r", "\n").split("\n")
            ]
            lines = [line for line in lines if line]
            value = "\n".join(lines)

            if not value or len(value) > 500:
                return "Unknown"

            return value

        value = re.sub(r"\s+", " ", raw).strip()

        if field == "Job Level":
            return SearchWorker._canonical_job_level(value)

        if field == "Job Function":
            return SearchWorker._canonical_job_function(value)

        if field == "Company Website":
            return (
                value
                if SearchWorker._is_valid_official_website(value)
                else "Unknown"
            )

        if field == "Country":
            return SearchWorker._canonical_country(value)

        if field == "Zip Code":
            return (
                value
                if SearchWorker._is_plausible_zip(value)
                else "Unknown"
            )

        if field == "Revenue":
            return (
                value
                if SearchWorker._is_plausible_revenue(value)
                else "Unknown"
            )

        if field == "SIC Code":
            digits = re.sub(r"\D", "", value)
            return digits if re.fullmatch(r"\d{4}", digits) else "Unknown"

        if field == "NAICS Code":
            digits = re.sub(r"\D", "", value)
            return digits if re.fullmatch(r"\d{6}", digits) else "Unknown"

        if field == "SIC/NAICS Code Link":
            return (
                value
                if SearchWorker._is_valid_http_url(value)
                else "Unknown"
            )

        if len(value) > 180:
            return "Unknown"

        return value

    @staticmethod
    def _append_enrichment_note(
        record: dict[str, str],
        message: str,
    ) -> None:
        existing = clean_cell_value(record.get("Notes", "-"))

        if existing == "-":
            record["Notes"] = message
        elif message not in existing:
            record["Notes"] = f"{existing} | {message}"

    def _collect_official_website_evidence(
        self,
        record: dict[str, str],
    ) -> str:
        """
        Collect factual address/location evidence from the official company site.

        Only runs when Address/Zip enrichment is needed. It checks:
        1. the homepage
        2. up to two same-domain Contact/About/Locations pages discovered on
           the homepage.

        Results are cached per website for the remainder of the run.
        """
        website = clean_cell_value(
            record.get("Company Website", "-")
        )

        if not self._is_valid_official_website(
            website
        ):
            return "-"

        try:
            website_key = (
                (urlsplit(website).hostname or "")
                .lower()
                .removeprefix("www.")
            )
        except Exception:
            website_key = website.lower()

        if website_key in self.website_evidence_cache:
            return self.website_evidence_cache[
                website_key
            ]

        session = requests.Session()
        session.headers.update(
            {
                "User-Agent": (
                    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
                    "AppleWebKit/537.36 (KHTML, like Gecko) "
                    "Chrome/120 Safari/537.36"
                )
            }
        )

        def fetch_html(url: str) -> tuple[str, str]:
            response = session.get(
                url,
                timeout=(3, 7),
                allow_redirects=True,
            )
            response.raise_for_status()

            content_type = (
                response.headers
                .get("content-type", "")
                .lower()
            )

            if (
                content_type
                and "html" not in content_type
            ):
                return "", response.url

            return (
                response.text[:650_000],
                response.url,
            )

        def html_to_text(html: str) -> str:
            if not html:
                return ""

            html = re.sub(
                r"(?is)<script.*?>.*?</script>",
                " ",
                html,
            )
            html = re.sub(
                r"(?is)<style.*?>.*?</style>",
                " ",
                html,
            )
            html = re.sub(
                r"(?is)<noscript.*?>.*?</noscript>",
                " ",
                html,
            )
            html = re.sub(
                r"(?is)<!--.*?-->",
                " ",
                html,
            )
            plain = re.sub(
                r"(?s)<[^>]+>",
                " ",
                html,
            )

            replacements = {
                "&nbsp;": " ",
                "&amp;": "&",
                "&quot;": '"',
                "&#39;": "'",
                "&apos;": "'",
            }

            for old_value, new_value in replacements.items():
                plain = plain.replace(
                    old_value,
                    new_value,
                )

            return re.sub(
                r"\s+",
                " ",
                plain,
            ).strip()

        def address_windows(
            plain: str,
        ) -> list[str]:
            if not plain:
                return []

            keywords = (
                "address",
                "headquarters",
                "head office",
                "registered office",
                "corporate office",
                "contact us",
                "contact",
                "location",
                "locations",
                "postal",
                "postcode",
                "zip code",
                "street",
                "road",
                "avenue",
                "boulevard",
                "suite",
                "floor",
            )

            lower = plain.lower()
            windows: list[str] = []

            for keyword in keywords:
                start_pos = 0

                while len(windows) < 16:
                    pos = lower.find(
                        keyword,
                        start_pos,
                    )

                    if pos < 0:
                        break

                    left = max(
                        0,
                        pos - 300,
                    )
                    right = min(
                        len(plain),
                        pos + 650,
                    )

                    snippet = plain[
                        left:right
                    ].strip()

                    if (
                        snippet
                        and snippet not in windows
                    ):
                        windows.append(snippet)

                    start_pos = (
                        pos + len(keyword)
                    )

            return windows

        pages: list[
            tuple[str, str, str]
        ] = []

        try:
            homepage_html, homepage_url = fetch_html(
                website
            )

            if homepage_html:
                pages.append(
                    (
                        "HOME",
                        homepage_url,
                        homepage_html,
                    )
                )

            candidate_links: list[
                tuple[int, str]
            ] = []

            if homepage_html:
                anchor_pattern = re.compile(
                    r"(?is)<a\b[^>]*href=[\"']([^\"'#]+)[\"'][^>]*>(.*?)</a>"
                )

                for href, label in anchor_pattern.findall(
                    homepage_html
                ):
                    label_text = re.sub(
                        r"(?s)<[^>]+>",
                        " ",
                        label,
                    )
                    label_text = re.sub(
                        r"\s+",
                        " ",
                        label_text,
                    ).strip()

                    combined = (
                        href + " " + label_text
                    ).lower()

                    score = 0

                    if "contact" in combined:
                        score += 100
                    if "location" in combined:
                        score += 90
                    if "office" in combined:
                        score += 75
                    if "about" in combined:
                        score += 50

                    if score <= 0:
                        continue

                    absolute = urljoin(
                        homepage_url or website,
                        href,
                    )

                    try:
                        host = (
                            urlsplit(absolute)
                            .hostname
                            or ""
                        ).lower().removeprefix(
                            "www."
                        )
                    except Exception:
                        continue

                    if host != website_key:
                        continue

                    candidate_links.append(
                        (
                            score,
                            absolute,
                        )
                    )

            seen_urls = {
                item[1]
                for item in pages
            }

            for _score, url in sorted(
                candidate_links,
                reverse=True,
            ):
                if len(pages) >= 3:
                    break

                if url in seen_urls:
                    continue

                seen_urls.add(url)

                try:
                    page_html, final_url = (
                        fetch_html(url)
                    )

                    if page_html:
                        pages.append(
                            (
                                "DETAIL",
                                final_url,
                                page_html,
                            )
                        )
                except Exception:
                    continue

        except Exception as exc:
            self.log(
                "WEBSITE EVIDENCE WARNING | "
                f"{type(exc).__name__}: {exc}"
            )

        evidence_chunks: list[str] = []

        for page_type, page_url, html in pages:
            plain = html_to_text(html)
            windows = address_windows(plain)

            if not windows and plain:
                windows = [
                    plain[-2200:]
                ]

            if windows:
                evidence_chunks.append(
                    f"[{page_type}] {page_url}\n"
                    + "\n---\n".join(
                        windows[:8]
                    )
                )

        evidence = "\n=====\n".join(
            evidence_chunks
        )[:9000]

        if not evidence:
            evidence = "-"

        self.website_evidence_cache[
            website_key
        ] = evidence

        if evidence != "-":
            self.log(
                "WEBSITE EVIDENCE | "
                f"company='{record.get('Company Name', '-')}' | "
                f"chars={len(evidence)} | "
                f"pages={len(pages)}"
            )

        return evidence

    @staticmethod
    def _google_maps_cache_key(record: dict[str, str]) -> str:
        website = clean_cell_value(record.get('Company Website', '-'))
        host = ''
        if SearchWorker._is_valid_official_website(website):
            host = (urlsplit(website).hostname or '').lower().removeprefix('www.')
        # Fixed slots and the existing office prevent cross-city/branch collisions.
        values = [SearchWorker._normalize_company_for_match(record.get('Company Name', '-'))]
        values.extend(SearchWorker._maps_geo_text(record.get(field, '-'))
                      for field in ('City', 'State', 'Country'))
        values.extend([host, '|'.join(SearchWorker._maps_address_parts(record.get('Address', '-')))])
        return json.dumps(values, ensure_ascii=False, separators=(',', ':'))

    @staticmethod
    def _extract_postal_code_from_address(
        full_address: str,
    ) -> str:
        """
        Extract common postal-code formats from a Google Maps address.

        This is deliberately conservative: it only returns text visibly present
        in the Maps address; it does not infer a postal code from the city.
        """
        value = re.sub(
            r"\s+",
            " ",
            full_address or "",
        ).strip()

        if not value:
            return "-"

        patterns = [
            # US ZIP / ZIP+4
            r"(?<!\d)\d{5}(?:-\d{4})?(?!\d)",
            # Singapore / India / many 6-digit systems
            r"(?<!\d)\d{6}(?!\d)",
            # Canada
            r"\b[A-Z]\d[A-Z][ -]?\d[A-Z]\d\b",
            # UK
            r"\b[A-Z]{1,2}\d[A-Z\d]?\s*\d[A-Z]{2}\b",
            # Netherlands
            r"\b\d{4}\s?[A-Z]{2}\b",
            # Australia / NZ and several 4-digit systems
            r"(?<!\d)\d{4}(?!\d)",
        ]

        for pattern in patterns:
            matches = re.findall(
                pattern,
                value,
                flags=re.IGNORECASE,
            )

            if matches:
                return str(matches[-1]).strip()

        return "-"

    @staticmethod
    def _street_address_from_maps_address(
        full_address: str,
        record: dict[str, str],
    ) -> str:
        """
        Convert a full Google Maps address to a street-level Address value.

        Example:
            160 Robinson Rd, #14-04, Singapore 068914
        becomes:
            160 Robinson Rd, #14-04
        """
        full_address = re.sub(
            r"\s+",
            " ",
            full_address or "",
        ).strip(" ,")

        if not full_address:
            return "-"

        parts = [
            part.strip()
            for part in full_address.split(",")
            if part.strip()
        ]

        if not parts:
            return "-"

        known_location_tokens = {
            clean_cell_value(
                record.get("City", "-")
            ).lower().strip(" ."),
            clean_cell_value(
                record.get("State", "-")
            ).lower().strip(" ."),
            clean_cell_value(
                record.get("Country", "-")
            ).lower().strip(" ."),
        }
        known_location_tokens.discard("-")
        known_location_tokens.discard("")

        postal = (
            SearchWorker
            ._extract_postal_code_from_address(
                full_address
            )
        )

        keep: list[str] = []

        for part in parts:
            low = part.lower().strip(" .")
            remove = False

            if any(
                token
                and (
                    low == token
                    or low.startswith(
                        token + " "
                    )
                    or low.endswith(
                        " " + token
                    )
                )
                for token in known_location_tokens
            ):
                remove = True

            if (
                postal != "-"
                and postal.lower() in low
            ):
                # A trailing locality/state/postal component is not part of
                # the street address. Keep it only when it clearly also
                # contains a street/unit marker.
                street_markers = (
                    " st",
                    " street",
                    " rd",
                    " road",
                    " ave",
                    " avenue",
                    " blvd",
                    " boulevard",
                    " dr",
                    " drive",
                    " lane",
                    " ln",
                    " hwy",
                    " highway",
                    " #",
                    " suite",
                    " floor",
                    " fl ",
                    " unit",
                )

                if not any(
                    marker in " " + low
                    for marker in street_markers
                ):
                    remove = True

            if not remove:
                keep.append(part)

        if keep:
            return ", ".join(keep)

        # If stripping was too aggressive, never return nothing when Maps gave
        # a factual address.
        return full_address

    @staticmethod
    def _maps_geo_text(value: str) -> str:
        """Normalize spelling, without guessing a different city or region."""
        if SearchWorker._is_unknown_value(value):
            return ""
        value = unicodedata.normalize('NFKD', str(value)).casefold()
        value = ''.join(ch for ch in value if not unicodedata.combining(ch))
        value = re.sub(r'[^\w]+', ' ', value).strip()
        value = re.sub(r'^st\s+', 'saint ', value)
        aliases = {'bangalore': 'bengaluru', 'bombay': 'mumbai', 'madras': 'chennai',
                   'calcutta': 'kolkata', 'gurgaon': 'gurugram', 'new york city': 'new york'}
        value = re.sub(r'\s+(?:district|metropolitan area|metro area)$', '', value)
        return aliases.get(value, value)

    @staticmethod
    def _maps_country_name(value: str) -> str:
        """Recognize explicit country labels; an omitted label stays unknown."""
        value = SearchWorker._maps_geo_text(value)
        groups = {
            'united states': ('united states', 'united states of america', 'usa', 'us', 'u s a'),
            'united kingdom': ('united kingdom', 'uk', 'u k', 'great britain', 'england', 'scotland', 'wales', 'northern ireland'),
            'united arab emirates': ('united arab emirates', 'uae', 'u a e'),
            'switzerland': ('switzerland', 'ch', 'schweiz', 'suisse'),
            'germany': ('germany', 'de', 'deutschland'),
            'netherlands': ('netherlands', 'the netherlands', 'nl'),
            'south korea': ('south korea', 'republic of korea', 'korea'),
            'china': ('china', 'people s republic of china'),
        }
        for canonical, aliases in groups.items():
            if value in aliases:
                return canonical
        # These labels can be identified without country-code ambiguity.
        countries = {
            'india', 'canada', 'australia', 'new zealand', 'singapore', 'france', 'ireland',
            'spain', 'portugal', 'italy', 'austria', 'belgium', 'luxembourg', 'denmark',
            'norway', 'sweden', 'finland', 'iceland', 'poland', 'czechia', 'czech republic',
            'hungary', 'romania', 'bulgaria', 'greece', 'turkey', 'israel', 'saudi arabia',
            'qatar', 'bahrain', 'kuwait', 'oman', 'japan', 'taiwan', 'hong kong', 'macao',
            'malaysia', 'indonesia', 'thailand', 'vietnam', 'philippines', 'pakistan',
            'bangladesh', 'sri lanka', 'nepal', 'brazil', 'mexico', 'argentina', 'chile',
            'colombia', 'peru', 'ecuador', 'panama', 'costa rica', 'uruguay', 'paraguay',
            'venezuela', 'south africa', 'egypt', 'kenya', 'nigeria', 'morocco', 'ghana',
            'russia', 'ukraine', 'belarus', 'estonia', 'latvia', 'lithuania', 'croatia',
            'slovenia', 'slovakia', 'serbia', 'cyprus', 'malta', 'mauritius', 'bermuda',
        }
        return value if value in countries else ''

    @staticmethod
    def _maps_state_aliases() -> dict[str, str]:
        # Full US state names are frequently replaced by two-letter Maps labels.
        pairs = ('AL:Alabama|AK:Alaska|AZ:Arizona|AR:Arkansas|CA:California|CO:Colorado|'
                 'CT:Connecticut|DE:Delaware|DC:District of Columbia|FL:Florida|GA:Georgia|'
                 'HI:Hawaii|ID:Idaho|IL:Illinois|IN:Indiana|IA:Iowa|KS:Kansas|KY:Kentucky|'
                 'LA:Louisiana|ME:Maine|MD:Maryland|MA:Massachusetts|MI:Michigan|MN:Minnesota|'
                 'MS:Mississippi|MO:Missouri|MT:Montana|NE:Nebraska|NV:Nevada|NH:New Hampshire|'
                 'NJ:New Jersey|NM:New Mexico|NY:New York|NC:North Carolina|ND:North Dakota|'
                 'OH:Ohio|OK:Oklahoma|OR:Oregon|PA:Pennsylvania|RI:Rhode Island|SC:South Carolina|'
                 'SD:South Dakota|TN:Tennessee|TX:Texas|UT:Utah|VT:Vermont|VA:Virginia|'
                 'WA:Washington|WV:West Virginia|WI:Wisconsin|WY:Wyoming|'
                 'ON:Ontario|QC:Quebec|BC:British Columbia|AB:Alberta|MB:Manitoba|'
                 'NB:New Brunswick|NS:Nova Scotia|NL:Newfoundland and Labrador|'
                 'PE:Prince Edward Island|SK:Saskatchewan|NT:Northwest Territories|NU:Nunavut|YT:Yukon|'
                 'KA:Karnataka|MH:Maharashtra|TN:Tamil Nadu|TS:Telangana|AP:Andhra Pradesh|'
                 'KL:Kerala|GJ:Gujarat|RJ:Rajasthan|UP:Uttar Pradesh|MP:Madhya Pradesh|'
                 'WB:West Bengal|BR:Bihar|HR:Haryana|PB:Punjab|OD:Odisha|AS:Assam|'
                 'DL:Delhi|JH:Jharkhand|CG:Chhattisgarh|UK:Uttarakhand|HP:Himachal Pradesh|'
                 'NSW:New South Wales|VIC:Victoria|QLD:Queensland|TAS:Tasmania')
        return {name.casefold(): code.casefold() for code, name in
                (pair.split(':') for pair in pairs.split('|'))}

    @staticmethod
    def _maps_address_parts(value: str) -> list[str]:
        """Canonical components for same-office comparisons, preserving numbers."""
        if SearchWorker._is_unknown_value(value):
            return []
        states = SearchWorker._maps_state_aliases()
        replacements = {'st': 'street', 'rd': 'road', 'ave': 'avenue', 'av': 'avenue',
                        'blvd': 'boulevard', 'dr': 'drive', 'ln': 'lane', 'hwy': 'highway',
                        'pkwy': 'parkway', 'ste': 'suite', 'fl': 'floor',
                        'e': 'east', 'w': 'west', 'n': 'north', 's': 'south',
                        'ne': 'northeast', 'nw': 'northwest', 'se': 'southeast', 'sw': 'southwest'}
        parts = []
        for raw in re.split(r'[,\n]+', str(value)):
            if SearchWorker._maps_geo_text(raw) in {
                'corporate office', 'corporate headquarters', 'headquarters',
                'head office', 'primary', 'primary location', 'registered office',
                'postbox', 'po box', 'p o box',
            }:
                # A LinkedIn section label or unnumbered mailbox label is not
                # a street. Numbered mailboxes and suite numbers remain intact.
                continue
            raw = re.sub(r'\b(\d{5})-\d{4}\b', r'\1', raw)
            country = SearchWorker._maps_country_name(raw)
            part = country or SearchWorker._maps_geo_text(raw)
            if not part:
                continue
            for state, code in sorted(states.items(), key=lambda item: -len(item[0])):
                part = re.sub(r'(?<!\w)' + re.escape(state) + r'(?!\w)', code, part)
            # A standalone state code must not become a compass direction.
            if not country and not re.fullmatch(r'[a-z]{2}(?:\s+\d{4,6})?', part):
                part = ' '.join(replacements.get(token, token) for token in part.split())
            parts.append(part)
        return parts

    @staticmethod
    def _maps_address_compatible(existing: str, candidate: str) -> bool:
        """Every known old component must occur in the new office address."""
        old = SearchWorker._maps_address_parts(existing)
        new = SearchWorker._maps_address_parts(candidate)
        if not old or not new:
            return False
        if old == new:
            return True
        def contains(part, known):
            return re.search(r'(?<!\w)' + re.escape(known) + r'(?!\w)', part) is not None
        street = re.compile(r'\b(?:street|road|avenue|boulevard|drive|lane|highway|parkway|way|court|circle|suite|floor)\b')
        unit = re.compile(r'^(?:suite|unit|floor|apartment|apt)\s+\w+$')
        for index, part in enumerate(old):
            if (unit.fullmatch(part) and not any(unit.fullmatch(item) for item in new)
                    and any(re.match(r'^\d+\s+\w', item) for item in old)):
                # Maps often omits the suite. A different explicit unit still
                # fails, and every known street/locality component must match.
                continue
            if SearchWorker._maps_country_name(part):
                actual_country = SearchWorker._maps_country_name(new[-1])
                if not actual_country:
                    continue
            is_street = bool(street.search(part) or re.match(r'^\d', part))
            # A city must not match a street named after another city.
            choices = new if is_street or len(new) == 1 else new[1:]
            if not any(contains(other, part) for other in choices):
                # A short city/state entry can omit the separating comma.
                if is_street or not contains(' '.join(choices), part):
                    return False
        return True

    @staticmethod
    def _maps_addresses_equivalent(first: str, second: str) -> bool:
        """Same complete office address, allowing common display abbreviations."""
        units = lambda value: {part for part in SearchWorker._maps_address_parts(value)
                               if re.match(r'^(?:suite|unit|floor|apartment|apt)\s+', part)}
        if units(first) != units(second):
            return False
        return (SearchWorker._maps_address_compatible(first, second)
                and SearchWorker._maps_address_compatible(second, first))

    @staticmethod
    def _maps_address_matches_location(record: dict[str, str], full_address: str) -> bool:
        """Validate the office location without requiring an omitted country label."""
        if SearchWorker._is_unknown_value(full_address):
            return False
        existing = clean_cell_value(record.get('Address', '-'))
        if not SearchWorker._is_unknown_value(existing):
            if not SearchWorker._maps_address_compatible(existing, full_address):
                return False
            # LinkedIn's company address can differ from the prospect's city.
            # A street-only old value still needs the record's locality check.
            old_parts = SearchWorker._maps_address_parts(existing)
            if len(old_parts) > 1 or not re.search(r'\d|\b(?:street|road|avenue|drive|lane)\b', old_parts[0]):
                return SearchWorker._maps_address_matches_country(record, full_address)
        parts = [SearchWorker._maps_geo_text(part) for part in re.split(r'[,\n]+', full_address)]
        parts = [part for part in parts if part]
        # Exclude a street component so "Paris Rd, Lyon" cannot validate Paris.
        local_parts = parts[1:] if len(parts) > 1 and re.search(r'\d|\b(?:st|street|rd|road|ave|avenue|drive|lane|center|centre)\b', parts[0]) else parts
        locality = ' '.join(local_parts)
        city = SearchWorker._maps_geo_text(record.get('City', '-'))
        state = SearchWorker._maps_geo_text(record.get('State', '-'))
        aliases = SearchWorker._maps_state_aliases()
        def present(term):
            return bool(term and re.search(r'(?<!\w)' + re.escape(term) + r'(?!\w)', locality))
        if city and not present(city):
            return False
        if state:
            state_code = aliases.get(state, state)
            state_names = [state, state_code] + [name for name, code in aliases.items() if code == state_code]
            state_found = any(present(name) for name in state_names)
            # Detect explicit conflicting state labels, but not a city named
            # New York or a street containing a two-letter word.
            explicit_codes = set()
            for part in local_parts:
                for name, code in aliases.items():
                    if part == name or re.fullmatch(re.escape(name) + r'\s+\d[\w -]*', part):
                        explicit_codes.add(code)
                    if re.fullmatch(re.escape(code) + r'(?:\s+\d[\w -]*)?', part):
                        explicit_codes.add(code)
            if state_code in aliases.values() and explicit_codes and state_code not in explicit_codes:
                return False
            if not city and not state_found:
                return False
        return SearchWorker._maps_address_matches_country(record, full_address)

    @staticmethod
    def _maps_address_matches_country(record: dict[str, str], full_address: str) -> bool:
        # Prefer the existing company's office country over a prospect's country.
        old_parts = re.split(r'[,\n]+', clean_cell_value(record.get('Address', '-')))
        office_country = SearchWorker._maps_country_name(old_parts[-1]) if old_parts else ''
        expected_raw = office_country or record.get('Country', '-')
        if SearchWorker._is_unknown_value(expected_raw):
            return True
        expected = SearchWorker._maps_country_name(expected_raw) or SearchWorker._maps_geo_text(expected_raw)
        parts = [part.strip() for part in re.split(r'[,\n]+', full_address) if part.strip()]
        if not parts:
            return False
        tail = SearchWorker._maps_geo_text(parts[-1])
        words = tail.split()
        actual = SearchWorker._maps_country_name(parts[-1])
        if not actual:
            for size in (4, 3, 2, 1):
                actual = SearchWorker._maps_country_name(' '.join(words[-size:]))
                if actual:
                    break
        if actual:
            return actual == expected
        # Country text can share the postal-code component (Singapore 068914).
        if re.search(r'(?<!\w)' + re.escape(expected) + r'(?!\w)', tail):
            return True
        # A country is often omitted on domestic Maps results. Location/office
        # validation is performed separately; omission is not a contradiction.
        return (not SearchWorker._is_unknown_value(record.get('Address', '-'))
                or any(not SearchWorker._is_unknown_value(record.get(field, '-'))
                       for field in ('City', 'State')))

    @staticmethod
    def _maps_listing_is_subunit(expected: str, actual: str) -> bool:
        """A shared institutional domain does not identify its general contact."""
        expected = SearchWorker._maps_geo_text(expected)
        actual = SearchWorker._maps_geo_text(actual)
        markers = (r'\b(?:department|dept)\b', r'\bparking (?:lot|garage)\b',
                   r'\bbookstore\b', r'\blibrary\b', r'\bcafeteria\b',
                   r'\bdining hall\b', r'\bstudent union\b', r'\bfire station\b')
        return any(re.search(marker, actual) and not re.search(marker, expected) for marker in markers)

    @staticmethod
    def _maps_result_matches_company(
        expected_company: str,
        actual_name: str,
        official_website: str,
        maps_website: str,
    ) -> bool:
        if SearchWorker._maps_listing_is_subunit(expected_company, actual_name):
            return False
        name_score = SearchWorker._maps_company_match_score(expected_company, actual_name)
        same_domain = False
        if is_public_source_url(official_website) and is_public_source_url(maps_website):
            expected_host, actual_host = source_host(official_website), source_host(maps_website)
            same_domain = (expected_host == actual_host
                           or actual_host.endswith('.' + expected_host)
                           or expected_host.endswith('.' + actual_host))
            # An explicit conflicting website outweighs a similar display name.
            if not same_domain:
                return False
        if name_score > 0:
            return True
        # A parent domain cannot validate an explicitly different named unit.
        # Generic office suffixes were already accepted by the name scorer.
        expected = SearchWorker._normalize_company_for_match(expected_company)
        actual = SearchWorker._normalize_company_for_match(actual_name)
        if expected and actual != expected and re.search(r'(?<!\w)' + re.escape(expected) + r'(?!\w)', actual):
            return False
        return same_domain

    def _dismiss_google_consent(
        self,
        page,
    ) -> None:
        """
        Best-effort handling for Google's consent screen.

        This does not bypass CAPTCHA or access controls.
        """
        candidates = (
            "Accept all",
            "Reject all",
            "I agree",
        )

        for label in candidates:
            try:
                button = page.get_by_role(
                    "button",
                    name=label,
                    exact=True,
                )

                if (
                    button.count()
                    and button.first.is_visible()
                ):
                    button.first.click(
                        timeout=2_000
                    )
                    page.wait_for_timeout(500)
                    return
            except Exception:
                continue

    def _nudge_google_maps_details_after_click(self, page, source_label: str, query: str):
        try:
            movement = page.evaluate(PLACE_SCRIPT, {"scroll": GOOGLE_MAPS_POST_CLICK_SCROLL_PX})
            self.trace("GOOGLE_MAPS", "post_click_small_scroll", source=source_label,
                       query=query, movement=movement)
            page.wait_for_timeout(GOOGLE_MAPS_POST_CLICK_WAIT_MS)
            return movement
        except Exception as exc:
            self.trace("GOOGLE_MAPS", "post_click_scroll_error", error=str(exc))
            return {}

    def _wait_and_scroll_google_maps_place(self, page, source_label: str, query: str,
                                          need_address: bool, need_phone: bool) -> dict[str, Any]:
        """Wait for lazy details and scroll the selected place, excluding result cards."""
        page.wait_for_timeout(GOOGLE_MAPS_DETAIL_INITIAL_WAIT_MS)
        scans = []
        stagnant = 0
        state = {}
        for index in range(GOOGLE_MAPS_DETAIL_MAX_SCROLLS + 1):
            if self.stop_event.is_set():
                break
            state = page.evaluate(PLACE_SCRIPT, {})
            scans.append({"round": index, **state})
            if ((not need_address or state.get("addressFound")) and
                    (not need_phone or state.get("phoneFound"))):
                break
            movement = page.evaluate(PLACE_SCRIPT, {"scroll": GOOGLE_MAPS_DETAIL_SCROLL_STEP})
            stagnant = stagnant + 1 if movement.get("after", 0) <= movement.get("before", 0) else 0
            # Always give asynchronously inserted fields a final observation.
            page.wait_for_timeout(GOOGLE_MAPS_DETAIL_SCROLL_WAIT_MS)
            if stagnant >= GOOGLE_MAPS_DETAIL_STAGNANT_LIMIT:
                state = page.evaluate(PLACE_SCRIPT, {})
                break
        self.trace("GOOGLE_MAPS", "details_scroll_complete", source=source_label,
                   query=query, need_address=need_address, need_phone=need_phone,
                   final_state=state, scans=scans)
        return {"address_found": bool(state.get("addressFound")),
                "phone_found": bool(state.get("phoneFound")), "scans": scans}

    def _extract_google_maps_place(self, page) -> dict[str, str]:
        """Use the same scoped selectors for readiness checks and final extraction."""
        data = page.evaluate(PLACE_SCRIPT, {})
        return {field: clean_cell_value(data.get(key) or "-") for field, key in (
            ("Place Name", "title"), ("Full Address", "address"),
            ("Phone Number", "phone"), ("Company Website", "website"),
            ("Address Link", "url"),
        )}

    @staticmethod
    def _build_google_maps_query(
        company_name: str,
        city: str,
    ) -> str:
        """
        Build the default Maps query exactly as:
            Company Name+City

        Spaces are normalized to "+" so:
            Anyscale + San Francisco
        becomes:
            Anyscale+San+Francisco

        When City is missing/invalid, search only by Company Name.
        """
        company = SearchWorker._normalize_ws(
            clean_cell_value(company_name)
        )
        city_value = SearchWorker._normalize_ws(
            clean_cell_value(city)
        )

        if SearchWorker._is_unknown_value(company):
            return ""

        if SearchWorker._is_unknown_value(city_value):
            raw = company
        else:
            raw = f"{company}+{city_value}"

        raw = raw.strip()
        raw = re.sub(r"\s+", "+", raw)
        raw = re.sub(r"\++", "+", raw)
        return raw.strip("+")

    @staticmethod
    def _normalize_maps_phone(
        value: str,
    ) -> str:
        """
        Normalize a Google Maps phone number.

        Keeps an existing leading "+" but never invents one.
        Removes spaces, dashes, parentheses and other punctuation.
        """
        raw = re.sub(
            r"(?i)^\s*(?:phone\s*:?\s*|tel\s*:?\s*)",
            "",
            str(value or ""),
        ).strip()

        if SearchWorker._is_unknown_value(raw):
            return "-"

        raw = re.split(r'(?:;ext=|\bext\.?\s*|\bx\s*)(?=\d)', raw, maxsplit=1,
                       flags=re.I)[0].strip()
        has_plus = raw.lstrip().startswith("+")
        digits = re.sub(r"\D", "", raw)

        if not 7 <= len(digits) <= 15:
            return "-"

        return ("+" if has_plus else "") + digits

    @staticmethod
    def _maps_company_match_score(expected: str, actual: str) -> int:
        """Rank exact/fuzzy names and narrowly defined office descriptors."""
        expected_raw = SearchWorker._normalize_ws(expected)
        actual_raw = SearchWorker._normalize_ws(actual)
        if SearchWorker._maps_listing_is_subunit(expected_raw, actual_raw):
            return 0
        if SearchWorker._is_unknown_value(expected_raw) or SearchWorker._is_unknown_value(actual_raw):
            return 0
        if expected_raw == actual_raw:
            return 500
        if expected_raw.casefold() == actual_raw.casefold():
            return 450
        expected_normalized = SearchWorker._normalize_company_for_match(expected_raw)
        actual_normalized = SearchWorker._normalize_company_for_match(actual_raw)
        if expected_normalized and expected_normalized == actual_normalized:
            return 400
        fuzzy, similarity = SearchWorker._fuzzy_company_match(expected_raw, actual_raw)
        if fuzzy:
            return 300 + int(similarity / 2)
        # Maps adds office labels that LinkedIn company identity matching must
        # not accept globally. Only these generic suffixes can be removed.
        descriptor = (r'(?:corporate (?:headquarters|office)|headquarters|head office|registered office|'
                      r'corporate campus|main office|hq|(?:local|branch|regional|sales|district|satellite) office|'
                      r'regional (?:headquarters|head office|hq))')
        trimmed = re.sub(r'(?i)[\s,()–—-]+' + descriptor + r'\s*[)]?$', '', actual_raw).strip()
        if trimmed != actual_raw and SearchWorker._link_company_names_match(expected_raw, trimmed):
            return 275
        return 0

    def _get_google_maps_page(self):
        """
        Reuse a single Google Maps tab for the whole run.

        The main LinkedIn page and existing persistent browser context remain
        unchanged.
        """
        if not self.context:
            return None

        try:
            if (
                self.google_maps_page
                and not self.google_maps_page.is_closed()
            ):
                self.google_maps_page.set_default_timeout(
                    GOOGLE_MAPS_PAGE_TIMEOUT_MS
                )
                return self.google_maps_page
        except Exception:
            self.google_maps_page = None

        self.google_maps_page = self.context.new_page()
        self.google_maps_page.set_default_timeout(
            GOOGLE_MAPS_PAGE_TIMEOUT_MS
        )
        return self.google_maps_page

    def _find_google_maps_company_result(
        self,
        page,
        company_name: str,
        source_label: str,
        query: str,
        excluded_urls: set[str] | None = None,
        allow_website_verification: bool = False,
    ) -> dict[str, Any]:
        """
        Inspect visible Google Maps search results and choose the best matching
        company instead of blindly clicking the first result.

        Returns:
            {
                "found": bool,
                "name": str,
                "score": int,
                "index": int,
            }
        """
        best: dict[str, Any] | None = None
        excluded_urls = excluded_urls or set()

        for scan_round in range(
            GOOGLE_MAPS_RESULT_SCAN_ROUNDS
        ):
            # Prefer the visible results feed. Fall back to all Maps place links
            # for layouts where Google does not expose role="feed".
            result_links = page.locator(
                'div[role="feed"] a[href*="/maps/place/"]'
            )

            try:
                count = result_links.count()
            except Exception:
                count = 0

            if count == 0:
                result_links = page.locator(
                    'a[href*="/maps/place/"]'
                )
                try:
                    count = result_links.count()
                except Exception:
                    count = 0

            candidates: list[dict[str, Any]] = []

            for result_index in range(
                min(
                    count,
                    GOOGLE_MAPS_RESULT_SCAN_LIMIT,
                )
            ):
                try:
                    candidate = result_links.nth(
                        result_index
                    )

                    if not candidate.is_visible():
                        continue

                    href = candidate.get_attribute("href") or ""
                    if href in excluded_urls:
                        continue

                    label = self._normalize_ws(
                        candidate.get_attribute(
                            "aria-label"
                        )
                        or candidate.get_attribute(
                            "title"
                        )
                        or ""
                    )

                    if not label:
                        try:
                            raw_text = str(
                                candidate.inner_text()
                                or ""
                            ).strip()
                        except Exception:
                            raw_text = ""

                        if raw_text:
                            # Keep line boundaries until the company-name line
                            # is isolated; lower lines may contain category,
                            # rating, phone, or address text.
                            first_line = next(
                                (
                                    line.strip()
                                    for line in raw_text.splitlines()
                                    if line.strip()
                                ),
                                "",
                            )
                            label = self._normalize_ws(
                                first_line
                            )

                    if not label:
                        continue

                    score = self._maps_company_match_score(
                        company_name,
                        label,
                    )
                    # A domain search may display a trading name. Open it for
                    # verification; only the final name/domain + office checks
                    # can accept it. The query itself is never identity proof.
                    if score == 0 and allow_website_verification:
                        score = 1

                    item = {
                        "index": result_index,
                        "name": label,
                        "score": score,
                        "url": href,
                    }
                    candidates.append(item)

                    if (
                        score > 0
                        and (
                            best is None
                            or score > best["score"]
                        )
                    ):
                        best = {
                            **item,
                            "locator": candidate,
                        }

                except Exception as exc:
                    self.trace(
                        "GOOGLE_MAPS",
                        "result_candidate_error",
                        source=source_label,
                        query=query,
                        candidate_index=result_index,
                        error_type=type(exc).__name__,
                        error=str(exc),
                    )

            self.trace(
                "GOOGLE_MAPS",
                "result_scan",
                source=source_label,
                query=query,
                scan_round=scan_round,
                expected_company=company_name,
                candidates=candidates,
                best=(
                    {
                        key: value
                        for key, value in best.items()
                        if key != "locator"
                    }
                    if best
                    else None
                ),
            )

            # Exact/normalized match: no need to load more result cards.
            if best and best["score"] >= 400:
                break

            # For a strong partial match, scan one extra round if possible so an
            # exact company lower in the visible list can still win.
            if (
                best
                and best["score"] >= 250
                and scan_round
                >= GOOGLE_MAPS_RESULT_SCAN_ROUNDS - 2
            ):
                break

            movement = page.evaluate(
                f"""
                () => {{
                    const feed =
                        document.querySelector(
                            'div[role="feed"]'
                        );

                    if (!feed) {{
                        return {{
                            moved: false,
                            before: 0,
                            after: 0,
                        }};
                    }}

                    const before = feed.scrollTop;

                    feed.scrollTop = Math.min(
                        feed.scrollTop
                            + {GOOGLE_MAPS_RESULT_SCROLL_PX},
                        feed.scrollHeight
                            - feed.clientHeight
                    );

                    feed.dispatchEvent(
                        new Event(
                            'scroll',
                            {{ bubbles: true }}
                        )
                    );

                    return {{
                        moved:
                            feed.scrollTop > before,
                        before:
                            Math.round(before),
                        after:
                            Math.round(feed.scrollTop),
                    }};
                }}
                """
            )

            self.trace(
                "GOOGLE_MAPS",
                "result_scan_scroll",
                source=source_label,
                query=query,
                scan_round=scan_round,
                movement=movement,
            )

            if not (movement or {}).get("moved", False) and best:
                break

            page.wait_for_timeout(
                GOOGLE_MAPS_DETAIL_SCROLL_WAIT_MS
            )

        if not best:
            return {
                "found": False,
                "name": "-",
                "score": 0,
                "index": -1,
            }

        # Result feeds can reorder while scrolling. Re-find by URL, not nth().
        locator = page.locator('a[href=' + json.dumps(best['url']) + ']:visible').first

        try:
            locator.scroll_into_view_if_needed(
                timeout=3_000
            )
        except Exception:
            pass

        locator.click(
            timeout=5_000
        )

        self.log(
            "[MAPS] Result opened: "
            f"{best['name']}"
        )

        self.trace(
            "GOOGLE_MAPS",
            "matched_result_clicked",
            source=source_label,
            query=query,
            expected_company=company_name,
            result_name=best["name"],
            match_score=best["score"],
            result_index=best["index"],
        )

        return {
            "found": True,
            "name": best["name"],
            "score": best["score"],
            "index": best["index"],
            "url": best["url"],
        }

    def _get_google_maps_company_details(
        self,
        record: dict[str, str],
        source_label: str,
    ) -> dict[str, Any]:
        """
        Search Google Maps and return validated Address / Phone details.

        Uses the existing Playwright browser context and the application's
        existing wait/scroll helpers. The application is Playwright-based, so
        Playwright explicit locator/function waits are used rather than adding a
        second Selenium driver stack.
        """
        expected_company = clean_cell_value(
            record.get("Company Name", "-")
        )

        failure = {
            "company_matched": False,
            "company_name": "-",
            "address": "-",
            "zip_code": "-",
            "phone": "-",
            "phone_link": "-",
            "address_link": "-",
            "maps_query": "-",
            "maps_query_type": "-",
        }

        if self._is_unknown_value(
            expected_company
        ):
            return failure

        queries = self._google_maps_queries(
            record
        )

        if not queries:
            return failure

        page = self._get_google_maps_page()

        if page is None:
            return failure

        last_reason = ""
        best_result = None
        excluded_by_query: dict[str, set[str]] = {}
        # Revisit the feed after rejecting a branch; bound attempts per query.
        pending_queries = list(queries)

        for query_index, query_info in enumerate(
            pending_queries,
            1,
        ):
            query_type = query_info.get(
                "type",
                "UNKNOWN",
            )
            query = query_info.get(
                "query",
                "",
            )

            if not query:
                continue
            if self.stop_event.is_set():
                break
            excluded_urls = excluded_by_query.setdefault(query, set())
            matched_result = {}

            def retry_next_candidate():
                candidate_url = matched_result.get("url")
                if candidate_url and candidate_url not in excluded_urls and len(excluded_urls) < 2:
                    excluded_urls.add(candidate_url)
                    pending_queries.insert(query_index, dict(query_info))

            search_url = (
                "https://www.google.com/maps/search/"
                + quote_plus(
                    query,
                    safe="+",
                )
                + "?hl=en"
            )

            self.log(
                f"[MAPS] Searching: {query}"
            )

            self.trace(
                "GOOGLE_MAPS",
                "search_start",
                source=source_label,
                record=self._record_log_context(
                    record
                ),
                query_index=query_index,
                query_type=query_type,
                query=query,
                search_url=search_url,
            )

            try:
                page.goto(
                    search_url,
                    wait_until="domcontentloaded",
                    timeout=GOOGLE_MAPS_PAGE_TIMEOUT_MS,
                )
            except PlaywrightTimeoutError as exc:
                last_reason = (
                    f"Google Maps navigation timeout: {exc}"
                )
                self.trace(
                    "GOOGLE_MAPS",
                    "navigation_timeout",
                    source=source_label,
                    query=query,
                    error=str(exc),
                )
                continue
            except Exception as exc:
                last_reason = (
                    f"Google Maps navigation failed: "
                    f"{type(exc).__name__}: {exc}"
                )
                continue

            self._dismiss_google_consent(
                page
            )

            # Explicit wait for either a detailed place title or a result card.
            try:
                page.wait_for_function(
                    """
                    () => {
                        const title =
                            document.querySelector('h1.DUwDvf');

                        const result =
                            document.querySelector(
                                'div[role="feed"] '
                                + 'a[href*="/maps/place/"]'
                            )
                            || document.querySelector(
                                'a[href*="/maps/place/"]'
                            );

                        return !!title || !!result;
                    }
                    """,
                    timeout=GOOGLE_MAPS_RESULT_WAIT_MS,
                    polling=200,
                )
            except Exception:
                pass

            direct_place = (
                "/maps/place/" in page.url
                or page.locator("h1.DUwDvf:visible").count() > 0
            )

            if not direct_place:
                try:
                    matched_result = (
                        self._find_google_maps_company_result(
                            page=page,
                            company_name=expected_company,
                            source_label=source_label,
                            query=query,
                            excluded_urls=excluded_urls,
                            allow_website_verification=query_type == 'FALLBACK_WEBSITE',
                        )
                    )
                except Exception as exc:
                    matched_result = {
                        "found": False,
                        "name": "-",
                        "score": 0,
                        "index": -1,
                    }
                    last_reason = (
                        "Google Maps result selection failed: "
                        f"{type(exc).__name__}: {exc}"
                    )

                if not matched_result["found"]:
                    last_reason = (
                        "No Google Maps result reasonably matched "
                        f"'{expected_company}'."
                    )

                    self.log(
                        "[MAPS] Company match failed"
                    )
                    self.log(
                        "[MAPS] Google Maps data ignored"
                    )

                    self.trace(
                        "GOOGLE_MAPS",
                        "no_matching_result",
                        source=source_label,
                        query=query,
                        expected_company=expected_company,
                    )

                    # Company Website+City remains a fallback only if the
                    # Company Name query did not find a usable company.
                    continue

                try:
                    page.wait_for_function(
                        "() => !!(" + PLACE_SCRIPT + ")({}).title",
                        timeout=GOOGLE_MAPS_RESULT_WAIT_MS,
                        polling=200,
                    )
                except Exception:
                    pass

                self._nudge_google_maps_details_after_click(
                    page,
                    source_label=source_label,
                    query=query,
                )

            # Whether Google opened the company directly or we clicked a result,
            # wait/scroll the second/details panel before extraction.
            try:
                details_state = (
                    self._wait_and_scroll_google_maps_place(
                        page,
                        source_label=source_label,
                        query=query,
                        need_address=True,
                        need_phone=self._is_unknown_value(
                            record.get(
                                "Phone Number",
                                "-",
                            )
                        ),
                    )
                )

                self.trace(
                    "GOOGLE_MAPS",
                    "details_ready_for_extract",
                    source=source_label,
                    query=query,
                    details_state=details_state,
                )
            except Exception as exc:
                # A failed scroll should not crash the record; attempt extraction
                # once from whatever DOM Google rendered.
                self.trace(
                    "GOOGLE_MAPS",
                    "details_wait_scroll_error",
                    source=source_label,
                    query=query,
                    error_type=type(exc).__name__,
                    error=str(exc),
                )

            try:
                place = self._extract_google_maps_place(
                    page
                )
            except Exception as exc:
                last_reason = (
                    "Google Maps details extraction failed: "
                    f"{type(exc).__name__}: {exc}"
                )
                continue

            # One short second pass catches cards inserted immediately after the
            # last scroll event.
            first_address = clean_cell_value(
                place.get("Full Address", "-")
            )
            first_phone = clean_cell_value(
                place.get("Phone Number", "-")
            )

            if (
                (
                    self._is_unknown_value(
                        record.get("Address", "-")
                    )
                    and self._is_unknown_value(
                        first_address
                    )
                )
                or (
                    self._is_unknown_value(
                        record.get(
                            "Phone Number",
                            "-",
                        )
                    )
                    and self._is_unknown_value(
                        first_phone
                    )
                )
            ):
                try:
                    page.wait_for_timeout(
                        GOOGLE_MAPS_DETAIL_SCROLL_WAIT_MS
                    )
                    second_place = (
                        self._extract_google_maps_place(
                            page
                        )
                    )

                    for key, value in second_place.items():
                        if (
                            self._is_unknown_value(
                                place.get(key, "-")
                            )
                            and not self._is_unknown_value(
                                value
                            )
                        ):
                            place[key] = value

                except Exception:
                    pass

            opened_name = clean_cell_value(
                place.get("Place Name", "-")
            )
            maps_website = clean_cell_value(
                place.get(
                    "Company Website",
                    "-",
                )
            )
            full_address = clean_cell_value(
                place.get("Full Address", "-")
            )

            company_match_score = (
                self._maps_company_match_score(
                    expected_company,
                    opened_name,
                )
            )

            company_matched = (
                self._maps_result_matches_company(
                    expected_company=expected_company,
                    actual_name=opened_name,
                    official_website=clean_cell_value(
                        record.get(
                            "Company Website",
                            "-",
                        )
                    ),
                    maps_website=maps_website,
                )
            )

            if not company_matched:
                last_reason = (
                    "Opened Google Maps company did not match "
                    f"'{expected_company}'."
                )

                self.log(
                    f"[MAPS] Result opened: {opened_name}"
                )
                self.log(
                    "[MAPS] Company match failed"
                )
                self.log(
                    "[MAPS] Google Maps data ignored"
                )

                self.trace(
                    "GOOGLE_MAPS",
                    "opened_company_mismatch",
                    source=source_label,
                    query=query,
                    expected_company=expected_company,
                    opened_company=opened_name,
                    match_score=company_match_score,
                    place=place,
                )
                retry_next_candidate()
                continue

            # Preserve the existing wrong-location protection. When City is
            # unavailable, this check gracefully falls back to State/Country or
            # simply accepts company validation if no geography is available.
            location_matched = True

            if not self._is_unknown_value(
                full_address
            ):
                location_matched = (
                    self._maps_address_matches_location(
                        record,
                        full_address,
                    )
                )

            country_matched = (
                True
                if self._is_unknown_value(
                    full_address
                )
                else self._maps_address_matches_country(
                    record,
                    full_address,
                )
            )

            # A shared corporate website cannot validate a different branch.
            accepted_location = location_matched and country_matched

            if not accepted_location:
                last_reason = (
                    "Google Maps company matched, but the opened "
                    "address did not match the expected location."
                )

                self.log(
                    f"[MAPS] Result opened: {opened_name}"
                )
                self.log(
                    "[MAPS] Company matched, location validation failed"
                )
                self.log(
                    "[MAPS] Google Maps data ignored"
                )
                self.trace("GOOGLE_MAPS", "opened_location_mismatch", source=source_label,
                           query=query, expected_location={key:record.get(key, '-') for key in
                           ('Address', 'City', 'State', 'Country')}, place=place)
                retry_next_candidate()
                continue

            normalized_phone = (
                self._normalize_maps_phone(
                    place.get(
                        "Phone Number",
                        "-",
                    )
                )
            )

            # Maps may expose a different office's international phone even
            # when its visible street address matches the requested location.
            place_country = (self._maps_country_name(re.split(r'[,\n]+', full_address)[-1])
                             or self._maps_country_name(re.split(r'[,\n]+', record.get('Address', '-'))[-1])
                             or self._maps_country_name(record.get('Country', '-')))
            if (normalized_phone.startswith('+') and not normalized_phone.startswith('+1')
                    and place_country in {'united states', 'canada'}):
                self.log('[MAPS] Phone ignored: international prefix conflicts with place country')
                self.trace('GOOGLE_MAPS', 'phone_country_conflict', source=source_label,
                           address=full_address, phone=normalized_phone)
                normalized_phone = '-'

            zip_code = (
                self._extract_postal_code_from_address(
                    full_address
                )
                if not self._is_unknown_value(
                    full_address
                )
                else "-"
            )

            phone_link = (
                f"tel:{normalized_phone}"
                if not self._is_unknown_value(
                    normalized_phone
                )
                else "-"
            )

            if self._is_unknown_value(full_address) and self._is_unknown_value(normalized_phone):
                last_reason = "Matched Maps place contains no usable address or phone."
                retry_next_candidate()
                continue

            result = {
                "company_matched": True,
                "company_name": opened_name,
                "company_website": maps_website,
                "location_matched": accepted_location,
                "company_address_matched": (
                    not self._is_unknown_value(record.get('Address', '-'))
                    and not self._maps_address_is_locality_only(record)
                    and self._maps_address_compatible(record.get('Address', '-'), full_address)),
                # Keep the COMPLETE readable Google Maps business address.
                "address": (
                    full_address
                    if not self._is_unknown_value(
                        full_address
                    )
                    else "-"
                ),
                "zip_code": zip_code,
                "phone": normalized_phone,
                "phone_link": phone_link,
                "address_link": (
                    clean_cell_value(place.get("Address Link", "-"))
                    if not self._is_unknown_value(full_address) else "-"
                ),
                "maps_query": query,
                "maps_query_type": query_type,
            }

            self.log(
                f"[MAPS] Result opened: {opened_name}"
            )
            self.log(
                "[MAPS] Company matched"
            )
            self.log(
                f"[MAPS] Address: {result['address']}"
            )
            self.log(
                f"[MAPS] Phone: {result['phone']}"
            )

            self.trace(
                "GOOGLE_MAPS",
                "validated_company_details",
                source=source_label,
                query=query,
                query_type=query_type,
                expected_company=expected_company,
                result=result,
                place=place,
            )

            if best_result and self._maps_addresses_equivalent(best_result['address'], result['address']):
                # Keep both sources attached to their own fields for one office.
                for key in ('phone', 'phone_link', 'zip_code'):
                    if self._is_unknown_value(best_result[key]):
                        best_result[key] = result[key]
            elif best_result is None:
                best_result = result
            else:
                # Choose a complete place instead of mixing two different offices.
                completeness = lambda item: sum(not self._is_unknown_value(item[key])
                                                for key in ('address', 'zip_code', 'phone'))
                if completeness(result) > completeness(best_result):
                    best_result = result
            need_phone = self._is_unknown_value(record.get('Phone Number', '-'))
            need_address = self._is_unknown_value(record.get('Address', '-'))
            need_zip = self._is_unknown_value(record.get('Zip Code', '-'))
            if ((not need_phone or not self._is_unknown_value(best_result['phone'])) and
                    (not need_zip or not self._is_unknown_value(best_result['zip_code'])) and
                    (not need_address or not self._is_unknown_value(best_result['address']))):
                return best_result
            retry_next_candidate()

        if best_result:
            return best_result

        self.trace(
            "GOOGLE_MAPS",
            "company_details_unavailable",
            source=source_label,
            record=self._record_log_context(
                record
            ),
            reason=last_reason,
            queries=queries,
        )

        failure['reason'] = last_reason or 'No usable Maps place found'
        return failure

    def _google_maps_queries(self, record: dict[str, str]) -> list[dict[str, str]]:
        """Target a known office; otherwise disambiguate the prospect's city."""
        company = clean_cell_value(record.get('Company Name', '-'))
        if self._is_unknown_value(company):
            return []
        existing = clean_cell_value(record.get('Address', '-'))
        if not self._is_unknown_value(existing):
            labels = {'corporate office', 'corporate headquarters', 'headquarters',
                      'head office', 'primary', 'primary location', 'registered office',
                      'postbox', 'po box', 'p o box'}
            location = ', '.join(part.strip() for part in re.split(r'[,\n]+', existing)
                                 if part.strip() and self._maps_geo_text(part) not in labels)
        else:
            location = ' '.join(clean_cell_value(record.get(field, '-')) for field in
                                ('City', 'State', 'Country')
                                if not self._is_unknown_value(record.get(field, '-')))
        primary = self._build_google_maps_query(company, location)
        queries = [{'type': 'PRIMARY_COMPANY_NAME', 'query': primary}]
        website = clean_cell_value(record.get('Company Website', '-'))
        if self._is_valid_official_website(website):
            host = (urlsplit(website).hostname or '').lower().removeprefix('www.')
            fallback = self._build_google_maps_query(host, location)
            if fallback and fallback != primary:
                queries.append({'type': 'FALLBACK_WEBSITE', 'query': fallback})
        return queries[:GOOGLE_MAPS_MAX_QUERIES_PER_COMPANY]


    @staticmethod
    def _maps_address_is_locality_only(record: dict[str, str]) -> bool:
        """Identify a saved city/state label that still needs a street address."""
        parts = SearchWorker._maps_address_parts(record.get('Address', '-'))
        if not parts or any(re.search(r'\d', part) for part in parts):
            return False
        known = set()
        for field in ('City', 'State', 'Country'):
            known.update(SearchWorker._maps_address_parts(record.get(field, '-')))
        return bool(known) and all(part in known for part in parts)

    @staticmethod
    def _maps_phones_equivalent(first: str, second: str, record: dict[str, str]) -> bool:
        a, b = (SearchWorker._normalize_maps_phone(value) for value in (first, second))
        if a == '-' or b == '-':
            return False
        if a == b:
            return True
        tail = re.split(r'[,\n]+', record.get('Address', '-'))[-1]
        country = (SearchWorker._maps_country_name(tail)
                   or SearchWorker._maps_country_name(record.get('Country', '-')))
        code = {'united states':'1', 'canada':'1', 'india':'91', 'united kingdom':'44',
                'switzerland':'41', 'australia':'61', 'new zealand':'64',
                'france':'33', 'germany':'49'}.get(country)
        if not code:
            return False
        def national(value):
            if value.startswith('+' + code):
                return value[len(code) + 1:]
            if value.startswith('+'):
                return None
            if code == '1':
                return value[1:] if len(value) == 11 and value.startswith('1') else value
            return value.removeprefix('0')
        return national(a) is not None and national(a) == national(b)

    def _merge_google_maps_fields(self, record: dict[str, str], proposed: dict[str, str]) -> dict[str, str]:
        """Fill missing facts from the same office despite display abbreviations."""
        applied = {}
        missing = self._is_unknown_value
        proposed = dict(proposed)
        if not is_public_source_url(proposed.get('Address Link', '-')):
            proposed.pop('Address Link', None)
        phone_link = proposed.get('Phone Number Link', '-')
        if isinstance(phone_link, str) and phone_link.lower().startswith('tel:'):
            linked_phone = proposed.get('Phone Number', record.get('Phone Number', '-'))
            valid_phone_link = (not missing(linked_phone)
                                and self._maps_phones_equivalent(phone_link[4:], linked_phone, record))
        else:
            valid_phone_link = is_public_source_url(phone_link)
        if not valid_phone_link:
            proposed.pop('Phone Number Link', None)
        old_address, new_address = record.get('Address', '-'), proposed.get('Address', '-')
        address_allowed = missing(old_address) or self._maps_address_compatible(old_address, new_address)
        old_phone, new_phone = record.get('Phone Number', '-'), proposed.get('Phone Number', '-')
        same_phone = (not missing(old_phone) and not missing(new_phone)
                      and self._maps_phones_equivalent(old_phone, new_phone, record))
        new_parts = self._maps_address_parts(new_address)
        if (address_allowed and self._maps_address_is_locality_only(record) and len(new_parts) > 1
                and re.search(r'\d|\b(?:street|road|avenue|drive|lane|highway)\b', new_parts[0])
                and not missing(proposed.get('Address Link', '-'))):
            # A saved locality is incomplete. Keep full existing street addresses;
            # upgrade only this positively identified city/state placeholder.
            record['Address'] = applied['Address'] = new_address
            record['Address Link'] = applied['Address Link'] = proposed['Address Link']
        for field, value in proposed.items():
            if missing(value) or not missing(record.get(field, '-')):
                continue
            if field in {'Address', 'Zip Code', 'Address Link', 'Phone Number'} and not address_allowed:
                continue
            if field == 'Phone Number Link' and not (same_phone or (missing(old_phone) and address_allowed)):
                continue
            record[field] = value
            applied[field] = value
        return applied

    @staticmethod
    def _google_maps_needs_enrichment(record: dict[str, str]) -> bool:
        """Retry missing Maps values and source links for values already saved."""
        missing = SearchWorker._is_unknown_value
        if SearchWorker._maps_address_is_locality_only(record):
            return True
        if any(missing(record.get(field, '-'))
               for field in ('Address', 'Zip Code', 'Phone Number')):
            return True
        return any(not missing(record.get(value_field, '-'))
                   and missing(record.get(link_field, '-'))
                   for value_field, link_field in (
                       ('Address', 'Address Link'),
                       ('Phone Number', 'Phone Number Link'),
                   ))

    def _apply_google_maps_details(self, record: dict, details: dict) -> dict:
        """Revalidate cached/fresh company contacts and record their office scope."""
        from maps_office_scope import classify_maps_location
        details = dict(details)
        address = clean_cell_value(details.get('address'))
        has_address = not self._is_unknown_value(address)
        details['location_matched'] = (details.get('location_matched') is not False
            and has_address and self._maps_address_matches_location(record, address)
            and self._maps_address_matches_country(record, address))
        if not self._is_unknown_value(details.get('company_name')):
            details['company_matched'] = (details.get('company_matched') is True
                and self._maps_result_matches_company(record.get('Company Name', '-'),
                    details['company_name'], record.get('Company Website', '-'),
                    details.get('company_website', '-')))
        classification_record = dict(record)
        if self._maps_address_is_locality_only(record):
            classification_record['Address'] = '-'
        existing = classification_record.get('Address', '-')
        details['company_address_matched'] = (
            not self._is_unknown_value(existing) and has_address
            and self._maps_address_compatible(existing, address))
        scope = classify_maps_location(classification_record, details)
        if details.get('reason') and not scope['is_mergeable']:
            scope['reason'] += '; ' + clean_cell_value(details['reason'])
        applied = {}
        if scope['is_mergeable']:
            values = {field: clean_cell_value(details.get(key)) for field, key in (
                ('Address', 'address'), ('Zip Code', 'zip_code'), ('Phone Number', 'phone'),
                ('Phone Number Link', 'phone_link'), ('Address Link', 'address_link'))}
            applied = self._merge_google_maps_fields(record, values)
        record['_maps_evidence'] = {
            **scope, 'profile': record.get('Prospect LinkedIn Link', '-'),
            'company': record.get('Company Name', '-'),
            'applied_fields': ', '.join(applied) or 'None',
            'checked_at': datetime.now(timezone.utc).isoformat(),
        }
        return applied

    @timed('maps_lookup')
    def _google_maps_enrich_record(
        self,
        record: dict[str, str],
        source_label: str,
    ) -> dict[str, str]:
        """
        Fill only missing Address / Zip / Phone values from Google Maps.

        Existing reliable values are never overwritten.

        The actual Maps lookup is delegated to
        _get_google_maps_company_details(), which selects and validates the
        correct company result before extracting the detailed panel.
        """
        # A prospect city is a search target, not evidence that the person works
        # at this particular office. Preserve another office as a complete group.
        if self.prefer_local_office and any(not self._is_unknown_value(record.get(k))
                                           for k in ('City', 'State', 'Country')):
            target = dict(record)
            target['Address'] = '-'
            address = record.get('Address', '-')
            if (not self._is_unknown_value(address) and
                    not self._maps_address_matches_location(target, address)):
                local = dict(record)
                contact_keys = ('Address', 'Zip Code', 'Phone Number', 'Address Link', 'Phone Number Link')
                for key in contact_keys:
                    local[key] = '-'
                # Recursive call has no saved office, so validation uses prospect geography.
                local = self._google_maps_enrich_record(local, source_label)
                evidence = local.get('_maps_evidence', {})
                if evidence.get('is_mergeable') and not self._is_unknown_value(local.get('Address')):
                    record['_previous_company_location'] = record.get('_previous_company_location') or {
                        'profile': record.get('Prospect LinkedIn Link', '-'),
                        'company': record.get('Company Name', '-'),
                        'address': address, 'zip': record.get('Zip Code', '-'),
                        'phone': record.get('Phone Number', '-'),
                        'source_url': record.get('Address Link', '-'),
                        'phone_source': record.get('Phone Number Link', '-'),
                        'basis': 'Another company location; local Maps contact selected using prospect geography',
                    }
                    for key in (*contact_keys, '_maps_evidence'):
                        record[key] = local.get(key, '-')
                    self._append_enrichment_note(record, 'Maps company location selected near prospect; previous company address preserved')
                elif evidence:
                    record['_maps_evidence'] = evidence
                return record

        needs_address = self._is_unknown_value(
            record.get("Address", "-")
        )
        needs_zip = self._is_unknown_value(
            record.get("Zip Code", "-")
        )
        needs_phone = self._is_unknown_value(
            record.get("Phone Number", "-")
        )

        if not self._google_maps_needs_enrichment(record):
            return record

        if self._is_unknown_value(
            record.get("Company Name", "-")
        ):
            self.log(
                "[MAPS] Skipped: Company Name is missing"
            )
            return record

        cache_key = (
            self._google_maps_cache_key(
                record
            )
        )

        if (
            cache_key
            and cache_key
            in self.google_maps_cache
        ):
            cached = dict(
                self.google_maps_cache[
                    cache_key
                ]
            )

            applied_from_cache = self._apply_google_maps_details(record, cached)

            self.log(
                "[MAPS] Cache hit: "
                f"{record.get('Company Name', '-')}"
            )

            self.trace(
                "GOOGLE_MAPS",
                "cache_hit",
                source=source_label,
                record=self._record_log_context(
                    record
                ),
                cached=cached,
                applied=applied_from_cache,
            )

            return record

        try:
            details = (
                self._get_google_maps_company_details(
                    record,
                    source_label,
                )
            )
        except Exception as exc:
            self.log(
                "[MAPS] Lookup failed: "
                f"{type(exc).__name__}: {exc}"
            )

            self.trace(
                "GOOGLE_MAPS",
                "lookup_exception",
                source=source_label,
                record=self._record_log_context(
                    record
                ),
                error_type=type(exc).__name__,
                error=str(exc),
            )
            return record

        applied = self._apply_google_maps_details(record, details)

        if cache_key:
            self.google_maps_cache[
                cache_key
            ] = dict(
                details
            )

        if applied:
            self._append_enrichment_note(
                record,
                "Google Maps enriched: "
                + ", ".join(
                    applied.keys()
                ),
            )

            self.log(
                "MAPS ENRICHMENT APPLIED | "
                f"company='{record.get('Company Name', '-')}' | "
                f"address='{record.get('Address', '-')}' | "
                f"zip='{record.get('Zip Code', '-')}' | "
                f"phone='{record.get('Phone Number', '-')}'"
            )

        self.trace(
            "GOOGLE_MAPS",
            "enrichment_result",
            source=source_label,
            record=self._record_log_context(
                record
            ),
            details=details,
            applied=applied,
            final_address=record.get(
                "Address",
                "-",
            ),
            final_zip=record.get(
                "Zip Code",
                "-",
            ),
            final_phone=record.get(
                "Phone Number",
                "-",
            ),
        )

        return record

    @timed('smart_validation')
    def _smart_validate_record(
        self,
        record: dict[str, str],
        source_label: str,
        validate_existing: bool = True,
    ) -> dict[str, str]:
        """
        SMART-mode validation + enrichment.

        1. Preserve direct scraped values.
        2. Identify required missing/invalid fields.
        3. Send only those fields (plus explicitly validated fields) to Qwen.
        4. Apply conservative confidence thresholds.
        5. Keep valid scraped values unless the field is intentionally being
           validated/corrected.
        6. If Qwen is unavailable or returns invalid JSON, continue scraping and
           mark unresolved fields in Notes for later enrichment.
        """
        requested_fields = self._required_enrichment_fields(
            record,
            validate_existing=validate_existing,
        )

        self.trace(
            "ENRICHMENT",
            "field_check",
            source=source_label,
            validate_existing=validate_existing,
            record=self._record_log_context(record),
            requested_fields=list(requested_fields),
            current_values={
                field: clean_cell_value(
                    record.get(field, "-")
                )
                for field in ENRICHMENT_AUDIT_FIELDS
            },
        )

        requested_fields, cache_hits = self._apply_smart_cache(
            record,
            requested_fields,
        )

        if cache_hits:
            self.log(
                "SMART CACHE HIT | "
                f"source='{source_label}' | "
                + ", ".join(cache_hits)
            )

            self.trace(
                "ENRICHMENT",
                "cache_hit",
                source=source_label,
                record=self._record_log_context(record),
                fields=cache_hits,
                cached_values={
                    field: clean_cell_value(
                        record.get(field, "-")
                    )
                    for field in cache_hits
                },
            )

            cached_unresolved = [
                field
                for field in cache_hits
                if self._field_missing_or_invalid(record, field)
            ]

            if cached_unresolved:
                self._append_enrichment_note(
                    record,
                    "Unresolved after cached LLM enrichment: "
                    + ", ".join(cached_unresolved),
                )

        if not requested_fields:
            self.log(
                "SMART ENRICHMENT | "
                f"source='{source_label}' | satisfied from scraped/cache data"
            )
            return record

        # Check/recheck the local LLM. A failed check is only temporary:
        # another record can retry after LLM_RECHECK_INTERVAL_SECONDS.
        if requested_fields:
            llm_ok, model = self._ensure_llm_available(
                force=(self.llm_available is None)
            )

            if llm_ok:
                self.log(
                    f"LLM FALLBACK CONNECTED | source='{source_label}' | "
                    f"model='{model}'"
                )
            else:
                self.log(
                    f"LLM FALLBACK OFFLINE | source='{source_label}' | "
                    f"{model}"
                )

        if self.llm_available is False:
            unresolved_now = [
                field
                for field in requested_fields
                if self._field_missing_or_invalid(record, field)
            ]

            if unresolved_now:
                self._append_enrichment_note(
                    record,
                    "LLM enrichment pending: " + ", ".join(unresolved_now),
                )

            self.log(
                "SMART ENRICHMENT OFFLINE | "
                f"source='{source_label}' | "
                "scraped values preserved"
            )

            self.trace(
                "ENRICHMENT",
                "offline",
                source=source_label,
                record=self._record_log_context(record),
                unresolved_fields=unresolved_now,
            )

            return record

        originally_missing_or_invalid = {
            field
            for field in requested_fields
            if self._field_missing_or_invalid(record, field)
        }

        before = {
            field: clean_cell_value(record.get(field, "-"))
            for field in requested_fields
        }

        temporary_website_evidence_added = False

        if (
            "Address" in requested_fields
            or "Zip Code" in requested_fields
        ):
            website_evidence = (
                self._collect_official_website_evidence(
                    record
                )
            )

            if website_evidence != "-":
                record[
                    "Official Website Evidence"
                ] = website_evidence
                temporary_website_evidence_added = True

                self.log(
                    "SMART WEBSITE EVIDENCE READY | "
                    f"source='{source_label}' | "
                    f"chars={len(website_evidence)}"
                )

                self.trace(
                    "ENRICHMENT",
                    "website_evidence",
                    source=source_label,
                    record=self._record_log_context(record),
                    website=clean_cell_value(
                        record.get("Company Website", "-")
                    ),
                    evidence_chars=len(website_evidence),
                    website_evidence=website_evidence,
                )

        try:
            started = time.perf_counter()

            self.trace(
                "ENRICHMENT",
                "llm_call_start",
                source=source_label,
                record=self._record_log_context(record),
                requested_fields=requested_fields,
            )

            result = self.llm.enrich_missing_fields(
                record=record,
                requested_fields=requested_fields,
            )

            self.qwen_calls += 1
            self.qwen_seconds += time.perf_counter() - started

            values = result.get("values", {})
            confidence = result.get("confidence", {})

            if not isinstance(values, dict):
                raise ValueError("LLM enrichment returned invalid values object.")

            applied = []
            unresolved = []
            rejected: list[dict[str, Any]] = []

            for field in requested_fields:
                old_value = before[field]

                proposed = self._normalize_llm_enrichment_value(
                    field,
                    clean_cell_value(values.get(field, "Unknown")),
                )

                try:
                    score = int(confidence.get(field, 0))
                except Exception:
                    score = 0

                score = max(0, min(100, score))
                threshold = self._confidence_threshold(field)

                old_is_valid = (
                    field not in originally_missing_or_invalid
                    and not self._is_unknown_value(old_value)
                )

                # Never replace a valid scraped value with Unknown or with a
                # low-confidence inference.
                if old_is_valid:
                    if proposed == "Unknown":
                        rejected.append(
                            {
                                "field": field,
                                "reason": "LLM returned Unknown; kept valid scraped value",
                                "scraped_value": old_value,
                                "proposed_value": proposed,
                                "confidence": score,
                                "threshold": threshold,
                            }
                        )
                        continue

                    if score < threshold:
                        rejected.append(
                            {
                                "field": field,
                                "reason": "confidence below threshold; kept valid scraped value",
                                "scraped_value": old_value,
                                "proposed_value": proposed,
                                "confidence": score,
                                "threshold": threshold,
                            }
                        )
                        continue

                    if proposed != old_value:
                        record[field] = proposed
                        applied.append(
                            f"{field}: '{old_value}' -> '{proposed}' ({score}%)"
                        )

                    continue

                # Missing/invalid scraped field:
                # Apply only sufficiently supported values. Otherwise mark it
                # Unknown so it is visibly unresolved rather than fabricated.
                if proposed != "Unknown" and score >= threshold:
                    record[field] = proposed
                    applied.append(
                        f"{field}: '{old_value}' -> '{proposed}' ({score}%)"
                    )
                else:
                    record[field] = "Unknown"
                    unresolved.append(field)

                    rejected.append(
                        {
                            "field": field,
                            "reason": (
                                "LLM returned Unknown"
                                if proposed == "Unknown"
                                else "confidence below threshold"
                            ),
                            "scraped_value": old_value,
                            "proposed_value": proposed,
                            "confidence": score,
                            "threshold": threshold,
                        }
                    )

            self.trace(
                "ENRICHMENT",
                "decision",
                source=source_label,
                record=self._record_log_context(record),
                requested_fields=requested_fields,
                raw_values=values,
                confidence=confidence,
                applied=applied,
                unresolved=unresolved,
                rejected=rejected,
                final_values={
                    field: clean_cell_value(
                        record.get(field, "-")
                    )
                    for field in requested_fields
                },
                thresholds={
                    field: self._confidence_threshold(field)
                    for field in requested_fields
                },
            )

            # Cache the final decision (including Unknown) so another lead
            # from the same company/title does not trigger the same LLM work.
            self._store_smart_cache(
                record,
                requested_fields,
            )

            if applied:
                applied_fields = [
                    item.split(":", 1)[0].strip()
                    for item in applied
                ]

                self._append_enrichment_note(
                    record,
                    "LLM enriched: "
                    + ", ".join(applied_fields),
                )

                self.log(
                    "SMART ENRICHMENT APPLIED | "
                    f"source='{source_label}' | "
                    + " | ".join(applied)
                )

            if unresolved:
                self._append_enrichment_note(
                    record,
                    "Unresolved after LLM enrichment: "
                    + ", ".join(unresolved),
                )

                self.log(
                    "SMART ENRICHMENT UNRESOLVED | "
                    f"source='{source_label}' | "
                    + ", ".join(unresolved)
                )

            if not applied and not unresolved:
                self.log(
                    "SMART ENRICHMENT VALIDATED | "
                    f"source='{source_label}' | no corrections required"
                )

            if temporary_website_evidence_added:
                record.pop(
                    "Official Website Evidence",
                    None,
                )

        except Exception as exc:
            if temporary_website_evidence_added:
                record.pop(
                    "Official Website Evidence",
                    None,
                )

            self.trace(
                "ENRICHMENT",
                "exception",
                source=source_label,
                record=self._record_log_context(record),
                requested_fields=requested_fields,
                error_type=type(exc).__name__,
                error=str(exc),
            )

            # A local-server/network failure should not trigger another long
            # timeout for every subsequent record in the same run.
            if isinstance(exc, requests.RequestException):
                self.llm_available = False
                self.llm_last_check_ts = time.time()

            # Keep all scraped values exactly as they were and mark fields for a
            # future enrichment pass. Scraping must continue.
            pending = sorted(
                originally_missing_or_invalid
                or set(requested_fields)
            )

            self._append_enrichment_note(
                record,
                "LLM enrichment pending: " + ", ".join(pending),
            )

            self.log(
                "SMART ENRICHMENT WARNING | "
                f"source='{source_label}' | "
                f"{type(exc).__name__}: {exc} | "
                "scraped values preserved; unresolved fields marked for later"
            )

        return record

    def _ensure_llm_available(
        self,
        force: bool = False,
    ) -> tuple[bool, str]:
        """
        Recheck a previously unavailable local model periodically.

        This prevents a single startup/network hiccup from causing every record
        in the run to be written as "LLM enrichment pending".
        """
        now = time.time()

        if self.llm_available is True and not force:
            return True, self.llm.model_id

        if (
            self.llm_available is False
            and not force
            and (now - self.llm_last_check_ts)
            < LLM_RECHECK_INTERVAL_SECONDS
        ):
            return False, "Local LLM retry interval has not elapsed."

        ok, detail = self.llm.health_check()

        self.llm_last_check_ts = time.time()
        self.llm_available = ok

        return ok, detail

    @staticmethod
    def _remove_llm_pending_note(
        record: dict[str, str],
    ) -> None:
        notes = clean_cell_value(record.get("Notes", "-"))

        if notes == "-":
            return

        parts = [
            part.strip()
            for part in notes.split("|")
            if part.strip()
        ]

        parts = [
            part
            for part in parts
            if not part.lower().startswith(
                "llm enrichment pending:"
            )
        ]

        record["Notes"] = " | ".join(parts) if parts else "-"

    def _retry_pending_enrichment_rows(
        self,
        workbook: WorkbookManager,
        pending_rows: list[tuple[int, dict[str, str]]],
        source_label: str,
        validate_existing: bool,
    ) -> None:
        """
        Final recovery pass.

        Records are still checkpointed during scraping. If the local model was
        temporarily unavailable, retry pending rows once before the final Excel
        save and update them in place when enrichment succeeds.
        """
        if not pending_rows or self.stop_event.is_set():
            return

        self.log(
            "LLM FINAL RETRY | "
            f"source='{source_label}' | "
            f"pending_rows={len(pending_rows)}"
        )

        ok, detail = self._ensure_llm_available(force=True)

        if not ok:
            self.log(
                "LLM FINAL RETRY OFFLINE | "
                f"source='{source_label}' | {detail}"
            )
            return

        self.log(
            "LLM FINAL RETRY CONNECTED | "
            f"source='{source_label}' | model='{detail}'"
        )

        updated = 0

        for row_idx, original in pending_rows:
            if self.stop_event.is_set():
                break

            record = dict(original)
            self._remove_llm_pending_note(record)

            record = self._smart_validate_record(
                record,
                source_label=f"{source_label} final retry",
                validate_existing=validate_existing,
            )

            notes = clean_cell_value(
                record.get("Notes", "-")
            ).lower()

            if "llm enrichment pending:" not in notes:
                workbook.update_record(row_idx, record)
                updated += 1

        self.log(
            "LLM FINAL RETRY COMPLETE | "
            f"source='{source_label}' | "
            f"updated_rows={updated}/{len(pending_rows)}"
        )

    @staticmethod
    def _record_log_context(
        record: dict[str, str],
    ) -> dict[str, str]:
        return {
            "name": " ".join(
                value
                for value in (
                    clean_cell_value(record.get("First Name", "-")),
                    clean_cell_value(record.get("Last Name", "-")),
                )
                if value != "-"
            ).strip() or "-",
            "job_title": clean_cell_value(
                record.get("Job Title", "-")
            ),
            "company": clean_cell_value(
                record.get("Company Name", "-")
            ),
            "prospect_url": clean_cell_value(
                record.get("Prospect LinkedIn Link", "-")
            ),
            "company_url": clean_cell_value(
                record.get("Company LinkedIn URL", "-")
            ),
        }

    def _local_llm_trace(
        self,
        event: str,
        payload: dict[str, Any],
    ) -> None:
        callback = self.callbacks.get("trace")

        if callback:
            callback(
                "LOCAL_LLM",
                {
                    "event": event,
                    **payload,
                },
            )

    def trace(
        self,
        category: str,
        event: str,
        **payload,
    ) -> None:
        callback = self.callbacks.get("trace")

        if callback:
            callback(
                category,
                {
                    "event": event,
                    **payload,
                },
            )

    def _log_enrichment_audit(
        self,
        workbook: WorkbookManager,
        source_label: str,
    ) -> None:
        audit = workbook.enrichment_audit()

        self.log(
            "ENRICHMENT AUDIT | "
            f"source='{source_label}' | "
            f"rows={audit['total_rows']} | "
            f"rows_with_missing={audit['rows_with_unresolved_fields']}"
        )

        self.trace(
            "OUTPUT_AUDIT",
            "final_enrichment_audit",
            source=source_label,
            **audit,
        )

        missing_counts = [
            f"{field}={count}"
            for field, count in audit[
                "field_missing_counts"
            ].items()
            if count
        ]

        if missing_counts:
            self.log(
                "ENRICHMENT MISSING COUNTS | "
                + " | ".join(missing_counts)
            )
            self.log(
                "ENRICHMENT DIAGNOSTIC | "
                "Open the *_activity.log file to see Local LLM health, "
                "requested fields, evidence, raw response, confidence, "
                "rejection reasons, and final row audit."
            )
        else:
            self.log(
                "ENRICHMENT AUDIT COMPLETE | "
                "all required enrichment fields are populated"
            )

    def send(self, name: str, payload: Any = None):
        # These events must reach the active loop, not queue behind it.
        if name == "pause":
            self.pause_event.set()
            return
        if name == "resume":
            self.pause_event.clear()
            return
        if name in {"stop", "shutdown", "close_browser"}:
            self.stop_event.set()
            self.pause_event.clear()
            if name == "stop":
                return
        self.commands.put((name, payload))

    def log(self, message: str):
        cb = self.callbacks.get("log")
        if cb:
            cb(message)

    def status(self, message: str):
        cb = self.callbacks.get("status")
        if cb:
            cb(message)

    @timed('browser')
    def _launch_browser(self):
        if self.context and (self._browser_headless is None or self._browser_headless == self.headless):
            return
        if self.context:
            self._close_browser(close_enrichment=False)

        CHROME_PROFILE_DIR.mkdir(parents=True, exist_ok=True)

        self.status("Launching Chrome...")
        self.log(f"Chrome session folder: {CHROME_PROFILE_DIR}")

        try:
            self.context = self.pw.chromium.launch_persistent_context(
                user_data_dir=str(CHROME_PROFILE_DIR),
                channel="chrome",
                headless=self.headless,
                chromium_sandbox=True,
                viewport={'width': 1365, 'height': 900} if self.headless else None,
                args=[
                    "--start-minimized",
                    "--disable-background-timer-throttling",
                    "--disable-backgrounding-occluded-windows",
                    "--disable-renderer-backgrounding",
                ],
            )
        except Exception as chrome_exc:
            self.log(f"Chrome channel launch failed: {chrome_exc}")
            self.log("Trying Playwright Chromium fallback...")

            self.context = self.pw.chromium.launch_persistent_context(
                user_data_dir=str(CHROME_PROFILE_DIR),
                headless=self.headless,
                chromium_sandbox=True,
                viewport={'width': 1365, 'height': 900} if self.headless else None,
                args=[
                    "--start-minimized",
                    "--disable-background-timer-throttling",
                    "--disable-backgrounding-occluded-windows",
                    "--disable-renderer-backgrounding",
                ],
            )

        pages = self.context.pages
        self.page = pages[0] if pages else self.context.new_page()

        self.page.set_default_timeout(PAGE_TIMEOUT_MS)
        self._browser_headless = self.headless
        self.context.on('request', lambda req: self.metrics.count('linkedin_navigations')
            if req.is_navigation_request() and (urlsplit(req.url).hostname or '').endswith('linkedin.com') else None)

        self.status("Chrome ready")
        self.log(
            "Chrome is ready. Sign in manually if LinkedIn requests authentication."
        )

    def _open_login(self):
        self.headless = False
        self._launch_browser()

        self.page.goto(
            "https://www.linkedin.com/sales/",
            wait_until="domcontentloaded",
            timeout=PAGE_TIMEOUT_MS,
        )

        self.status("Waiting for sign-in")
        self.log(
            "Use the Chrome window to sign in normally. "
            "Complete any legitimate verification yourself."
        )

    def _wait_if_paused(self) -> bool:
        while self.pause_event.is_set() and not self.stop_event.is_set():
            self.status("Paused")
            time.sleep(0.25)

        return not self.stop_event.is_set()

    def _wait_for_manual_login_if_needed(self):
        current = (self.page.url or "").lower()
        indicators = (
            "/login",
            "/checkpoint",
            "/authwall",
            "uas/login",
        )

        if not any(token in current for token in indicators):
            return

        if self.headless:
            self.stop_event.set()
            raise ProfileAccessError('Sign-in or verification is required. Stop the run, use Open Chrome / Sign In, then restart. Headless mode cannot complete interactive sign-in.')

        self.log(
            "LinkedIn requires sign-in or verification. "
            "Waiting for you to complete it in Chrome."
        )
        self.status("Complete sign-in in Chrome")

        deadline = time.time() + 300

        while time.time() < deadline and not self.stop_event.is_set():
            current = (self.page.url or "").lower()

            if not any(token in current for token in indicators):
                self.log("Sign-in detected. Continuing.")
                return

            time.sleep(1)

        raise TimeoutError("Sign-in was not completed within 5 minutes.")

    @staticmethod
    def _clean_card_text(text: str) -> str:
        text = (text or "").replace("\xa0", " ")
        text = re.sub(r"[ \t]+", " ", text)
        text = re.sub(r"\n{3,}", "\n\n", text)
        return text.strip()

    @timed('search_results')
    def _collect_current_cards(self) -> list[dict[str, str]]:
        """
        VERSION 10 - SAME-CARD COMPANY EXTRACTION

        Sales Navigator usually renders several anchors for one prospect
        (photo, name, accessibility/profile links). The previous logic counted
        anchors, so it could stop at a small header container before reaching
        the company link.

        This version:
        1. Canonicalizes each lead URL.
        2. Walks upward while the ancestor still belongs to only ONE distinct
           prospect.
        3. Reads /sales/company/ or /sales/account/ only from that same card.
        4. Falls back to parsing a visible "Job Title · Company" line.
        """
        started = time.perf_counter()

        try:
            cards = self.page.evaluate(
                """
                () => {
                    const LEAD_SELECTOR =
                        'a[href*="/sales/lead/"], a[href*="/sales/people/"]';

                    const COMPANY_SELECTOR =
                        'a[href*="/sales/company/"], a[href*="/sales/account/"]';

                    function cleanText(value) {
                        return (value || '')
                            .replace(/\\u00a0/g, ' ')
                            .replace(/[ \\t]+/g, ' ')
                            .trim();
                    }

                    function normalizeLeadUrl(href) {
                        if (!href) return '';

                        try {
                            const u = new URL(href, window.location.origin);

                            // Sales Navigator can put tokens after the lead id in
                            // the pathname, separated by commas. Keep the stable
                            // lead identity portion so all anchors for one person
                            // collapse to the same key.
                            let path = u.pathname.replace(/\\/$/, '');

                            const leadMatch = path.match(
                                /^(\\/sales\\/(?:lead|people)\\/[^,/?#]+)/
                            );

                            if (leadMatch) {
                                path = leadMatch[1];
                            }

                            return u.origin + path;
                        } catch (e) {
                            return href
                                .split('?')[0]
                                .split('#')[0]
                                .replace(/\\/$/, '');
                        }
                    }

                    function cleanNameCandidate(value) {
                        let s = cleanText(value);
                        if (!s) return '';

                        const match = s.match(
                            /^view\\s+(.+?)(?:'s|’s)\\s+profile$/i
                        );

                        if (match) {
                            s = match[1].trim();
                        }

                        s = s
                            .replace(/^view\\s+profile$/i, '')
                            .trim();

                        s = s
                            .replace(/\\s+is\\s+reachable.*$/i, '')
                            .trim();

                        s = s
                            .replace(
                                /\\s*[·•]\\s*(?:1st|2nd|3rd|4th|5th|6th|7th|8th|9th)\\b.*$/i,
                                ''
                            )
                            .trim();

                        return s;
                    }

                    function nameScore(value) {
                        const s = cleanNameCandidate(value);

                        if (!s) return -1000;

                        const lower = s.toLowerCase();

                        const badValues = [
                            'view profile',
                            'degree connection',
                            'linkedin premium',
                            'save',
                            'message',
                            'select',
                            'show more'
                        ];

                        if (badValues.some(item => lower.includes(item))) {
                            return -500;
                        }

                        const words = s.split(/\\s+/).filter(Boolean);

                        if (words.length < 1 || words.length > 8) {
                            return -100;
                        }

                        let score = 20;

                        if (words.length >= 2) score += 30;
                        if (/^[\\p{L}.'’,\\- ]+$/u.test(s)) score += 20;
                        if (s.length >= 3 && s.length <= 90) score += 10;

                        return score;
                    }

                    function distinctLeadUrls(node) {
                        const urls = new Set();

                        if (!node) return urls;

                        for (const anchor of node.querySelectorAll(LEAD_SELECTOR)) {
                            const key = normalizeLeadUrl(anchor.href || '');

                            if (key) {
                                urls.add(key);
                            }
                        }

                        return urls;
                    }

                    function findProspectCard(anchor, targetUrl) {
                        let node = anchor;
                        let bestNode = null;

                        /*
                        Keep expanding while this ancestor still contains only
                        the target prospect. Stop as soon as the ancestor starts
                        containing a second distinct prospect.
                        */
                        for (let depth = 0; depth < 18 && node; depth++) {
                            const text = cleanText(node.innerText || '');

                            if (text.length >= 20 && text.length <= 5000) {
                                const prospectUrls = distinctLeadUrls(node);

                                if (
                                    prospectUrls.size === 1 &&
                                    prospectUrls.has(targetUrl)
                                ) {
                                    bestNode = node;
                                } else if (
                                    prospectUrls.size > 1 &&
                                    bestNode
                                ) {
                                    break;
                                }
                            }

                            node = node.parentElement;
                        }

                        // Semantic fallback when the page uses list items.
                        if (!bestNode) {
                            const li = anchor.closest('li');

                            if (li) {
                                const urls = distinctLeadUrls(li);

                                if (
                                    urls.size === 1 &&
                                    urls.has(targetUrl)
                                ) {
                                    bestNode = li;
                                }
                            }
                        }

                        return bestNode || anchor.parentElement || anchor;
                    }

                    function companyFromCard(card) {
                        if (!card) return '';

                        // Preferred: exact Sales Navigator company/account link.
                        const companyLinks = Array.from(
                            card.querySelectorAll(COMPANY_SELECTOR)
                        );

                        for (const link of companyLinks) {
                            const possibleValues = [
                                link.innerText,
                                link.textContent,
                                link.getAttribute('aria-label'),
                                link.getAttribute('title')
                            ];

                            for (const value of possibleValues) {
                                const company = cleanText(value);

                                if (
                                    company &&
                                    company.length >= 2 &&
                                    company.length <= 180
                                ) {
                                    return company;
                                }
                            }
                        }

                        /*
                        Fallback:
                            Intern · Josh Software, Inc.
                            Professor · University of Northern Iowa
                        */
                        const lines = (card.innerText || '')
                            .split('\\n')
                            .map(cleanText)
                            .filter(Boolean);

                        for (const line of lines) {
                            const match = line.match(
                                /^(.+?)\\s*[·•]\\s*(.+)$/
                            );

                            if (!match) continue;

                            const left = cleanText(match[1]);
                            const right = cleanText(match[2]);

                            if (!left || !right) continue;

                            // Ignore person-name relationship lines:
                            // "Manasi Mane · 2nd"
                            if (
                                /^(?:1st|2nd|3rd|4th|5th|6th|7th|8th|9th)\\b/i
                                .test(right)
                            ) {
                                continue;
                            }

                            if (
                                /degree connection/i.test(right) ||
                                /viewed/i.test(right) ||
                                /linkedin premium/i.test(right)
                            ) {
                                continue;
                            }

                            if (
                                /^(save|message|select|show more)$/i.test(right)
                            ) {
                                continue;
                            }

                            if (right.length >= 2 && right.length <= 180) {
                                return right;
                            }
                        }

                        return '';
                    }

                    function bestPersonName(card, targetLeadUrl) {
                        let bestName = '';
                        let bestScore = -10000;

                        const candidateAnchors = Array.from(
                            card.querySelectorAll(LEAD_SELECTOR)
                        );

                        for (const anchor of candidateAnchors) {
                            const candidateUrl = normalizeLeadUrl(
                                anchor.href || ''
                            );

                            if (candidateUrl !== targetLeadUrl) {
                                continue;
                            }

                            const values = [
                                anchor.innerText,
                                anchor.textContent,
                                anchor.getAttribute('aria-label'),
                                anchor.getAttribute('title'),
                                anchor.querySelector('img')?.getAttribute('alt')
                            ];

                            for (const value of values) {
                                const candidate = cleanNameCandidate(value);
                                const score = nameScore(value);

                                if (score > bestScore) {
                                    bestScore = score;
                                    bestName = candidate;
                                }
                            }
                        }

                        return {
                            name: bestName,
                            score: bestScore
                        };
                    }

                    const allLeadAnchors = Array.from(
                        document.querySelectorAll(LEAD_SELECTOR)
                    );

                    const results = new Map();

                    for (const leadAnchor of allLeadAnchors) {
                        const rawUrl = cleanText(leadAnchor.href || '');

                        if (!rawUrl) continue;

                        const normalizedUrl = normalizeLeadUrl(rawUrl);

                        if (!normalizedUrl) continue;

                        const card = findProspectCard(
                            leadAnchor,
                            normalizedUrl
                        );

                        if (!card) continue;

                        const cardText = cleanText(card.innerText || '');

                        if (!cardText || cardText.length < 20) {
                            continue;
                        }

                        const person = bestPersonName(
                            card,
                            normalizedUrl
                        );

                        const company = companyFromCard(card);

                        const previous = results.get(normalizedUrl);

                        if (!previous) {
                            results.set(
                                normalizedUrl,
                                {
                                    url: rawUrl,
                                    name: person.name,
                                    nameScore: person.score,
                                    company: company,
                                    text: cardText
                                }
                            );
                            continue;
                        }

                        if (
                            person.score >
                            (previous.nameScore ?? -10000)
                        ) {
                            previous.name = person.name;
                            previous.nameScore = person.score;
                        }

                        if (
                            cardText.length >
                            (previous.text || '').length
                        ) {
                            previous.text = cardText;
                        }

                        if (
                            !previous.company &&
                            company
                        ) {
                            previous.company = company;
                        }
                    }

                    return Array.from(results.values());
                }
                """
            )

        except Exception as exc:
            self.log(
                "CARD DOM EXTRACTION ERROR | "
                f"{type(exc).__name__}: {exc}"
            )
            cards = []

        elapsed = time.perf_counter() - started
        self.dom_calls += 1
        self.dom_seconds += elapsed

        results: dict[str, dict[str, str]] = {}

        for item in cards or []:
            href = str(item.get("url", "")).strip()
            name = str(item.get("name", "")).strip()
            company = str(item.get("company", "")).strip()
            card_text = self._clean_card_text(
                str(item.get("text", ""))
            )

            if not href or not card_text:
                continue

            if (
                "/sales/lead/" not in href
                and "/sales/people/" not in href
            ):
                continue

            key = normalize_url(href)

            if not key:
                continue

            candidate = {
                "url": href,
                "name": name,
                "company": company,
                "text": card_text[:MAX_CARD_TEXT_CHARS],
            }

            existing = results.get(key)

            if existing is None:
                results[key] = candidate
            else:
                if len(candidate["text"]) > len(existing["text"]):
                    existing["text"] = candidate["text"]

                if (
                    (
                        not existing.get("name")
                        or self._is_bad_name_candidate(
                            existing["name"]
                        )
                    )
                    and candidate.get("name")
                    and not self._is_bad_name_candidate(
                        candidate["name"]
                    )
                ):
                    existing["name"] = candidate["name"]

                if (
                    not existing.get("company")
                    and candidate.get("company")
                ):
                    existing["company"] = candidate["company"]

        return list(results.values())

    @staticmethod
    def _is_bad_name_candidate(value: str) -> bool:
        s = (value or "").strip().lower()

        if not s:
            return True

        bad_fragments = (
            "view profile",
            "degree connection",
            "linkedin premium",
            "show more",
            "select",
            "save",
            "message",
        )

        return any(fragment in s for fragment in bad_fragments)

    def _wait_for_new_loaded_result(self, previous_urls: list[str]) -> None:
        """
        After a scroll, return as soon as LinkedIn renders a lead URL that was
        not present in the previous visible batch.
        """
        try:
            self.page.wait_for_function(
                """
                (previous) => {
                    const oldSet = new Set(previous);

                    const anchors = document.querySelectorAll(
                        'a[href*="/sales/lead/"], a[href*="/sales/people/"]'
                    );

                    for (const a of anchors) {
                        const href = (a.href || '').trim();

                        if (href && !oldSet.has(href)) {
                            return true;
                        }
                    }

                    return false;
                }
                """,
                arg=previous_urls[-80:],
                timeout=NEW_RESULT_TIMEOUT_MS,
                polling=NEW_RESULT_POLL_MS,
            )
        except Exception:
            pass

    @staticmethod
    def _clean_person_name(display_name: str) -> str:
        name = re.sub(r"\s+", " ", display_name or "").strip()

        if not name:
            return ""

        # Accessibility/status text appended by Sales Navigator.
        name = re.sub(
            r"\s+is\s+reachable.*$",
            "",
            name,
            flags=re.IGNORECASE,
        ).strip()

        # Degree suffix: "Matt Decoskey · 3rd"
        name = re.sub(
            r"\s*[·•]\s*(?:1st|2nd|3rd|4th|5th|6th|7th|8th|9th)\b.*$",
            "",
            name,
            flags=re.IGNORECASE,
        ).strip()

        # Accessibility label: "View Matt Decoskey's profile"
        match = re.match(
            r"^view\s+(.+?)(?:'s|’s)\s+profile$",
            name,
            flags=re.IGNORECASE,
        )
        if match:
            name = match.group(1).strip()

        if name.lower() == "view profile":
            return ""

        # Remove common honorifics from First Name.
        name = re.sub(
            r"^(?:mr|mrs|ms|miss|dr|prof)\.?\s+",
            "",
            name,
            flags=re.IGNORECASE,
        ).strip()

        # Credentials after a comma are not part of Last Name:
        # "Alexa Jawlakian, M.S., LMFT" -> "Alexa Jawlakian"
        if "," in name:
            name = name.split(",", 1)[0].strip()

        return name

    @staticmethod
    def _split_name(display_name: str) -> tuple[str, str]:
        name = SearchWorker._clean_person_name(display_name)

        if not name:
            return "-", "-"

        parts = name.split()

        if len(parts) == 1:
            return parts[0], "-"

        # Keep compound/multi-part surnames together.
        return parts[0], " ".join(parts[1:])

    @staticmethod
    def _classify_job_level(title: str) -> str:
        t = (title or "").lower()
        rules = [
            (("chief ", "ceo", "cfo", "coo", "cto", "cio", "cmo", "cro"), "C-Level"),
            (("executive vice president", "evp"), "EVP"),
            (("senior vice president", "svp"), "SVP"),
            (("vice president", "vp"), "VP"),
            (("managing director",), "Managing Director"),
            (("director",), "Director"),
            (("head of", "head,"), "Head"),
            (("senior manager", "sr. manager", "sr manager"), "Senior Manager"),
            (("manager",), "Manager"),
            (("lead",), "Lead"),
            (("senior ", "sr. "), "Senior"),
        ]
        for needles, level in rules:
            if any(re.search(r"(?<!\w)" + re.escape(n.strip()) + r"(?!\w)", t) for n in needles):
                return level
        return "-"

    @staticmethod
    def _classify_job_function(title: str) -> str:
        t = (title or "").lower()
        mapping = [
            (("finance", "financial", "accounting", "accountant", "tax", "treasury", "controller"), "Finance / Accounting"),
            (("marketing", "brand", "growth"), "Marketing"),
            (("sales", "business development", "revenue"), "Sales / Business Development"),
            (("human resources", "hr ", "talent", "people"), "Human Resources"),
            (("information technology", "chief information officer", "cio", " it ", "technology", "software", "engineering", "developer"), "IT / Engineering"),
            (("procurement", "purchasing", "sourcing", "supply chain"), "Procurement / Supply Chain"),
            ((
                "operations",
                "manufacturing",
                "operational",
                "operating officer",
                "chief operating officer",
                " coo ",
            ), "Operations"),
            (("legal", "counsel", "attorney"), "Legal"),
            (("security", "cyber"), "Security"),
        ]
        padded = f" {t} "
        for needles, function in mapping:
            if any(re.search(r"(?<!\w)" + re.escape(n.strip()) + r"(?!\w)", padded) for n in needles):
                return function
        return "-"

    @staticmethod
    def _looks_like_control_line(line: str) -> bool:
        l = re.sub(r"\s+", " ", line or "").strip().lower()

        if not l:
            return True

        exact_controls = {
            "save",
            "saved",
            "view profile",
            "show more",
            "message",
            "select",
            "select all",
            "unsave",
            "viewed",
        }

        if l in exact_controls:
            return True

        # These are relationship/status lines, not job titles/locations.
        if "degree connection" in l:
            return True

        if "linkedin premium member" in l:
            return True

        # Sales Navigator can render the profile state as its own line:
        #     Viewed
        # or alongside relationship metadata:
        #     · Viewed
        # This is UI/status text and must never become Job Title.
        if re.fullmatch(r"[·•]?\s*viewed\s*", l, flags=re.IGNORECASE):
            return True

        if re.fullmatch(r"[·•]?\s*(?:1st|2nd|3rd|4th|5th|6th|7th|8th|9th)", l):
            return True

        return False

    @staticmethod
    def _looks_like_country_token(value: str) -> bool:
        token = re.sub(
            r"\s+",
            " ",
            (value or "").strip(" ."),
        ).lower()

        countries = {
            "united states", "usa", "u.s.a", "u.s.", "us",
            "united kingdom", "uk", "u.k.", "great britain",
            "canada", "india", "singapore", "australia",
            "new zealand", "ireland", "norway", "malaysia",
            "czech republic", "czechia", "poland", "turkey",
            "mexico", "netherlands", "france", "italy",
            "germany", "denmark", "china", "japan", "sweden",
            "austria", "indonesia", "saudi arabia", "south africa",
            "hong kong", "switzerland", "spain", "iceland",
            "argentina", "chile", "peru", "uruguay", "finland",
            "romania", "philippines", "thailand", "vietnam",
            "united arab emirates", "uae", "brazil", "colombia",
            "panama", "portugal", "belgium", "luxembourg",
            "israel", "south korea", "korea", "taiwan",
            "pakistan", "bangladesh", "sri lanka", "nepal",
            "egypt", "kenya", "nigeria", "ghana", "morocco",
            "qatar", "kuwait", "bahrain", "oman", "jordan",
            "greece", "hungary", "croatia", "serbia", "slovakia",
            "slovenia", "bulgaria", "estonia", "latvia",
            "lithuania", "ukraine", "cyprus", "malta",
        }

        return token in countries

    @staticmethod
    def _looks_like_revenue_line(value: str) -> bool:
        s = re.sub(r"\s+", " ", value or "").strip()

        if not s:
            return False

        return bool(
            re.search(
                r"(?i)\brevenue\b",
                s,
            )
            or re.search(
                r"(?i)(?:[$€£¥₹]|USD|EUR|GBP|JPY|INR)\s*"
                r"\d[\d,.]*\s*(?:K|M|B|T|million|billion)?",
                s,
            )
        )

    @staticmethod
    def _looks_like_industry_line(value: str) -> bool:
        s = re.sub(r"\s+", " ", value or "").strip().lower()

        if not s:
            return False

        industry_terms = (
            "services",
            "consulting",
            "software",
            "technology",
            "information",
            "internet",
            "management",
            "manufacturing",
            "retail",
            "wholesale",
            "education",
            "healthcare",
            "medical",
            "financial",
            "banking",
            "insurance",
            "oil and gas",
            "telecommunications",
            "construction",
            "real estate",
            "hospitality",
            "food",
            "beverage",
            "human resources",
            "marketing",
            "advertising",
            "media",
            "transportation",
            "logistics",
            "supply chain",
            "consumer",
            "electronics",
            "biotechnology",
            "pharmaceutical",
            "research",
            "legal",
            "law practice",
            "security",
            "outsourcing",
            "training",
            "coaching",
            "investment",
        )

        return any(term in s for term in industry_terms)

    @staticmethod
    def _looks_like_state_region_token(value: str) -> bool:
        """
        Conservative region/state detector used only to prevent values such as
        "Maharashtra, India" from becoming City="Maharashtra".

        The list covers common markets used by the scraper, while suffix rules
        handle many other international region labels.
        """
        token = re.sub(
            r"\s+",
            " ",
            (value or "").strip(" ."),
        ).lower()

        if not token:
            return False

        suffixes = (
            " state",
            " province",
            " region",
            " prefecture",
            " emirate",
            " county",
            " territory",
            " governorate",
            " department",
            " oblast",
            " canton",
        )

        if any(token.endswith(suffix) for suffix in suffixes):
            return True

        known_regions = {
            # India
            "andhra pradesh",
            "arunachal pradesh",
            "assam",
            "bihar",
            "chhattisgarh",
            "goa",
            "gujarat",
            "haryana",
            "himachal pradesh",
            "jharkhand",
            "karnataka",
            "kerala",
            "madhya pradesh",
            "maharashtra",
            "manipur",
            "meghalaya",
            "mizoram",
            "nagaland",
            "odisha",
            "punjab",
            "rajasthan",
            "sikkim",
            "tamil nadu",
            "telangana",
            "tripura",
            "uttar pradesh",
            "uttarakhand",
            "west bengal",
            "delhi",
            "jammu and kashmir",
            "ladakh",
            "puducherry",
            # United States
            "alabama",
            "alaska",
            "arizona",
            "arkansas",
            "california",
            "colorado",
            "connecticut",
            "delaware",
            "florida",
            "georgia",
            "hawaii",
            "idaho",
            "illinois",
            "indiana",
            "iowa",
            "kansas",
            "kentucky",
            "louisiana",
            "maine",
            "maryland",
            "massachusetts",
            "michigan",
            "minnesota",
            "mississippi",
            "missouri",
            "montana",
            "nebraska",
            "nevada",
            "new hampshire",
            "new jersey",
            "new mexico",
            "new york",
            "north carolina",
            "north dakota",
            "ohio",
            "oklahoma",
            "oregon",
            "pennsylvania",
            "rhode island",
            "south carolina",
            "south dakota",
            "tennessee",
            "texas",
            "utah",
            "vermont",
            "virginia",
            "washington",
            "west virginia",
            "wisconsin",
            "wyoming",
            # Canada
            "alberta",
            "british columbia",
            "manitoba",
            "new brunswick",
            "newfoundland and labrador",
            "nova scotia",
            "ontario",
            "prince edward island",
            "quebec",
            "saskatchewan",
            # Australia
            "new south wales",
            "queensland",
            "south australia",
            "tasmania",
            "victoria",
            "western australia",
            "australian capital territory",
            "northern territory",
            # Common UK profile regions
            "england",
            "scotland",
            "wales",
            "northern ireland",
        }

        # LinkedIn frequently abbreviates US states in Experience locations
        # (for example ``Greenville, NC``). Treat the standard postal codes as
        # regions so the city/state split remains usable; the country can then
        # be filled from the profile location when it is omitted here.
        us_state_codes = {
            "al", "ak", "az", "ar", "ca", "co", "ct", "de", "fl", "ga",
            "hi", "id", "il", "in", "ia", "ks", "ky", "la", "me", "md",
            "ma", "mi", "mn", "ms", "mo", "mt", "ne", "nv", "nh", "nj",
            "nm", "ny", "nc", "nd", "oh", "ok", "or", "pa", "ri", "sc",
            "sd", "tn", "tx", "ut", "vt", "va", "wa", "wv", "wi", "wy",
            "dc",
        }
        return token in known_regions or token in us_state_codes

    @staticmethod
    def _profile_location_to_fields(
        location: str,
    ) -> tuple[str, str, str]:
        """
        Parse a profile/search-card location into City / State / Country.

        Priority is positional only when the final component is a recognized
        country. This prevents industry text and region-only text from being
        mapped into the wrong geography column.

        Examples:
            Pune Division, Maharashtra, India
                -> Pune Division / Maharashtra / India

            Maharashtra, India
                -> - / Maharashtra / India

            Pune, India
                -> Pune / - / India

            India
                -> - / - / India
        """
        value = SearchWorker._normalize_ws(location)

        if not value or value == "-":
            return "-", "-", "-"

        parts = [
            SearchWorker._normalize_ws(part).strip(" .")
            for part in value.split(",")
            if SearchWorker._normalize_ws(part).strip(" .")
        ]

        if not parts:
            return "-", "-", "-"

        # One-component values need semantic classification.
        if len(parts) == 1:
            only = parts[0]

            if SearchWorker._looks_like_country_token(only):
                return (
                    "-",
                    "-",
                    SearchWorker._canonical_country(only),
                )

            if SearchWorker._looks_like_state_region_token(only):
                return "-", only, "-"

            # Metro/Division/Area labels shown by LinkedIn are valid City-like
            # profile location values when no separate state/country exists.
            if re.search(
                r"(?i)\b(area|division|metro|metropolitan|city)\b",
                only,
            ):
                return only, "-", "-"

            return only, "-", "-"

        last = parts[-1]

        # Never assume a final component is Country unless it looks like one.
        if not SearchWorker._looks_like_country_token(last):
            # A two-part value such as "Pune Division, Maharashtra" contains
            # City + State but no Country.
            if len(parts) == 2:
                first, second = parts

                if SearchWorker._looks_like_state_region_token(second):
                    return first, second, "-"

                if SearchWorker._looks_like_state_region_token(first):
                    return "-", first, "-"

            return "-", "-", "-"

        country = SearchWorker._canonical_country(last)

        if len(parts) >= 3:
            # LinkedIn profile locations normally use:
            # City/Area, State/Region, Country.
            city = parts[0]
            state = parts[-2]

            return (
                city or "-",
                state or "-",
                country or "-",
            )

        # Exactly two parts: determine whether the first token is a state/region
        # or a city so State is not incorrectly written into City.
        first = parts[0]

        if SearchWorker._looks_like_state_region_token(first):
            return "-", first, country

        return first, "-", country

    @staticmethod
    def _fill_missing_location_fields(
        record: dict[str, str],
        source_fields: dict[str, str],
    ) -> list[str]:
        """
        Field-level location fallback.

        Existing valid City / State / Country values are never replaced.
        Only missing fields are copied from source_fields.
        """
        filled: list[str] = []

        for field in ("City", "State", "Country"):
            existing = clean_cell_value(
                record.get(field, "-")
            )
            proposed = clean_cell_value(
                source_fields.get(field, "-")
            )

            if not SearchWorker._is_unknown_value(existing):
                continue

            if SearchWorker._is_unknown_value(proposed):
                continue

            if field == "Country":
                proposed = SearchWorker._canonical_country(
                    proposed
                )

            record[field] = proposed
            filled.append(field)

        return filled

    @staticmethod
    def _is_location_candidate(line: str) -> bool:
        l = (line or "").strip()

        if not l or SearchWorker._looks_like_control_line(l):
            return False

        low = l.lower()

        if (
            " in role" in low
            or " in company" in low
            or low.startswith("about:")
            or "is reachable" in low
            or SearchWorker._looks_like_revenue_line(l)
        ):
            return False

        parts = [
            p.strip()
            for p in l.split(",")
            if p.strip()
        ]

        if len(parts) < 2 or len(l) > 160:
            return False

        # The final component should look like a country. This prevents
        # LinkedIn industry strings such as:
        #   "Technology, Information and Internet"
        # from becoming City="Technology" / Country="Information and Internet".
        return (
            SearchWorker._looks_like_country_token(parts[-1])
            or (
                len(parts) == 2
                and SearchWorker._looks_like_state_region_token(parts[-1])
            )
        )


    def _parse_card_fast(
        self,
        card: dict[str, str],
        extraction_mode: str = "FAST",
    ) -> dict[str, str]:
        """
        Accuracy-first deterministic parser.

        Fixes observed in the supplied output:
        - "View profile" is no longer treated as a person's name.
        - "Matt Decoskey is reachable" -> First Name=Matt, Last Name=Decoskey.
        - "2nd/3rd degree connection" is not treated as Job Title.
        - "· 2nd" is not treated as Country.
        - Company is read from the company's Sales Navigator link when present.
        - "Viewed" and similar UI/status labels can never become Job Title.
        - Explicit "Job Title · Company" text is preferred over generic fallback lines.
        """
        lead_url = card["url"]
        raw_text = self._clean_card_text(card.get("text", ""))
        display_name = self._clean_person_name(card.get("name", ""))
        dom_company = re.sub(
            r"\s+",
            " ",
            card.get("company", "") or "",
        ).strip()

        raw_lines = []
        for line in raw_text.splitlines():
            line = re.sub(r"\s+", " ", line).strip()

            if not line:
                continue

            raw_lines.append(line)

        # Preserve useful lines, but remove obvious UI/control-only lines.
        lines = []
        seen_lines = set()

        for line in raw_lines:
            if self._looks_like_control_line(line):
                continue

            if line not in seen_lines:
                seen_lines.add(line)
                lines.append(line)

        # ------------------------------------------------------
        # NAME
        # ------------------------------------------------------
        if not display_name or self._is_bad_name_candidate(display_name):
            # Strongest fallback: a line explicitly carrying "is reachable".
            for line in raw_lines:
                if re.search(r"\bis reachable\b", line, flags=re.IGNORECASE):
                    candidate = self._clean_person_name(line)

                    if candidate and not self._is_bad_name_candidate(candidate):
                        display_name = candidate
                        break

        if not display_name or self._is_bad_name_candidate(display_name):
            # Second fallback: "Name · 2nd/3rd".
            for line in raw_lines:
                match = re.match(
                    r"^(.+?)\s*[·•]\s*(?:1st|2nd|3rd|4th|5th|6th|7th|8th|9th)\b",
                    line,
                    flags=re.IGNORECASE,
                )

                if match:
                    candidate = self._clean_person_name(match.group(1))

                    if candidate and not self._is_bad_name_candidate(candidate):
                        display_name = candidate
                        break

        first_name, last_name = self._split_name(display_name)

        # ------------------------------------------------------
        # TENURE / DESCRIPTION
        # ------------------------------------------------------
        tenure = "-"
        description = "-"

        for line in raw_lines:
            lower = line.lower()

            if " in role" in lower or " in company" in lower:
                tenure = line

            if lower.startswith("about:"):
                description = re.sub(
                    r"^about:\s*",
                    "",
                    line,
                    flags=re.IGNORECASE,
                ).strip() or "-"

        # ------------------------------------------------------
        # COMPANY
        # ------------------------------------------------------
        company = dom_company or "-"

        # ------------------------------------------------------
        # JOB TITLE
        # ------------------------------------------------------
        title = "-"

        # Identify where the person's name appears so we can prioritize
        # the text immediately below it.
        name_idx = -1
        cleaned_name_for_compare = self._clean_person_name(display_name).lower()

        for i, line in enumerate(raw_lines):
            cleaned_line = self._clean_person_name(line).lower()

            if (
                cleaned_name_for_compare
                and cleaned_line
                and cleaned_line == cleaned_name_for_compare
            ):
                name_idx = i
                break

            if (
                cleaned_name_for_compare
                and cleaned_name_for_compare in cleaned_line
                and "is reachable" in line.lower()
            ):
                name_idx = i
                break

        search_lines = (
            raw_lines[name_idx + 1 : name_idx + 12]
            if name_idx >= 0
            else raw_lines[:16]
        )

        def valid_title_candidate(line: str) -> bool:
            lower = line.lower()

            if self._looks_like_control_line(line):
                return False

            if (
                " in role" in lower
                or " in company" in lower
                or lower.startswith("about:")
                or "is reachable" in lower
            ):
                return False

            if self._is_location_candidate(line):
                return False

            # Extra defensive UI/status filters.
            if re.fullmatch(
                r"(?:viewed|recently hired|shared groups?|shared education|"
                r"\d+\s+mutual connections?)",
                lower.strip(),
                flags=re.IGNORECASE,
            ):
                return False

            return True

        # ======================================================
        # PASS 1 - DEFINITIVE TITLE + COMPANY PAIR
        # ======================================================
        # Prefer a line such as:
        #
        #   Professor · University of Northern Iowa
        #   Intern · Josh Software, Inc.
        #
        # This is much stronger evidence than taking the first
        # text line below the prospect name.
        normalized_company = self._normalize_company_for_match(company)

        for line in search_lines:
            if not valid_title_candidate(line):
                continue

            # Support both middle-dot characters and minor spacing changes.
            match = re.match(
                r"^(.+?)\s*[·•]\s*(.+)$",
                line,
            )

            if not match:
                continue

            left = self._normalize_ws(match.group(1))
            right = self._normalize_ws(match.group(2))

            if not left or not right:
                continue

            # Ignore relationship text such as:
            #   Syed Kirmani · 3rd
            if re.fullmatch(
                r"(?:1st|2nd|3rd|4th|5th|6th|7th|8th|9th).*",
                right,
                flags=re.IGNORECASE,
            ):
                continue

            # Ignore status/control text on the right.
            if (
                "viewed" in right.lower()
                or "degree connection" in right.lower()
                or "linkedin premium" in right.lower()
            ):
                continue

            right_company = self._normalize_company_for_match(right)

            # If the DOM company link was captured, require the right side
            # of the title line to refer to that same company.
            if normalized_company:
                if (
                    right_company
                    and (
                        right_company == normalized_company
                        or right_company in normalized_company
                        or normalized_company in right_company
                    )
                ):
                    title = left
                    break
            else:
                # Company DOM link was unavailable, so use this line to
                # populate both title and company.
                title = left
                company = right
                normalized_company = self._normalize_company_for_match(company)
                break

        # ======================================================
        # PASS 1B - FLATTENED TITLE + KNOWN COMPANY
        # ======================================================
        # Some DOM layouts flatten:
        #   "Professor University of Northern Iowa"
        # rather than preserving the middle dot.
        if title == "-" and company != "-":
            for line in search_lines:
                if not valid_title_candidate(line):
                    continue

                if line.lower().endswith(company.lower()):
                    candidate = line[: len(line) - len(company)].strip(" ·|-")

                    if (
                        candidate
                        and not self._looks_like_control_line(candidate)
                    ):
                        title = candidate
                        break

        # ======================================================
        # PASS 2 - GENERIC FALLBACK
        # ======================================================
        # Only use the first plausible text line if no explicit
        # Title · Company relationship could be found.
        if title == "-":
            for line in search_lines:
                if not valid_title_candidate(line):
                    continue

                # Do not accept a bare company name as the title.
                if (
                    company != "-"
                    and self._company_names_match(line, company)
                ):
                    continue

                if 2 <= len(line) <= 220:
                    title = line
                    break

        # ------------------------------------------------------
        # LOCATION
        # ------------------------------------------------------
        location = "-"

        for line in raw_lines:
            if self._is_location_candidate(line):
                location = line
                break

        city, state, country = self._profile_location_to_fields(
            location
        )

        result = {
            "Tenure": tenure,
            "First Name": first_name,
            "Last Name": last_name,
            "Job Title": title,
            "Job Level": self._classify_job_level(title),
            "Job Function": self._classify_job_function(title),
            "Company Name": company,
            "Industry": "-",
            "City": city,
            "State": state,
            "Country": country,
            "Description": description,
        }

        # SMART mode: Qwen is only a fallback for genuinely missing
        # core fields. Existing deterministic values are not overwritten.
        needs_llm = (
            result["First Name"] == "-"
            or result["Last Name"] == "-"
            or result["Job Title"] == "-"
            or result["Company Name"] == "-"
        )

        if (
            extraction_mode.upper() == "SMART"
            and USE_LLM_FALLBACK
            and self.llm_available is not False
            and needs_llm
        ):
            try:
                qwen_started = time.perf_counter()

                llm_result = self.llm.extract_search_card(
                    card_text=raw_text,
                    lead_url=lead_url,
                )

                self.qwen_calls += 1
                self.qwen_seconds += time.perf_counter() - qwen_started

                # Clean LLM name fields as well.
                if llm_result.get("First Name") != "-":
                    llm_result["First Name"] = self._clean_person_name(
                        llm_result["First Name"]
                    ) or "-"

                if llm_result.get("Last Name") != "-":
                    llm_result["Last Name"] = re.sub(
                        r"\s+is\s+reachable.*$",
                        "",
                        llm_result["Last Name"],
                        flags=re.IGNORECASE,
                    ).strip() or "-"

                for key, value in llm_result.items():
                    if (
                        key in result
                        and result[key] == "-"
                        and value != "-"
                    ):
                        result[key] = value

            except Exception as exc:
                if isinstance(exc, requests.RequestException):
                    self.llm_available = False
                    self.llm_last_check_ts = time.time()

                self.log(
                    f"Qwen fallback skipped for one card: {exc}"
                )

        return result

    @staticmethod
    def _normalize_ws(value: str) -> str:
        return re.sub(r"\s+", " ", value or "").strip()

    @staticmethod
    def _location_to_fields(
        preferred_location: str,
        fallback_location: str = "-",
    ) -> tuple[str, str, str]:
        """
        Map the preferred location first.

        Examples:
        - "United States" -> Country
        - "St Augustine, Florida, United States"
              -> City / State / Country
        - "Raleigh-Durham-Chapel Hill Area"
              -> City-like area value; missing Country can be filled from the
                 top-profile fallback location.
        """
        preferred = SearchWorker._normalize_ws(preferred_location)
        fallback = SearchWorker._normalize_ws(fallback_location)

        city = state = country = "-"

        def parse(value: str):
            if not value or value == "-":
                return "-", "-", "-"

            parts = [p.strip() for p in value.split(",") if p.strip()]

            if len(parts) >= 3:
                return parts[0], parts[1], parts[-1]

            if len(parts) == 2:
                if SearchWorker._looks_like_country_token(parts[1]):
                    return parts[0], "-", SearchWorker._canonical_country(parts[1])
                if SearchWorker._looks_like_state_region_token(parts[1]):
                    return parts[0], parts[1], "-"
                return parts[0], "-", parts[1]

            # A single value can be either a country or a metro/area.
            if value.lower().endswith(" area"):
                return value, "-", "-"

            return "-", "-", value

        city, state, country = parse(preferred)

        f_city, f_state, f_country = parse(fallback)

        if city == "-":
            city = f_city
        if state == "-":
            state = f_state
        if country == "-":
            country = f_country

        return city, state, country

    def _open_lead_panel(self, lead_url: str, person_name: str) -> None:
        """
        Click the prospect name/link so Sales Navigator opens the right-side panel.
        """
        target_path = urlsplit(lead_url).path.rstrip("/")

        loc = self.page.locator(f'a[href*="{target_path}"]')

        # Prefer the name-labelled lead link rather than photo/control links.
        if person_name:
            named = loc.filter(has_text=person_name)
            try:
                if named.count() > 0:
                    named.first.click(timeout=5000)
                else:
                    loc.first.click(timeout=5000)
            except Exception:
                loc.first.click(timeout=5000)
        else:
            loc.first.click(timeout=5000)

        deadline = time.time() + PROFILE_PANEL_OPEN_TIMEOUT_MS / 1000.0

        while time.time() < deadline and not self.stop_event.is_set():
            try:
                found = self.page.evaluate(
                    """
                    (personName) => {
                        const vw = window.innerWidth;

                        return Array.from(
                            document.querySelectorAll('aside, section, div')
                        ).some(el => {
                            const r = el.getBoundingClientRect();
                            const txt = (el.innerText || '').trim();

                            if (
                                r.width < 320 ||
                                r.height < 220 ||
                                r.left < vw * 0.55 ||
                                r.right < vw * 0.88 ||
                                !txt
                            ) return false;

                            return personName
                                ? txt.toLowerCase().includes(personName.toLowerCase())
                                : true;
                        });
                    }
                    """,
                    person_name or "",
                )

                if found:
                    return

            except Exception:
                pass

            time.sleep(0.18)

        raise RuntimeError("The prospect side panel did not open in time.")

    @timed('linkedin_profile_panel')
    def _capture_profile_panel(self, person_name: str) -> dict[str, Any]:
        """
        Scroll the right-side profile panel and aggregate its text and links.

        We do not depend on one LinkedIn class name because those change often.
        """
        chunks: list[str] = []
        links_by_href: dict[str, dict[str, str]] = {}

        previous_scroll_top = None
        stagnant = 0

        for _ in range(PROFILE_PANEL_MAX_SCROLLS):
            if self.stop_event.is_set():
                break

            if not self._wait_if_paused():
                break

            snapshot = self.page.evaluate(
                """
                (args) => {
                    const personName = args.personName || '';
                    const step = args.step || 850;
                    const vw = window.innerWidth;
                    const candidates = [];

                    for (const el of Array.from(document.querySelectorAll('aside, section, div'))) {
                        const r = el.getBoundingClientRect();
                        if (
                            r.width < 320 ||
                            r.height < 220 ||
                            r.left < vw * 0.55 ||
                            r.right < vw * 0.88
                        ) continue;

                        const txt = (el.innerText || '').trim();
                        if (!txt || txt.length < 80) continue;

                        const style = getComputedStyle(el);
                        const delta = Math.max(0, el.scrollHeight - el.clientHeight);
                        const hasName = personName
                            ? txt.toLowerCase().includes(personName.toLowerCase())
                            : false;

                        let score = 0;
                        if (hasName) score += 250;
                        if (delta > 120) score += 120;
                        if (style.overflowY === 'auto' || style.overflowY === 'scroll') score += 60;
                        if (r.left > vw * 0.62) score += 40;
                        score += Math.min(delta / 20, 100);

                        candidates.push({el, score});
                    }

                    candidates.sort((a, b) => b.score - a.score);
                    const target = candidates[0]?.el;

                    if (!target) {
                        return {
                            text: '',
                            links: [],
                            scrollTop: 0,
                            scrollHeight: 0,
                            clientHeight: 0
                        };
                    }

                    const result = {
                        text: (target.innerText || '').trim(),
                        links: Array.from(target.querySelectorAll('a[href]'))
                            .slice(0, 600)
                            .map(a => ({
                                text: (a.innerText || a.textContent || '').replace(/\\s+/g, ' ').trim(),
                                href: a.href || ''
                            })),
                        scrollTop: target.scrollTop,
                        scrollHeight: target.scrollHeight,
                        clientHeight: target.clientHeight
                    };

                    target.scrollTop = Math.min(
                        target.scrollHeight,
                        target.scrollTop + step
                    );

                    return result;
                }
                """,
                {
                    "personName": person_name or "",
                    "step": PROFILE_PANEL_SCROLL_STEP,
                },
            )

            chunk = self._clean_card_text(str(snapshot.get("text", "")))

            if chunk and chunk not in chunks:
                chunks.append(chunk)

            for item in snapshot.get("links", []) or []:
                href = str(item.get("href", "")).strip()
                label = self._normalize_ws(str(item.get("text", "")))

                if href:
                    links_by_href[href] = {
                        "text": label,
                        "href": href,
                    }

            scroll_top = float(snapshot.get("scrollTop", 0) or 0)
            scroll_height = float(snapshot.get("scrollHeight", 0) or 0)
            client_height = float(snapshot.get("clientHeight", 0) or 0)

            if previous_scroll_top == scroll_top:
                stagnant += 1
            else:
                stagnant = 0

            previous_scroll_top = scroll_top

            if stagnant >= 2:
                break

            if (
                scroll_height > 0
                and scroll_top + client_height >= scroll_height - 15
            ):
                break

            self.page.wait_for_timeout(PROFILE_PANEL_SCROLL_WAIT_MS)

        # Aggregate unique lines from every captured scroll position.
        unique_lines = []
        seen = set()

        for chunk in chunks:
            for line in chunk.splitlines():
                line = self._normalize_ws(line)

                if line and line not in seen:
                    seen.add(line)
                    unique_lines.append(line)

        return {
            "text": "\n".join(unique_lines),
            "links": list(links_by_href.values()),
        }

    @staticmethod
    def _extract_top_profile_location(
        panel_text: str,
        person_name: str,
    ) -> str:
        """
        Extract the location shown at the top of the prospect panel.
        """
        lines = [
            SearchWorker._normalize_ws(x)
            for x in (panel_text or "").splitlines()
            if SearchWorker._normalize_ws(x)
        ]

        # Examine only the first portion of the profile.
        for line in lines[:30]:
            low = line.lower()

            if person_name and line.lower() == person_name.lower():
                continue

            if (
                "connection" in low
                or low.startswith("viewed:")
                or " yrs " in low
                or " mos" in low
                or "present" in low
                or line.startswith("$")
            ):
                continue

            # Structured profile location. Reuse the same safeguards as the
            # search-card parser so title/industry text cannot become location.
            if (
                "," in line
                and len(line) <= 140
                and SearchWorker._is_location_candidate(line)
            ):
                return line

        # Single-country fallback, e.g. "United States".
        country_like = {
            "united states",
            "united kingdom",
            "canada",
            "india",
            "singapore",
            "australia",
            "france",
            "germany",
            "italy",
            "spain",
            "netherlands",
            "ireland",
            "malaysia",
            "philippines",
            "indonesia",
            "thailand",
            "vietnam",
            "japan",
            "china",
            "hong kong",
            "switzerland",
            "sweden",
            "denmark",
            "norway",
            "finland",
            "south africa",
            "united arab emirates",
            "saudi arabia",
        }

        for line in lines[:35]:
            if line.lower() in country_like:
                return line

        return "-"

    @staticmethod
    def _is_duration_only_line(value: str) -> bool:
        """
        Examples that are company-level duration summaries rather than titles:
            7 yrs 2 mos
            5 yrs 3 mos
            1 yr 1 mo
        """
        s = SearchWorker._normalize_ws(value).lower()

        if not s:
            return False

        return bool(
            re.fullmatch(
                r"(?:\d+\s*yrs?\s*)?(?:\d+\s*mos?\s*)?(?:\d+\s*yrs?)?",
                s,
                flags=re.IGNORECASE,
            )
        )

    @staticmethod
    def _is_experience_noise_line(value: str) -> bool:
        s = SearchWorker._normalize_ws(value)
        low = s.lower()

        if not s:
            return True

        if SearchWorker._looks_like_control_line(s):
            return True

        if SearchWorker._is_duration_only_line(s):
            return True

        noise_prefixes = (
            "summarized by ai",
            "account insights",
            "strategic priorities",
            "business challenges",
            "view relationship map",
            "lead iq",
            "sources:",
            "was this helpful",
        )

        if any(low.startswith(prefix) for prefix in noise_prefixes):
            return True

        return False

    @staticmethod
    def _parse_experience_blocks(
        panel_text: str,
        panel_links: list[dict[str, str]],
    ) -> list[dict[str, str]]:
        """
        VERSION 11 CURRENT-ROLE RULE

        Search the captured side-panel text for the literal keyword "Present".

        Example:
            Professor
            University of Northern Iowa
            1984–Present · 42 yrs
            Cedar Falls, Iowa, U.S.A.

        Result:
            Job Title            = Professor
            Company Name         = University of Northern Iowa
            Tenure               = 1984–Present · 42 yrs
            Location             = Cedar Falls, Iowa, U.S.A.
            Company Profile Link = /sales/company/... link

        The parser intentionally does not require a month before the year.
        """
        lines = [
            SearchWorker._normalize_ws(x)
            for x in (panel_text or "").splitlines()
            if SearchWorker._normalize_ws(x)
        ]

        company_link_map: dict[str, str] = {}

        for item in panel_links or []:
            href = str(item.get("href", "")).strip()
            label = SearchWorker._normalize_ws(str(item.get("text", "")))

            if (
                href
                and label
                and (
                    "/sales/company/" in href
                    or "/sales/account/" in href
                )
            ):
                company_link_map[label.lower()] = href

        present_re = re.compile(r"\bpresent\b", flags=re.IGNORECASE)

        date_range_re = re.compile(
            r"(?i)"
            r"(?:"
            r"(?:jan|feb|mar|apr|may|jun|jul|aug|sep|sept|oct|nov|dec)\s+"
            r")?"
            r"\d{4}\s*[–—-]\s*"
            r"(?:present|"
            r"(?:(?:jan|feb|mar|apr|may|jun|jul|aug|sep|sept|oct|nov|dec)\s+)?\d{4})"
        )

        blocks: list[dict[str, str]] = []

        for i, line in enumerate(lines):
            if not present_re.search(line):
                continue

            if SearchWorker._is_experience_noise_line(line):
                continue

            tenure = line

            # Find the nearest linked company around/above the Present line.
            company = "-"
            company_url = "-"
            company_idx = -1

            for idx in range(i - 1, max(-1, i - 12), -1):
                candidate = lines[idx]
                href = company_link_map.get(candidate.lower())

                if href:
                    company = candidate
                    company_url = href
                    company_idx = idx
                    break

            # Find the nearest meaningful title above Present, excluding company.
            title = "-"
            title_idx = -1

            for idx in range(i - 1, max(-1, i - 10), -1):
                candidate = lines[idx]

                if idx == company_idx:
                    continue
                if candidate.lower() in company_link_map:
                    continue
                if SearchWorker._is_experience_noise_line(candidate):
                    continue
                if date_range_re.search(candidate):
                    continue
                if "experience" in candidate.lower():
                    continue

                if 2 <= len(candidate) <= 240:
                    title = candidate
                    title_idx = idx
                    break

            # Fallback if company is text-only, not a company hyperlink.
            if company == "-":
                # Common layout: Title -> Company -> Present.
                if title_idx >= 0:
                    for idx in range(title_idx + 1, i):
                        candidate = lines[idx]

                        if SearchWorker._is_experience_noise_line(candidate):
                            continue
                        if date_range_re.search(candidate):
                            continue
                        if candidate == title:
                            continue
                        if "experience" in candidate.lower():
                            continue

                        if 1 < len(candidate) <= 180:
                            company = candidate
                            company_idx = idx
                            break

                # Grouped-company layout fallback.
                if company == "-":
                    for idx in range(i - 1, max(-1, i - 10), -1):
                        candidate = lines[idx]

                        if idx == title_idx:
                            continue
                        if SearchWorker._is_experience_noise_line(candidate):
                            continue
                        if date_range_re.search(candidate):
                            continue
                        if "experience" in candidate.lower():
                            continue

                        if 1 < len(candidate) <= 180:
                            company = candidate
                            company_idx = idx
                            break

            # Closest location after the Present line.
            location = "-"

            for idx in range(i + 1, min(len(lines), i + 7)):
                candidate = lines[idx]

                if SearchWorker._is_experience_noise_line(candidate):
                    continue
                if present_re.search(candidate):
                    break
                if date_range_re.search(candidate):
                    break

                if SearchWorker._is_location_candidate(candidate):
                    location = candidate
                    break

                if (
                    location == "-"
                    and 2 <= len(candidate) <= 180
                    and candidate.lower() not in company_link_map
                ):
                    location = candidate

            blocks.append(
                {
                    "Job Title": title or "-",
                    "Company Name": company or "-",
                    "Tenure": tenure or "-",
                    "Location": location or "-",
                    "Company Profile Link": company_url or "-",
                    "Is Current": "Present",
                }
            )

        return blocks

    @staticmethod
    def _select_current_experience(
        blocks: list[dict[str, str]],
    ) -> dict[str, str]:
        """
        Strict current-role rule.

        If at least one parsed block contains Present:
            use the first/current block.

        If NO Present block can be found:
            Tenure = Dead
            Job Title / Company / Location = "-"
            Move on to the next prospect without using a past role as current.
        """
        for block in blocks:
            if "present" in block.get("Tenure", "").lower():
                return block

        return {
            "Job Title": "-",
            "Company Name": "-",
            "Tenure": "Dead",
            "Location": "-",
            "Company Profile Link": "-",
            "Is Current": "Dead",
        }

    def _close_lead_panel(self) -> None:
        """
        Close the right-side Sales Navigator lead panel after one prospect.

        LinkedIn changes labels/classes, so this tries a few conservative
        approaches. Failure to close is non-fatal.
        """
        try:
            clicked = self.page.evaluate(
                """
                () => {
                    const vw = window.innerWidth;

                    const buttons = Array.from(
                        document.querySelectorAll('button')
                    ).filter(btn => {
                        const r = btn.getBoundingClientRect();
                        if (r.right < vw * 0.85 || r.top > 180) return false;

                        const label = (
                            btn.getAttribute('aria-label') ||
                            btn.getAttribute('title') ||
                            btn.innerText ||
                            ''
                        ).trim().toLowerCase();

                        return (
                            label === 'close' ||
                            label.includes('close') ||
                            label === 'dismiss'
                        );
                    });

                    if (buttons[0]) {
                        buttons[0].click();
                        return true;
                    }

                    return false;
                }
                """
            )

            if not clicked:
                # Escape often closes the right-side detail panel.
                self.page.keyboard.press("Escape")

            self.page.wait_for_timeout(200)

        except Exception:
            pass

    @staticmethod
    def _normalize_company_for_match(value: str) -> str:
        """
        Normalize company names for comparison without trying to invent aliases.

        Examples:
            "Deloitte Tax LLP" -> "deloitte tax"
            "Deloitte Tax, LLP" -> "deloitte tax"
            "EY" -> "ey"

        This intentionally does NOT claim that different brand names are the same
        (for example EY vs Ernst & Young) unless the visible text actually matches.
        """
        value = SearchWorker._normalize_ws(value)

        if not value or value == "-":
            return ""

        value = value.lower()

        # Replace punctuation with spaces.
        value = re.sub(r"[^a-z0-9&]+", " ", value)

        # Remove common legal entity suffixes only from the end.
        value = re.sub(
            r"\b(?:llp|llc|ltd|limited|inc|incorporated|corp|corporation|plc|pvt|private|co|company|sa|srl|sl|pty|pte|bv|gmbh|ag|nv|spa|sas|oy|ab|as|kk|sro|sarl|pbc)\b\.?$",
            "",
            value,
            flags=re.IGNORECASE,
        )

        # Collapse whitespace again.
        value = re.sub(r"\s+", " ", value).strip()

        return value

    @staticmethod
    def _company_names_match(
        search_company: str,
        present_company: str,
    ) -> bool:
        """
        Company match rule:
        - compare normalized case/punctuation/legal suffixes
        - allow containment for meaningful multi-character names
        - missing company on either side = not matched
        """
        a = SearchWorker._normalize_company_for_match(search_company)
        b = SearchWorker._normalize_company_for_match(present_company)

        if not a or not b:
            return False

        if a == b:
            return True

        # Helpful for layouts such as:
        # "Deloitte Tax" vs "Deloitte Tax US"
        # but avoid containment for very short names.
        if len(a) >= 4 and len(b) >= 4:
            if a in b or b in a:
                # Do not collapse a one-token brand into a different named
                # subsidiary (for example Cognizant vs Cognizant Moment).
                if len(a.split()) == 1 or len(b.split()) == 1:
                    return False
                return True

        return False


    def _open_company_page_from_panel_click(
        self,
        company_url: str,
        company_name: str,
    ):
        """
        Click the matched company inside the right-side prospect panel.

        The click opens a new tab so the Sales Navigator search stays intact.
        A direct URL fallback is used only if LinkedIn changes the click behavior.
        """
        if not company_url or company_url == "-":
            raise RuntimeError("Matched company does not have a company link.")

        target_path = urlsplit(company_url).path.rstrip("/")

        try:
            with self.context.expect_page(timeout=7000) as page_info:
                clicked = self.page.evaluate(
                    """
                    (args) => {
                        const targetPath = args.targetPath || '';
                        const companyName = (args.companyName || '').trim().toLowerCase();
                        const vw = window.innerWidth;

                        const links = Array.from(
                            document.querySelectorAll('a[href]')
                        ).filter(a => {
                            const href = a.href || '';
                            if (!href.includes(targetPath)) return false;

                            const text = (
                                a.innerText ||
                                a.textContent ||
                                a.getAttribute('aria-label') ||
                                ''
                            ).replace(/\\s+/g, ' ').trim().toLowerCase();

                            if (companyName && text && !text.includes(companyName)) {
                                return false;
                            }

                            const r = a.getBoundingClientRect();
                            return (
                                r.width > 0 &&
                                r.height > 0 &&
                                r.left > vw * 0.52
                            );
                        });

                        const target = links[0];
                        if (!target) return false;

                        target.setAttribute('target', '_blank');
                        target.click();
                        return true;
                    }
                    """,
                    {
                        "targetPath": target_path,
                        "companyName": company_name or "",
                    },
                )

                if not clicked:
                    raise RuntimeError(
                        "Matched company link was not found in the prospect panel."
                    )

            company_page = page_info.value
            company_page.set_default_timeout(PAGE_TIMEOUT_MS)

            try:
                company_page.wait_for_load_state(
                    "domcontentloaded",
                    timeout=COMPANY_PAGE_TIMEOUT_MS,
                )
            except Exception as exc:
                # LinkedIn can occasionally create a transient popup/tab and
                # close it during redirect. Let the caller recover directly.
                if company_page.is_closed():
                    self.log(
                        "COMPANY CLICK TRANSIENT TAB CLOSED | "
                        f"company='{company_name}' | "
                        f"{type(exc).__name__}: {exc}"
                    )
                    return self._open_company_page_direct(
                        company_url,
                        reason="clicked company tab closed during load",
                    )

            if company_page.is_closed():
                return self._open_company_page_direct(
                    company_url,
                    reason="clicked company tab closed before extraction",
                )

            self.log(
                "COMPANY CLICK | "
                f"clicked='{company_name}' | "
                f"url='{company_url}'"
            )

            return company_page

        except Exception as exc:
            self.log(
                "COMPANY CLICK WARNING | "
                f"{type(exc).__name__}: {exc} | "
                "using direct company URL fallback"
            )

            return self._open_company_page_direct(
                company_url,
                reason="company click failed",
            )


    def _open_company_page_direct(
        self,
        company_url: str,
        reason: str = "",
    ):
        """
        Open the company URL in a fresh tab.

        Used as a recovery path when LinkedIn closes the tab created by the
        company-link click before extraction completes.
        """
        if not self.context:
            raise RuntimeError(
                "Browser context is not available for company-page recovery."
            )

        if reason:
            self.log(
                "COMPANY PAGE RECOVERY | "
                f"reason='{reason}' | "
                f"url='{company_url}'"
            )

        page = self.context.new_page()
        page.set_default_timeout(PAGE_TIMEOUT_MS)

        page.goto(
            company_url,
            wait_until="domcontentloaded",
            timeout=COMPANY_PAGE_TIMEOUT_MS,
        )

        return page


    @timed('linkedin_company')
    def _read_company_details(
        self,
        company_url: str,
        company_name: str,
        extraction_mode: str,
    ) -> dict[str, str]:
        """
        VERSION 12 COMPANY-PAGE VALIDATION

        Workflow:
        1. Open/click the matched company from the prospect panel.
        2. Give the Sales Navigator company page time to render.
        3. Read the company identity from the company page itself.
        4. Compare it with the company copied from the search-result card.
        5. If the names do not match, raise CompanyPageMismatchError.
           The caller closes/skips this prospect and moves to the next person.
        6. Only after the company name matches, copy:
              Industry
              City / State / Country
              Company Size
           from the company summary block.

        Example:
            University of Northern Iowa
            Higher Education
            Cedar Falls, Iowa, United States
            3K+ employees
        """
        fields = {
            "Company Name": company_name or "-",
            "Company Website": "-",
            "Industry": "-",
            "Company Size": "-",
            "See All Company Size": "-",
            "Industry Link": "-",
            "Address": "-",
            "City": "-",
            "State": "-",
            "Zip Code": "-",
            "Address Link": "-",
            "Country": "-",
            "Phone Number": "-",
            "Phone Number Link": "-",
            "Revenue": "-",
            "Revenue Link": "-",
            "SIC Code": "-",
            "NAICS Code": "-",
            "SIC/NAICS Code Link": "-",
        }

        if not company_url or company_url == "-":
            raise CompanyPageMismatchError(
                f"No company page link was available for '{company_name}'."
            )

        cache_key = normalize_url(company_url)

        if cache_key in self.company_cache:
            cached = dict(self.company_cache[cache_key])

            cached_name = cached.get("Company Name", "-")
            if not self._company_names_match(company_name, cached_name):
                raise CompanyPageMismatchError(
                    f"Cached company page '{cached_name}' does not match "
                    f"search company '{company_name}'."
                )

            return cached

        company_page = self._open_company_page_from_panel_click(
            company_url=company_url,
            company_name=company_name,
        )

        try:
            # --------------------------------------------------
            # WAIT FOR THE COMPANY PAGE TO RENDER PROPERLY
            # --------------------------------------------------
            recovery_used = False

            def recover_closed_company_page(reason: str):
                nonlocal company_page, recovery_used

                if recovery_used:
                    raise RuntimeError(
                        "Company page closed again after recovery: "
                        + reason
                    )

                recovery_used = True

                try:
                    if company_page and not company_page.is_closed():
                        company_page.close()
                except Exception:
                    pass

                company_page = self._open_company_page_direct(
                    company_url,
                    reason=reason,
                )

            if company_page.is_closed():
                recover_closed_company_page(
                    "company page was already closed before extraction"
                )

            try:
                company_page.wait_for_load_state(
                    "domcontentloaded",
                    timeout=COMPANY_PAGE_TIMEOUT_MS,
                )
            except Exception as exc:
                if company_page.is_closed():
                    recover_closed_company_page(
                        "company page closed while waiting for DOM content"
                    )
                else:
                    self.log(
                        "COMPANY PAGE LOAD WARNING | "
                        f"{type(exc).__name__}: {exc}"
                    )

            try:
                company_page.wait_for_function(
                    """
                    () => {
                        const body = document.body;
                        if (!body) return false;

                        const text = (body.innerText || '')
                            .replace(/\\s+/g, ' ')
                            .trim();

                        if (text.length < 250) return false;

                        return (
                            /\\bemployees\\b/i.test(text) ||
                            /\\babout\\b/i.test(text) ||
                            /\\bindustry\\b/i.test(text)
                        );
                    }
                    """,
                    timeout=COMPANY_PAGE_ELEMENT_WAIT_MS,
                    polling=250,
                )
            except Exception as exc:
                if company_page.is_closed():
                    recover_closed_company_page(
                        "company page closed while waiting for company elements"
                    )
                else:
                    # LinkedIn can render the summary without satisfying every
                    # heuristic above. Continue to the final settle/evaluate.
                    self.log(
                        "COMPANY ELEMENT WAIT WARNING | "
                        f"{type(exc).__name__}: {exc}"
                    )

            try:
                company_page.wait_for_timeout(COMPANY_PAGE_SETTLE_MS)
            except Exception as exc:
                if company_page.is_closed():
                    recover_closed_company_page(
                        "company page closed during final settle"
                    )
                    # The directly opened page already completed DOMContentLoaded.
                    company_page.wait_for_timeout(800)
                else:
                    raise

            if company_page.is_closed():
                recover_closed_company_page(
                    "company page closed immediately before data extraction"
                )
                company_page.wait_for_timeout(800)

            def evaluate_company_page():
                return company_page.evaluate(
                                """
                                () => {
                                    function clean(value) {
                                        return (value || '')
                                            .replace(/\\u00a0/g, ' ')
                                            .replace(/[ \\t]+/g, ' ')
                                            .trim();
                                    }
                
                                    function visible(el) {
                                        if (!el) return false;
                                        const r = el.getBoundingClientRect();
                                        const s = getComputedStyle(el);
                
                                        return (
                                            r.width > 0 &&
                                            r.height > 0 &&
                                            s.display !== 'none' &&
                                            s.visibility !== 'hidden'
                                        );
                                    }
                
                                    const bodyText = clean(document.body?.innerText || '');
                
                                    const links = Array.from(
                                        document.querySelectorAll('a[href]')
                                    )
                                    .filter(visible)
                                    .slice(0, 1000)
                                    .map(a => ({
                                        text: clean(
                                            a.innerText ||
                                            a.textContent ||
                                            a.getAttribute('aria-label') ||
                                            ''
                                        ),
                                        href: a.href || ''
                                    }));
                
                                    // Headings are useful for validating the company identity.
                                    const headings = Array.from(
                                        document.querySelectorAll(
                                            'h1, h2, [role="heading"]'
                                        )
                                    )
                                    .filter(visible)
                                    .map(el => {
                                        const r = el.getBoundingClientRect();
                                        return {
                                            text: clean(
                                                el.innerText ||
                                                el.textContent ||
                                                el.getAttribute('aria-label') ||
                                                ''
                                            ),
                                            left: r.left,
                                            top: r.top,
                                            width: r.width,
                                            height: r.height
                                        };
                                    })
                                    .filter(x => x.text);
                
                                    // Capture the compact left-side company identity block.
                                    const vw = window.innerWidth;
                                    const blocks = [];
                
                                    for (const el of Array.from(
                                        document.querySelectorAll('aside, section, div')
                                    )) {
                                        if (!visible(el)) continue;
                
                                        const r = el.getBoundingClientRect();
                
                                        if (
                                            r.left > vw * 0.34 ||
                                            r.top < 120 ||
                                            r.width < 180 ||
                                            r.width > vw * 0.40 ||
                                            r.height < 100 ||
                                            r.height > 650
                                        ) {
                                            continue;
                                        }
                
                                        const txt = clean(el.innerText || '');
                                        if (!txt || txt.length < 30 || txt.length > 1200) {
                                            continue;
                                        }
                
                                        let score = 0;
                                        if (/\\bemployees\\b/i.test(txt)) score += 120;
                                        if (/\\bvisit website\\b/i.test(txt)) score += 60;
                                        if (/\\bsave\\b/i.test(txt)) score += 20;
                                        if (r.left < vw * 0.25) score += 30;
                
                                        blocks.push({
                                            text: txt,
                                            score,
                                            left: r.left,
                                            top: r.top
                                        });
                                    }
                
                                    blocks.sort((a, b) => b.score - a.score);
                
                                    return {
                                        text: bodyText,
                                        links,
                                        headings,
                                        summaryBlock: blocks[0]?.text || ''
                                    };
                                }
                                """
                            )

            try:
                data = evaluate_company_page()
            except Exception as exc:
                if company_page.is_closed() and not recovery_used:
                    recover_closed_company_page(
                        "company page closed during data extraction"
                    )
                    company_page.wait_for_timeout(800)
                    data = evaluate_company_page()
                else:
                    raise

            body_text = self._clean_card_text(
                str(data.get("text", ""))
            )

            summary_text = self._clean_card_text(
                str(data.get("summaryBlock", ""))
            )

            links = data.get("links", []) or []
            headings = data.get("headings", []) or []

            # --------------------------------------------------
            # BUILD SUMMARY LINES
            # --------------------------------------------------
            summary_lines = [
                self._normalize_ws(x)
                for x in summary_text.splitlines()
                if self._normalize_ws(x)
            ]

            body_lines = [
                self._normalize_ws(x)
                for x in body_text.splitlines()
                if self._normalize_ws(x)
            ]

            # If compact summary detection failed, body text remains a fallback.
            lines = summary_lines if summary_lines else body_lines

            employee_re = re.compile(
                r"(?i)\b\d[\d,.]*\s*[KMB]?\s*\+?\s+employees\b"
            )

            # --------------------------------------------------
            # FIND COMPANY PAGE NAME
            # --------------------------------------------------
            actual_company_name = "-"

            # First prefer left/top headings because they represent the account.
            heading_candidates: list[tuple[float, str]] = []

            for item in headings:
                label = self._normalize_ws(
                    str(item.get("text", ""))
                )

                if not label:
                    continue

                low = label.lower()

                if low in {
                    "about",
                    "account iq",
                    "relationship explorer",
                    "common searches",
                    "your personas",
                }:
                    continue

                if "employees" in low:
                    continue

                left = float(item.get("left", 99999) or 99999)
                top = float(item.get("top", 99999) or 99999)

                # Prefer headings in the left company identity area.
                score = top + (left * 0.35)
                heading_candidates.append((score, label))

            heading_candidates.sort(key=lambda x: x[0])

            # Prefer a heading that already resembles the expected company,
            # but only among actual visible heading elements.
            for _score, label in heading_candidates[:15]:
                if self._company_names_match(company_name, label):
                    actual_company_name = label
                    break

            if actual_company_name == "-" and heading_candidates:
                actual_company_name = heading_candidates[0][1]

            # Sequence fallback around the employee-count line:
            # Company -> Industry -> Location -> Employees
            employee_idx = -1
            employee_value = "-"

            for idx, line in enumerate(lines):
                match = employee_re.search(line)

                if match:
                    employee_idx = idx
                    employee_value = match.group(0)
                    break

            if actual_company_name == "-" and employee_idx >= 3:
                # Most common account summary layout.
                actual_company_name = lines[employee_idx - 3]

            actual_company_name = self._normalize_ws(actual_company_name)

            # --------------------------------------------------
            # VERIFY COMPANY NAME BEFORE COPYING ANY COMPANY DATA
            # --------------------------------------------------
            if (
                actual_company_name == "-"
                or not self._company_names_match(
                    company_name,
                    actual_company_name,
                )
            ):
                raise CompanyPageMismatchError(
                    "Company page mismatch: "
                    f"search card='{company_name}' | "
                    f"company page='{actual_company_name}'"
                )

            self.log(
                "COMPANY PAGE MATCH | "
                f"search='{company_name}' | "
                f"page='{actual_company_name}'"
            )

            # Keep the original search-card spelling as the canonical output.
            fields["Company Name"] = company_name

            # --------------------------------------------------
            # COMPANY SIZE + LINK
            # --------------------------------------------------
            if employee_value != "-":
                fields["Company Size"] = employee_value

            for item in links:
                label = self._normalize_ws(
                    str(item.get("text", ""))
                )
                href = str(item.get("href", "")).strip()

                match = employee_re.search(label)

                if match:
                    fields["Company Size"] = match.group(0)
                    fields["See All Company Size"] = href or "-"
                    break

            # --------------------------------------------------
            # REVENUE FROM COMPANY PAGE
            # --------------------------------------------------
            revenue_re = re.compile(
                r"(?i)(?:[$€£¥₹]|USD|EUR|GBP|JPY|INR)\s*"
                r"\d[\d,.]*\s*(?:K|M|B|T|million|billion)?"
                r"(?:\s*[-–—]\s*"
                r"(?:[$€£¥₹]|USD|EUR|GBP|JPY|INR)?\s*"
                r"\d[\d,.]*\s*(?:K|M|B|T|million|billion)?)?"
                r"(?:\s+in\s+revenue)?"
            )

            for line in body_lines:
                if "revenue" not in line.lower():
                    continue

                match = revenue_re.search(line)

                if match:
                    fields["Revenue"] = self._normalize_ws(
                        match.group(0)
                    )
                    break

            for item in links:
                label = self._normalize_ws(
                    str(item.get("text", ""))
                )
                href = str(item.get("href", "")).strip()

                if "revenue" in label.lower() and href:
                    fields["Revenue Link"] = href
                    break

            # --------------------------------------------------
            # INDUSTRY + LOCATION FROM VERIFIED COMPANY SUMMARY
            # --------------------------------------------------
            #
            # Do NOT depend on the employee-count line being visible.
            #
            # Sales Navigator normally renders the left account summary as:
            #
            #   Capgemini
            #   IT Services and IT Consulting
            #   Paris, Île-de-France, France
            #   300K+ employees
            #
            # or:
            #
            #   University of Northern Iowa
            #   Higher Education
            #   Cedar Falls, Iowa, United States
            #   3K+ employees
            #
            # Since the company name has already been verified, anchor on that
            # company line and read the meaningful lines immediately after it.
            # --------------------------------------------------

            summary_company_idx = -1

            # First search the compact left-side company summary.
            for idx, line in enumerate(summary_lines):
                if self._company_names_match(
                    actual_company_name,
                    line,
                ):
                    summary_company_idx = idx
                    break

            # If the compact block was not detected, use the page body.
            industry_location_lines = summary_lines

            if summary_company_idx < 0:
                industry_location_lines = body_lines

                for idx, line in enumerate(body_lines):
                    if self._company_names_match(
                        actual_company_name,
                        line,
                    ):
                        summary_company_idx = idx
                        break

            def is_company_summary_noise(value: str) -> bool:
                value = self._normalize_ws(value)

                if not value:
                    return True

                low = value.lower()

                if self._looks_like_control_line(value):
                    return True

                if employee_re.search(value):
                    return True

                if self._looks_like_revenue_line(value):
                    return True

                if self._company_names_match(
                    actual_company_name,
                    value,
                ):
                    return True

                exact_noise = {
                    "about",
                    "account iq",
                    "relationship explorer",
                    "relationship map",
                    "save",
                    "saved",
                    "visit website",
                    "common searches",
                    "your personas",
                }

                if low in exact_noise:
                    return True

                if low.startswith(
                    (
                        "decision makers",
                        "cxo ",
                        "director+",
                        "create new",
                    )
                ):
                    return True

                return False

            # --------------------------------------------------
            # PRIMARY INDUSTRY EXTRACTION
            # --------------------------------------------------
            # The first meaningful NON-location value after the verified
            # company name is normally the industry.
            if summary_company_idx >= 0:
                scan_end = min(
                    len(industry_location_lines),
                    summary_company_idx + 8,
                )

                for idx in range(
                    summary_company_idx + 1,
                    scan_end,
                ):
                    candidate = self._normalize_ws(
                        industry_location_lines[idx]
                    )

                    if is_company_summary_noise(candidate):
                        continue

                    # The location follows the industry, so do not mistake
                    # a City/State/Country line for Industry.
                    if self._is_location_candidate(candidate):
                        continue

                    if (
                        2 <= len(candidate) <= 180
                        and not self._looks_like_revenue_line(candidate)
                    ):
                        fields["Industry"] = candidate
                        fields["Industry Link"] = company_url

                        self.log(
                            "INDUSTRY | "
                            f"company='{company_name}' | "
                            f"industry='{candidate}'"
                        )
                        break

            # --------------------------------------------------
            # COMPANY LOCATION EXTRACTION
            # --------------------------------------------------
            # Company-page geography is extracted here for fallback use.
            # The caller preserves any valid profile geography first.
            if summary_company_idx >= 0:
                scan_end = min(
                    len(industry_location_lines),
                    summary_company_idx + 10,
                )

                for idx in range(
                    summary_company_idx + 1,
                    scan_end,
                ):
                    candidate = self._normalize_ws(
                        industry_location_lines[idx]
                    )

                    if is_company_summary_noise(candidate):
                        continue

                    if not self._is_location_candidate(candidate):
                        continue

                    parts = [
                        p.strip()
                        for p in candidate.split(",")
                        if p.strip()
                    ]

                    if len(parts) >= 3:
                        fields["City"] = parts[0].strip(" .")
                        fields["State"] = parts[1].strip(" .")
                        fields["Country"] = self._canonical_country(parts[-1].strip(" ."))

                    elif len(parts) == 2:
                        fields["City"] = parts[0].strip(" .")
                        fields["State"] = "-"
                        fields["Country"] = self._canonical_country(parts[-1].strip(" ."))

                    break

            # --------------------------------------------------
            # SECONDARY FALLBACK AROUND EMPLOYEE SIZE
            # --------------------------------------------------
            # Keep the previous structural fallback for pages where the
            # compact company-name anchor is unavailable.
            if employee_idx >= 0:
                if (
                    fields["City"] == "-"
                    and employee_idx - 1 >= 0
                ):
                    location_line = self._normalize_ws(
                        lines[employee_idx - 1]
                    )

                    if (
                        location_line
                        and self._is_location_candidate(location_line)
                        and "employee" not in location_line.lower()
                    ):
                        parts = [
                            p.strip()
                            for p in location_line.split(",")
                            if p.strip()
                        ]

                        if len(parts) >= 3:
                            fields["City"] = parts[0].strip(" .")
                            fields["State"] = parts[1].strip(" .")
                            fields["Country"] = self._canonical_country(parts[-1].strip(" ."))
                        elif len(parts) == 2:
                            fields["City"] = parts[0].strip(" .")
                            fields["State"] = "-"
                            fields["Country"] = self._canonical_country(parts[-1].strip(" ."))

                if (
                    fields["Industry"] == "-"
                    and employee_idx - 2 >= 0
                ):
                    industry_line = self._normalize_ws(
                        lines[employee_idx - 2]
                    )

                    if (
                        industry_line
                        and not self._looks_like_control_line(industry_line)
                        and not employee_re.search(industry_line)
                        and not self._looks_like_revenue_line(industry_line)
                        and not self._company_names_match(
                            actual_company_name,
                            industry_line,
                        )
                        and not self._is_location_candidate(industry_line)
                        and len(industry_line) <= 180
                    ):
                        fields["Industry"] = industry_line
                        fields["Industry Link"] = company_url

            # --------------------------------------------------
            # FALLBACK: explicit Industry label/value
            # --------------------------------------------------
            if fields["Industry"] == "-":
                for i, line in enumerate(body_lines):
                    low = line.lower()

                    if low in {"industry", "industries"} and i + 1 < len(body_lines):
                        candidate = body_lines[i + 1]

                        if 1 < len(candidate) <= 160:
                            fields["Industry"] = candidate
                            fields["Industry Link"] = company_url
                            break

                    match = re.match(
                        r"(?i)^industry\s*:?\s*(.+)$",
                        line,
                    )

                    if match and match.group(1).strip():
                        fields["Industry"] = match.group(1).strip()
                        fields["Industry Link"] = company_url
                        break

            # --------------------------------------------------
            # EXTERNAL WEBSITE
            # --------------------------------------------------
            for item in links:
                label = self._normalize_ws(
                    str(item.get("text", ""))
                )
                href = str(item.get("href", "")).strip()

                if (
                    href.startswith(("http://", "https://"))
                    and "linkedin.com" not in href.lower()
                    and (
                        "website" in label.lower()
                        or "visit website" in label.lower()
                    )
                ):
                    fields["Company Website"] = href
                    break

            # Cache only verified company-page data.
            self.company_cache[cache_key] = dict(fields)

            self.log(
                "COMPANY DETAILS | "
                f"company='{fields['Company Name']}' | "
                f"industry='{fields['Industry']}' | "
                f"location='{fields['City']}, {fields['State']}, {fields['Country']}' | "
                f"size='{fields['Company Size']}'"
            )

            if fields["Industry"] == "-":
                self.log(
                    "INDUSTRY WARNING | "
                    f"Verified company='{company_name}', but no industry "
                    "was found in the company summary."
                )

            return fields

        finally:
            # Always close the company tab before the next prospect.
            try:
                company_page.close()
            except Exception:
                pass


    def _scroll_search_results(self):
        """
        Scroll the CENTER search-results column, not the right-side profile panel.
        """
        try:
            self.page.evaluate(
                """
                () => {
                    const vw = window.innerWidth;
                    const candidates = [];

                    for (const el of Array.from(document.querySelectorAll('*'))) {
                        const r = el.getBoundingClientRect();
                        const style = getComputedStyle(el);
                        const delta = el.scrollHeight - el.clientHeight;

                        if (
                            delta < 250 ||
                            !(style.overflowY === 'auto' || style.overflowY === 'scroll')
                        ) continue;

                        // Sales Navigator result list is normally in the central
                        // column, to the right of filters and left of side panel.
                        if (
                            r.left < vw * 0.24 ||
                            r.left > vw * 0.67 ||
                            r.width < 360 ||
                            r.height < 300
                        ) continue;

                        candidates.push({
                            el,
                            score: delta + r.height + r.width
                        });
                    }

                    candidates.sort((a, b) => b.score - a.score);

                    if (candidates[0]) {
                        candidates[0].el.scrollTop += 1100;
                        return true;
                    }

                    window.scrollBy(0, 1100);
                    return false;
                }
                """
            )
        except Exception:
            try:
                self.page.mouse.wheel(0, 1100)
            except Exception:
                pass

    def _go_to_next_results_page(self) -> bool:
        """
        VERSION 15:
        Sales Navigator search results are commonly paginated in batches
        (often 25 profiles per page). If some profiles are skipped because of
        company mismatch / missing company link, processing the first 25
        profiles is not enough to SAVE 25 valid rows.

        This helper moves to the next result page when available.
        It deliberately searches near the bottom of the search-results area so
        it does not confuse the prospect side-panel "Next" control with the
        search-results pagination control.
        """
        try:
            # First push the results area/page to the bottom so pagination is
            # rendered and visible.
            self.page.evaluate(
                """
                () => {
                    const vw = window.innerWidth;
                    const candidates = [];

                    for (const el of Array.from(document.querySelectorAll('*'))) {
                        const r = el.getBoundingClientRect();
                        const style = getComputedStyle(el);
                        const delta = el.scrollHeight - el.clientHeight;

                        if (
                            delta < 200 ||
                            !(style.overflowY === 'auto' || style.overflowY === 'scroll')
                        ) continue;

                        if (
                            r.left < vw * 0.20 ||
                            r.left > vw * 0.78 ||
                            r.width < 350 ||
                            r.height < 250
                        ) continue;

                        candidates.push({
                            el,
                            score: delta + r.height + r.width
                        });
                    }

                    candidates.sort((a, b) => b.score - a.score);

                    if (candidates[0]) {
                        candidates[0].el.scrollTop =
                            candidates[0].el.scrollHeight;
                    }

                    window.scrollTo(0, document.body.scrollHeight);
                }
                """
            )

            self.page.wait_for_timeout(500)

            before_urls = set(
                normalize_url(card.get("url", ""))
                for card in self._collect_current_cards()
                if card.get("url")
            )

            clicked = self.page.evaluate(
                """
                () => {
                    const vw = window.innerWidth;
                    const vh = window.innerHeight;

                    function clean(value) {
                        return (value || '')
                            .replace(/\\s+/g, ' ')
                            .trim()
                            .toLowerCase();
                    }

                    function isVisible(el) {
                        if (!el) return false;

                        const r = el.getBoundingClientRect();
                        const s = getComputedStyle(el);

                        return (
                            r.width > 0 &&
                            r.height > 0 &&
                            r.bottom > 0 &&
                            r.top < vh &&
                            s.display !== 'none' &&
                            s.visibility !== 'hidden'
                        );
                    }

                    const candidates = Array.from(
                        document.querySelectorAll(
                            'button, a[role="button"], a[href]'
                        )
                    ).filter(el => {
                        if (!isVisible(el)) return false;

                        const r = el.getBoundingClientRect();

                        // Search-results pagination should be low on the page,
                        // not in the right-side profile panel header.
                        if (r.top < vh * 0.45) return false;

                        // Avoid far-left filter controls and extreme right-side
                        // panel controls.
                        if (
                            r.left < vw * 0.20 ||
                            r.right > vw * 0.92
                        ) return false;

                        const label = clean(
                            el.getAttribute('aria-label') ||
                            el.getAttribute('title') ||
                            el.innerText ||
                            el.textContent ||
                            ''
                        );

                        const disabled =
                            el.disabled ||
                            el.getAttribute('aria-disabled') === 'true' ||
                            el.classList.contains('disabled');

                        if (disabled) return false;

                        return (
                            label === 'next' ||
                            label === 'next page' ||
                            label.includes('next page')
                        );
                    });

                    // Prefer the lowest visible candidate.
                    candidates.sort((a, b) => {
                        const ra = a.getBoundingClientRect();
                        const rb = b.getBoundingClientRect();
                        return rb.top - ra.top;
                    });

                    const target = candidates[0];

                    if (!target) {
                        return false;
                    }

                    target.click();
                    return true;
                }
                """
            )

            if not clicked:
                self.log("PAGINATION | no enabled Next results-page control found")
                return False

            self.log("PAGINATION | clicked Next results page")
            self.status("Loading next Sales Navigator results page...")

            # Wait for navigation/rendering. Sales Navigator may use either a
            # route change or client-side replacement, so check both URL/card
            # changes rather than depending only on a navigation event.
            deadline = time.time() + 15

            while time.time() < deadline and not self.stop_event.is_set():
                self.page.wait_for_timeout(350)

                cards = self._collect_current_cards()

                after_urls = {
                    normalize_url(card.get("url", ""))
                    for card in cards
                    if card.get("url")
                }

                if after_urls and after_urls != before_urls:
                    self.log(
                        "PAGINATION | next page loaded | "
                        f"visible_cards={len(cards)}"
                    )
                    return True

            self.log(
                "PAGINATION WARNING | Next was clicked, but a new result "
                "batch was not detected in time."
            )
            return False

        except Exception as exc:
            self.log(
                "PAGINATION ERROR | "
                f"{type(exc).__name__}: {exc}"
            )
            return False


    @staticmethod
    def _fuzzy_company_match(expected: str, actual: str) -> tuple[bool, float]:
        """Conservative name similarity; no partial-brand or acronym expansion."""
        def tokens(value: str) -> list[str]:
            value = unicodedata.normalize('NFKD', str(value or '')).casefold()
            value = ''.join(c for c in value if not unicodedata.combining(c))
            value = value.replace('&', ' and ')
            value = ''.join(c if c.isalnum() else ' ' for c in value)
            parts = value.split()
            suffixes = {
                'llp', 'llc', 'ltd', 'limited', 'inc', 'incorporated', 'corp',
                'corporation', 'plc', 'pvt', 'private', 'co', 'company',
                'sa', 'srl', 'sl', 'pty', 'pte', 'bv', 'gmbh', 'ag', 'nv',
                'spa', 'sas', 'oy', 'ab', 'as', 'kk', 'sro', 'sarl', 'pbc',
            }
            while parts and parts[-1] in suffixes:
                parts.pop()
            return parts

        a, b = tokens(expected), tokens(actual)
        if not a or not b:
            return False, 0.0
        normalized_a = ' '.join(a)
        normalized_b = ' '.join(b)
        if frozenset({normalized_a, normalized_b}) in COMPANY_NAME_ALIASES:
            return True, 100.0
        if a == b or sorted(a) == sorted(b):
            return True, 100.0
        if ''.join(a) == ''.join(b):
            return True, 100.0
        score = 100 * difflib.SequenceMatcher(None, ' '.join(sorted(a)), ' '.join(sorted(b))).ratio()
        # Added brand qualifiers can identify another subsidiary, not a typo.
        if len(a) != len(b) or min(len(''.join(a)), len(''.join(b))) <= 4:
            return False, round(score, 1)
        # Long shared business words must not hide a different short brand.
        for left, right in zip(sorted(a), sorted(b)):
            if left == right:
                continue
            if min(len(left), len(right)) <= 4 or difflib.SequenceMatcher(None, left, right).ratio() < 0.75:
                return False, round(score, 1)
        return score >= COMPANY_NAME_MATCH_THRESHOLD, round(score, 1)

    @staticmethod
    def _link_company_names_match(expected: str, actual: str) -> bool:
        """Apply the same conservative fuzzy identity check throughout."""
        fuzzy, _score = SearchWorker._fuzzy_company_match(expected, actual)
        return fuzzy


    @timed('import')
    def _load_link_inputs(self, input_path):
        from link_import import load_link_inputs
        return load_link_inputs(input_path,
            progress=lambda text: self.metrics.notify('import', 'running', text),
            cancelled=self.stop_event.is_set)

    def _load_profile_links(self, input_path):
        # Compatibility helper; the run uses the full audit-preserving inputs.
        return list(dict.fromkeys(item.canonical for item in self._load_link_inputs(input_path)
                                  if item.canonical))

    @timed('link_resolution')
    def _resolve_profile_link(self, item):
        if item.kind == 'LinkedIn profile':
            return item.canonical, 'Direct'
        if not item.canonical:
            raise ValueError(item.reason)
        if self.refresh_profiles and self.link_resolutions:
            self.link_resolutions.invalidate(item.canonical)
        cached = self.link_resolutions.get(item.canonical) if self.link_resolutions else None
        if cached:
            self.metrics.count('link_resolution_cache_hits')
            return cached, 'Cached'
        if not self._wait_if_paused():
            raise RuntimeError('Link resolution stopped.')
        self._launch_browser()
        self.status('Resolving Sales Navigator link')
        self.metrics.count('sales_link_resolution_visits')
        self.page.goto(item.canonical, wait_until='domcontentloaded', timeout=PAGE_TIMEOUT_MS)
        self._wait_for_manual_login_if_needed()
        self.page.wait_for_timeout(LINK_PROFILE_SETTLE_MS)

        def inspect():
            if not self._wait_if_paused():
                raise RuntimeError('Link resolution stopped.')
            state = self.page.evaluate(RESOLUTION_DOM)
            if state.get('auth'):
                raise ProfileAccessError('Sales Navigator requires sign-in or verification. Use Open Chrome / Sign In, then retry.')
            if state.get('unavailable'):
                raise ValueError('Sales Navigator lead is unavailable.')
            current = classify_link(self.page.url)
            if current.kind == 'LinkedIn profile':
                return current.canonical
            candidates = {classify_link(url).canonical for url in state.get('links', [])
                          if classify_link(url).kind == 'LinkedIn profile'}
            if len(candidates) > 1:
                raise ValueError('Multiple public profile links were exposed; no mapping was guessed.')
            return next(iter(candidates), None)

        resolved = inspect()
        if not resolved:
            # Only an explicitly labelled, unique action menu is eligible.
            menus = self.page.get_by_role('button', name='Open actions overflow menu', exact=True)
            visible = [menus.nth(i) for i in range(menus.count()) if menus.nth(i).is_visible()]
            if not visible:
                menus = self.page.get_by_role('button', name=re.compile(
                    r'^(More|More actions|More options|More lead actions|More profile actions)(?: for .+)?$', re.I))
                visible = [menus.nth(i) for i in range(menus.count()) if menus.nth(i).is_visible()]
            if len(visible) == 1:
                visible[0].click(timeout=3000)
                deadline = time.monotonic() + 2
                while time.monotonic() < deadline:
                    self.page.wait_for_timeout(150)
                    resolved = inspect()
                    if resolved:
                        break
        if not resolved:
            # Some layouts expose the action as a button/menuitem, not an anchor.
            actions = self.page.get_by_role('button', name=re.compile(r'^View LinkedIn profile$', re.I)).or_(
                self.page.get_by_role('menuitem', name=re.compile(r'^View LinkedIn profile$', re.I)))
            visible = [actions.nth(i) for i in range(actions.count()) if actions.nth(i).is_visible()]
            if len(visible) == 1:
                before = set(self.page.context.pages)
                try:
                    visible[0].click(timeout=3000)
                    deadline = time.monotonic() + 5
                    while time.monotonic() < deadline:
                        if not self._wait_if_paused():
                            raise RuntimeError('Link resolution stopped.')
                        resolved = inspect()
                        popups = [p for p in self.page.context.pages if p not in before]
                        popup_urls = {classify_link(p.url).canonical for p in popups
                                      if classify_link(p.url).kind == 'LinkedIn profile'}
                        if len(popup_urls) > 1:
                            raise ValueError('Ambiguous public profile popup; no mapping saved.')
                        resolved = resolved or next(iter(popup_urls), None)
                        if resolved:
                            break
                        self.page.wait_for_timeout(150)
                finally:
                    for popup in list(self.page.context.pages):
                        if popup not in before:
                            popup.close()
        if not resolved:
            raise ValueError('No explicit View LinkedIn profile link was found. Check Sales Navigator access or supply the normal profile URL.')
        if self.link_resolutions:
            self.link_resolutions.put(item.canonical, resolved)
        self.metrics.count('sales_links_resolved')
        return resolved, 'Resolved'

    def _scroll_standard_page(self, max_scrolls: int) -> None:
        """Progressively render a standard LinkedIn page."""
        last_height = -1
        stagnant = 0

        for _ in range(max_scrolls):
            if self.stop_event.is_set():
                break
            if not self._wait_if_paused():
                break

            try:
                values = self.page.evaluate(
                    """
                    () => {
                        const h = Math.max(
                            document.body?.scrollHeight || 0,
                            document.documentElement?.scrollHeight || 0
                        );
                        window.scrollBy(0, Math.max(700, window.innerHeight * 0.75));
                        return {height: h, y: window.scrollY};
                    }
                    """
                )
                height = int(values.get('height', 0) or 0)
            except Exception:
                height = 0

            if height == last_height:
                stagnant += 1
            else:
                stagnant = 0

            last_height = height
            self.page.wait_for_timeout(180)

            try:
                at_bottom = self.page.evaluate(
                    """
                    () => (
                        window.scrollY + window.innerHeight >=
                        (document.documentElement?.scrollHeight || document.body?.scrollHeight || 0) - 20
                    )
                    """
                )
            except Exception:
                at_bottom = False

            if at_bottom and stagnant >= 1:
                break

    def _standard_profile_name(self, timeout_ms: int = 12000) -> str:
        """Read rendered profile identity without requiring a particular H1 tag."""
        deadline = time.monotonic() + timeout_ms / 1000
        last_state = {}
        while True:
            if self.stop_event.is_set():
                raise RuntimeError('Profile capture stopped.')
            last_state = self.page.evaluate(r"""() => {
                const main = document.querySelector('main');
                const visible = el => !!el && el.getBoundingClientRect().width > 0 &&
                    el.getBoundingClientRect().height > 0 && getComputedStyle(el).visibility !== 'hidden';
                const text = (document.body?.innerText || '').slice(0, 12000);
                const auth = /\/(login|checkpoint|authwall|uas\/login)(\/|\?|$)/i.test(location.href) ||
                    Array.from(document.querySelectorAll('input[type="password"]')).some(visible) ||
                    /sign in to (view|see)|join (now|linkedin) to (view|see)|verify your identity|security verification/i.test(text);
                const unavailable = /this (page|profile) (doesn.t exist|is(n.t| not) available)|profile not found|page not found/i.test(text);
                const selectors = ['main h1', 'h1', '.pv-text-details__left-panel h2',
                    'main [data-anonymize="person-name"]', 'main .top-card-layout__title',
                    'main section:first-of-type h2'];
                const bad = /^(experience|about|education|activity|contact info|sign in|join now|linkedin|welcome|page not found)$/i;
                let name = '';
                for (const selector of selectors) {
                    for (const el of document.querySelectorAll(selector)) {
                        const value = (el.innerText || '').trim();
                        if (visible(el) && value && !bad.test(value) && value.length < 150) {
                            name = value; break;
                        }
                    }
                    if (name) break;
                }
                // Page titles are a fallback only on a rendered profile page.
                if (!name && main && /\/in\//i.test(location.pathname)) {
                    const title = document.title.replace(/^\(\d+\)\s*/, '').trim();
                    if (/\s[|–-]\sLinkedIn$/i.test(title)) {
                        const candidate = title.replace(/\s[|–-]\sLinkedIn$/i, '').split(/\s[-–]\s/)[0].trim();
                        if (candidate && !bad.test(candidate) && candidate.length < 150) name = candidate;
                    }
                }
                const locationCandidates = [];
                // Current LinkedIn layouts use generated classes. The location
                // still sits beside the profile's Contact info link.
                for (const contact of document.querySelectorAll('main a[href*="/overlay/contact-info/"]')) {
                    const block = contact.closest('p')?.parentElement || contact.parentElement;
                    for (const el of block?.querySelectorAll('p,span') || []) {
                        const value = (el.innerText || '').replace(/\s+/g, ' ').trim();
                        if (visible(el) && value.includes(',') && value.length < 160) locationCandidates.push(value);
                    }
                }
                for (const selector of [
                    '[data-anonymize="location"]',
                    '.top-card-layout__first-subline',
                    '.pv-text-details__left-panel .text-body-small',
                    'main section:first-of-type .text-body-small'
                ]) {
                    for (const el of document.querySelectorAll(selector)) {
                        const value = (el.innerText || el.textContent || '').replace(/\s+/g, ' ').trim();
                        if (visible(el) && value && !bad.test(value)) locationCandidates.push(value);
                    }
                }
                return {name, auth, unavailable, url: location.href, title: document.title,
                    locations: [...new Set(locationCandidates)]};
            }""")
            if last_state.get('auth'):
                raise ProfileAccessError('LinkedIn requires sign-in or verification. Complete it in the scraper Chrome window, then restart. URL=' + last_state.get('url', ''))
            if last_state.get('unavailable'):
                raise RuntimeError('LinkedIn reports this profile unavailable: ' + last_state.get('url', ''))
            if last_state.get('name'):
                self._last_profile_locations = last_state.get('locations', []) or []
                return self._clean_person_name(last_state['name'])
            if time.monotonic() >= deadline:
                raise RuntimeError(f"Profile name was not rendered after {timeout_ms} ms; URL={last_state.get('url')}; title={last_state.get('title')!r}. The page may be unavailable or its layout unsupported.")
            self.page.wait_for_timeout(250)

    def _read_standard_experience_dom(self) -> dict[str, Any]:
        return self.page.evaluate(r"""() => {
            const clean = v => (v || '').replace(/\u00a0/g, ' ').replace(/[ \t]+/g, ' ').trim();
            const main = document.querySelector('main') || document.body;
            const date = /\b\d{4}\s*[-–—]\s*Present\b/i;
            const heading = Array.from(main.querySelectorAll('h2,h3,[role="heading"],span'))
                .find(el => /^Experience$/i.test(clean(el.innerText)));
            const marker = main.querySelector('#experience');
            let section = heading?.closest('section') || marker?.closest('section');
            if (!section) {
                let node = heading?.parentElement || marker?.parentElement;
                while (node && node !== main) {
                    if (date.test(node.innerText || '') && node.querySelector('a[href*="/company/"], a[href*="/school/"]')) {
                        section = node; break;
                    }
                    node = node.parentElement;
                }
            }
            if (!section && /\/details\/experience/.test(location.pathname)) section = main;
            if (!section) return {experienceText:'',currentText:'',companyLinks:[],sectionFound:false};
            const candidates = Array.from(section.querySelectorAll('li, [role="listitem"], article, div'))
                .filter(el => date.test(el.innerText || '') &&
                    clean(el.innerText).split('\n').filter(Boolean).length >= 3);
            // Prefer a focused complete card with an employer link. For grouped
            // employers, prefer the leaf role and recover its enclosing company.
            const leafRoles = candidates.filter(el => !candidates.some(other =>
                other !== el && el.contains(other) && date.test(other.innerText || '') &&
                clean(other.innerText).split('\n').filter(Boolean).length >= 3));
            // Preserve page order when several current roles are displayed.
            // Expand metadata-only wrappers to include the role's title.
            let role = leafRoles[0];
            const status = /^(?:full[- ]?time|part[- ]?time|contract|temporary|self-employed|freelance|internship)$/i;
            while (role && role.parentElement && role.parentElement !== section) {
                const beforeDate = clean(role.innerText).split('\n').map(clean).filter(Boolean);
                const index = beforeDate.findIndex(s => date.test(s));
                if (beforeDate.slice(0,index).some(s => !status.test(s))) break;
                role = role.parentElement;
            }
            const linkData = a => {
                const alt = a.querySelector('img')?.alt || a.querySelector('svg[aria-label]')?.getAttribute('aria-label');
                const label = clean(alt || a.innerText || a.getAttribute('aria-label'));
                // A whole role can itself link to the company. Its title/date
                // text is not an employer label; continue up to the group header.
                if (/\b\d{4}\s*[-–—]/.test(label)) return null;
                return {text:label, href:a.href};
            };
            let node = role, companyLinks = [];
            while (node && node !== section) {
                companyLinks = Array.from(node.querySelectorAll('a[href*="/company/"], a[href*="/school/"]')).map(linkData).filter(Boolean);
                if (companyLinks.length) break;
                node = node.parentElement;
            }
            // Some public/detail pages render the employer as plain text while
            // the profile page still contains the company link elsewhere. Return
            // all remaining company links so the Python parser can select the
            // one whose label matches the employer line.
            const companyLinksScoped = companyLinks.length > 0;
            if (!companyLinks.length) {
                companyLinks = Array.from(main.querySelectorAll('a[href*="/company/"], a[href*="/school/"]')).map(linkData).filter(Boolean);
            }
            return {experienceText:clean(section.innerText), currentText:clean(role?.innerText),
                companyLinks, companyLinksScoped, sectionFound:true};
        }""") or {}

    def _capture_standard_profile(self, profile_url: str) -> dict[str, Any]:
        profile_url = canonical_profile_url(profile_url)
        self.page.goto(profile_url, wait_until='domcontentloaded', timeout=PAGE_TIMEOUT_MS)
        self._wait_for_manual_login_if_needed()
        self.page.wait_for_timeout(LINK_PROFILE_SETTLE_MS)
        profile_name = self._standard_profile_name()
        data = {}
        for attempt in range(2):
            # The dedicated Experience page is a normal LinkedIn UI route and
            # avoids dependence on the top-profile page's lazy section wrapper.
            if attempt:
                self.page.goto(profile_url + '/details/experience/',
                               wait_until='domcontentloaded', timeout=PAGE_TIMEOUT_MS)
                self._wait_for_manual_login_if_needed()
                self.page.wait_for_timeout(LINK_PROFILE_SETTLE_MS)
            for _ in range(LINK_PROFILE_MAX_SCROLLS):
                if self.stop_event.is_set() or not self._wait_if_paused():
                    raise RuntimeError('Profile capture stopped.')
                data = self._read_standard_experience_dom()
                if data.get('currentText') and (data.get('companyLinks') or attempt > 0):
                    data['name'] = profile_name
                    data['profileLocations'] = list(getattr(self, '_last_profile_locations', []))
                    parsed = self._parse_standard_current_experience(data)
                    if all(parsed.get(field) not in (None, '', '-', 'Dead')
                           for field in ('Job Title', 'Company Name', 'Tenure')):
                        return data
                    self.trace('LINK_PROFILE', 'role_parse_unresolved', url=self.page.url,
                               data=data, parsed=parsed)
                    break
                self.page.evaluate(r"""() => {
                    const marker = document.querySelector('#experience');
                    if (marker && marker.getBoundingClientRect().top > innerHeight) marker.scrollIntoView();
                    else window.scrollBy(0, 600);
                    for (const el of document.querySelectorAll('main, [role="main"]')) {
                        if (el.scrollHeight > el.clientHeight && /auto|scroll/.test(getComputedStyle(el).overflowY)) el.scrollTop += 600;
                    }
                }""")
                self.page.wait_for_timeout(500)
        diagnostic = self.page.evaluate("""() => ({
            headings: Array.from(document.querySelectorAll('h1,h2,h3,[role="heading"]')).map(el => el.innerText).slice(0,40),
            renderedText: (document.querySelector('main')?.innerText || '').slice(0,12000)
        })""")
        self.trace('LINK_PROFILE', 'experience_unresolved', url=self.page.url,
                   name=profile_name, data=data, diagnostic=diagnostic)
        raise RuntimeError('Current Experience could not be extracted. Profile was not saved as Dead. '
                           'Check the Experience page in the scraper browser; see experience_unresolved in the activity log.')

    def _parse_standard_current_experience(
        self,
        profile_data: dict[str, Any],
    ) -> dict[str, str]:
        """Parse Name + current Experience from a regular LinkedIn profile."""
        display_name = self._clean_person_name(
            str(profile_data.get('name', '') or '')
        )
        first_name, last_name = self._split_name(display_name)

        raw = self._clean_card_text(
            str(profile_data.get('currentText', '') or '')
        )
        lines = [
            self._normalize_ws(x)
            for x in raw.splitlines()
            if self._normalize_ws(x)
        ]
        # LinkedIn can expose both visible and accessibility copies of each line.
        lines = [line for i, line in enumerate(lines) if i == 0 or line != lines[i - 1]]

        present_idx = next(
            (i for i, line in enumerate(lines)
             if re.search(r'\b\d{4}\s*[-–—]\s*Present\b', line, re.I)),
            -1,
        )

        if present_idx < 0:
            return {
                'First Name': first_name,
                'Last Name': last_name,
                'Job Title': '-',
                'Company Name': '-',
                'Company Profile Link': '-',
                'Tenure': 'Dead',
                'Location': '-',
            }

        tenure = lines[present_idx]

        # LinkedIn may expose the employer as text with no anchor on the
        # Experience detail route. Infer that employer from the line directly
        # above the current date before using any company links. A role card
        # often puts ``Full-time`` immediately above the date, so that status
        # line must never be accepted as the company name.
        company_url = '-'
        company = '-'

        def is_employment_status(value: str) -> bool:
            return bool(re.fullmatch(
                r'(?i)(?:full[- ]?time|part[- ]?time|contract|temporary|'
                r'self[- ]?employed|freelance|internship|volunteer|'
                r'apprenticeship|seasonal|on[- ]?call)',
                self._normalize_ws(value),
            ))

        if present_idx >= 2:
            company_candidate = lines[present_idx - 1]
            company_candidate = re.sub(
                r'\s*[·•]\s*(?:full[- ]?time|part[- ]?time|contract|self-employed|freelance|internship).*$',
                '', company_candidate, flags=re.I,
            ).strip()
            if (
                company_candidate
                and not is_employment_status(company_candidate)
                and not re.search(r'\b(?:present|\d+\s+(?:yrs?|mos?))\b', company_candidate, re.I)
            ):
                company = company_candidate

        # Grouped Experience cards may render the employer and its total
        # tenure before the role title, e.g. ``East Carolina University`` /
        # ``19 yrs 4 mos`` / ``Chief Information Officer`` / ``Full-time`` /
        # ``Dec 2020 - Present``. Recover the employer from the line before
        # that total-tenure marker when the direct line was a status label.
        if company == '-':
            for idx in range(present_idx - 1, 0, -1):
                if not re.search(r'\b\d+\s+(?:yrs?|years?|mos?|months?)\b', lines[idx], re.I):
                    continue
                candidate = lines[idx - 1]
                if candidate and not is_employment_status(candidate):
                    company = candidate
                    break

        # Find the company URL whose visible label best matches the inferred
        # employer. Do not blindly use the first link, because detail pages can
        # include unrelated recommendation/company links.
        link_items = profile_data.get('companyLinks', []) or []
        # A link recovered from the current role's enclosing employer card is
        # stronger evidence than the line above its date (often a job title).
        scoped_links = {
            normalize_url(str(item.get('href', '') or '')): item
            for item in link_items
            if '/company/' in str(item.get('href', '') or '')
            or '/school/' in str(item.get('href', '') or '')
        }
        if profile_data.get('companyLinksScoped') and len(scoped_links) == 1:
            company = '-'
        best_link: tuple[float, dict[str, Any]] | None = None
        for item in link_items:
            href = str(item.get('href', '') or '').strip()
            label = self._normalize_ws(str(item.get('text', '') or ''))

            if '/company/' not in href and '/school/' not in href:
                continue

            label = re.sub(r'\s+logo\s*$', '', label, flags=re.I)
                # Remove common duration/status suffixes from grouped-company links.
            label = re.sub(r'\s+\d+\s+(?:yrs?|years?|mos?|months?).*$', '', label, flags=re.I).strip(' ·|-')
            label = re.sub(r'\s*[·•]\s*(?:full[- ]?time|part[- ]?time|contract).*$', '', label, flags=re.I).strip()
            if company != '-' and label:
                matched, score = self._fuzzy_company_match(company, label)
                candidate_score = 100.0 if matched else score
            elif label:
                candidate_score = 40.0
            else:
                candidate_score = 0.0
            if best_link is None or candidate_score > best_link[0]:
                best_link = (candidate_score, item)

        valid_link_count = len(scoped_links)
        if best_link and (
            best_link[0] >= COMPANY_NAME_MATCH_THRESHOLD
            or (company == '-' and valid_link_count == 1 and profile_data.get('companyLinksScoped'))
        ):
            company_url = str(best_link[1].get('href', '') or '-')
            label = self._normalize_ws(str(best_link[1].get('text', '') or ''))
            label = re.sub(r'\s+logo\s*$', '', label, flags=re.I)
            label = re.sub(r'\s+\d+\s+(?:yrs?|years?|mos?|months?).*$', '', label, flags=re.I).strip(' ·|-')
            label = re.sub(r'\s*[·•]\s*(?:full[- ]?time|part[- ]?time|contract).*$', '', label, flags=re.I).strip()
            if label and company == '-':
                company = label

        company = re.sub(
            r'\s*[·•]\s*(?:full[- ]?time|part[- ]?time|contract|self-employed|freelance|internship).*$',
            '', company, flags=re.I,
        ).strip()

        def is_noise(value: str) -> bool:
            low = value.lower().strip()
            if not value:
                return True
            if is_employment_status(value):
                return True
            if re.search(r'\bpresent\b', low):
                return True
            if re.fullmatch(
                r'(?:\d+\s+yrs?\s*)?(?:\d+\s+mos?\s*)?',
                low,
            ):
                return True
            return False

        title = '-'
        for idx in range(present_idx - 1, max(-1, present_idx - 6), -1):
            candidate = lines[idx]
            if is_noise(candidate):
                continue
            employer_candidate = re.sub(
                r'\s*[·•]\s*(?:full[- ]?time|part[- ]?time|contract|self-employed|freelance|internship).*$',
                '', candidate, flags=re.I,
            ).strip()
            if company != '-' and self._link_company_names_match(company, employer_candidate):
                continue
            if 2 <= len(candidate) <= 220:
                title = candidate
                break

        location = '-'
        for idx in range(present_idx + 1, min(len(lines), present_idx + 5)):
            candidate = lines[idx]
            low = candidate.lower()
            if not candidate:
                continue
            if low.startswith(('skills:', 'skill:', 'show all')):
                continue
            if 'present' in low:
                break
            candidate = re.sub(r'\s*[·•]\s*(?:on-site|hybrid|remote)\s*$', '', candidate, flags=re.I)
            if 1 < len(candidate) <= 180 and self._is_location_candidate(candidate):
                location = candidate
                break

        profile_location = '-'
        for candidate in profile_data.get('profileLocations', []) or []:
            candidate = self._normalize_ws(str(candidate))
            if self._is_location_candidate(candidate):
                profile_location = candidate
                break

        if location == '-':
            location = profile_location

        # Keep the role-card location for traceability while exposing the
        # richer top-profile location as a fallback for a missing country
        # (for example ``Greenville, NC`` plus ``Greenville, North Carolina,
        # United States`` in the profile header).
        if profile_location == '-' and location != '-':
            profile_location = location

        return {
            'First Name': first_name,
            'Last Name': last_name,
            'Job Title': title,
            'Company Name': company,
            'Company Profile Link': company_url,
            'Tenure': tenure,
            'Location': location,
            'Profile Location': profile_location,
        }

    def _open_standard_company_about(self, company_url: str) -> str:
        """Open the regular LinkedIn company page and then its About tab."""
        company_url = (company_url or '').strip()
        if not company_url or company_url == '-':
            raise RuntimeError('Current Experience does not contain a company link.')

        # First try to click the company link/icon from the current profile,
        # matching the user's Link Scraper workflow. Force same-tab navigation
        # so the worker keeps ownership of one browser page.
        target_path = urlsplit(company_url).path.rstrip('/')
        clicked_company = False

        try:
            clicked_company = self.page.evaluate(
                """
                (targetPath) => {
                    const links = Array.from(
                        document.querySelectorAll('a[href*="/company/"], a[href*="/school/"]')
                    );
                    const target = links.find(a => {
                        try {
                            return new URL(a.href).pathname.replace(/\\/$/, '') === targetPath;
                        } catch (e) {
                            return false;
                        }
                    });

                    if (!target) return false;
                    target.removeAttribute('target');
                    target.click();
                    return true;
                }
                """,
                target_path,
            )
        except Exception:
            clicked_company = False

        if clicked_company:
            try:
                self.page.wait_for_url(
                    re.compile(r'linkedin\.com/(?:company|school)/', re.I),
                    timeout=8_000,
                )
            except Exception:
                pass

        # Fallback if the company icon/link was not clickable in the current DOM.
        if (
            not clicked_company
            or not re.search(r'/(?:company|school)/', (self.page.url or '').lower())
        ):
            self.page.goto(
                company_url,
                wait_until='domcontentloaded',
                timeout=PAGE_TIMEOUT_MS,
            )

        self._wait_for_manual_login_if_needed()
        self.page.wait_for_timeout(1_800)

        # Follow the visible About tab when possible.
        clicked_about = False
        try:
            about_links = self.page.locator('a').filter(
                has_text=re.compile(r'^\s*About\s*$', re.I)
            )
            if about_links.count() > 0:
                about_links.first.click(timeout=5_000)
                clicked_about = True
        except Exception:
            clicked_about = False

        if not clicked_about:
            base = company_url.split('?', 1)[0].split('#', 1)[0].rstrip('/')
            if not base.lower().endswith('/about'):
                base += '/about'
            self.page.goto(
                base + '/',
                wait_until='domcontentloaded',
                timeout=PAGE_TIMEOUT_MS,
            )

        try:
            self.page.wait_for_function(
                """
                () => {
                    const text = (document.body?.innerText || '')
                        .replace(/\\s+/g, ' ')
                        .trim();

                    return (
                        /\\bIndustry\\b/i.test(text) ||
                        /\\bCompany size\\b/i.test(text) ||
                        /\\bWebsite\\b/i.test(text) ||
                        /\\bHeadquarters\\b/i.test(text)
                    );
                }
                """,
                timeout=10_000,
                polling=250,
            )
        except Exception:
            pass

        self.page.wait_for_timeout(LINK_COMPANY_SETTLE_MS)
        return self.page.url

    @staticmethod
    def _company_slug_fallback_url(company_name: str) -> str:
        """Build one safe LinkedIn company-slug candidate from visible text."""
        value = unicodedata.normalize('NFKD', company_name or '')
        value = ''.join(c for c in value if not unicodedata.combining(c)).casefold()
        value = re.sub(r'[^a-z0-9]+', '-', value).strip('-')
        return f'https://www.linkedin.com/company/{value}/about/' if value else '-'

    @staticmethod
    def _company_slug_fallback_urls(company_name: str) -> list[str]:
        """Return company and school slug candidates for text-only employers.

        LinkedIn represents universities and many other education employers as
        ``/school/<slug>/`` pages, while ordinary employers use ``/company``.
        Trying both candidates lets a text-only Experience card reach the
        correct About page without guessing from the employer's industry.
        """
        value = unicodedata.normalize('NFKD', company_name or '')
        value = ''.join(c for c in value if not unicodedata.combining(c)).casefold()
        value = re.sub(r'[^a-z0-9]+', '-', value).strip('-')
        if not value:
            return []
        company_url = f'https://www.linkedin.com/company/{value}/about/'
        school_url = f'https://www.linkedin.com/school/{value}/about/'
        if re.search(r'\b(?:university|college|school|institute)\b', company_name or '', re.I):
            return [school_url, company_url]
        return [company_url, school_url]

    @timed('linkedin_company')
    def _read_standard_company_about(
        self,
        company_url: str,
        expected_company: str,
    ) -> dict[str, str]:
        """
        Read standard LinkedIn company About details:
        Company Name, Industry, Company Size, Website, Phone and Primary Address.
        """
        cache_key = "standard|" + normalize_url(company_url)

        def same_company(actual: str) -> bool:
            matched, score = self._fuzzy_company_match(expected_company, actual)
            self.log(f'LINK COMPANY MATCH | expected={expected_company!r} | actual={actual!r} | '
                     f'similarity={score:.1f}% | threshold={COMPANY_NAME_MATCH_THRESHOLD:.1f}% | accepted={matched}')
            return matched

        if cache_key in self.company_cache:
            cached = dict(self.company_cache[cache_key])
            cached_name = cached.get("Company Name", "-")

            if not same_company(cached_name):
                raise CompanyPageMismatchError(
                    "Cached standard company mismatch: "
                    f"experience='{expected_company}' | "
                    f"page='{cached_name}'"
                )

            self.log(
                "LINK COMPANY CACHE HIT | "
                f"company='{cached_name}'"
            )
            return cached

        if self.page is None or (self._browser_headless is not None and self._browser_headless != self.headless):
            self._launch_browser()
        about_url = self._open_standard_company_about(company_url)

        # Progressive scroll loads the lower Locations section.
        self._scroll_standard_page(LINK_COMPANY_MAX_SCROLLS)
        self.page.wait_for_timeout(600)

        data = self.page.evaluate(
            """
            () => {
                function clean(value) {
                    return (value || '')
                        .replace(/\\u00a0/g, ' ')
                        .replace(/[ \\t]+/g, ' ')
                        .trim();
                }

                function visible(el) {
                    if (!el) return false;
                    const r = el.getBoundingClientRect();
                    const s = getComputedStyle(el);
                    return r.width > 0 && r.height > 0 &&
                        s.display !== 'none' && s.visibility !== 'hidden';
                }

                const h1s = Array.from(document.querySelectorAll('h1'))
                    .filter(visible)
                    .map(el => clean(el.innerText || el.textContent || ''))
                    .filter(Boolean);

                const main = document.querySelector('main') || document.body;
                const about = Array.from(main.querySelectorAll('section'))
                    .find(s => s.querySelector('dl') && /Website|Industry|Company size/i.test(s.innerText))
                    || main;
                const definitions = {};
                for (const term of about.querySelectorAll('dt')) {
                    const value = term.nextElementSibling;
                    if (value?.tagName === 'DD') definitions[clean(term.innerText).toLowerCase()] = clean(value.innerText);
                }
                const links = Array.from(about.querySelectorAll('a[href]'))
                    .filter(visible)
                    .map(a => ({
                        text: clean(
                            a.innerText || a.textContent ||
                            a.getAttribute('aria-label') || ''
                        ),
                        href: a.href || ''
                    }));

                // Find the smallest rendered container that looks like the
                // Locations area and contains the Primary location.
                const locationCandidates = Array.from(
                    document.querySelectorAll('section, div')
                ).filter(visible).map(el => ({
                    el,
                    text: clean(el.innerText || '')
                })).filter(x =>
                    /^Locations\\b/i.test(x.text) &&
                    /\\bPrimary\\b/i.test(x.text) &&
                    /Get directions/i.test(x.text) &&
                    x.text.length <= 5000
                );

                locationCandidates.sort((a, b) => a.text.length - b.text.length);
                let locationNode = locationCandidates[0]?.el || null;
                const primaryCard = Array.from(main.querySelectorAll('li, [role="listitem"]'))
                    .find(el => Array.from(el.querySelectorAll('span')).some(s =>
                        /^Primary$/i.test(clean(s.innerText))) && el.querySelector('p'));
                let primaryText = '';
                if (primaryCard) {
                    // Restrict the address and direction link to one office.
                    // Descriptive headings and adjacent branches are not addresses.
                    locationNode = primaryCard;
                    primaryText = 'Primary\\n' + Array.from(primaryCard.querySelectorAll('p'))
                        .filter(visible).map(el => clean(el.innerText)).filter(Boolean).join('\\n');
                }

                return {
                    companyName: h1s[0] || '',
                    companyNames: h1s,
                    bodyText: clean(about.innerText || ''),
                    definitions,
                    links,
                    locationsText: primaryText || clean(locationNode?.innerText || ''),
                    locationLinks: locationNode ?
                        Array.from(locationNode.querySelectorAll('a[href]')).map(a => ({
                            text: clean(a.innerText || a.textContent || ''),
                            href: a.href || ''
                        })) : []
                };
            }
            """
        )

        body_text = self._clean_card_text(str(data.get('bodyText', '') or ''))
        lines = [
            self._normalize_ws(x)
            for x in body_text.splitlines()
            if self._normalize_ws(x)
        ]

        page_company = ''
        page_company_candidates = [
            self._normalize_ws(str(value or ''))
            for value in (data.get('companyNames') or [])
        ]
        if not page_company_candidates:
            page_company_candidates = [
                self._normalize_ws(str(data.get('companyName', '') or ''))
            ]

        # Prefer the visible heading that fuzzy-matches the employer from the
        # Experience card. This avoids treating a secondary H1 such as
        # ``About`` or an unrelated navigation heading as the company.
        for candidate in page_company_candidates:
            if candidate and same_company(candidate):
                page_company = candidate
                break
        if not page_company:
            page_company = next(
                (candidate for candidate in page_company_candidates if candidate),
                '',
            )

        # Title fallback for company pages where the H1 is not exposed.
        if not page_company:
            try:
                title = re.sub(r'^\(\d+\)\s*', '', self.page.title())
                page_company = re.sub(
                    r'\s*[:|]\s*(?:About\s*\|\s*)?LinkedIn.*$',
                    '',
                    title,
                    flags=re.I,
                ).strip()
            except Exception:
                page_company = ''

        if not same_company(page_company):
            raise CompanyPageMismatchError(
                'Standard LinkedIn company mismatch: '
                f"experience='{expected_company}' | page='{page_company or '-'}'"
            )

        def value_after(*labels: str) -> str:
            label_set = {x.lower() for x in labels}
            for label in label_set:
                value = (data.get('definitions') or {}).get(label)
                if value:
                    return self._normalize_ws(value)
            for i, line in enumerate(lines):
                low = line.lower().strip()
                for label in label_set:
                    if low == label or low.rstrip(':') == label:
                        if i + 1 < len(lines) and lines[i + 1]:
                            return lines[i + 1]
                    # Some LinkedIn layouts render the label and value in one
                    # text node (for example ``Website: ecu.edu``).
                    prefix = label + ':'
                    if low.startswith(prefix):
                        inline = line[len(prefix):].strip()
                        if inline:
                            return inline
                    # A few responsive layouts omit the colon but keep the
                    # label and value in the same rendered line.
                    match = re.match(
                        rf'^{re.escape(label)}\s+(.+)$',
                        line,
                        flags=re.I,
                    )
                    if match and match.group(1).strip():
                        return match.group(1).strip()
            return '-'

        industry = value_after('Industry')
        company_size = value_after('Company size', 'Company Size')
        website = value_after('Website')
        phone = value_after('Phone')

        links = data.get('links', []) or []
        phone_link = '-'
        company_size_link = '-'

        employee_re = EMPLOYEE_SIZE_RE

        # Prefer actual clickable link targets for website, phone and employee
        # count when LinkedIn exposes them as anchors.
        for item in links:
            label = self._normalize_ws(str(item.get('text', '') or ''))
            href = str(item.get('href', '') or '').strip()

            if href.lower().startswith('tel:'):
                phone_link = href
                phone = self._normalize_maps_phone(href[4:]) or '-'

            employee_match = employee_re.search(label)
            if employee_match:
                if company_size == '-':
                    company_size = employee_match.group(0)
                if href:
                    company_size_link = href

            if website != '-' and label == website and href.startswith(('http://', 'https://')):
                if 'linkedin.com' not in href.lower():
                    website = href

        if website == '-':
            for item in links:
                label = self._normalize_ws(str(item.get('text', '') or ''))
                href = str(item.get('href', '') or '').strip()
                if (
                    href.startswith(('http://', 'https://'))
                    and 'linkedin.com' not in href.lower()
                    and ('website' in label.lower() or label.lower().startswith('http'))
                ):
                    website = href
                    break

        # Primary address from Locations section.
        address = '-'
        address_link = '-'
        city = state = country = '-'

        loc_text = self._clean_card_text(
            str(data.get('locationsText', '') or '')
        )
        loc_lines = [
            self._normalize_ws(x)
            for x in loc_text.splitlines()
            if self._normalize_ws(x)
        ]

        primary_idx = next(
            (i for i, x in enumerate(loc_lines) if x.lower() == 'primary'),
            -1,
        )

        def geography(parts: list[str]) -> tuple[str, str, str]:
            """Parse address components without treating a state as a country."""
            parts = [p.strip(" .") for p in parts if p.strip(" .")]
            if not parts:
                return '-', '-', '-'
            last = self._canonical_country(parts[-1])
            if len(parts) >= 4 and self._looks_like_country_token(last):
                region = parts[-2]
                if re.fullmatch(r'\d{4,6}(?:-\d{4})?', region):
                    region = '-'
                return parts[-3], region, last
            if len(parts) == 3:
                if self._looks_like_country_token(last):
                    if re.fullmatch(r'\d{4,6}(?:-\d{4})?', parts[-2]):
                        return parts[0], '-', self._canonical_country(last)
                    if self._looks_like_state_region_token(parts[-2]):
                        return parts[-3], parts[-2], self._canonical_country(last)
                    return parts[-2], '-', self._canonical_country(last)
                if self._looks_like_state_region_token(last):
                    return parts[-2], last, '-'
            if len(parts) == 2:
                if self._looks_like_country_token(last):
                    return parts[0], '-', self._canonical_country(last)
                if self._looks_like_state_region_token(last):
                    return parts[0], last, '-'
            return '-', '-', '-'

        if primary_idx >= 0:
            # LinkedIn can render the "Primary" badge either before or after
            # the building/address text. Capture a small window around it so:
            #
            #   Torre Madero
            #   Macacha Guemes 515, Capital Federal, Buenos Aires ., AR
            #
            # is preserved as a two-line Address.
            address_lines: list[str] = []

            start_idx = max(0, primary_idx - 4)
            end_idx = min(len(loc_lines), primary_idx + 7)

            for value in loc_lines[start_idx:end_idx]:
                low = value.lower().strip()

                if low == 'primary':
                    continue
                if re.match(r'^locations?\b', low):
                    continue
                if low.startswith('get directions'):
                    continue
                if low in {'show all locations', 'see all locations'}:
                    continue

                if value and value not in address_lines:
                    address_lines.append(value)

            if address_lines:
                address = '\n'.join(address_lines)

                # Use the line with the richest comma-separated geography for
                # City / State / Country instead of assuming where the building
                # name appears.
                geo_source = max(
                    address_lines,
                    key=lambda item: item.count(','),
                )

                parts = [
                    p.strip()
                    for p in geo_source.split(',')
                    if p.strip()
                ]

                city, state, country = geography(parts)

        for item in data.get('locationLinks', []) or []:
            label = self._normalize_ws(str(item.get('text', '') or ''))
            href = str(item.get('href', '') or '').strip()
            if 'get directions' in label.lower() and href:
                if href.rstrip().endswith('where='):
                    continue
                address_link = href
                break

        # Headquarters is a fallback only when Locations/Primary was not loaded.
        if address == '-':
            headquarters = value_after('Headquarters')
            if headquarters != '-':
                address = headquarters

                hq_parts = [
                    p.strip(" .")
                    for p in headquarters.split(",")
                    if p.strip(" .")
                ]

                city, state, country = geography(hq_parts)

        fields = {
            'Company Name': page_company or expected_company or '-',
            'Company Website': website,
            'Industry': industry,
            'Company Size': company_size,
            'See All Company Size': company_size_link,
            'Industry Link': about_url,
            'Address': address,
            'City': city,
            'State': state,
            'Zip Code': '-',
            'Address Link': address_link,
            'Country': country,
            'Phone Number': phone,
            'Phone Number Link': phone_link,
            'Revenue': '-',
            'Revenue Link': '-',
            'SIC Code': '-',
            'NAICS Code': '-',
            'SIC/NAICS Code Link': '-',
        }

        # US About addresses often combine region and postal code in one
        # component, e.g. NC 27858-4353. Keep those in separate columns.
        if country == 'United States':
            postal_match = re.fullmatch(r'(.+?)\s+(\d{5}(?:-\d{4})?)', state)
            if postal_match:
                fields['State'] = postal_match.group(1)
                fields['Zip Code'] = postal_match.group(2)
        elif country == 'Switzerland' and address != '-':
            postal_match = re.search(r'\b(\d{4})\s*,\s*CH\b', address, re.I)
            if postal_match:
                fields['Zip Code'] = postal_match.group(1)

        self.company_cache[cache_key] = dict(fields)

        self.log(
            'LINK COMPANY | '
            f"experience='{expected_company}' | page='{fields['Company Name']}' | "
            f"industry='{industry}' | size='{company_size}'"
        )

        return fields

    def _process_link_file(self, payload: dict[str, Any]):
        """Mixed-link workflow: resolve, deduplicate, then use the standard scraper."""
        self._configure_run(payload)
        input_path: str = payload['input_path']
        requested_max: int = payload['max_records']
        workbook: WorkbookManager = payload['workbook']
        self.email_domain_checks_enabled = bool(payload.get('email_domain_checks', False))
        extraction_mode: str = payload.get(
            'extraction_mode',
            'FAST',
        ).upper()

        self.dom_calls = 0
        self.dom_seconds = 0.0
        self.qwen_calls = 0
        self.qwen_seconds = 0.0
        self._reset_public_company_enricher()
        self.company_cache = {}
        self.smart_company_cache.clear()
        self.smart_job_cache.clear()
        self.website_evidence_cache.clear()
        self.google_maps_cache = {}
        self._install_caches()
        self.stop_event.clear()

        self.llm_available = None
        self.llm_last_check_ts = 0.0

        if extraction_mode == "SMART":
            llm_ok, model = self._ensure_llm_available(
                force=True
            )

            if llm_ok:
                self.log(
                    f"LINK SMART mode: Qwen connected: {model}"
                )
            else:
                self.log(
                    "LINK SMART mode: initial Qwen check failed. "
                    "The app will retry automatically during the run and "
                    "again before the final Excel save."
                )
        else:
            self.llm_available = None

        links = self._load_link_inputs(input_path)
        self.link_resolutions = LinkResolutionStore(
            workbook.output_path.parent / '.linkedin_link_mappings.sqlite3')
        if not links:
            raise ValueError(
                'No URLs were found. Select Excel/CSV containing LinkedIn /in/ or /sales/lead/ links.'
            )

        unique_count = len({item.canonical for item in links if item.canonical})
        target = min(requested_max, unique_count) if requested_max is not None else unique_count

        self.log(
            f"LINK SCRAPER | input='{input_path}' | links_found={len(links)} | "
            f"target_saved={target}"
        )

        processed = saved = duplicates = failed = skipped = 0
        consecutive_errors = 0
        pending_enrichment_rows: list[
            tuple[int, dict[str, str]]
        ] = []

        self.callbacks['progress'](
            processed, saved, duplicates, failed, target
        )

        handled_profiles = set()
        for index, item in enumerate(links, 1):
            if self.stop_event.is_set() or (target > 0 and saved >= target):
                break
            if not self._wait_if_paused():
                break

            profile_url, resolution = '', 'Unresolved'
            def audit(processing, detail=''):
                workbook.record_link_resolution(item.original, profile_url, item.kind,
                                                resolution, processing, detail)
            try:
                profile_url, resolution = self._resolve_profile_link(item)
            except ProfileAccessError as exc:
                resolution = 'Access required'
                audit('Stopped', str(exc))
                self.log(f'LINK ACCESS REQUIRED | {exc}')
                self.stop_event.set()
                break
            except Exception as exc:
                failed += 1
                processed += 1
                resolution = 'Invalid' if not item.canonical else 'Unresolved'
                audit('Stopped' if self.stop_event.is_set() else 'Not processed', str(exc))
                self.log(f"LINK RESOLUTION FAILED | url='{item.original}' | {exc}")
                self.callbacks['progress'](processed, saved, duplicates, failed, target)
                continue
            audit('Ready')
            if profile_url in handled_profiles:
                duplicates += 1
                processed += 1
                audit('Duplicate', 'This public profile was already handled in this run.')
                self.callbacks['progress'](processed, saved, duplicates, failed, target)
                continue
            retry_row = workbook.find_link_row(profile_url)
            if retry_row is None and item.kind == 'Sales Navigator lead':
                retry_row = workbook.find_link_row(item.canonical)
                if retry_row is not None:
                    previous = workbook.read_record(retry_row)
                    previous['Prospect LinkedIn Link'] = profile_url
                    # Old companion data remains attributable to its original URL.
                    previous.pop('_profile_state', None)
                    workbook.update_record(retry_row, previous)
                    workbook.checkpoint()

            if (retry_row is not None and not self.refresh_profiles
                    and not workbook.profile_needs_visit(profile_url)):
                handled_profiles.add(profile_url)
                self.metrics.count('profile_visits_avoided')
                try:
                    record = self._enrich_saved_company_row(
                        workbook, profile_url, source_label="Link Scraper existing row",
                    )
                except ProfileAccessError as exc:
                    audit('Access required', str(exc))
                    self.stop_event.set()
                    break
                except Exception as exc:
                    audit('Failed', str(exc))
                    failed += 1
                    processed += 1
                    self.log(f'LINK ENRICHMENT ERROR | {exc}')
                    self.callbacks['progress'](processed, saved, duplicates, failed, target)
                    continue
                if record is not None:
                    audit('Updated')
                    saved += 1
                    processed += 1
                    self.callbacks['record'](record)
                    self.callbacks['harvest'](saved, target)
                    self.callbacks['progress'](
                        processed, saved, duplicates, failed, target,
                    )
                    continue
                audit('Already saved', 'Existing profile reused; no changed enrichment.')
                duplicates += 1
                processed += 1
                self.callbacks['progress'](
                    processed, saved, duplicates, failed, target
                )
                continue

            self.status(
                f"Link Scraper: profile {index}/{len(links)} | saved {saved}/{target}"
            )

            record = self._base_record(profile_url)
            record['Source File Name'] = Path(input_path).name

            try:
                profile_data = self._capture_profile_once(profile_url)
                current = self._parse_standard_current_experience(profile_data)
                if any(current.get(field) in (None, '', '-', 'Dead')
                       for field in ('Tenure', 'Job Title', 'Company Name')):
                    raise RuntimeError('Current role, tenure, or employer is unresolved; incomplete profile was not saved.')

                record['First Name'] = current.get('First Name', '-')
                record['Last Name'] = current.get('Last Name', '-')
                record['Job Title'] = current.get('Job Title', '-')
                record['Job Level'] = self._classify_job_level(record['Job Title'])
                record['Job Function'] = self._classify_job_function(record['Job Title'])
                record['Company Name'] = current.get('Company Name', '-')
                record['Tenure'] = current.get('Tenure', 'Dead')

                role_location = current.get('Location', '-')
                if role_location and role_location != '-':
                    c, s, country = self._location_to_fields(
                        role_location,
                        current.get('Profile Location', '-'),
                    )
                    record['City'] = c
                    record['State'] = s
                    record['Country'] = country

                self.log(
                    'LINK PROFILE | '
                    f"name='{record['First Name']} {record['Last Name']}' | "
                    f"title='{record['Job Title']}' | "
                    f"company='{record['Company Name']}' | "
                    f"tenure='{record['Tenure']}' | "
                    f"location='{role_location}'"
                )

                if record['Tenure'].lower() != 'dead':
                    company_url = current.get('Company Profile Link', '-')
                    expected_company = current.get('Company Name', '-')
                    record['Company LinkedIn URL'] = company_url

                    if company_url == '-' or expected_company == '-':
                        # Public/detail pages sometimes expose the employer as
                        # text but omit its company hyperlink. Keep the reliable
                        # role data and try one name-derived company URL.
                        self.log(
                            'LINK COMPANY LINK MISSING | trying a name-derived company URL | '
                            f"profile='{profile_url}' | company='{expected_company}'"
                        )

                    company_candidates = []
                    if company_url != '-':
                        company_candidates.append(company_url)
                    elif expected_company != '-':
                        company_candidates.extend(
                            self._company_slug_fallback_urls(expected_company)
                        )

                    company_fields = None
                    for candidate_company_url in company_candidates:
                        try:
                            company_fields = self._read_standard_company_about(
                                company_url=candidate_company_url,
                                expected_company=expected_company,
                            )
                            company_url = candidate_company_url
                            record['Company LinkedIn URL'] = company_url
                            break
                        except ProfileAccessError:
                            raise
                        except CompanyPageMismatchError as exc:
                            self.log(
                                f"LINK COMPANY CANDIDATE REJECTED | {exc}"
                            )
                        except Exception as exc:
                            self.log(
                                'LINK COMPANY CANDIDATE FAILED | '
                                f"url='{candidate_company_url}' | {type(exc).__name__}: {exc}"
                            )

                    if company_fields:
                        for field, value in company_fields.items():
                            if field in record and value != '-':
                                if field in {'City', 'State', 'Country'} and record[field] != '-':
                                    continue
                                # Company page is authoritative for company fields,
                                # including Name/Industry/Size/Website/Phone/Address.
                                record[field] = value

                record['_profile_state'] = {
                    'profile': profile_url, 'company': record.get('Company Name', '-'),
                    'company_url': record.get('Company LinkedIn URL', '-'),
                    'captured_at': profile_data.get('_captured_at', datetime.now(timezone.utc).isoformat()),
                }
                if retry_row is not None:
                    previous = dict(zip(HEADERS, [cell.value for cell in workbook.ws[retry_row]]))
                    record = self._retain_same_employer_email(record, previous)
                record = self._enrich_offsite(record, "Link Scraper")

                if extraction_mode == "SMART":
                    record = self._smart_validate_record(
                        record,
                        source_label="Link Scraper",
                        validate_existing=True,
                    )
                else:
                    self.trace(
                        "ENRICHMENT",
                        "fast_mode_llm_skipped",
                        source="Link Scraper",
                        record=self._record_log_context(record),
                        reason=(
                            "FAST mode uses deterministic scraping + "
                            "Google Maps + public company sources; Local LLM enrichment is disabled."
                        ),
                    )

                if retry_row is not None:
                    previous = dict(zip(HEADERS, [cell.value for cell in workbook.ws[retry_row]]))
                    employer_fields = {
                        "Business Email ID", "Email Status", "Bounce Email Status",
                        'Company Name', 'Company Website', 'Industry', 'Company Size',
                        'See All Company Size', 'Industry Link', 'Address', 'Address Link',
                        'City', 'State', 'Country', 'Zip Code', 'Phone Number',
                        'Phone Number Link', 'Revenue', 'Revenue Link', 'SIC Code',
                        'NAICS Code', 'SIC/NAICS Code Link',
                    }
                    for field, old_value in previous.items():
                        # A repair must not resurrect stale enrichment merely
                        # because the employer name stayed the same. Its old
                        # address/phone may belong to a different branch.
                        if field in employer_fields:
                            continue
                        if field == 'Notes' and str(old_value).lower().startswith(
                            ('google maps enriched:', 'llm enrichment pending:')
                        ):
                            continue
                        if field == 'Source File Name' and old_value not in (None, '', '-'):
                            record[field] = old_value
                            continue
                        if record.get(field) in (None, '', '-') and old_value not in (None, '', '-', 'Dead'):
                            record[field] = old_value
                    workbook.update_record(retry_row, record)
                    workbook.checkpoint()
                    row_idx = retry_row
                else:
                    row_idx = workbook.add_record(record)
                self._remember_enrichment(record)
                handled_profiles.add(profile_url)
                audit('Updated' if retry_row is not None else 'Saved')
                consecutive_errors = 0

                if "llm enrichment pending:" in clean_cell_value(
                    record.get("Notes", "-")
                ).lower():
                    pending_enrichment_rows.append(
                        (row_idx, dict(record))
                    )

                saved += 1
                processed += 1

                self.callbacks['record'](record)
                self.callbacks['harvest'](saved, target)
                self.callbacks['progress'](
                    processed, saved, duplicates, failed, target
                )

                self.log(
                    f"LINK SAVED {saved}/{target} | "
                    f"{record['First Name']} {record['Last Name']} | "
                    f"{record['Company Name']}"
                )

            except ProfileAccessError as exc:
                audit('Access required', str(exc))
                self.log(f'LINK ACCESS REQUIRED | {exc}')
                self.stop_event.set()
                break
            except Exception as exc:
                failed += 1
                processed += 1
                consecutive_errors += 1
                audit('Failed', str(exc))
                self.log(
                    'LINK PROFILE ERROR | '
                    f"url='{profile_url}' | {type(exc).__name__}: {exc}"
                )
                self.callbacks['progress'](
                    processed, saved, duplicates, failed, target
                )
                if consecutive_errors >= 3:
                    self.log('LINK STOPPED | Three consecutive extraction failures. Inspect the browser and activity log before retrying.')
                    self.stop_event.set()
                    break

        if extraction_mode == "SMART":
            self._retry_pending_enrichment_rows(
                workbook=workbook,
                pending_rows=pending_enrichment_rows,
                source_label="Link Scraper",
                validate_existing=True,
            )

        self._log_enrichment_audit(
            workbook,
            "Link Scraper",
        )

        workbook.save(force=True)

        self.log(
            'LINK FINAL | '
            f"processed={processed} | saved={saved}/{target} | "
            f"skipped={skipped} | duplicates={duplicates} | failed={failed} | "
            f"Qwen calls={self.qwen_calls} | "
            f"Qwen time={self.qwen_seconds:.1f}s | "
            f"company cache={len(self.company_cache)} | "
            f"smart company cache={len(self.smart_company_cache)} | "
            f"smart job cache={len(self.smart_job_cache)}"
        )
        self.status('Stopped' if self.stop_event.is_set() else 'Finished')
        self._finish_enrichment_sessions()
        self._report_performance(workbook)
        self.callbacks['done']()

    def _base_record(self, lead_url: str) -> dict[str, str]:
        record = {header: "-" for header in HEADERS}
        record["Date"] = time.strftime("%d-%b-%y")
        record["Prospect LinkedIn Link"] = lead_url
        return record

    def _process_search(self, payload: dict[str, Any]):
        """
        V11 STRICT PRESENT / COMPANY-MATCH WORKFLOW

        Important change:
        A prospect row is ALWAYS written when the search-card data was read,
        even if:
        - side panel fails to open,
        - Present block cannot be parsed,
        - company page times out,
        - company enrichment fails.

        This prevents an empty Excel workbook when one enrichment stage fails.
        """
        self._configure_run(payload)
        search_url: str = payload["search_url"]
        max_records: int = payload["max_records"]
        workbook: WorkbookManager = payload["workbook"]
        self.email_domain_checks_enabled = bool(payload.get('email_domain_checks', False))
        extraction_mode: str = payload.get(
            "extraction_mode",
            "FAST",
        ).upper()

        self._launch_browser()

        self.dom_calls = 0
        self.dom_seconds = 0.0
        self.qwen_calls = 0
        self.qwen_seconds = 0.0
        self._reset_public_company_enricher()
        self.company_cache = {}
        self.smart_company_cache.clear()
        self.smart_job_cache.clear()
        self.website_evidence_cache.clear()
        self.google_maps_cache = {}
        self._install_caches()

        self.llm_available = None
        self.llm_last_check_ts = 0.0

        if extraction_mode == "SMART":
            llm_ok, model = self._ensure_llm_available(
                force=True
            )

            if llm_ok:
                self.log(
                    f"SMART mode: Qwen connected: {model}"
                )
            else:
                self.log(
                    "SMART mode: initial Qwen check failed. "
                    "The app will retry automatically during the run and "
                    "again before the final Excel save."
                )
        else:
            self.llm_available = None
            self.log(
                "FAST mode: LinkedIn + Google Maps + public company sources. "
                "Local Qwen/LLM is completely disabled for enrichment."
            )

        self.stop_event.clear()

        self.status("Opening Sales Navigator search")
        self.log(f"Search URL: {search_url}")

        self.page.goto(
            search_url,
            wait_until="domcontentloaded",
            timeout=PAGE_TIMEOUT_MS,
        )

        self._wait_for_manual_login_if_needed()
        self.page.wait_for_timeout(PAGE_SETTLE_MS)

        seen: set[str] = set()

        processed = 0
        saved = 0
        duplicates = 0
        failed = 0
        skipped = 0
        stagnant = 0
        pending_enrichment_rows: list[
            tuple[int, dict[str, str]]
        ] = []

        run_started = time.perf_counter()

        self.callbacks["progress"](
            processed,
            saved,
            duplicates,
            failed,
            max_records,
        )

        for _scroll_round in range(MAX_SCROLL_ROUNDS):
            if self.stop_event.is_set():
                break

            if not self._wait_if_paused():
                break

            cards = self._collect_current_cards()

            self.log(
                f"BATCH | visible lead cards detected={len(cards)}"
            )

            current_urls = [
                card.get("url", "")
                for card in cards
                if card.get("url")
            ]

            new_cards = 0

            for card in cards:
                if self.stop_event.is_set() or saved >= max_records:
                    break

                if not self._wait_if_paused():
                    break

                lead_url = card["url"]
                key = normalize_url(lead_url)

                if not key or key in seen:
                    continue

                seen.add(key)
                new_cards += 1

                existing_row = workbook.find_link_row(lead_url)
                if (existing_row is not None and not self.refresh_profiles
                        and not workbook.profile_needs_visit(lead_url)):
                    self.metrics.count('profile_visits_avoided')
                    record = self._enrich_saved_company_row(
                        workbook, lead_url, source_label="Sales Navigator existing row",
                    )
                    if record is not None:
                        saved += 1
                        processed += 1
                        self.callbacks["record"](record)
                        self.callbacks["harvest"](saved, max_records)
                        self.callbacks["progress"](
                            processed, saved, duplicates, failed, max_records,
                        )
                        continue
                    duplicates += 1
                    processed += 1

                    self.callbacks["progress"](
                        processed,
                        saved,
                        duplicates,
                        failed,
                        max_records,
                    )
                    continue

                # ---------------------------------------------------
                # STEP 1: SEARCH-CARD FALLBACK RECORD
                # ---------------------------------------------------
                try:
                    card_data = self._parse_card_fast(
                        card,
                        extraction_mode=extraction_mode,
                    )
                except Exception as exc:
                    # Even if the detailed parser fails, preserve the URL.
                    card_data = {
                        "First Name": "-",
                        "Last Name": "-",
                        "Job Title": "-",
                        "Job Level": "-",
                        "Job Function": "-",
                        "Company Name": card.get("company", "-") or "-",
                        "Tenure": "-",
                        "City": "-",
                        "State": "-",
                        "Country": "-",
                        "Description": "-",
                    }
                    self.log(
                        f"CARD PARSE WARNING | {type(exc).__name__}: {exc}"
                    )

                first_name = card_data.get("First Name", "-")
                last_name = card_data.get("Last Name", "-")
                search_company = card_data.get("Company Name", "-")

                person_name = self._normalize_ws(
                    " ".join(
                        x for x in (first_name, last_name)
                        if x and x != "-"
                    )
                )

                self.log(
                    "CARD | "
                    f"name='{person_name or '-'}' | "
                    f"company='{search_company}' | "
                    f"title='{card_data.get('Job Title', '-')}'"
                )

                # Build the row NOW, before side-panel/company enrichment.
                # Therefore a later failure cannot produce an empty workbook.
                record = self._base_record(lead_url)

                record["First Name"] = first_name
                record["Last Name"] = last_name
                record["Prospect LinkedIn Link"] = lead_url

                # Search-card values are provisional fallbacks.
                record["Job Title"] = card_data.get("Job Title", "-")
                record["Job Level"] = card_data.get("Job Level", "-")
                record["Job Function"] = card_data.get("Job Function", "-")
                record["Company Name"] = search_company
                record["Tenure"] = card_data.get("Tenure", "-")

                # PRIMARY LOCATION SOURCE:
                # Use the location shown directly on the current Sales
                # Navigator profile/result card first.
                record["City"] = card_data.get("City", "-")
                record["State"] = card_data.get("State", "-")
                record["Country"] = card_data.get("Country", "-")

                primary_card_location = {
                    "City": record["City"],
                    "State": record["State"],
                    "Country": record["Country"],
                }

                if any(
                    not self._is_unknown_value(value)
                    for value in primary_card_location.values()
                ):
                    self.log(
                        "PROFILE LOCATION PRIMARY | "
                        "source='search card' | "
                        f"city='{record['City']}' | "
                        f"state='{record['State']}' | "
                        f"country='{record['Country']}'"
                    )

                    self.trace(
                        "LOCATION",
                        "profile_primary_search_card",
                        record=self._record_log_context(record),
                        values=primary_card_location,
                    )

                record["Description"] = card_data.get("Description", "-")

                company_match = False
                present_found = False
                panel_error = None
                company_error = None
                skip_record = False

                # ---------------------------------------------------
                # STEP 2: SIDE-PANEL CURRENT ROLE
                # ---------------------------------------------------
                try:
                    self.status(
                        f"Opening {person_name or 'prospect'} "
                        f"({saved + 1}/{max_records})"
                    )

                    self._open_lead_panel(
                        lead_url=lead_url,
                        person_name=person_name,
                    )

                    panel = self._capture_profile_panel(person_name)
                    self.profile_cache['sales|' + key] = {
                        **panel, '_captured_at': datetime.now(timezone.utc).isoformat()}

                    panel_text = panel.get("text", "")
                    panel_links = panel.get("links", [])

                    self.log(
                        "PANEL | "
                        f"name='{person_name or '-'}' | "
                        f"text_chars={len(panel_text)} | "
                        f"links={len(panel_links)}"
                    )

                    top_location = self._extract_top_profile_location(
                        panel_text=panel_text,
                        person_name=person_name,
                    )

                    # The opened prospect panel is still a direct profile
                    # source, so it may fill only geography fields that the
                    # search card did not provide.
                    panel_city, panel_state, panel_country = (
                        self._profile_location_to_fields(
                            top_location
                        )
                    )

                    panel_location_fields = {
                        "City": panel_city,
                        "State": panel_state,
                        "Country": panel_country,
                    }

                    panel_filled = (
                        self._fill_missing_location_fields(
                            record,
                            panel_location_fields,
                        )
                    )

                    if panel_filled:
                        self.log(
                            "PROFILE LOCATION PRIMARY | "
                            "source='profile panel' | "
                            f"filled={panel_filled} | "
                            f"city='{record['City']}' | "
                            f"state='{record['State']}' | "
                            f"country='{record['Country']}'"
                        )

                        self.trace(
                            "LOCATION",
                            "profile_primary_panel",
                            record=self._record_log_context(record),
                            raw_location=top_location,
                            parsed=panel_location_fields,
                            filled=panel_filled,
                        )

                    experience_blocks = self._parse_experience_blocks(
                        panel_text=panel_text,
                        panel_links=panel_links,
                    )

                    self.log(
                        "PRESENT SEARCH | "
                        f"name='{person_name or '-'}' | "
                        f"present_blocks={len(experience_blocks)}"
                    )

                    # ---------------------------------------------------
                    # STRICT USER FLOW
                    # ---------------------------------------------------
                    # Search card already supplied:
                    #   Name / Job Title / Company
                    #
                    # Side panel now supplies:
                    #   Present tenure and current-company validation.
                    # ---------------------------------------------------
                    if not experience_blocks:
                        # Literal Present keyword not found.
                        # Keep card Name/Title/Company and only mark Tenure Dead.
                        record["Tenure"] = "Dead"

                        self.log(
                            "DEAD | "
                            f"name='{person_name or '-'}' | "
                            "no Present keyword found"
                        )

                        existing_notes = (
                            record["Notes"]
                            if record["Notes"] != "-"
                            else ""
                        )

                        record["Notes"] = (
                            (existing_notes + " | " if existing_notes else "")
                            + "No Present keyword found; Tenure=Dead"
                        )

                    else:
                        present_found = True

                        # More than one current role can exist. Select the
                        # Present block whose company matches the company copied
                        # from this prospect's search-result card.
                        current = None

                        for block in experience_blocks:
                            panel_company = block.get("Company Name", "-")

                            if self._company_names_match(
                                search_company,
                                panel_company,
                            ):
                                current = block
                                break

                        if current is None:
                            present_companies = [
                                block.get("Company Name", "-")
                                for block in experience_blocks
                            ]

                            self.log(
                                "SKIPPED COMPANY MISMATCH | "
                                f"name='{person_name or '-'}' | "
                                f"search_company='{search_company}' | "
                                f"present_companies={present_companies}"
                            )

                            # User rule: mismatch -> move directly to next lead.
                            skip_record = True

                        else:
                            present_company = current.get("Company Name", "-")
                            current_tenure = current.get("Tenure", "-")
                            company_match = True

                            # Copy the whole line containing Present.
                            record["Tenure"] = current_tenure

                            # Keep the Job Title copied from the search card.
                            # Use the panel title only when card extraction failed.
                            if (
                                record.get("Job Title", "-") == "-"
                                and current.get("Job Title", "-") != "-"
                            ):
                                record["Job Title"] = current["Job Title"]
                                record["Job Level"] = self._classify_job_level(
                                    record["Job Title"]
                                )
                                record["Job Function"] = self._classify_job_function(
                                    record["Job Title"]
                                )

                            record["Company Name"] = search_company

                            self.log(
                                "COMPANY MATCH | "
                                f"search='{search_company}' | "
                                f"present='{present_company}' | "
                                f"tenure='{current_tenure}'"
                            )

                            # Profile/search-card geography remains the
                            # primary source. The company page is used only as
                            # a field-level fallback for missing City / State /
                            # Country values.

                            # Matched company: click its exact company link.
                            company_url = current.get(
                                "Company Profile Link",
                                "-",
                            )
                            record["Company LinkedIn URL"] = company_url

                            if company_url != "-":
                                try:
                                    self.status(
                                        f"Opening matched company: {search_company}"
                                    )

                                    company_fields = self._read_company_details(
                                        company_url=company_url,
                                        company_name=search_company,
                                        extraction_mode=extraction_mode,
                                    )

                                    # COMPANY LOCATION FALLBACK:
                                    # Fill only missing profile geography. Never
                                    # replace a valid City/State/Country already
                                    # found on the current profile/page.
                                    company_location_fields = {
                                        "City": company_fields.get(
                                            "City",
                                            "-",
                                        ),
                                        "State": company_fields.get(
                                            "State",
                                            "-",
                                        ),
                                        "Country": company_fields.get(
                                            "Country",
                                            "-",
                                        ),
                                    }

                                    company_location_filled = (
                                        self._fill_missing_location_fields(
                                            record,
                                            company_location_fields,
                                        )
                                    )

                                    if company_location_filled:
                                        self.log(
                                            "COMPANY LOCATION FALLBACK | "
                                            f"filled={company_location_filled} | "
                                            f"city='{record['City']}' | "
                                            f"state='{record['State']}' | "
                                            f"country='{record['Country']}'"
                                        )

                                        self.trace(
                                            "LOCATION",
                                            "company_page_fallback",
                                            record=self._record_log_context(
                                                record
                                            ),
                                            company_location=(
                                                company_location_fields
                                            ),
                                            filled=(
                                                company_location_filled
                                            ),
                                        )

                                    for field, value in company_fields.items():
                                        if field not in record or value == "-":
                                            continue

                                        if field in {
                                            "City",
                                            "State",
                                            "Country",
                                        }:
                                            # Already handled above with
                                            # field-level fallback.
                                            continue

                                        if field in {
                                            "Company Name",
                                            "Company Website",
                                            "Industry",
                                            "Company Size",
                                            "See All Company Size",
                                        }:
                                            record[field] = value
                                        elif record[field] == "-":
                                            record[field] = value

                                except CompanyPageMismatchError as exc:
                                    company_error = str(exc)
                                    skip_record = True

                                    self.log(
                                        "SKIPPED COMPANY PAGE MISMATCH | "
                                        f"name='{person_name or '-'}' | "
                                        f"{company_error}"
                                    )

                                except Exception as exc:
                                    company_error = (
                                        f"{type(exc).__name__}: {exc}"
                                    )

                                    self.log(
                                        "COMPANY ENRICHMENT WARNING | "
                                        f"{company_error}"
                                    )
                            else:
                                skip_record = True
                                self.log(
                                    "SKIPPED NO COMPANY LINK | "
                                    f"name='{person_name or '-'}' | "
                                    f"company='{search_company}' | "
                                    "matched company text found but no company URL was captured, "
                                    "so the company page could not be verified"
                                )

                except Exception as exc:
                    panel_error = f"{type(exc).__name__}: {exc}"

                    # DO NOT discard the search-card record.
                    # The row will still be written below.
                    self.log(
                        "PANEL EXTRACTION WARNING | "
                        f"name='{person_name or '-'}' | "
                        f"{panel_error}"
                    )

                    existing_notes = (
                        record["Notes"]
                        if record["Notes"] != "-"
                        else ""
                    )

                    record["Notes"] = (
                        (existing_notes + " | " if existing_notes else "")
                        + f"Panel extraction failed: {panel_error}"
                    )

                finally:
                    # Reset UI state for the next prospect.
                    self._close_lead_panel()

                # ---------------------------------------------------
                # STEP 4: SAVE OR MOVE TO NEXT
                # ---------------------------------------------------
                if skip_record:
                    skipped += 1
                    processed += 1

                    self.log(
                        f"SKIP COUNT | skipped={skipped} | "
                        f"saved={saved}/{max_records} | "
                        f"processed={processed}"
                    )

                    self.callbacks["progress"](
                        processed,
                        saved,
                        duplicates,
                        failed,
                        max_records,
                    )

                    continue

                if record.get("Revenue", "-") != "-" and record.get("Revenue Link", "-") == "-":
                    record["Revenue Link"] = record.get("Industry Link", "-")
                if company_error:
                    self._append_enrichment_note(record, "Company extraction failed: " + company_error)

                try:
                    record['_profile_state'] = {
                        'profile': lead_url, 'company': record.get('Company Name', '-'),
                        'company_url': record.get('Company LinkedIn URL', '-'),
                        'captured_at': datetime.now(timezone.utc).isoformat(),
                    }
                    record = self._enrich_offsite(record, "Sales Navigator")

                    if extraction_mode == "SMART":
                        record = self._smart_validate_record(
                            record,
                            source_label="Sales Navigator",
                            validate_existing=True,
                        )
                    else:
                        self.trace(
                            "ENRICHMENT",
                            "fast_mode_llm_skipped",
                            source="Sales Navigator",
                            record=self._record_log_context(record),
                            reason=(
                                "FAST mode uses deterministic scraping + "
                                "Google Maps + public company sources; Local LLM enrichment is disabled."
                            ),
                        )

                    if existing_row is not None:
                        previous = workbook.read_record(existing_row)
                        for field in HEADERS[:8]:
                            if self._is_unknown_value(record.get(field)):
                                record[field] = previous.get(field, '-')
                        self._append_enrichment_note(record, previous.get('Notes', '-'))
                        workbook.update_record(existing_row, record)
                        workbook.checkpoint()
                        row_idx = existing_row
                    else:
                        row_idx = workbook.add_record(record)
                    self._remember_enrichment(record)

                    if "llm enrichment pending:" in clean_cell_value(
                        record.get("Notes", "-")
                    ).lower():
                        pending_enrichment_rows.append(
                            (row_idx, dict(record))
                        )

                    saved += 1
                    processed += 1

                    self.log(
                        f"SAVED {saved}/{max_records} | "
                        f"{record['First Name']} {record['Last Name']} | "
                        f"title='{record['Job Title']}' | "
                        f"company='{record['Company Name']}' | "
                        f"tenure='{record['Tenure']}' | "
                        f"present={'YES' if present_found else 'NO'} | "
                        f"company_match={'YES' if company_match else 'NO'}"
                    )

                    self.callbacks["record"](record)
                    self.callbacks["harvest"](saved, max_records)

                except Exception as exc:
                    failed += 1
                    processed += 1

                    self.log(
                        "EXCEL SAVE FAILED | "
                        f"{type(exc).__name__}: {exc}"
                    )

                self.callbacks["progress"](
                    processed,
                    saved,
                    duplicates,
                    failed,
                    max_records,
                )

            if saved >= max_records:
                break

            if new_cards == 0:
                stagnant += 1
            else:
                stagnant = 0

            if stagnant >= MAX_STAGNANT_SCROLLS:
                # We have exhausted the currently visible/paginated result set.
                # The target is SAVED rows, not merely profiles attempted.
                if saved < max_records:
                    self.log(
                        "CURRENT RESULT PAGE EXHAUSTED | "
                        f"processed={processed} | saved={saved}/{max_records} | "
                        f"skipped={skipped} | trying next results page"
                    )

                    if self._go_to_next_results_page():
                        stagnant = 0
                        continue

                self.log(
                    "No additional search-result page was available. "
                    f"Finished with {saved}/{max_records} saved rows."
                )
                break

            self.status(
                f"Saved {saved}/{max_records}; loading more prospects..."
            )

            self._scroll_search_results()
            self._wait_for_new_loaded_result(current_urls)

        # Before the final save, give any rows that were written while Qwen
        # was temporarily unavailable one final enrichment opportunity.
        if extraction_mode == "SMART":
            self._retry_pending_enrichment_rows(
                workbook=workbook,
                pending_rows=pending_enrichment_rows,
                source_label="Sales Navigator",
                validate_existing=True,
            )

        self._log_enrichment_audit(
            workbook,
            "Sales Navigator",
        )

        # Always force-save whatever was collected.
        workbook.save(force=True)

        elapsed = max(
            time.perf_counter() - run_started,
            0.001,
        )

        self.log(
            "FINAL PERF | "
            f"processed={processed} | "
            f"saved={saved}/{max_records} | "
            f"skipped={skipped} | "
            f"failed={failed} | "
            f"elapsed={elapsed:.1f}s | "
            f"rate={saved / elapsed:.2f} leads/sec | "
            f"company cache={len(self.company_cache)} | "
            f"Qwen calls={self.qwen_calls} | "
            f"Qwen time={self.qwen_seconds:.1f}s | "
            f"smart company cache={len(self.smart_company_cache)} | "
            f"smart job cache={len(self.smart_job_cache)}"
        )

        self.status(
            "Stopped"
            if self.stop_event.is_set()
            else "Finished"
        )

        self.callbacks["done"]()

    def _close_browser(self, close_enrichment=True):
        if close_enrichment:
            self._finish_enrichment_sessions()
        self._reset_public_company_enricher()
        try:
            if self.context:
                self.context.close()
        finally:
            self.context = None
            self.page = None
            self.google_maps_page = None

    def run(self):
        self.pw = sync_playwright().start()

        try:
            while True:
                name, payload = self.commands.get()

                if name == "shutdown":
                    self.stop_event.set()
                    self._close_browser()
                    return

                if name == "open_browser":
                    try:
                        self._open_login()
                    except Exception as exc:
                        self.status("Browser error")
                        self.log(f"Browser error: {exc}")

                elif name == "process":
                    self.pause_event.clear()
                    self.stop_event.clear()

                    try:
                        self._process_search(payload)
                    except Exception as exc:
                        try:
                            payload["workbook"].save(force=True)
                        except Exception as save_exc:
                            self.log(f"Recovery save failed: {save_exc}")
                        self.status("Error")
                        self.log(
                            f"Processing error: "
                            f"{type(exc).__name__}: {exc}"
                        )
                        self.callbacks["done"]()
                    finally:
                        self._finish_enrichment_sessions()
                        self._report_performance(payload['workbook'])

                elif name == "process_links":
                    self.pause_event.clear()
                    self.stop_event.clear()

                    try:
                        self._process_link_file(payload)
                    except Exception as exc:
                        try:
                            payload["workbook"].save(force=True)
                        except Exception as save_exc:
                            self.log(f"Recovery save failed: {save_exc}")
                        self.status("Error")
                        self.log(
                            f"Link Scraper error: "
                            f"{type(exc).__name__}: {exc}"
                        )
                        self.callbacks["done"]()
                    finally:
                        self._finish_enrichment_sessions()
                        self._report_performance(payload['workbook'])

                elif name == "pause":
                    self.pause_event.set()
                    self.status("Paused")

                elif name == "resume":
                    self.pause_event.clear()
                    self.status("Running")

                elif name == "stop":
                    self.stop_event.set()
                    self.pause_event.clear()

                elif name == "close_browser":
                    self.stop_event.set()
                    self.pause_event.clear()
                    self._close_browser()
                    self.status("Chrome closed")

        finally:
            try:
                self._close_browser()
            except Exception:
                pass

            try:
                self.pw.stop()
            except Exception:
                pass


class LeadExtractorApp(ctk.CTk):
    def __init__(self):
        super().__init__()

        ctk.set_appearance_mode("dark")
        ctk.set_default_color_theme("blue")

        self.title(f"{APP_NAME} v{APP_VERSION}")

        # VERSION 18 UI:
        # Fit the application inside the actual monitor instead of forcing a
        # taller window than some laptop screens can display.
        self.update_idletasks()
        screen_w = self.winfo_screenwidth()
        screen_h = self.winfo_screenheight()

        usable_w = max(screen_w - 70, 760)
        usable_h = max(screen_h - 110, 560)

        window_w = min(1420, usable_w)
        window_h = min(900, usable_h)

        x = max((screen_w - window_w) // 2, 0)
        y = max((screen_h - window_h) // 2, 0)

        self.geometry(f"{window_w}x{window_h}+{x}+{y}")

        min_w = min(980, max(screen_w - 40, 760))
        min_h = min(640, max(screen_h - 80, 560))
        self.minsize(min_w, min_h)

        self.output_path = ""
        self.link_input_path = ""
        self.workbook: WorkbookManager | None = None
        self.processing = False
        self.importing = False
        self.worksheet_window = None

        # Every run gets a persistent activity log beside the Excel workbook.
        self.activity_log_path = ""
        self.activity_log_lock = threading.Lock()

        self.callbacks = {
            "stage": lambda *args: self.after(0, self._update_stage, *args),
            "log": lambda msg: self.after(
                0,
                self._log,
                msg,
            ),
            "trace": self._write_detailed_trace,
            "status": lambda msg: self.after(
                0,
                self._set_status,
                msg,
            ),
            "progress": lambda *args: self.after(
                0,
                self._update_progress,
                *args,
            ),
            "record": lambda rec: self.after(
                0,
                self._show_record,
                rec,
            ),
            "harvest": lambda found, maximum: self.after(
                0,
                self._update_harvest,
                found,
                maximum,
            ),
            "done": lambda: self.after(
                0,
                self._processing_done,
            ),
        }

        self.worker = SearchWorker(self.callbacks)
        self.worker.start()

        self._build_ui()
        self._check_llm_async()

        self.protocol(
            "WM_DELETE_WINDOW",
            self._on_close,
        )

    def _open_worksheet(self):
        from tkinter import messagebox
        if self.processing:
            messagebox.showinfo('Worksheet', 'Stop the run before editing the output worksheet.')
            return
        if self.worksheet_window is not None and self.worksheet_window.winfo_exists():
            self.worksheet_window.lift()
            return
        if self._prepare_workbook():
            from worksheet_editor import WorksheetEditor
            self.worksheet_window = WorksheetEditor(self, self.workbook)

    def _reset_stages(self):
        from scraper_ui import reset_stages
        reset_stages(self)

    def _update_stage(self, name, state, detail=''):
        from scraper_ui import update_stage
        update_stage(self, name, state, detail)

    def _preview_link_input(self, path):
        # Parsing stays off the Tk thread so large uploads do not freeze the UI.
        from link_import import load_link_inputs
        self.importing = True
        self.start_btn.configure(state='disabled')
        self.link_file_button.configure(state='disabled')
        self._update_stage('import', 'running', 'Reading file')
        result_queue = queue.Queue()
        def read():
            try:
                links = load_link_inputs(path, progress=lambda text: result_queue.put(('progress', text)))
                valid = len({link.canonical for link in links if link.canonical})
                invalid = sum(not link.canonical for link in links)
                result_queue.put(('done', f'{valid:,} unique profile links · {invalid:,} invalid URLs'))
            except Exception as exc:
                result_queue.put(('error', str(exc)))
        def poll():
            try:
                while True:
                    state, text = result_queue.get_nowait()
                    self._update_stage('import', 'running' if state == 'progress' else state, text)
                    if state != 'progress':
                        self.importing = False
                        self.start_btn.configure(state='normal')
                        self.link_file_button.configure(state='normal')
                        self.link_count_label.configure(text=text)
                        self._log('File import: ' + text)
                        if state == 'error':
                            self.link_input_path = ''
                        return
            except queue.Empty:
                pass
            self.after(100, poll)
        threading.Thread(target=read, daemon=True, name='link-file-preview').start()
        self.after(100, poll)

    def _build_ui(self):
        from scraper_ui import build_ui
        build_ui(self, APP_VERSION)

    def _configure_activity_log(
        self,
        output_path: str,
        reset: bool = False,
    ) -> None:
        output = Path(output_path)
        log_path = output.with_name(
            output.stem + "_activity.log"
        )

        changed = (
            str(log_path) != self.activity_log_path
        )

        self.activity_log_path = str(log_path)

        if reset or changed:
            try:
                log_path.parent.mkdir(
                    parents=True,
                    exist_ok=True,
                )

                header = {
                    "event": "RUN_LOG_INITIALIZED",
                    "timestamp": time.strftime(
                        "%Y-%m-%d %H:%M:%S"
                    ),
                    "app": APP_NAME,
                    "version": APP_VERSION,
                    "output_workbook": str(output),
                    "local_llm_base_url": LLM_BASE_URL,
                    "configured_llm_model": LLM_FALLBACK_MODEL,
                }

                with self.activity_log_lock:
                    with log_path.open(
                        "a",
                        encoding="utf-8",
                    ) as handle:
                        handle.write(
                            json.dumps(
                                header,
                                ensure_ascii=False,
                            )
                            + "\n"
                        )

            except Exception:
                # The GUI must remain usable even if diagnostics cannot be
                # written due to folder permissions.
                pass

    def _append_activity_line(
        self,
        line: str,
    ) -> None:
        if not self.activity_log_path:
            return

        try:
            path = Path(
                self.activity_log_path
            )
            path.parent.mkdir(
                parents=True,
                exist_ok=True,
            )

            with self.activity_log_lock:
                with path.open(
                    "a",
                    encoding="utf-8",
                ) as handle:
                    handle.write(line + "\n")

        except Exception:
            pass

    def _write_detailed_trace(
        self,
        category: str,
        payload: dict[str, Any],
    ) -> None:
        entry = {
            "timestamp": time.strftime(
                "%Y-%m-%d %H:%M:%S"
            ),
            "category": category,
            **payload,
        }

        try:
            line = json.dumps(
                entry,
                ensure_ascii=False,
                default=str,
            )
        except Exception:
            line = json.dumps(
                {
                    "timestamp": entry["timestamp"],
                    "category": category,
                    "event": "TRACE_SERIALIZATION_ERROR",
                    "payload_repr": repr(payload),
                },
                ensure_ascii=False,
            )

        self._append_activity_line(line)

    def _log(self, message: str):
        stamp = time.strftime("%H:%M:%S")

        self._append_activity_line(
            json.dumps(
                {
                    "timestamp": time.strftime(
                        "%Y-%m-%d %H:%M:%S"
                    ),
                    "category": "ACTIVITY",
                    "message": message,
                },
                ensure_ascii=False,
            )
        )

        self.log_box.configure(state="normal")
        self.log_box.insert(
            "end",
            f"[{stamp}] {message}\n",
        )
        self.log_box.see("end")
        self.log_box.configure(state="disabled")

    def _set_status(self, message: str):
        self.status_label.configure(text=message)

        if (
            "Chrome" in message
            or "Browser" in message
            or "sign-in" in message.lower()
        ):
            self.browser_status.configure(
                text=f"Browser: {message}"
            )

    def _check_llm_async(self):
        def work():
            client = LocalLLMClient()
            ok, detail = client.health_check()

            self.after(
                0,
                self._llm_check_done,
                ok,
                detail,
            )

        threading.Thread(
            target=work,
            daemon=True,
        ).start()

    def _llm_check_done(
        self,
        ok: bool,
        detail: str,
    ):
        if ok:
            self.llm_status.configure(
                text=f"LLM: Connected ({detail})"
            )
            self._log(
                f"Local LLM connected: {detail}"
            )
        else:
            self.llm_status.configure(
                text="LLM: Offline"
            )
            self._log(
                f"Local LLM not reachable at {LLM_BASE_URL}. "
                "Start your llamafile server before processing."
            )

    def _open_browser(self):
        self.worker.send("open_browser")
        self.browser_status.configure(
            text="Browser: Launching..."
        )

    def _on_process_mode_change(self, mode: str):
        """
        Show only the input controls relevant to the selected process.
        This keeps the left panel compact and avoids disabled controls taking
        up vertical space.
        """
        if mode == "Link Scraper":
            if hasattr(self, "sales_input_frame"):
                self.sales_input_frame.pack_forget()

            if hasattr(self, "link_input_frame"):
                self.link_input_frame.pack(
                    fill="x",
                    padx=12,
                    pady=(0, 10),
                )

            if hasattr(self, "start_btn"):
                self.start_btn.configure(
                    text="START LINK SCRAPER"
                )

            if hasattr(self, "harvest_label"):
                self.harvest_label.configure(
                    text="Link Scraper saved: 0"
                )

            if hasattr(self, "workspace_note"):
                self.workspace_note.configure(
                    text="Link Scraper activity and extraction details"
                )

        else:
            if hasattr(self, "link_input_frame"):
                self.link_input_frame.pack_forget()

            if hasattr(self, "sales_input_frame"):
                self.sales_input_frame.pack(
                    fill="x",
                    padx=12,
                    pady=(0, 10),
                )

            if hasattr(self, "start_btn"):
                self.start_btn.configure(
                    text="START SALES NAVIGATOR"
                )

            if hasattr(self, "harvest_label"):
                self.harvest_label.configure(
                    text="Leads copied from search: 0"
                )

            if hasattr(self, "workspace_note"):
                self.workspace_note.configure(
                    text="Sales Navigator activity and extraction details"
                )

    def _choose_link_input(self):
        from tkinter import filedialog

        if self.processing or self.importing:
            return

        path = filedialog.askopenfilename(
            title='Choose profile-link input file',
            filetypes=[
                ('Excel / CSV', '*.xlsx *.xlsm *.csv'),
                ('Excel workbook', '*.xlsx *.xlsm'),
                ('CSV file', '*.csv'),
                ('All files', '*.*'),
            ],
        )

        if not path:
            return

        self.link_input_path = path
        self.link_file_label.configure(text=path)
        self._log(f'Link Scraper input selected: {path}')
        self._preview_link_input(path)

    def _choose_output(self):
        from tkinter import filedialog

        if self.processing:
            return

        path = filedialog.asksaveasfilename(
            title="Select workbook to append, or enter a new filename",
            defaultextension=".xlsx",
            filetypes=[
                ("Excel workbook", "*.xlsx"),
            ],
            initialfile=(
                "Link_Scraper_Results.xlsx"
                if (self.process_mode.get() or "Sales Navigator") == "Link Scraper"
                else "Sales_Navigator_Results.xlsx"
            ),
        )

        if not path:
            return

        self.output_path = path
        self.output_label.configure(text=path)

        self._configure_activity_log(
            path,
            reset=True,
        )

        self._log(
            f"Output selected: {path}"
        )
        self._log(
            "Detailed activity log: "
            f"{self.activity_log_path}"
        )

    def _prepare_workbook(self) -> bool:
        from tkinter import messagebox

        if not self.output_path:
            self._choose_output()

            if not self.output_path:
                return False

        path = Path(self.output_path)

        expected_log = str(
            path.with_name(
                path.stem + "_activity.log"
            )
        )

        if self.activity_log_path != expected_log:
            self._configure_activity_log(
                self.output_path,
                reset=True,
            )

        if path.exists():
            self._log("Appending to existing workbook; existing prospects will be skipped.")

        try:
            self.workbook = WorkbookManager(
                self.output_path, checkpoint_every=EXCEL_CHECKPOINT_EVERY
            )
            return True

        except PermissionError:
            messagebox.showerror(
                "Excel File Is Open",
                "Close the output workbook in Excel and try again.",
            )
            return False

        except Exception as exc:
            messagebox.showerror(
                "Output Error",
                str(exc),
            )
            return False

    def _validate_search_url(self, value: str) -> bool:
        value = (value or "").strip().lower()

        return (
            value.startswith(("http://", "https://"))
            and "linkedin.com/sales/" in value
        )

    def _build_offline_benchmark(self):
        from tkinter import filedialog, messagebox
        if self.processing:
            messagebox.showinfo('Benchmark', 'Stop the active run before benchmarking its saved workbook.')
            return
        try:
            size = int(self.benchmark_sample_entry.get().strip())
            if not 1 <= size <= 200:
                raise ValueError
        except ValueError:
            messagebox.showwarning('Benchmark', 'Enter a sample size from 1 to 200 (150 recommended).')
            return
        source = filedialog.askopenfilename(title='Choose saved results for offline benchmark',
            initialdir=str(Path(self.output_path).parent if self.output_path else Path(__file__).resolve().parent / 'Output'),
            filetypes=[('Saved results', '*.xlsx *.csv *.tsv')])
        if not source:
            return
        self.benchmark_button.configure(state='disabled')
        self._update_stage('benchmark', 'running')
        self._log('Building offline benchmark from ' + source)

        def work():
            try:
                from benchmark import build_benchmark
                result = build_benchmark(source, Path(source).parent / 'benchmarks', sample_size=size)
                error = None
            except Exception as exc:
                result, error = None, str(exc)
            try:
                self.after(0, self._benchmark_finished, result, error)
            except Exception:
                # Closing the application does not invalidate the offline report.
                pass
        threading.Thread(target=work, daemon=True).start()

    def _benchmark_finished(self, result, error):
        from tkinter import messagebox
        self.benchmark_button.configure(state='normal')
        self._update_stage('benchmark', 'error' if error else 'done')
        if error:
            self._log('Benchmark failed: ' + error)
            messagebox.showerror('Benchmark', error)
            return
        self._log('Offline benchmark saved: ' + result['output_dir'])
        messagebox.showinfo('Benchmark complete',
            'Saved the sample, review template and field report to:\n\n' + result['output_dir']
            + '\n\nAccuracy remains unknown until the review template is completed.')

    def _start_processing(self):
        from tkinter import messagebox

        if self.processing or self.importing:
            return

        mode = self.process_mode.get() or 'Sales Navigator'

        if mode == 'Sales Navigator':
            search_url = self.search_url_entry.get().strip()

            if not self._validate_search_url(search_url):
                messagebox.showwarning(
                    'Invalid Search URL',
                    (
                        'Paste one LinkedIn Sales Navigator URL, for example:\n\n'
                        'https://www.linkedin.com/sales/search/people?...'
                    ),
                )
                return
        else:
            search_url = ''

            if not self.link_input_path:
                self._choose_link_input()

            if not self.link_input_path or self.importing:
                return

        try:
            max_records = (None if mode == 'Link Scraper' and self.all_links_option.get()
                           else int(self.max_records_entry.get().strip()))
            if max_records is not None and not 1 <= max_records <= 1048575:
                raise ValueError
        except ValueError:
            messagebox.showwarning(
                'Invalid Maximum',
                'Enter a target between 1 and 1,048,575, or choose All imported links.',
            )
            return

        if not self._prepare_workbook():
            return

        if self.worksheet_window is not None and self.worksheet_window.winfo_exists():
            self.worksheet_window.destroy()
        self._reset_stages()
        self.processing = True
        self.start_btn.configure(state='disabled')
        self.link_file_button.configure(state='disabled')
        self.worksheet_button.configure(state='disabled')
        self.process_mode.configure(state='disabled')
        for control in self.run_controls.values():
            control.configure(state='normal')
        self.run_controls['Resume'].configure(state='disabled')
        self.progress.set(0)
        self.harvest_label.configure(text='Leads saved: 0')
        self._update_progress(0, 0, 0, 0, 0)

        if mode == 'Link Scraper':
            payload = {
                'input_path': self.link_input_path,
                'max_records': max_records,
                'workbook': self.workbook,
                'extraction_mode': self.extraction_mode.get() or 'FAST',
                'email_domain_checks': bool(self.email_domain_checks.get()),
                'headless': bool(self.headless_option.get()),
                'enrichment_sessions': 2 if self.parallel_enrichment.get() else 1,
                'prefer_local_office': bool(self.local_office_option.get()),
                'persistent_cache': True,
                'refresh_profiles': bool(self.refresh_profiles_option.get()),
            }
            self.worker.send('process_links', payload)
            self._log(
                f'Link Scraper started in {self.extraction_mode.get()} mode. '
                f"Saved-lead target: {format(max_records, ',') if max_records is not None else 'All imported links'}"
            )
        else:
            payload = {
                'search_url': search_url,
                'max_records': max_records,
                'workbook': self.workbook,
                'extraction_mode': self.extraction_mode.get() or 'FAST',
                'email_domain_checks': bool(self.email_domain_checks.get()),
                'headless': bool(self.headless_option.get()),
                'enrichment_sessions': 2 if self.parallel_enrichment.get() else 1,
                'prefer_local_office': bool(self.local_office_option.get()),
                'persistent_cache': True,
                'refresh_profiles': bool(self.refresh_profiles_option.get()),
            }
            self.worker.send('process', payload)
            self._log(
                f'Processing started in {self.extraction_mode.get()} mode. '
                f'Maximum leads: {max_records:,}'
            )

    def _pause(self):
        if self.processing:
            self.worker.send("pause")
            self.run_controls["Pause"].configure(state="disabled")
            self.run_controls["Resume"].configure(state="normal")
            self._log("Pause requested.")

    def _resume(self):
        if self.processing:
            self.worker.send("resume")
            self.run_controls["Pause"].configure(state="normal")
            self.run_controls["Resume"].configure(state="disabled")
            self._log("Resume requested.")

    def _stop(self):
        if self.processing:
            self.worker.send("stop")
            self._log(
                "Stop requested. Current Excel progress will be preserved."
            )

    def _update_harvest(
        self,
        found: int,
        maximum: int,
    ):
        mode = self.process_mode.get() if hasattr(self, "process_mode") else "Sales Navigator"
        prefix = "Link Scraper saved" if mode == "Link Scraper" else "Leads copied from search"
        self.harvest_label.configure(
            text=f"{prefix}: {found:,} / {maximum:,} maximum"
        )

    def _update_progress(
        self,
        processed: int,
        saved: int,
        duplicates: int,
        failed: int,
        total: int,
    ):
        # Maximum is a saved-lead target. Some examined prospects can be
        # skipped because their Present company does not match the card company.
        remaining = max(
            total - saved,
            0,
        )

        self.counter_labels[
            "Processed"
        ].configure(text=str(processed))

        self.counter_labels[
            "Saved"
        ].configure(text=str(saved))

        self.counter_labels[
            "Duplicates"
        ].configure(text=str(duplicates))

        self.counter_labels[
            "Failed"
        ].configure(text=str(failed))

        self.counter_labels[
            "Remaining"
        ].configure(text=str(remaining))

        self.progress.set(
            (saved / total)
            if total
            else 0
        )

    def _show_record(
        self,
        record: dict[str, str],
    ):
        keys = [
            "First Name",
            "Last Name",
            "Business Email ID",
            "Email Status",
            "Job Title",
            "Job Level",
            "Job Function",
            "Company Name",
            "Tenure",
            "City",
            "State",
            "Country",
            "Industry",
            "Description",
            "Prospect LinkedIn Link",
        ]

        text = "\n".join(
            f"{key}: {record.get(key, '-')}"
            for key in keys
        )

        self.preview.configure(state="normal")
        self.preview.delete("1.0", "end")
        self.preview.insert("1.0", text)
        self.preview.configure(state="disabled")

    def _processing_done(self):
        for name, widgets in self.stage_widgets.items():
            if widgets['running']:
                self._update_stage(name, 'stopped', 'Run ended')
        self.processing = False
        self.start_btn.configure(state="normal")
        self.link_file_button.configure(state="normal")
        self.worksheet_button.configure(state="normal")
        self.process_mode.configure(state="normal")
        for control in self.run_controls.values():
            control.configure(state="disabled")

        if self.output_path:
            self._log(
                f"Workbook saved: {self.output_path}"
            )

        if self.activity_log_path:
            self._log(
                f"Detailed activity log saved: {self.activity_log_path}"
            )

        self._log("Processing finished.")

    def _clear_session(self):
        from tkinter import messagebox

        if self.processing:
            messagebox.showwarning(
                "Busy",
                "Stop processing before clearing the browser session.",
            )
            return

        confirm = messagebox.askyesno(
            "Clear browser session?",
            (
                "This removes the dedicated Chrome profile used by this app, "
                "including its saved website session cookies.\n\nContinue?"
            ),
        )

        if not confirm:
            return

        self.worker.send("close_browser")
        time.sleep(0.5)

        try:
            if CHROME_PROFILE_DIR.exists():
                shutil.rmtree(
                    CHROME_PROFILE_DIR
                )

            self._log(
                "Dedicated browser session cleared."
            )

            self.browser_status.configure(
                text="Browser: Session cleared"
            )

        except Exception as exc:
            messagebox.showerror(
                "Could Not Clear Session",
                (
                    "Close the Chrome window opened by the app, "
                    f"then try again.\n\n{exc}"
                ),
            )

    def _on_close(self):
        try:
            self.worker.send("shutdown")
        except Exception:
            pass

        # Cancel pending progress animations and UI callbacks before destroying Tk.
        for callback in self.tk.call('after', 'info'):
            self.tk.call('after', 'cancel', callback)
        self.destroy()


if __name__ == "__main__":
    app = LeadExtractorApp()
    app.mainloop()
